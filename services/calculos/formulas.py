"""Deteccion de patrones de formula: valores implicitos, patrones por fila/columna, hardcodes, promedios simples."""
from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

from services.calculos.excel_utils import (
    MAX_HALLAZGOS,
    _casi_igual,
    _col_label,
    _header_col,
    _norm_label,
    _parece_fila_resultado,
    _rango_contiguo,
    _rango_ref,
    _row_label,
    _short,
)


def _hallazgo_implicito(
    sheet: str,
    target: dict[str, Any],
    col: int,
    actual: float,
    esperado: float,
    regla: str,
    componentes: list[str],
) -> dict:
    return {
        "modo": "formula_implicita",
        "descripcion": (
            f"En la hoja '{sheet}', {target['label']} parece seguir la regla '{regla}', "
            f"pero en { _header_col(target['ws'], col) } declara {actual:g} y el patrón esperaría {esperado:g}. "
            "El archivo no necesita tener fórmula: la relación se infiere por los valores de la serie."
        ),
        "severidad": "media",
        "evidencia": {
            "regla": "formula_implicita_rota",
            "hoja": sheet,
            "celda": f"{get_column_letter(col)}{target['row']}",
            "fila_detalle": target["label"],
            "columna_detalle": _header_col(target["ws"], col),
            "patron_inferido": regla,
            "componentes": componentes,
            "valor_actual": round(actual, 2),
            "valor_esperado": round(esperado, 2),
        },
    }


def _evaluar_patron(target: dict[str, Any], esperado_por_col: dict[int, float]) -> tuple[int, list[tuple[int, float, float]]]:
    comunes = sorted(set(target["values"]) & set(esperado_por_col))
    if len(comunes) < 4:
        return 0, []
    ok = 0
    fallas = []
    for col in comunes:
        actual = target["values"][col]
        esperado = esperado_por_col[col]
        if _casi_igual(actual, esperado):
            ok += 1
        else:
            fallas.append((col, actual, esperado))
    return ok, fallas


def _analizar_patrones_pares(ws, rows: list[dict[str, Any]]) -> list[dict]:
    hallazgos: list[dict] = []
    sheet = ws.title
    vistos: set[tuple[int, str]] = set()
    for target in rows:
        target["ws"] = ws
        # Una coincidencia numérica parcial no define la composición de un
        # total. Con más de dos componentes posibles se requiere una regla
        # explícita o el control de bloque completo, no adivinar un par.
        if re.match(r"^(?:sub)?total(?:\b|_)", _norm_label(ws, target["row"])):
            componentes = [r for r in rows if r["row"] < target["row"] and not _parece_fila_resultado(ws, r["row"])]
            if len(componentes) > 2:
                continue
        for i, a in enumerate(rows):
            if a["row"] == target["row"] or abs(a["row"] - target["row"]) > 8:
                continue
            for b in rows[i + 1:]:
                if b["row"] == target["row"] or abs(b["row"] - target["row"]) > 8:
                    continue
                if not _parece_fila_resultado(target["ws"], target["row"]):
                    continue
                operaciones = [
                    ("{target} = {a} + {b}", {c: a["values"][c] + b["values"][c] for c in set(a["values"]) & set(b["values"])}),
                    ("{target} = {a} - {b}", {c: a["values"][c] - b["values"][c] for c in set(a["values"]) & set(b["values"])}),
                    ("{target} = {b} - {a}", {c: b["values"][c] - a["values"][c] for c in set(a["values"]) & set(b["values"])}),
                ]
                for template, esperado in operaciones:
                    ok, fallas = _evaluar_patron(target, esperado)
                    total = ok + len(fallas)
                    if not fallas or total < 4 or ok < max(4, int(total * 0.7)):
                        continue
                    regla = template.format(target=target["label"], a=a["label"], b=b["label"])
                    key = (target["row"], regla)
                    if key in vistos:
                        continue
                    vistos.add(key)
                    col, actual, exp = fallas[0]
                    hallazgos.append(_hallazgo_implicito(
                        sheet, target, col, actual, exp, regla, [a["label"], b["label"]]
                    ))
                    if len(hallazgos) >= MAX_HALLAZGOS:
                        return hallazgos
    return hallazgos


def _hallazgo_patron(
    sheet: str,
    eje: str,
    celda: str,
    etiqueta: str,
    formula: str,
    esperado: str,
    severidad: str = "media",
    impacto_decision: bool = False,
) -> dict:
    return {
        "modo": "formula",
        "descripcion": (
            f"En la hoja '{sheet}', {celda} ({etiqueta}) tiene una fórmula que se aparta del patrón "
            f"de la {eje}. Conviene revisar si hubo referencia corrida, rango incompleto o edición manual."
        ),
        "severidad": severidad,
        "evidencia": {
            "regla": "formula_fuera_de_patron",
            "hoja": sheet,
            "celda": celda,
            "fila_detalle": etiqueta,
            "formula": formula,
            "patron_esperado": esperado,
            "eje": eje,
            "impacta_decision": impacto_decision,
        },
    }


def _hallazgo_hardcode(
    sheet: str,
    eje: str,
    celda: str,
    etiqueta: str,
    valor: Any,
    esperado: str,
    severidad: str = "media",
    impacto_decision: bool = False,
) -> dict:
    return {
        "modo": "formula",
        "descripcion": (
            f"En la hoja '{sheet}', {celda} ({etiqueta}) contiene un valor fijo dentro de una {eje} "
            "donde las celdas vecinas usan fórmulas. Puede ser una carga manual válida; conviene confirmar su origen y el criterio de cálculo."
        ),
        "severidad": severidad,
        "evidencia": {
            "regla": "valor_fijo_en_serie_de_formulas",
            "hoja": sheet,
            "celda": celda,
            "fila_detalle": etiqueta,
            "valor": _short(valor),
            "patron_esperado": esperado,
            "eje": eje,
            "impacta_decision": impacto_decision,
        },
    }


def _hallazgo_promedio_simple(ws, cell, formula: str) -> dict:
    celda = f"{get_column_letter(cell.column)}{cell.row}"
    rango = _rango_ref(formula)
    detalle_rango = f" sobre {rango}" if rango else ""
    return {
        "modo": "pertinencia_calculo",
        "descripcion": (
            f"En la hoja '{ws.title}', {celda} usa un promedio simple{detalle_rango} en un contexto "
            "donde se identificaron indicios de ponderación en el archivo. "
            "La fórmula puede estar bien escrita, pero conviene justificar si corresponde promedio simple o ponderado."
        ),
        "severidad": "media",
        "evidencia": {
            "regla": "promedio_simple_con_posibles_ponderadores",
            "hoja": ws.title,
            "celda": celda,
            "formula": formula,
            "rango": rango,
            "fila_detalle": _row_label(ws, cell.row),
            "columna_detalle": _col_label(ws, cell.column),
        },
    }


def _hallazgo_promedio_simple_grupo(ws, cells: list[tuple[Any, str]]) -> dict:
    ordenadas = sorted(cells, key=lambda item: item[0].row)
    first_cell, formula = ordenadas[0]
    ref = _rango_contiguo(ordenadas)
    rango_formula = _rango_ref(formula)
    detalle_rango = f" sobre {rango_formula}" if rango_formula else ""
    cantidad = len(ordenadas)
    filas_detalle = [_row_label(ws, cell.row) for cell, _ in ordenadas[:6]]
    return {
        "modo": "pertinencia_calculo",
        "descripcion": (
            f"En la hoja '{ws.title}', la columna {_col_label(ws, first_cell.column)} usa promedio simple "
            f"en {ref} ({cantidad} celdas){detalle_rango}, dentro de un contexto donde aparecen ponderadores "
            "posibles (créditos, horas, pesos o carga horaria). Conviene justificar una vez el criterio de cálculo, "
            "sin repetir la misma señal fila por fila."
        ),
        "severidad": "media",
        "evidencia": {
            "regla": "promedio_simple_con_posibles_ponderadores",
            "hoja": ws.title,
            "celda": ref,
            "formula": formula,
            "rango": rango_formula,
            "cantidad_celdas": cantidad,
            "filas_detalle": filas_detalle,
            "columna_detalle": _col_label(ws, first_cell.column),
            "agrupado": True,
        },
    }
