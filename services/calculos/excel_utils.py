"""Helpers de bajo nivel sobre celdas/filas/columnas de openpyxl, constantes regex compartidas."""
from __future__ import annotations

import re
from collections import Counter
from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

EXCEL_EXTS = {"xlsx", "xlsm"}
REF_RE = re.compile(r"(?<![A-Za-z0-9_])(\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z0-9_])")
MAX_HALLAZGOS = 8
TOLERANCIA_VALOR = 2.0
TOLERANCIA_RELATIVA = 0.001
COLUMNA_NO_CALCULO_IMPLICITO_RE = re.compile(
    r"(^|\b)(id|codigo|c[oó]digo|cod|dni|cuit|cuil|legajo|expediente|archivo|"
    r"cohorte|año|anio|year|periodo|per[ií]odo|mes|fecha|comisi[oó]n|curso|"
    r"estudiante|alumno|persona|nombre|apellido|proveedor(?:es)?|contratistas?|adjudicatari[oa]s?|beneficiari[oa]s?)($|\b|_)",
    re.I,
)
FILA_RESULTADO_IMPLICITO_RE = re.compile(
    r"total|subtotal|saldo|superavit|superávit|deficit|déficit|caja|resultado|"
    r"diferencia|brecha|neto|final|remanente|acumulado|balance",
    re.I,
)
DECISION_RE = re.compile(
    r"condici[oó]n|estado|resultado|decisi[oó]n|aprob|promoc|regular|libre|alerta|dictamen|clasificaci[oó]n|control|validaci[oó]n|consistencia",
    re.I,
)
CALCULO_CRITICO_RE = re.compile(
    r"promedio|ponderado|nota|calificaci[oó]n|total|saldo|resultado|puntaje|score|brecha|diferencia|control|validaci[oó]n|consistencia",
    re.I,
)
ACTOR_DETALLE_RE = re.compile(
    r"proveedor(?:es)?|contratistas?|adjudicatari[oa]s?|beneficiari[oa]s?|"
    r"vendors?|suppliers?|parties|raz[oó]n\s*social|cuit|cuil|dni",
    re.I,
)
ACTOR_AGREGADO_PERMITIDO_RE = re.compile(
    r"deudas?\s+con\s+proveedor(?:es)?|pagos?\s+a\s+proveedor(?:es)?|"
    r"acreedores?\s+comerciales?",
    re.I,
)
ESTADO_CONTROL_RE = re.compile(r"^(?:estado|status|resultado|validaci[oó]n|control)$", re.I)
ESTADO_REVISION_RE = re.compile(r"^(?:revisar|alerta|fail|failed|fall[oó]|error|inconsistente|observado|no ok)$", re.I)
CIERRE_ETIQUETA_RE = re.compile(r"total|promedio\s*general|subtotal|acumulado", re.I)
PONDERADOR_RE = re.compile(r"ponder|peso|credit|cr[eé]dito|hora|carga horaria|coef|unidad|cantidad|volumen|casos|respuestas?|respondidas?|encuestas?", re.I)
PROMEDIO_SIMPLE_RE = re.compile(r"\b(AVERAGE|PROMEDIO|MEDIA)\s*\(", re.I)


def _extension(filename: str) -> str:
    return (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _non_empty_value(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _short(value: Any, max_len: int = 80) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", " ", text)
    return text if len(text) <= max_len else text[: max_len - 3].rstrip() + "..."


def _row_label(ws, row: int) -> str:
    base = f"fila {row}"
    for col in range(1, min(ws.max_column, 4) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
            return f"{base} ({get_column_letter(col)}: {_short(value)})"
    return base


def _col_label(ws, col: int) -> str:
    letter = get_column_letter(col)
    for row in range(1, min(ws.max_row, 12) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
            return f"{letter} ({_short(value)})"
    return letter


def _es_actor_detalle(value: Any) -> bool:
    texto = str(value or "")
    if ACTOR_AGREGADO_PERMITIDO_RE.search(texto):
        return False
    return bool(ACTOR_DETALLE_RE.search(texto))


def _numero(value: Any) -> float | None:
    if value is None or isinstance(value, bool) or _is_formula(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or re.search(r"[A-Za-zÀ-ÿ]", text):
        return None
    text = re.sub(r"[^\d,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    comma_decimal = "," in text and ("." not in text or text.rfind(",") > text.rfind("."))
    if comma_decimal:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _casi_igual(a: float, b: float) -> bool:
    return abs(a - b) <= max(TOLERANCIA_VALOR, abs(b) * TOLERANCIA_RELATIVA)


def _header_col(ws, col: int) -> str:
    letter = get_column_letter(col)
    for row in range(1, min(ws.max_row, 12) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value):
            return f"{letter}: {_short(value)}"
    return letter


def _texto_columna(ws, col: int, max_rows: int = 12) -> str:
    partes: list[str] = []
    for row in range(1, min(ws.max_row, max_rows) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
            partes.append(str(value))
    return " ".join(partes)


def _columna_para_calculo_implicito(ws, col: int) -> bool:
    texto = _texto_columna(ws, col)
    if COLUMNA_NO_CALCULO_IMPLICITO_RE.search(texto):
        return False
    nums: list[float] = []
    for row in range(1, ws.max_row + 1):
        n = _numero(ws.cell(row=row, column=col).value)
        if n is not None:
            nums.append(n)
    if len(nums) >= 3:
        distintos = {round(n, 8) for n in nums}
        if len(distintos) <= 1:
            return False
        if all(float(n).is_integer() and 1900 <= n <= 2100 for n in nums):
            return False
    return True


def _parece_fila_resultado(ws, row: int) -> bool:
    return bool(FILA_RESULTADO_IMPLICITO_RE.search(_norm_label(ws, row)))


def _columna_marca(ws, col: int, patron: re.Pattern[str]) -> bool:
    return bool(patron.search(_texto_columna(ws, col)))


def _formula_refiere_celda(formula: str, row: int, col: int) -> bool:
    letra = get_column_letter(col)
    return bool(re.search(rf"(?<![A-Z0-9_])\$?{letra}\$?{row}(?![A-Z0-9_])", str(formula), re.I))


def _impacta_decision(ws, row: int, col: int) -> bool:
    if _columna_marca(ws, col, DECISION_RE):
        return True
    if not _columna_marca(ws, col, CALCULO_CRITICO_RE):
        return False
    for decision_col in range(1, ws.max_column + 1):
        if decision_col == col or not _columna_marca(ws, decision_col, DECISION_RE):
            continue
        value = ws.cell(row=row, column=decision_col).value
        if _is_formula(value) and _formula_refiere_celda(value, row, col):
            return True
    return False


def _numeric_rows(ws, min_vals: int = 4) -> list[dict[str, Any]]:
    rows = []
    for row in range(1, ws.max_row + 1):
        vals: dict[int, float] = {}
        for col in range(2, ws.max_column + 1):
            if not _columna_para_calculo_implicito(ws, col):
                continue
            n = _numero(ws.cell(row=row, column=col).value)
            if n is not None:
                vals[col] = n
        if len(vals) >= min_vals:
            rows.append({"row": row, "label": _row_label(ws, row), "values": vals})
    return rows


def _norm_label(ws, row: int) -> str:
    text = str(ws.cell(row=row, column=1).value or "").strip().lower()
    text = re.sub(r"[áà]", "a", text)
    text = re.sub(r"[éè]", "e", text)
    text = re.sub(r"[íì]", "i", text)
    text = re.sub(r"[óò]", "o", text)
    text = re.sub(r"[úù]", "u", text)
    return re.sub(r"\s+", " ", text)


def _normalize_formula(formula: str, row: int, col: int) -> str:
    """Fórmula en notación R/C: lo relativo, con signo (R+0C-4); lo fijo con «$», sin signo (R2C3).

    Una referencia fija ($C$2, la fila de pesos de un promedio ponderado) es la misma en todas las
    celdas del patrón: si se tratara como relativa, cambiaría de una fila a otra y el patrón no se
    reconocería nunca.
    """
    def repl(match: re.Match[str]) -> str:
        ref = match.group(1)
        m = re.fullmatch(r"(\$?)([A-Z]{1,3})(\$?)(\d+)", ref.upper())
        if not m:
            return ref.replace("$", "")
        try:
            rr, cc = coordinate_to_tuple(m.group(2) + m.group(4))
        except ValueError:
            return ref.replace("$", "")
        fila = f"R{rr}" if m.group(3) else f"R{rr - row:+d}"
        columna = f"C{cc}" if m.group(1) else f"C{cc - col:+d}"
        return fila + columna

    return REF_RE.sub(repl, formula.upper()).replace("$", "")


def _formula_cells_row(ws, row: int) -> list[tuple[int, str, str]]:
    cells = []
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row, column=col).value
        if _is_formula(value):
            cells.append((col, value, _normalize_formula(value, row, col)))
    return cells


def _formula_cells_col(ws, col: int) -> list[tuple[int, str, str]]:
    cells = []
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=col).value
        if _is_formula(value):
            cells.append((row, value, _normalize_formula(value, row, col)))
    return cells


def _formula_referencia_bloque_propio(formula: str, eje_fijo: int, min_pos: int, max_pos: int, es_columna: bool) -> bool:
    """True si la formula referencia, dentro de su propio eje (misma columna
    si es_columna, misma fila si no), un rango que cubre [min_pos, max_pos]
    -- las posiciones de las celdas del patron mayoritario."""
    posiciones = []
    for m in REF_RE.finditer(formula.upper()):
        try:
            rr, cc = coordinate_to_tuple(m.group(1).replace("$", ""))
        except ValueError:
            continue
        if es_columna and cc == eje_fijo:
            posiciones.append(rr)
        elif not es_columna and rr == eje_fijo:
            posiciones.append(cc)
    if not posiciones:
        return False
    return min(posiciones) <= min_pos and max(posiciones) >= max_pos


def _hay_espacio_de_separacion(ws, eje_fijo: int, pos_celda: int, min_patron: int, max_patron: int, es_columna: bool) -> bool:
    if pos_celda > max_patron:
        gap_pos = max_patron + 1
    elif pos_celda < min_patron:
        gap_pos = min_patron - 1
    else:
        return False
    if gap_pos == pos_celda:
        return False
    value = ws.cell(row=gap_pos, column=eje_fijo).value if es_columna else ws.cell(row=eje_fijo, column=gap_pos).value
    return not _non_empty_value(value)


def _es_cierre_legitimo(ws, formula: str, eje_fijo: int, pos_celda: int, posiciones_patron: list[int], es_columna: bool) -> bool:
    """Distingue una fila/columna de cierre (total, promedio general,
    subtotal, acumulado que resume el bloque de patron) de una celda
    genuinamente fuera de patron: la formula debe referenciar, dentro de su
    propio eje, un rango que cubra el bloque de celdas del patron mayoritario,
    y ademas estar separada de ese bloque por un hueco vacio o tener una
    etiqueta de cierre."""
    min_patron, max_patron = min(posiciones_patron), max(posiciones_patron)
    if not _formula_referencia_bloque_propio(formula, eje_fijo, min_patron, max_patron, es_columna):
        return False
    if _hay_espacio_de_separacion(ws, eje_fijo, pos_celda, min_patron, max_patron, es_columna):
        return True
    etiqueta = _row_label(ws, pos_celda) if es_columna else _col_label(ws, pos_celda)
    return bool(CIERRE_ETIQUETA_RE.search(etiqueta))


def _columna_estado(ws) -> tuple[int, int] | None:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    for row in range(1, min(max_row, 20) + 1):
        for col in range(1, max_col + 1):
            value = ws.cell(row=row, column=col).value
            if not _non_empty_value(value) or _is_formula(value):
                continue
            texto = _short(value).lower()
            if ESTADO_CONTROL_RE.search(texto):
                return row, col
    return None


def _fila_contiene_actor_detalle(ws, row: int, max_cols: int = 12) -> bool:
    max_col = ws.max_column or 0
    for col in range(1, min(max_col, max_cols) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and _es_actor_detalle(value):
            return True
    return False


def _texto_fila(ws, row: int, max_cols: int = 8) -> str:
    partes = []
    max_col = ws.max_column or 0
    for col in range(1, min(max_col, max_cols) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
            partes.append(_short(value, 60))
    return " | ".join(partes)


def _texto_fila_con_encabezados(ws, header_row: int, row: int, max_cols: int = 8) -> str:
    partes = []
    max_col = ws.max_column or 0
    for col in range(1, min(max_col, max_cols) + 1):
        value = ws.cell(row=row, column=col).value
        if not _non_empty_value(value) or _is_formula(value):
            continue
        if _es_actor_detalle(value):
            continue
        header = ws.cell(row=header_row, column=col).value
        if _es_actor_detalle(header):
            continue
        etiqueta = _short(header, 30) if _non_empty_value(header) and not _is_formula(header) else ""
        valor = _short(value, 60)
        partes.append(f"{etiqueta}: {valor}" if etiqueta else valor)
    return "; ".join(partes)


def _texto_hoja(ws, max_rows: int = 80, max_cols: int = 30) -> str:
    partes: list[str] = []
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    for row in range(1, min(max_row, max_rows) + 1):
        for col in range(1, min(max_col, max_cols) + 1):
            value = ws.cell(row=row, column=col).value
            if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
                partes.append(str(value))
    return " ".join(partes)


def _rango_ref(formula: str) -> str | None:
    match = re.search(r"([A-Z]{1,3}\d+\s*:\s*[A-Z]{1,3}\d+)", formula.upper())
    return match.group(1).replace(" ", "") if match else None


def _rango_contiguo(cells: list[tuple[Any, str]]) -> str:
    ordenadas = sorted(cells, key=lambda item: item[0].row)
    col = get_column_letter(ordenadas[0][0].column)
    filas = [item[0].row for item in ordenadas]
    if filas == list(range(min(filas), max(filas) + 1)):
        return f"{col}{min(filas)}:{col}{max(filas)}"
    primeras = [f"{col}{fila}" for fila in filas[:4]]
    if len(filas) > 4:
        primeras.append(f"+{len(filas) - 4} más")
    return ", ".join(primeras)


def _patron_mayoritario(items: list[tuple[int, str, str]]) -> tuple[str | None, int]:
    if len(items) < 3:
        return None, 0
    counts = Counter(pattern for _, _, pattern in items)
    patron, n = counts.most_common(1)[0]
    if n >= max(3, int(len(items) * 0.6)):
        return patron, n
    return None, 0


def _nombre_columna_raw(ws, col: int) -> str:
    for row in range(1, min(ws.max_row, 12) + 1):
        value = ws.cell(row=row, column=col).value
        if _non_empty_value(value) and not _is_formula(value) and not _es_actor_detalle(value):
            return str(value)
    return get_column_letter(col)
