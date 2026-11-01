"""Integridad estructural: cortes de bloque/totales, estados explicitos, valores implicitos agregados."""
from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

from services.calculos.excel_utils import (
    CALCULO_CRITICO_RE,
    DECISION_RE,
    ESTADO_REVISION_RE,
    MAX_HALLAZGOS,
    _col_label,
    _columna_estado,
    _es_actor_detalle,
    _es_cierre_legitimo,
    _fila_contiene_actor_detalle,
    _formula_cells_col,
    _formula_cells_row,
    _impacta_decision,
    _is_formula,
    _non_empty_value,
    _numeric_rows,
    _numero,
    _patron_mayoritario,
    _row_label,
    _short,
    _texto_fila,
    _texto_fila_con_encabezados,
)
from services.calculos.formulas import (
    _analizar_patrones_pares,
    _evaluar_patron,
    _hallazgo_hardcode,
    _hallazgo_implicito,
    _hallazgo_patron,
)
from services.calculos.relaciones import (
    _analizar_relaciones_por_nombre,
)


def _es_corte_bloque(ws, row: int) -> bool:
    label = ws.cell(row=row, column=1).value
    if not _non_empty_value(label):
        return True
    nums = [_numero(ws.cell(row=row, column=col).value) for col in range(2, ws.max_column + 1)]
    tiene_num = any(v is not None for v in nums)
    text = str(label).strip()
    return (text.isupper() and not tiene_num) or text.lower().startswith("total ")


def _analizar_totales_bloque(ws, rows: list[dict[str, Any]]) -> list[dict]:
    hallazgos: list[dict] = []
    sheet = ws.title
    for target in rows:
        raw_label = str(ws.cell(row=target["row"], column=1).value or "").strip().lower()
        if not raw_label.startswith("total "):
            continue
        start = target["row"] - 1
        while start >= 1 and not _es_corte_bloque(ws, start):
            start -= 1
        componentes = _numeric_rows(ws, min_vals=1)
        componentes = [r for r in componentes if start < r["row"] < target["row"]]
        if len(componentes) < 2:
            continue
        esperado: dict[int, float] = {}
        for comp in componentes:
            for col, val in comp["values"].items():
                esperado[col] = esperado.get(col, 0.0) + val
        target["ws"] = ws
        ok, fallas = _evaluar_patron(target, esperado)
        total = ok + len(fallas)
        if not fallas or total < 4 or ok < max(4, int(total * 0.7)):
            continue
        col, actual, exp = fallas[0]
        regla = f"{target['label']} = suma del bloque anterior"
        hallazgos.append(_hallazgo_implicito(
            sheet, target, col, actual, exp, regla, [c["label"] for c in componentes[:8]]
        ))
        if len(hallazgos) >= MAX_HALLAZGOS:
            return hallazgos
    return hallazgos


def _analizar_estados_explicitos(ws) -> list[dict]:
    ubicacion = _columna_estado(ws)
    if not ubicacion:
        return []
    header_row, estado_col = ubicacion
    max_row = ws.max_row or header_row
    hallazgos: list[dict] = []
    for row in range(header_row + 1, max_row + 1):
        value = ws.cell(row=row, column=estado_col).value
        estado = _short(value, 40)
        if not ESTADO_REVISION_RE.search(estado):
            continue
        if _fila_contiene_actor_detalle(ws, row):
            continue
        detalle = _texto_fila_con_encabezados(ws, header_row, row) or _texto_fila(ws, row) or _row_label(ws, row)
        if _es_actor_detalle(detalle):
            continue
        impacto = bool(DECISION_RE.search(detalle) or CALCULO_CRITICO_RE.search(detalle))
        severidad = "alta" if re.search(r"fail|failed|fall[oó]|error|inconsistente", estado, re.I) else "media"
        celda = f"{get_column_letter(estado_col)}{row}"
        hallazgos.append({
            "modo": "control_explicito",
            "descripcion": (
                f"En la hoja '{ws.title}', {celda} figura con estado '{estado}'. "
                f"El propio archivo pide revisar este control: {detalle}."
            ),
            "severidad": severidad,
            "evidencia": {
                "regla": "estado_explicito_revision",
                "hoja": ws.title,
                "celda": celda,
                "estado": estado,
                "fila_detalle": detalle,
                "columna_estado": _col_label(ws, estado_col),
                "impacta_decision": impacto,
            },
        })
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    return hallazgos


def _analizar_valores_implicitos(ws) -> list[dict]:
    rows = _numeric_rows(ws, min_vals=3)
    if len(rows) < 3:
        return []
    hallazgos = []
    for inicio, fin in _rangos_de_cuadros(ws):
        cuadro = [r for r in rows if inicio <= r['row'] <= fin]
        hallazgos.extend(_analizar_relaciones_por_nombre(ws, cuadro))
        hallazgos.extend(_analizar_totales_bloque(ws, cuadro))
        hallazgos.extend(_analizar_patrones_pares(ws, cuadro))
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    return hallazgos[:MAX_HALLAZGOS]


def _rangos_de_cuadros(ws):
    """Encabezados y rótulos separan cuadros; un hueco solo no invalida un cierre.

    Los años de un encabezado son estructura, nunca valores fijos intercalados.
    No se combinan patrones de distintos niveles de gobierno o cuadros apilados.
    """
    inicio = 1
    for row in range(1, ws.max_row + 1):
        valores = [c.value for c in ws[row] if _non_empty_value(c.value)]
        if not valores or any(_is_formula(v) for v in valores):
            continue
        numeros = [_numero(v) for v in valores]
        nums = [n for n in numeros if n is not None]
        textos = [v for v, n in zip(valores, numeros) if n is None]
        anios = len(nums) >= 2 and all(n.is_integer() and 1900 <= n <= 2100 for n in nums)
        etiqueta = str(ws.cell(row=row, column=1).value or '').strip().lower()
        encabezado_anual = anios and bool(re.fullmatch(r'indicador(?:es)?|clase|categor[ií]a|concepto|rubro|a[nñ]o|per[ií]odo', etiqueta))
        texto_estructural = bool(textos) and not nums
        if encabezado_anual or texto_estructural:
            if inicio < row:
                yield inicio, row - 1
            inicio = row + 1
    if inicio <= ws.max_row:
        yield inicio, ws.max_row


_REF = r"\$?([A-Z]{1,3})\$?(\d+)"


def _suma_relativa(formula: str, row: int, col: int):
    """Celdas que suma una fórmula de pura suma, relativas a su posición; None si no es una suma.

    «=+U69+SUM(U72:U80)» y «=+T69+T72+…+T80» suman las mismas celdas relativas: son equivalentes.
    """
    from openpyxl.utils import column_index_from_string

    cuerpo = str(formula or "").lstrip("=").replace(" ", "").upper()
    if not cuerpo:
        return None
    celdas = []
    for termino in [t for t in cuerpo.split("+") if t]:
        m = re.fullmatch(_REF, termino)
        if m:
            celdas.append((int(m.group(2)) - row, column_index_from_string(m.group(1)) - col))
            continue
        m = re.fullmatch(r"SUM\(" + _REF + r":" + _REF + r"\)", termino)
        if not m:
            return None
        c1, r1, c2, r2 = column_index_from_string(m.group(1)), int(m.group(2)), column_index_from_string(m.group(3)), int(m.group(4))
        if (r2 - r1 + 1) * (c2 - c1 + 1) > 500:
            return None
        celdas += [(r - row, c - col) for r in range(min(r1, r2), max(r1, r2) + 1) for c in range(min(c1, c2), max(c1, c2) + 1)]
    return frozenset(celdas) if len(celdas) == len(set(celdas)) else None


def _analizar_hoja(ws, detectar_formulas: bool = True, inferir_patrones: bool = False, ws_valores=None) -> list[dict]:
    from services.calculos.contraste_formula import enriquecer
    if (
        _es_actor_detalle(ws.title)
        or _fila_contiene_actor_detalle(ws, 1)
        or (ws_valores is not None and (_es_actor_detalle(ws_valores.title) or _fila_contiene_actor_detalle(ws_valores, 1)))
    ):
        return []
    hallazgos: list[dict] = []
    sheet = ws.title
    if ws_valores is not None:
        hallazgos.extend(_analizar_estados_explicitos(ws_valores))
        if len(hallazgos) >= MAX_HALLAZGOS:
            return hallazgos[:MAX_HALLAZGOS]

    if detectar_formulas:
        for row in range(1, ws.max_row + 1):
            formulas = _formula_cells_row(ws, row)
            patron, _ = _patron_mayoritario(formulas)
            if not patron:
                continue
            posiciones_patron = [col for col, _, normalizada in formulas if normalizada == patron]
            suma_patron = next((_suma_relativa(f, row, c) for c, f, nz in formulas if nz == patron), None)
            for col, formula, normalizada in formulas:
                if normalizada != patron:
                    # Misma suma escrita de otra forma (rango en vez de celdas): no se aparta del patrón.
                    if suma_patron is not None and _suma_relativa(formula, row, col) == suma_patron:
                        continue
                    if _es_cierre_legitimo(ws, formula, row, col, posiciones_patron, es_columna=False):
                        continue
                    impacto = _impacta_decision(ws, row, col)
                    hallazgos.append(_hallazgo_patron(
                        sheet, "fila", f"{get_column_letter(col)}{row}",
                        f"{_row_label(ws, row)} / columna {_col_label(ws, col)}", formula, patron,
                        "alta" if impacto else "media", impacto,
                    ))
                    if len(hallazgos) >= MAX_HALLAZGOS:
                        return hallazgos
            cols = [col for col, _, _ in formulas]
            for col in range(min(cols), max(cols) + 1):
                cell = ws.cell(row=row, column=col)
                if not _is_formula(cell.value) and _non_empty_value(cell.value):
                    left = any(c < col for c in cols)
                    right = any(c > col for c in cols)
                    if left and right:
                        impacto = _impacta_decision(ws, row, col)
                        hallazgos.append(_hallazgo_hardcode(
                            sheet, "fila", f"{get_column_letter(col)}{row}",
                            f"{_row_label(ws, row)} / columna {_col_label(ws, col)}", cell.value, patron,
                            "alta" if impacto else "media", impacto,
                        ))
                        if len(hallazgos) >= MAX_HALLAZGOS:
                            return hallazgos

        for inicio, fin in _rangos_de_cuadros(ws):
            for col in range(1, ws.max_column + 1):
                formulas = [f for f in _formula_cells_col(ws, col) if inicio <= f[0] <= fin]
                patron, _ = _patron_mayoritario(formulas)
                if not patron:
                    continue
                posiciones_patron = [row for row, _, normalizada in formulas if normalizada == patron]
                suma_patron = next((_suma_relativa(f, r, col) for r, f, nz in formulas if nz == patron), None)
                for row, formula, normalizada in formulas:
                    if normalizada != patron:
                        if suma_patron is not None and _suma_relativa(formula, row, col) == suma_patron:
                            continue
                        if _es_cierre_legitimo(ws, formula, col, row, posiciones_patron, es_columna=True):
                            continue
                        impacto = _impacta_decision(ws, row, col)
                        hallazgos.append(_hallazgo_patron(
                            sheet, "columna", f"{get_column_letter(col)}{row}",
                            f"{_row_label(ws, row)} / columna {_col_label(ws, col)}", formula, patron,
                            "alta" if impacto else "media", impacto,
                        ))
                        if len(hallazgos) >= MAX_HALLAZGOS:
                            return hallazgos
                rows = [row for row, _, _ in formulas]
                for row in range(min(rows), max(rows) + 1):
                    cell = ws.cell(row=row, column=col)
                    if not _is_formula(cell.value) and _non_empty_value(cell.value):
                        above = any(r < row for r in rows)
                        below = any(r > row for r in rows)
                        if above and below:
                            impacto = _impacta_decision(ws, row, col)
                            hallazgos.append(_hallazgo_hardcode(
                                sheet, "columna", f"{get_column_letter(col)}{row}",
                                f"{_row_label(ws, row)} / columna {_col_label(ws, col)}", cell.value, patron,
                                "alta" if impacto else "media", impacto,
                            ))
                            if len(hallazgos) >= MAX_HALLAZGOS:
                                return hallazgos

    if inferir_patrones:
        hallazgos.extend(_analizar_valores_implicitos(ws))
    return enriquecer(ws, hallazgos[:MAX_HALLAZGOS])
