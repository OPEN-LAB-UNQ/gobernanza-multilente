"""Lectura completa de libros y particiones trazables, sin comparar cuadros distintos."""

from __future__ import annotations

import datetime as dt
import io
import json
import math
import re
from copy import deepcopy
from types import SimpleNamespace

import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter

from services.excel_estructura import (
    CODIGO_SERIE,
    encabezados_aislados,
    encabezados_apilados,
    filas_encabezado_anual,
    paneles_anuales,
)
from services.plural import contar, muestra

ANIO = re.compile(r"^(?:19|20)\d{2}(?:\.0)?$")
UNIDAD_FILA = re.compile(r"^en\s+(?:%|\$|millones|miles|pesos|porcentaje)", re.I)
UNIDAD_CELDA = re.compile(
    r"(?:n[uú]meros? [ií]ndice|[ií]ndice|%|porcentajes?|(?:miles|millones) de pesos|pesos|unidades|m[²2]"
    # Unidades físicas de los cuadros de energía, gas y producción («GWh», «MW», «°C», «Miles Tns.», «dam3»).
    r"|[kmgt]wh|[kmg]w|°\s?c|(?:miles (?:de )?)?tns?\.?|(?:miles (?:de )?)?toneladas|(?:dam|hm|m)[3³])", re.I
)
# Parecido a un código de serie: sin el prefijo numérico («_TASA_DE_DEANO_0_0_42_71») o con puntos en un tramo
# («364.1_DISTRIB._GGAS__30»). Sólo cuenta dentro de una fila donde la mayoría son códigos completos.
CODIGO_SERIE_AMPLIO = re.compile(r"^(?:\d{1,4}\.\d{1,3})?(?:_+[A-Z0-9.]+){2,}$")
def es_fila_codigos_serie(valores) -> bool:
    """Fila con los códigos de serie de cada columna: metadato del cuadro, ni encabezado ni datos."""
    from services.carga_archivos import _compactar

    presentes = [_compactar(v) for v in valores if _compactar(v)]
    if len(presentes) < 2 or not all(isinstance(v, str) and CODIGO_SERIE_AMPLIO.match(v) for v in presentes):
        return False
    return 2 * sum(bool(CODIGO_SERIE.match(v)) for v in presentes) > len(presentes)


def _es_numero_tabular(valor):
    """Reconocer números regionales sin convertir rótulos, fechas ni la fuente."""
    if isinstance(valor, bool):
        return False
    n = pd.to_numeric(valor, errors="coerce")
    if not pd.isna(n):
        return math.isfinite(float(n))
    if not isinstance(valor, str) or not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)*", valor.strip()):
        return False
    from services.perfilado import numero

    n = numero(valor)
    return n is not None and math.isfinite(n)


MAX_REFERENCIAS = 200
# Con más hojas que éstas en formato de porcentaje, el aviso no se repite hoja por hoja.
MAX_HOJAS_AVISO_PORCENTAJE = 3


def _rangos(numeros):
    """[2, 3, 4, 7] -> "2-4, 7": para declarar filas sin listas interminables."""
    partes, inicio, previo = [], None, None
    for n in sorted(numeros):
        if inicio is None:
            inicio = previo = n
        elif n == previo + 1:
            previo = n
        else:
            partes.append(f"{inicio}-{previo}" if previo != inicio else str(inicio))
            inicio = previo = n
    if inicio is not None:
        partes.append(f"{inicio}-{previo}" if previo != inicio else str(inicio))
    return ", ".join(partes)


def _completar_calculadas(crudo, calculadas):
    """Pone en la grilla leída el valor calculado de las fórmulas sin resultado guardado."""
    filas = max(r for r, _ in calculadas)
    columnas = max(c for _, c in calculadas)
    if crudo.shape[0] < filas or crudo.shape[1] < columnas:
        crudo = crudo.reindex(index=range(max(crudo.shape[0], filas)), columns=range(max(crudo.shape[1], columnas)))
    crudo = crudo.astype(object)
    for (r, c), v in calculadas.items():
        crudo.iat[r - 1, c - 1] = v
    return crudo


def info_libro(contenido):
    """Por hoja: visibilidad (hoja, filas y columnas ocultas) y celdas con fórmula sin resultado
    guardado. Sólo para .xlsx: el formato binario .xls no expone esos datos de la misma forma."""
    if contenido[:4] != b"PK\x03\x04":
        return {}
    import warnings

    from openpyxl import load_workbook

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            formulas = load_workbook(io.BytesIO(contenido), data_only=False)
            valores = load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception:  # noqa: BLE001 — sin este detalle la lectura sigue con pandas
        return {}
    from services.controles_excel.evaluador import Libro

    evaluador = Libro(valores, formulas)
    info = {}
    for ws in formulas.worksheets:
        wv = valores[ws.title]
        sin_valor = set()
        porcentajes = {}  # (fila, columna) → valor guardado
        con_numeros = False
        for fila in ws.iter_rows():
            for celda in fila:
                if celda.data_type == "f" and wv.cell(row=celda.row, column=celda.column).value is None:
                    sin_valor.add((celda.row, celda.column))
                elif celda.data_type == "n" and celda.value is not None:
                    con_numeros = True
                fmt = celda.number_format
                if fmt and "%" in fmt and es_formato_porcentaje(fmt):
                    v = wv.cell(row=celda.row, column=celda.column).value
                    # Un año del encabezado con formato % (2019 → 201900 %) no es un porcentaje.
                    if (isinstance(v, (int, float)) and not isinstance(v, bool)
                            and not (float(v).is_integer() and 1800 <= v <= 2200)):
                        porcentajes[(celda.row, celda.column)] = v
        calculadas = {}
        if sin_valor and not con_numeros and len(sin_valor) <= 20000:
            # Hoja hecha sólo de fórmulas guardadas sin resultado (libro generado sin recalcular):
            # sin calcularlas no habría nada que leer. Las sumas y la aritmética simple se calculan
            # desde sus celdas de origen; el resto queda sin valor y se declara. En una hoja con
            # números guardados, las fórmulas sin resultado se declaran sin completarlas.
            for r, c in sorted(sin_valor):
                v = evaluador.valor(ws.title, r, c)
                if v is not None:
                    calculadas[(r, c)] = v
            sin_valor -= set(calculadas)
        columnas = set()
        for clave, dim in ws.column_dimensions.items():
            if dim.hidden:
                desde = dim.min or column_index_from_string(clave)
                columnas.update(range(desde, (dim.max or desde) + 1))
        combinadas = {}
        for rango in ws.merged_cells.ranges:
            if rango.max_row > rango.min_row and rango.min_col == rango.max_col:
                v = wv.cell(row=rango.min_row, column=rango.min_col).value
                if isinstance(v, str) and v.strip():
                    for r in range(rango.min_row + 1, rango.max_row + 1):
                        combinadas[(r, rango.min_col)] = (rango.min_row, v)
        info[ws.title] = {
            "rangos_combinados": [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges],
            "celdas_combinadas": combinadas,
            "estado_hoja": ws.sheet_state,
            "filas_ocultas": sorted(r for r, d in ws.row_dimensions.items() if d.hidden),
            "columnas_ocultas": sorted(columnas),
            "formulas_sin_valor": sin_valor,
            "formulas_calculadas": calculadas,
            "porcentajes": porcentajes,
        }
    return info


def es_formato_porcentaje(fmt):
    """Formato de número que muestra la celda multiplicada por 100 («0.0%»). Un «%» entre comillas,
    escapado («\\%») o dentro de corchetes es un texto literal: no cambia la escala."""
    return "%" in re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', "", fmt or "")


def _completar_combinadas(crudo, combinadas):
    """Pone el rótulo de una celda combinada en vertical en cada fila de datos que abarca.

    Sólo entre filas con cifras (no en un encabezado de dos filas): «Interior» sin su grupo
    «Renovación de redes» es otra serie.
    """
    from services.cobertura_cifras import cifras_fila

    def con_cifras(r):
        return 0 < r <= crudo.shape[0] and bool(cifras_fila(list(crudo.iloc[r - 1])))

    completadas = []
    crudo = crudo.astype(object)
    for (r, c), (origen, valor) in sorted(combinadas.items()):
        if c > crudo.shape[1] or not con_cifras(r) or not con_cifras(origen):
            continue
        actual = crudo.iat[r - 1, c - 1]
        if actual is None or (isinstance(actual, float) and actual != actual):
            crudo.iat[r - 1, c - 1] = valor
            completadas.append((r, c, valor))
    return crudo, completadas


def _fecha_iso(valor):
    """Celda con fecha u hora → texto ISO; el resto queda igual."""
    if isinstance(valor, (dt.datetime, pd.Timestamp)):
        return valor.date().isoformat() if valor.time() == dt.time() else valor.isoformat()
    if isinstance(valor, (dt.date, dt.time)):
        return valor.isoformat()
    return valor


TOTAL_FILA = re.compile(r"^(?:sub)?total(?:es)?\b")


def _unir_total_suelto(crudo, grupos, activas):
    """Un «Total» que quedó solo en su grupo porque filas sin cifras lo separan de sus componentes.

    Una fila con rótulo y sin datos (un rubro vacío en ese cuadro) parte el cuadro como si fuera un
    título. Si el grupo que queda es una sola fila de total, se prueba unirlo con los grupos de arriba
    (mismo encabezado, mismas columnas), del más cercano al más lejano, y se une con el primero con el
    que CIERRA. Un título de sección entre dos grupos (texto sin cifras, sin sangría, que no es un
    período) corta la búsqueda: el total no sigue hacia otra sección; un rubro con sangría y sin
    cifras («   Asignaciones familiares») no la corta. Si no cierra con ninguno, los bloques no se fusionan:
    el total queda solo y se declara («no se armó su cierre»), sin afirmar que no cierra. Tampoco se une
    si arriba hay otro total (un bloque ya cerrado) o una fila de unidad («En %»).
    Devuelve (grupos, {tuple(filas): detalle}) con filas 0-based en los grupos y 1-based en el detalle.
    """
    from services.carga_archivos import _compactar
    from services.controles_excel.periodos import leer as leer_periodo
    from services.perfilado import numero

    def rotulo(i):
        for j in activas:
            v = crudo.iat[i, j]
            if _presente(v) and not _es_numero_tabular(v):
                return _compactar(v)
        return ""

    def sangria(i):
        for j in activas:
            v = crudo.iat[i, j]
            if _presente(v) and not _es_numero_tabular(v):
                t = str(v)
                return len(t) - len(t.lstrip())
        return 0

    def valor(i, j):
        v = crudo.iat[i, j]
        if not _presente(v) or not _es_numero_tabular(v):
            return None
        try:
            return float(v) if not isinstance(v, str) else numero(v)
        except (TypeError, ValueError):
            return None

    def con_cifras(i):
        return any(_presente(crudo.iat[i, j]) and _es_numero_tabular(crudo.iat[i, j]) for j in activas)

    def es_total(i):
        return bool(TOTAL_FILA.match(rotulo(i).lower()))

    def separadas(desde, hasta):
        """Filas entre dos grupos (exclusivas): None si alguna trae cifras o una unidad («En %»)."""
        filas = list(range(desde + 1, hasta))
        if any(con_cifras(i) or UNIDAD_FILA.match(rotulo(i)) for i in filas):
            return None
        return filas

    def es_titulo(i):
        """Texto sin cifras que no es un período ni un rubro con sangría (un rubro vacío del cuadro)."""
        t = rotulo(i)
        return bool(t) and not con_cifras(i) and leer_periodo(t) is None and not sangria(i)

    def coincide(a, b):
        return abs(a - b) <= 0.05 + 0.001 * max(abs(a), abs(b))

    def _sin_subtotales(partes):
        """Suma sin contar dos veces una fila que es la suma de las que la siguen (subtotal intercalado)."""
        excluidas, cambio = set(), True
        while cambio:
            cambio = False
            for k, v in enumerate(partes):
                if k in excluidas:
                    continue
                acumulado, usadas = 0.0, []
                for m in range(k + 1, len(partes)):
                    if m in excluidas:
                        continue
                    acumulado += partes[m]
                    usadas.append(m)
                    if len(usadas) >= 2 and coincide(v, acumulado):
                        excluidas.update(usadas)
                        cambio = True
                        break
        return sum(p for k, p in enumerate(partes) if k not in excluidas) if excluidas else None

    def cierra(filas, total):
        """El total es la suma de las filas (todas, las de menor sangría o sin contar subtotales)."""
        buenas = comparables = 0
        minimo = min(sangria(i) for i in filas)
        for j in activas:
            t = valor(total, j)
            partes = [valor(i, j) for i in filas]
            if t is None or any(p is None for p in partes) or len(partes) < 2:
                continue
            comparables += 1
            superiores = [p for i, p in zip(filas, partes) if sangria(i) == minimo]
            ajustada = _sin_subtotales(partes)
            if coincide(t, sum(partes)) or coincide(t, sum(superiores)) or (ajustada is not None and coincide(t, ajustada)):
                buenas += 1
        return comparables > 0 and buenas >= 0.8 * comparables

    salida, cortes = [], {}
    for g in grupos:
        if len(g) != 1 or not es_total(g[0]) or not salida:
            salida.append(g)
            continue
        unidos, sin_datos, elegido = [], [], None
        k = len(salida) - 1
        tope = g[0]
        while k >= 0:
            previo = salida[k]
            entre = separadas(previo[-1], tope)
            if entre is None or any(es_total(i) for i in previo) or UNIDAD_FILA.match(rotulo(previo[0])):
                break
            if unidos and any(es_titulo(i) for i in entre):
                break  # título de otra sección: el total no la abarca
            unidos = previo + unidos
            sin_datos = [i for i in entre if rotulo(i)] + sin_datos
            tope = previo[0]
            if cierra(unidos, g[0]):
                elegido = k
                break
            k -= 1
        if elegido is None or not sin_datos:
            # Sin un grupo de arriba con el que cierre (o separado de un bloque ya cerrado): queda solo y se declara.
            entre = separadas(salida[-1][-1], g[0])
            if entre is not None:
                cortes[tuple(g)] = {"fila_total": g[0] + 1, "rotulo": rotulo(g[0])[:80], "unido": False,
                                    "filas_sin_datos": [i + 1 for i in entre if rotulo(i)]}
            salida.append(g)
            continue
        del salida[elegido:]
        nuevo = unidos + g
        salida.append(nuevo)
        cortes[tuple(nuevo)] = {"fila_total": g[0] + 1, "rotulo": rotulo(g[0])[:80], "unido": True,
                                "filas": [nuevo[0] + 1, g[0] + 1], "filas_sin_datos": [i + 1 for i in sin_datos]}
    return salida, cortes


def leer_libro(contenido, archivo):
    from services.carga_archivos import _celda_numerica, _compactar, _detectar_filas_encabezado, _flatten_columnas
    from services.perfilado import interpretar_columna, preferencia_formato

    datos, hojas, bloques = [], [], []
    extra = info_libro(contenido)
    hojas_pct = []  # hojas con celdas en formato de porcentaje
    with pd.ExcelFile(io.BytesIO(contenido)) as libro:
        for hoja in libro.sheet_names:
            # dtype=object: sin inferencia de tipos. Si pandas convierte por su cuenta una columna de
            # textos «475.735» (la lee como 475,735 en formato inglés), la lectura por cuadro no se
            # entera y el punto de miles se pierde en silencio.
            crudo = libro.parse(hoja, header=None, dtype=object)
            detalle = {
                "archivo": archivo,
                "hoja": hoja,
                "estado": "sin_datos",
                "filas_datos": 0,
                "bloques": 0,
                "filas_estructura": [],
                "filas_no_interpretadas": [],
                "advertencias": [],
            }
            hojas.append(detalle)
            avisos = detalle["avisos_lectura"] = []
            # Formato de porcentaje («0.0%»): la celda guarda la fracción (0,304) y muestra 30,4 %. Se lee
            # el valor guardado, sin cambiarlo, y se avisa la escala. Con muchas hojas así, el aviso va en
            # las primeras y una línea final cuenta las demás: no se repite hoja por hoja.
            porcentajes = (extra.get(hoja) or {}).get("porcentajes") or {}
            pcts = sorted(porcentajes)
            if pcts:
                lista = [f"{get_column_letter(c)}{r}" for r, c in pcts]
                # Casi todos entre -1,5 y 1,5: la celda guarda la fracción. Si no, el formato multiplica por
                # 100 cifras que ya están en puntos (30,4 se ve 3.040 %): lo que se lee es lo guardado.
                fraccion = sum(abs(v) <= 1.5 for v in porcentajes.values()) >= 0.9 * len(porcentajes)
                detalle["formato_porcentaje"] = {"celdas": len(lista), "ejemplos": lista[:10], "fraccion": fraccion}
                hojas_pct.append(detalle)
                if len(hojas_pct) <= MAX_HOJAS_AVISO_PORCENTAJE:
                    if fraccion:
                        avisos.append(
                            f"{contar(len(lista), 'celda', 'celdas')} con formato de porcentaje ({muestra(lista)}) "
                            f"{'guarda' if len(lista) == 1 else 'guardan'} la fracción: se "
                            f"{'leyó' if len(lista) == 1 else 'leyeron'} como fracción (30,4 % → 0,304), no en puntos porcentuales.")
                    else:
                        avisos.append(
                            f"{contar(len(lista), 'celda', 'celdas')} con formato de porcentaje ({muestra(lista)}): el formato "
                            "multiplica por 100 valores que parecen ya expresados en puntos (30,4 se muestra como 3.040 %); "
                            "se leyó el valor guardado (30,4).")
            # Duraciones con formato [h]:mm llegan como timedelta: se leen como horas con decimales.
            duraciones = int(crudo.map(lambda v: isinstance(v, dt.timedelta)).values.sum())
            if duraciones:
                crudo = crudo.map(lambda v: v.total_seconds() / 3600 if isinstance(v, dt.timedelta) else v)
                avisos.append(f"{contar(duraciones, 'celda', 'celdas')} con formato de duración ([h]:mm) se "
                              f"{'leyó' if duraciones == 1 else 'leyeron'} como horas con decimales (9:45 → 9,75).")
            # Años guardados como fechas del mismo día y mes (01/01/2017 con formato «aaaa»): son años, no meses.
            for i in range(len(crudo)):
                fila_i = list(crudo.iloc[i])
                fechas = [(j, v) for j, v in enumerate(fila_i) if isinstance(v, (dt.datetime, dt.date))]
                otras_cifras = any(isinstance(v, (int, float)) and not isinstance(v, bool) and not pd.isna(v) for v in fila_i)
                if (len(fechas) >= 3 and not otras_cifras and {(v.month, v.day) for _, v in fechas} == {(1, 1)}
                        and len({v.year for _, v in fechas}) == len(fechas)):
                    for j, v in fechas:
                        crudo.iat[i, j] = v.year
                    avisos.append(f"La fila {i + 1} trae los años como fechas ({fechas[0][1]:%d/%m}/AAAA): se leyeron "
                                  f"como años, {fechas[0][1].year} a {fechas[-1][1].year}.")
            ex = extra.get(hoja) or {}
            sin_valor = ex.get("formulas_sin_valor") or set()
            filas_ocultas = set(ex.get("filas_ocultas") or [])
            columnas_ocultas = set(ex.get("columnas_ocultas") or [])
            hoja_oculta = ex.get("estado_hoja") in {"hidden", "veryHidden"}
            if hoja_oculta or filas_ocultas or columnas_ocultas:
                detalle["visibilidad"] = {
                    "estado_hoja": ex.get("estado_hoja"),
                    "filas_ocultas": sorted(filas_ocultas)[:MAX_REFERENCIAS],
                    "columnas_ocultas": [get_column_letter(c) for c in sorted(columnas_ocultas)][:MAX_REFERENCIAS],
                }
                partes = []
                if hoja_oculta:
                    partes.append("la hoja está oculta")
                if filas_ocultas:
                    partes.append(f"filas ocultas: {_rangos(filas_ocultas)}")
                if columnas_ocultas:
                    partes.append("columnas ocultas: " + ", ".join(get_column_letter(c) for c in sorted(columnas_ocultas)))
                avisos.append("Contenido no visible en la planilla (" + "; ".join(partes) + "). Sus datos se leen y se marcan como ocultos en la ubicación.")
            calculadas = ex.get("formulas_calculadas") or {}
            if calculadas:
                crudo = _completar_calculadas(crudo, calculadas)
                lista = [f"{get_column_letter(c)}{r}" for r, c in sorted(calculadas)]
                detalle["formulas_calculadas"] = {"celdas": len(lista), "ejemplos": lista[:10]}
                avisos.append(
                    f"{contar(len(lista), 'celda', 'celdas')} con fórmula sin resultado guardado se "
                    f"{'calculó' if len(lista) == 1 else 'calcularon'} desde sus celdas de origen (sumas y aritmética simple): "
                    f"{muestra(lista)}."
                )
            if ex.get("celdas_combinadas"):
                crudo, completadas = _completar_combinadas(crudo, ex["celdas_combinadas"])
                if completadas:
                    rotulos = list(dict.fromkeys(f"«{v}»" for _, _, v in completadas))
                    detalle["celdas_combinadas_completadas"] = len(completadas)
                    avisos.append(
                        f"{contar(len(completadas), 'fila toma', 'filas toman')} su rótulo de una celda combinada en vertical "
                        f"({muestra(rotulos)}): se completó en cada fila para no perder el grupo."
                    )
            if sin_valor:
                lista = [f"{get_column_letter(c)}{r}" for r, c in sorted(sin_valor)]
                ejemplos = lista[:10]
                detalle["formulas_sin_valor"] = {
                    "celdas": len(sin_valor),
                    "ejemplos": ejemplos,
                    "mensaje": f"{contar(len(sin_valor), 'celda', 'celdas')} con fórmula sin resultado guardado: no evaluables; guardá el archivo desde la planilla",
                }
                avisos.append(
                    f"{contar(len(sin_valor), 'celda', 'celdas')} con fórmula sin resultado guardado ({muestra(lista)}): no evaluables; guardá el archivo desde la planilla."
                )
            if crudo.empty:
                detalle["motivo"] = "Hoja vacía"
                detalle["advertencias"].extend(avisos)
                continue
            preferencia = preferencia_formato(v for v in crudo.to_numpy().flat if isinstance(v, str))
            # Los encabezados con años delimitan cuadros apilados. No se elige el mayor.
            anuales = filas_encabezado_anual(crudo.values.tolist())
            if anuales:
                encabezados = [[i] for i in anuales]
            else:
                header = _detectar_filas_encabezado(contenido, hoja, libro)
                if not header:
                    primera = next((int(i) for i, r in crudo.iterrows() if r.notna().any()), 0)
                    r = crudo.iloc[primera]
                    if sum(isinstance(v, str) and bool(v.strip()) for v in r) < 2:
                        if not any(_es_numero_tabular(v) for v in crudo.to_numpy().flat):
                            detalle.update(estado="sin_datos", texto_sin_cifras=True, motivo="Hoja de texto o índice; sin datos cuantitativos",
                                           filas_estructura=[int(i) + 1 for i, r in crudo.iterrows() if r.notna().any()])
                            detalle["advertencias"].extend(avisos)
                            continue
                        detalle["estado"] = "no_interpretable"
                        detalle["motivo"] = "No se reconoció un encabezado tabular; requiere preparar la hoja"
                        detalle["filas_no_interpretadas"] = [int(i) + 1 for i, r in crudo.iterrows() if r.notna().any()]
                        _declarar_no_revisada(detalle, crudo, _celda_numerica)
                        detalle["advertencias"].extend(avisos)
                        continue
                    header = [primera]
                encabezados = [header]
            # Otros cuadros apilados debajo, con su propio encabezado (períodos distintos,
            # trimestres, otra serie): cada uno se lee y se controla como cuadro propio.
            apilados = encabezados_apilados(crudo.values.tolist(), {i for h in encabezados for i in h})
            if apilados:
                encabezados = sorted(encabezados + [[i] for i in apilados], key=lambda h: h[0])
                detalle["cuadros_apilados"] = [i + 1 for i in apilados]
            usadas, estructura = set(), set()
            cubiertas = set()  # (fila, columna) 0-based de las celdas que quedaron en algún cuadro
            rotulos_previos = None  # (columnas, raw_header, filas del encabezado) del último cuadro con rótulos propios
            for numero, header in enumerate(encabezados):
                from services.excel_encabezados import combinar

                combinado = combinar(crudo.values.tolist(), ex.get("rangos_combinados", []), header)
                if combinado:
                    header, nombres_combinados = combinado
                inicio = header[-1] + 1
                fin = encabezados[numero + 1][0] if numero + 1 < len(encabezados) else len(crudo)
                estructura.update(i + 1 for i in header)
                # Un cuadro apilado que empieza con la fila de códigos de serie (planillas de un ministerio: anual arriba,
                # trimestral y mensual debajo) no repite los rótulos: son los del cuadro de arriba.
                codigos_encabezado = es_fila_codigos_serie(crudo.iloc[header[-1]].tolist())
                if codigos_encabezado:
                    detalle.setdefault("filas_codigos_serie", []).append(header[-1] + 1)
                # Una fila de unidades entre el encabezado y los datos es metadato, no otra observación; también la
                # fila de códigos de serie.
                unidades_fila = {}
                while inicio < fin:
                    valores = crudo.iloc[inicio].tolist()
                    if es_fila_codigos_serie(valores):
                        estructura.add(inicio + 1)
                        detalle.setdefault("filas_codigos_serie", []).append(inicio + 1)
                        inicio += 1
                        continue
                    presentes_unidad = [_compactar(v) for v in valores if _compactar(v)]
                    if not presentes_unidad or not all(UNIDAD_CELDA.fullmatch(v) for v in presentes_unidad):
                        break
                    unidades_fila.update({j: _compactar(v) for j, v in enumerate(valores) if _compactar(v)})
                    estructura.add(inicio + 1)
                    detalle.setdefault("filas_unidades", []).append(inicio + 1)
                    inicio += 1
                if combinado:
                    columnas = _flatten_columnas(nombres_combinados)
                elif len(header) == 1:
                    columnas = _flatten_columnas(crudo.iloc[header[0]].tolist())
                else:
                    niveles = [
                        crudo.iloc[i].map(_compactar).replace("", None).ffill().fillna("").tolist() for i in header[:-1]
                    ]
                    niveles.append(crudo.iloc[header[-1]].map(_compactar).tolist())
                    columnas = _flatten_columnas(list(zip(*niveles)))
                raw_header = [_compactar(v) for v in crudo.iloc[header[-1]]]
                if combinado:
                    raw_header = nombres_combinados
                rotulos_heredados = None
                if codigos_encabezado and rotulos_previos:
                    # Cada columna con su rótulo del cuadro de arriba; una columna que allí no tenía rótulo
                    # conserva su código.
                    previas, crudos_previos, rotulos_heredados = rotulos_previos

                    def con_rotulo(j):
                        return j < len(crudos_previos) and j < len(previas) and bool(crudos_previos[j])
                    columnas = [previas[j] if con_rotulo(j) else columnas[j] for j in range(len(columnas))]
                    raw_header = [crudos_previos[j] if con_rotulo(j) else raw_header[j] for j in range(len(raw_header))]
                elif not codigos_encabezado:
                    rotulos_previos = (columnas, raw_header, [i + 1 for i in header])
                activas = [
                    j
                    for j in range(len(columnas))
                    if ANIO.fullmatch(raw_header[j]) or crudo.iloc[inicio:fin, j].notna().any()
                ]
                aisladas = encabezados_aislados(crudo.iloc[header[-1]].tolist(), crudo.iloc[inicio:fin].values.tolist())
                if aisladas:
                    referencias = ", ".join(f"{get_column_letter(j + 1)}{header[-1] + 1}" for j in aisladas)
                    detalle["advertencias"].append(
                        f"Celdas aisladas fuera del cuadro: {referencias}; sin datos asociados y separadas por columnas vacías. No se incorporan como períodos."
                    )
                    activas = [j for j in activas if j not in aisladas]
                paneles = paneles_anuales(crudo.iloc[header[-1]].tolist(), crudo.iloc[inicio:fin].values.tolist())
                todas_activas = activas
                for panel in paneles:
                    activas = [j for j in todas_activas if j in panel]
                    if not activas:
                        continue
                    from services.controles_excel.periodos import es_auxiliar

                    # Las columnas auxiliares («CV (%)» después de cada año) se repiten a propósito: se declaran aparte.
                    repetidos = sorted({raw_header[j] for j in activas if raw_header[j] and not es_auxiliar(raw_header[j])
                                        and sum(raw_header[k] == raw_header[j] for k in activas) > 1})
                    if repetidos:
                        avisos.append(f"Encabezados repetidos en fila {header[-1] + 1}: "
                                      f"{', '.join(re.sub(r' 00:00:00$', '', r) for r in repetidos)}. Se conservan como columnas distintas.")
                    vertical = any(raw_header[j].lower() in {"año", "ano", "anio", "periodo", "year"} for j in activas)
                    grupos, actual = [], []
                    for i in range(inicio, fin):
                        # Una fila de códigos de serie dentro del cuadro separa bloques apilados (anual, trimestral,
                        # mensual en las planillas de un ministerio): es metadato y cada bloque se lee aparte.
                        if es_fila_codigos_serie(crudo.iloc[i].tolist()):
                            if actual:
                                grupos.append(actual)
                                actual = []
                            estructura.add(i + 1)
                            detalle.setdefault("filas_codigos_serie", []).append(i + 1)
                            continue
                        r = crudo.iloc[i, activas]
                        presentes = [v for v in r if _compactar(v)]
                        solo_rotulo = (
                            len(presentes) == 1
                            and isinstance(presentes[0], str)
                            and pd.isna(pd.to_numeric(presentes[0], errors="coerce"))
                        )
                        # «Total» con sus importes en fórmulas sin resultado guardado es una fila de
                        # datos no evaluable, no un título ni una nota.
                        if solo_rotulo and any((i + 1, j + 1) in sin_valor for j in activas):
                            solo_rotulo = False
                        if not presentes:
                            # El espacio visual dentro de un cuadro no separa un total
                            # de sus componentes. Un bloque ya cerrado por Total sí termina.
                            ultimo = _compactar(crudo.iloc[actual[-1], activas[0]]).lower() if actual else ""
                            if actual and not vertical and ultimo.split(" ")[0] in {"total", "subtotal"}:
                                grupos.append(actual)
                                actual = []
                            continue
                        if not presentes or solo_rotulo:
                            if actual:
                                grupos.append(actual)
                                actual = []
                            if presentes:
                                estructura.add(i + 1)
                            continue
                        if UNIDAD_FILA.match(_compactar(r.iloc[0])):
                            if actual:
                                grupos.append(actual)
                                actual = []
                            grupos.append([i])
                        else:
                            actual.append(i)
                    if actual:
                        grupos.append(actual)
                    cortes = {}
                    if not vertical:
                        # Un «Total» que quedó solo porque filas sin datos lo separan de sus componentes.
                        grupos, cortes = _unir_total_suelto(crudo, grupos, activas)
                    for indices in grupos:
                        frame = crudo.iloc[indices, activas].copy()
                        frame.columns = [columnas[j] for j in activas]
                        # Tablas de texto (guías, notas) quedan declaradas fuera del barrido cuantitativo.
                        if not any(_es_numero_tabular(v) for v in frame.to_numpy().flat):
                            estructura.update(i + 1 for i in indices)
                            continue
                        # Fechas como texto ISO («2022-06-01»): to_json las pasaría a milisegundos desde 1970.
                        frame = frame.apply(lambda col: col.map(_fecha_iso))
                        registros = json.loads(frame.to_json(orient="records", double_precision=15))
                        _leer_numeros_como_texto(
                            registros, [columnas[j] for j in activas], activas, indices, preferencia,
                            interpretar_columna, detalle, avisos,
                        )
                        start = len(datos)
                        datos.extend(registros)
                        usadas.update(i + 1 for i in indices)
                        cubiertas.update((i, j) for i in indices for j in activas)
                        bloques.append(
                            {
                                "archivo": archivo,
                                "hoja": hoja,
                                "bloque": len(bloques) + 1,
                                "inicio": start,
                                "fin": len(datos),
                                "filas": [i + 1 for i in indices],
                                "encabezados": [i + 1 for i in header],
                                # Filas de rótulos heredadas del cuadro de arriba (encabezado con códigos de serie).
                                **({"encabezados_rotulos": rotulos_heredados} if rotulos_heredados else {}),
                                "periodos_duplicados": sorted(
                                    {
                                        raw_header[j]
                                        for j in activas
                                        if ANIO.fullmatch(raw_header[j])
                                        and sum(raw_header[k] == raw_header[j] for k in activas) > 1
                                    }
                                ),
                                "columnas": {columnas[j]: get_column_letter(j + 1) for j in activas},
                            }
                        )
                        # Unidad anotada en el rótulo del período («2024 (miles de asistentes)»):
                        # metadato de la columna, comparable con la del resto del cuadro.
                        unidades_columna = _unidades_columna(crudo.iloc[header[-1]].tolist(), activas, columnas)
                        unidades_columna.update({columnas[j]: {"letra": get_column_letter(j + 1),
                                                              "unidad": unidades_fila[j], "rotulo": unidades_fila[j]}
                                                 for j in activas if j in unidades_fila})
                        if unidades_columna:
                            bloques[-1]["unidades_columna"] = unidades_columna
                        if tuple(indices) in cortes:
                            bloques[-1]["total_separado"] = cortes[tuple(indices)]
                        _declarar_encabezado_corrido(crudo, header[-1], indices, activas, detalle, avisos)
                        _declarar_columnas_fuera(crudo.iloc[header[-1]].tolist(), activas, columnas, frame, header[-1], detalle, avisos)
                        oculto = {
                            "hoja_oculta": hoja_oculta,
                            "filas_ocultas": [i + 1 for i in indices if i + 1 in filas_ocultas],
                            "columnas_ocultas": {
                                columnas[j]: get_column_letter(j + 1) for j in activas if j + 1 in columnas_ocultas
                            },
                        }
                        if hoja_oculta or oculto["filas_ocultas"] or oculto["columnas_ocultas"]:
                            bloques[-1]["origen_oculto"] = oculto
                        detalle["bloques"] += 1
            for i, r in crudo.iterrows():
                fila = int(i) + 1
                if fila in usadas or fila in estructura or not r.notna().any():
                    continue
                presentes = [v for v in r if _compactar(v)]
                if len(presentes) == 1 and isinstance(presentes[0], str):
                    estructura.add(fila)
                else:
                    detalle["filas_no_interpretadas"].append(fila)
            if detalle.get("filas_codigos_serie"):
                filas_cod = sorted(set(detalle["filas_codigos_serie"]))
                detalle["filas_codigos_serie"] = filas_cod
                avisos.append(f"{contar(len(filas_cod), 'fila', 'filas')} con códigos de serie ({muestra(filas_cod)}): "
                              "se leyeron como metadato, no como encabezado ni como datos; los cuadros de abajo llevan "
                              "los rótulos del cuadro de arriba.")
            _declarar_fuera_de_cuadros(crudo, cubiertas, estructura, detalle, avisos)
            _avisar_encabezado_con_datos(encabezados, crudo, detalle, avisos)
            _avisar_encabezado_entre_datos(crudo, usadas, detalle, avisos)
            detalle.update(
                filas_datos=len(usadas),
                filas_estructura=sorted(estructura),
                estado="incluida" if usadas else "sin_datos",
            )
            if not usadas:
                detalle["motivo"] = (
                    "Sin filas cuantitativas después del encabezado; texto conservado como estructura, fuera del barrido numérico"
                )
                _declarar_no_revisada(detalle, crudo, _celda_numerica)
            detalle["advertencias"].extend(avisos)
    if len(hojas_pct) > MAX_HOJAS_AVISO_PORCENTAJE:
        resto = hojas_pct[MAX_HOJAS_AVISO_PORCENTAJE:]
        celdas = sum(h["formato_porcentaje"]["celdas"] for h in resto)
        nombres = muestra(["«" + str(h["hoja"]) + "»" for h in resto], maximo=3)
        aviso = (f"Y {contar(len(resto), 'hoja más', 'hojas más')} ({nombres}) "
                 f"con {contar(celdas, 'celda', 'celdas')} en formato de porcentaje; se leyó el valor guardado.")
        ultima = hojas_pct[MAX_HOJAS_AVISO_PORCENTAJE - 1]
        ultima["avisos_lectura"].append(aviso)
        ultima["advertencias"].append(aviso)
    return datos, {
        "version": 1,
        "hojas": hojas,
        "bloques": bloques,
        "criterio": "Todas las hojas; controles por cuadro, sin comparar cuadros ni libros entre sí. Títulos, notas y encabezados no se tratan como observaciones numéricas.",
    }


def _unidades_columna(encabezado, activas, columnas):
    from services.controles_excel.periodos import leer

    salida = {}
    for j in activas:
        v = encabezado[j] if j < len(encabezado) else None
        p = leer(v) if isinstance(v, str) else None
        if p and p.get("unidad"):
            salida[columnas[j]] = {"letra": get_column_letter(j + 1), "unidad": p["unidad"], "rotulo": str(v).strip()}
    return salida


ANIO_EN_ROTULO = re.compile(r"(?<!\d)(?:19|20)\d{2}(?!\d)")
MEDIDA_DERIVADA = re.compile(r"var|%|partic|tasa|dif|prom|total|acum|crec|[ií]ndice|\bvs\b|/|\bi\.?a\b|interanual|cambio|brecha|ratio")


def _declarar_columnas_fuera(encabezado, activas, columnas, frame, fila_header, detalle, avisos):
    """Una columna con números cuyo rótulo menciona un año pero no se reconoce como período queda
    fuera de las series por período: se declara, no se pierde en silencio."""
    from services.controles_excel.periodos import leer

    periodos = [j for j in activas if j < len(encabezado) and leer(encabezado[j]) is not None]
    if len(periodos) < 2:
        return
    previas = {c["letra"] for c in detalle.get("columnas_no_controladas", [])}
    from services.controles_excel.periodos import es_auxiliar

    auxiliares = [j for j in activas if j not in periodos and j < len(encabezado) and es_auxiliar(encabezado[j])
                  and any(_es_numero_tabular(x) for x in frame[columnas[j]].tolist())]
    letras_aux = [get_column_letter(j + 1) for j in auxiliares]
    if auxiliares and ", ".join(letras_aux) not in previas:
        rotulos = list(dict.fromkeys(str(encabezado[j]).strip()[:40] for j in auxiliares))
        motivo = ("columnas auxiliares junto a los períodos: no son períodos ni componentes del cuadro; sus valores "
                  "no entran en las series por período ni en los cierres")
        detalle.setdefault("columnas_no_controladas", []).append(
            {"columna": ", ".join(rotulos[:3]), "letra": ", ".join(letras_aux),
             "celda": ", ".join(f"{x}{fila_header + 1}" for x in letras_aux), "motivo": motivo})
    for j in activas:
        v = encabezado[j] if j < len(encabezado) else None
        # Una columna con números y sin rótulo fuera del tramo de períodos tampoco entra en las series.
        sin_rotulo = not _presente(v) and not min(periodos) < j < max(periodos)
        if j in periodos or j in auxiliares or not (sin_rotulo or isinstance(v, str) and ANIO_EN_ROTULO.search(v)):
            continue
        # Medidas derivadas («Var. % 2023/22», «Participación 2023») no son un período con otro rótulo.
        if not sin_rotulo and MEDIDA_DERIVADA.search(v.lower()):
            continue
        if not any(_es_numero_tabular(x) for x in frame[columnas[j]].tolist()):
            continue
        letra = get_column_letter(j + 1)
        if letra in previas:
            continue
        motivo = ("columna sin rótulo fuera del tramo de períodos; sus valores no entran en las series por período" if sin_rotulo
                  else "el rótulo no se reconoce como período; sus valores no entran en las series por período")
        detalle.setdefault("columnas_no_controladas", []).append(
            {"columna": columnas[j], "letra": letra, "celda": f"{letra}{fila_header + 1}", "motivo": motivo})
        nombre = "sin rótulo" if sin_rotulo else f"«{str(v).strip()[:60]}»"
        avisos.append(f"Columna {nombre} ({letra}{fila_header + 1}): {motivo}.")


def _declarar_encabezado_corrido(crudo, fila_header, indices, activas, detalle, avisos):
    """Encabezado de períodos desplazado una columna respecto de los datos.

    La columna a la izquierda del primer período tiene números y ningún rótulo, y la del último
    período tiene rótulo y ningún dato: los períodos quedaron corridos una columna. No se
    realinea en silencio; se declara para que la persona verifique qué año es cada columna.
    """
    from services.controles_excel.periodos import leer

    encabezado = crudo.iloc[fila_header].tolist()
    periodos = [j for j in activas if j < len(encabezado) and _presente(encabezado[j]) and leer(encabezado[j]) is not None]
    if len(periodos) < 2:
        return
    izquierda, ultima = min(periodos) - 1, max(periodos)
    if izquierda < 0 or izquierda not in activas or _presente(encabezado[izquierda]):
        return
    cuerpo = crudo.iloc[indices]
    if not any(_es_numero_tabular(v) for v in cuerpo.iloc[:, izquierda].tolist()):
        return
    if any(_presente(v) for v in cuerpo.iloc[:, ultima].tolist()):
        return
    letra_izq, letra_ult = get_column_letter(izquierda + 1), get_column_letter(ultima + 1)
    motivo = (f"el encabezado de períodos parece corrido una columna: {letra_izq} tiene datos sin rótulo y "
              f"{letra_ult}{fila_header + 1} («{_rotulo_texto(encabezado[ultima])}») no tiene datos; "
              "los períodos de este cuadro pueden estar asignados a la columna equivocada")
    detalle.setdefault("columnas_no_controladas", []).append(
        {"columna": f"columna_{izquierda + 1}", "letra": letra_izq, "celda": f"{letra_izq}{fila_header + 1}", "motivo": motivo})
    avisos.append(f"Fila {fila_header + 1}: {motivo}.")


def _rotulo_texto(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()[:20]


def _declarar_fuera_de_cuadros(crudo, cubiertas, estructura, detalle, avisos):
    """Salvaguarda de cobertura: celdas numéricas que no quedaron en ningún cuadro leído.

    Un cuadro que no se reconoció (pegado a otro, sin encabezado, con otras columnas) deja sus
    números fuera de los controles; nunca se informa lectura completa con un bloque así.
    """
    if not cubiertas:
        return  # una hoja sin cuadros se declara entera como no revisada
    sueltas = {}
    for i in range(crudo.shape[0]):
        if i + 1 in estructura:
            continue
        for j in range(crudo.shape[1]):
            if (i, j) in cubiertas:
                continue
            v = crudo.iat[i, j]
            if _presente(v) and _es_numero_tabular(v):
                sueltas.setdefault(i + 1, []).append(j + 1)
    if sum(len(v) for v in sueltas.values()) < 2:
        return
    tramos, actual = [], []
    for fila in sorted(sueltas):
        if actual and fila > actual[-1] + 1:
            tramos.append(actual)
            actual = []
        actual.append(fila)
    tramos.append(actual)
    for tramo in tramos:
        n = sum(len(sueltas[f]) for f in tramo)
        columnas = sorted({c for f in tramo for c in sueltas[f]})
        filas = f"{tramo[0]}-{tramo[-1]}" if len(tramo) > 1 else str(tramo[0])
        letras = get_column_letter(columnas[0]) + (f"-{get_column_letter(columnas[-1])}" if len(columnas) > 1 else "")
        detalle.setdefault("fuera_de_cuadros", []).append({"filas": filas, "columnas": letras, "valores": n})
        avisos.append(
            f"Sin controlar: {'filas' if len(tramo) > 1 else 'fila'} {filas} (columnas {letras}): "
            f"{contar(n, 'valor numérico', 'valores numéricos')} fuera de los cuadros leídos."
        )


def _avisar_encabezado_entre_datos(crudo, usadas, detalle, avisos):
    """Una fila leída como datos que trae rótulos de período y ningún número es, casi siempre,
    el encabezado de otro cuadro que no se pudo separar: lo que sigue no se controló como
    cuadro propio y se declara."""
    from services.controles_excel.periodos import leer

    for fila in sorted(usadas):
        valores = [v for v in crudo.iloc[fila - 1].tolist() if _presente(v)]
        if any(_es_numero_tabular(v) and not isinstance(v, str) for v in valores):
            continue
        periodos = [str(v).strip() for v in valores if isinstance(v, str) and leer(v) is not None]
        if len(periodos) < 2:
            continue
        motivo = ("parece el encabezado de otro cuadro (" + ", ".join(f"«{p}»" for p in periodos[:3])
                  + ") que no se pudo separar; las filas siguientes no se controlaron como cuadro propio")
        detalle.setdefault("encabezados_no_separados", []).append({"fila": fila, "motivo": motivo})
        avisos.append(f"Fila {fila}: {motivo}.")


def _presente(v):
    return v is not None and not (isinstance(v, float) and math.isnan(v)) and str(v).strip() != ""


def _avisar_encabezado_con_datos(encabezados, crudo, detalle, avisos):
    """Una fila tomada como encabezado que trae importes (varios números con decimales, o
    números fuera de toda forma de calendario) probablemente era una fila de datos: no se
    pierde en silencio, se declara para que la persona verifique la lectura de la hoja."""
    from services.excel_estructura import anios_plausibles

    for header in encabezados:
        for i in header:
            nums = []
            for v in crudo.iloc[i].tolist():
                if isinstance(v, bool) or not isinstance(v, (int, float)) or (isinstance(v, float) and math.isnan(v)):
                    continue
                nums.append(float(v))
            if len(nums) < 2:
                continue
            decimales = sum(not n.is_integer() for n in nums)
            en_rango = all(1800 <= n <= 2200 for n in nums)
            if decimales >= 2 or (en_rango and not anios_plausibles(nums)):
                detalle.setdefault("encabezado_dudoso", []).append(i + 1)
                avisos.append(
                    f"La fila {i + 1} se leyó como encabezado pero contiene {contar(len(nums), 'valor numérico', 'valores numéricos')} "
                    "que parecen datos, no períodos ni rótulos: verificar la lectura de la hoja (sus valores no entran en los controles)."
                )


def _declarar_no_revisada(detalle, crudo, es_numerica):
    """Una hoja con números que no se incorporó queda declarada como no revisada, con su motivo."""
    celdas = sum(1 for v in crudo.to_numpy().flat if es_numerica(v))
    if not celdas:
        return
    motivo = detalle.get("motivo") or "No se reconoció un cuadro"
    detalle["no_revisada"] = {"celdas_numericas": int(celdas), "motivo": motivo}
    detalle.setdefault("avisos_lectura", []).append(
        f"Hoja no revisada: tiene {contar(celdas, 'celda numérica', 'celdas numéricas')} que no entraron en los controles ({motivo})."
    )


def _leer_numeros_como_texto(registros, nombres, indices_columna, filas_crudo, preferencia, interpretar, detalle, avisos):
    """Números guardados como texto («1.772.343.908», «1.234,5», «(300)») se leen con un único
    criterio por columna; «-» o «s/d» son sin dato; lo que no se puede leer se cuenta y se declara."""
    from services.perfilado import inferir_rol

    for nombre, j in zip(nombres, indices_columna):
        valores = [r.get(nombre) for r in registros]
        if not any(isinstance(v, str) for v in valores):
            continue
        # Códigos («1.10», «16.01»), clasificaciones y períodos escritos como texto se conservan:
        # convertirlos borraría ceros y distinciones que son parte del rótulo.
        if inferir_rol(nombre, valores) in {"identificador", "dimension", "temporal"}:
            continue
        nuevos, informe = interpretar(valores, preferencia)
        if not informe["numerica"]:
            continue
        from services.perfilado import tiene_llamada_nota

        con_nota = sum(1 for v in valores if tiene_llamada_nota(v))
        if con_nota:
            avisos.append(
                f"{contar(con_nota, 'celda', 'celdas')} de la columna «{nombre}» traen una llamada de nota pegada al número "
                "(por ejemplo «*» de dato provisorio); se leyeron sin la marca."
            )
        for r, v in zip(registros, nuevos):
            r[nombre] = v
        letra = get_column_letter(j + 1)
        if informe["convertidas"]:
            detalle.setdefault("numeros_como_texto", []).append(
                {"columna": nombre, "letra": letra, "celdas": informe["convertidas"], "formato": informe["formato"]}
            )
            avisos.append(
                f"{contar(informe['convertidas'], 'celda', 'celdas')} con números guardados como texto en columna «{nombre}» se leyeron como números "
                f"({'punto de miles y coma decimal' if informe['formato'] == 'es' else 'coma de miles y punto decimal'})."
            )
        suprimidos = set(informe.get("suprimidos") or ())
        if suprimidos:
            celdas = [f"{letra}{filas_crudo[k] + 1}" for k in sorted(suprimidos)]
            marcas = ", ".join(dict.fromkeys(f"«{str(valores[k]).strip()}»" for k in sorted(suprimidos)))
            avisos.append(
                f"{contar(len(celdas), 'celda', 'celdas')} con dato suprimido por confidencialidad estadística ({marcas}) "
                f"en columna «{nombre}» ({muestra(celdas)}): no son un número; los totales que las incluyen no se pueden comprobar."
            )
        menores = set(informe.get("menores") or ())
        if menores:
            celdas = [f"{letra}{filas_crudo[k] + 1}" for k in sorted(menores)]
            marcas = ", ".join(dict.fromkeys(f"«{str(valores[k]).strip()}»" for k in sorted(menores)))
            avisos.append(
                f"{contar(len(celdas), 'celda', 'celdas')} con valor menor que la unidad mínima publicada ({marcas}) "
                f"en columna «{nombre}» ({muestra(celdas)}): no se informa la cifra; los totales que las incluyen no se pueden comprobar."
            )
        suprimidos |= menores
        if [k for k in informe["sin_dato"] if k not in suprimidos]:
            celdas = [f"{letra}{filas_crudo[k] + 1}" for k in informe["sin_dato"] if k not in suprimidos]
            previo = detalle.setdefault("celdas_sin_dato", {"celdas": 0, "ejemplos": []})
            previo["celdas"] += len(celdas)
            previo["ejemplos"] = (previo["ejemplos"] + celdas)[:10]
            # Se cita la marca que usa la planilla («///», «s/i», «-»), no una de ejemplo.
            marcas = ", ".join(dict.fromkeys(f"«{str(valores[k]).strip()}»" for k in informe["sin_dato"]
                                             if k not in suprimidos and str(valores[k] or "").strip()))
            avisos.append(
                f"{contar(len(celdas), 'celda marcada', 'celdas marcadas')} sin dato ({marcas or '«-», «s/d»'}) en columna «{nombre}» ({muestra(celdas)}): se tratan como sin dato, no como cero ni como un valor."
            )
        if informe["no_interpretables"]:
            from services.carga_archivos import muestra_segura

            ejemplos = [
                f"{letra}{filas_crudo[k] + 1}: «{muestra_segura(nombre, valores, valores[k])}»"
                for k in informe["no_interpretables"][:5]
            ]
            detalle.setdefault("celdas_no_numericas", []).append(
                {"columna": nombre, "letra": letra, "celdas": len(informe["no_interpretables"]), "ejemplos": ejemplos}
            )
            avisos.append(
                f"{contar(len(informe['no_interpretables']), 'celda no interpretable', 'celdas no interpretables')} como número en columna «{nombre}» ({muestra(ejemplos, len(informe['no_interpretables']), 3, '; ')}): se conservan como texto."
            )


def preparar(datos, cobertura, normalizar, sanear):
    """Normaliza/sanea por cuadro. Ubicaciones sólo si cada registro se conserva."""
    cobertura = deepcopy(cobertura)
    finales, reportes, normalizaciones = [], [], []
    for bloque in cobertura["bloques"]:
        originales = datos[bloque["inicio"] : bloque["fin"]]
        nuevos, normalizacion = normalizar(originales)
        normalizados = nuevos
        nuevos, reporte = sanear(normalizados)
        # Cada registro conserva sólo sus columnas y en su orden (sin el BOM que limpia sanear).
        from services.sanitizacion import limpiar_nombre_columna

        nuevos = [
            {limpiar_nombre_columna(c): saneada[limpiar_nombre_columna(c)] for c in original if limpiar_nombre_columna(c) in saneada}
            for original, saneada in zip(normalizados, nuevos)
        ]
        mapa = {c: v for c, v in bloque.get("columnas", {}).items() if c not in reporte["columnas_excluidas"]}
        originales_sanos, _ = sanear(originales)
        indices, desde = [], 0
        for fila in nuevos:
            encontrado = next((i for i in range(desde, len(originales_sanos)) if originales_sanos[i] == fila), None)
            if encontrado is None:
                indices = []
                break
            indices.append(encontrado)
            desde = encontrado + 1
        filas = bloque.get("filas") or []
        bloque["filas"] = [filas[i] for i in indices] if filas and len(indices) == len(nuevos) else []
        bloque["columnas"] = mapa if bloque["filas"] else {}
        bloque["filas_entrada"] = len(originales)
        bloque["normalizacion_series"] = normalizacion
        bloque["columnas_excluidas"] = reporte["columnas_excluidas"]
        if nuevos and not any(any(v is not None for v in fila.values()) for fila in nuevos):
            nuevos = []
        bloque["inicio"] = len(finales)
        finales.extend(nuevos)
        bloque["fin"] = len(finales)
        bloque["filas_preparadas"] = len(nuevos)
        reportes.append(reporte)
        normalizaciones.append(normalizacion)
    for hoja in cobertura["hojas"]:
        bs = [b for b in cobertura["bloques"] if (b.get("archivo"), b.get("hoja")) == (hoja["archivo"], hoja["hoja"])]
        hoja["filas_preparadas"] = sum(b["fin"] - b["inicio"] for b in bs)
    excluidas = sorted({c for r in reportes for c in r["columnas_excluidas"]})
    redacciones = {}
    for r in reportes:
        for c, n in r.get("redacciones_por_columna", {}).items():
            redacciones[c] = redacciones.get(c, 0) + n
    reporte = {
        "columnas_excluidas": excluidas,
        "valores_redactados": sum(redacciones.values()),
        "redacciones_por_columna": redacciones,
        "excel": cobertura,
        "etapa_sanitizacion": {
            "ejecutada": True,
            "resultado": "informacion_agregada_disponible" if finales else "sin_informacion_agregada",
            "exclusiones": [
                {"columna": c, "motivo": "Identifica actores o documentos individuales"} for c in excluidas
            ],
            "valores_enmascarados": sum(redacciones.values()),
            "filas_restantes": len(finales),
            "columnas_restantes": sorted({c for fila in finales for c in fila}),
        },
    }
    return finales, reporte, {"formato": "por_cuadros", "bloques": len(normalizaciones)}


def ubicar(hallazgo, bloque):
    h = deepcopy(hallazgo)
    ev = h.setdefault("evidencia", {})
    ev.update(archivo_origen=bloque.get("archivo"), hoja=bloque.get("hoja"), bloque_excel=bloque["bloque"])
    columnas = ev.get("columnas") or [ev.get("columna"), ev.get("columna_mayor"), ev.get("columna_menor")]
    indices = ev.get("filas") or []
    letras = [bloque.get("columnas", {}).get(c) for c in columnas]
    filas = bloque.get("filas") or []
    ubicaciones = [
        f"{bloque['hoja']}!{col}{filas[i]}"
        for i in indices
        if isinstance(i, int) and 0 <= i < len(filas)
        for col in letras
        if col
    ]
    if ubicaciones:
        cols = sorted({column_index_from_string(c) for c in letras if c})
        if cols:
            ubicaciones = [
                f"{bloque['hoja']}!{get_column_letter(cols[0])}{filas[i]}"
                + (f":{get_column_letter(cols[-1])}{filas[i]}" if len(cols) > 1 else "")
                for i in indices
                if isinstance(i, int) and 0 <= i < len(filas)
            ]
        ev["ubicaciones_excel"] = ubicaciones
    if indices:
        ev["filas"] = [i + bloque["inicio"] for i in indices if isinstance(i, int)]
    if ev.get("detalle_series"):
        ev["detalle_series"] = [
            ubicar({"descripcion": "", "evidencia": d}, bloque)["evidencia"] for d in ev["detalle_series"]
        ]
        ev["ubicaciones_excel"] = list(
            dict.fromkeys(u for d in ev["detalle_series"] for u in d.get("ubicaciones_excel", []))
        )
    oculto = bloque.get("origen_oculto") or {}
    if oculto:
        referidas = {filas[i] for i in indices if isinstance(i, int) and 0 <= i < len(filas)}
        marcas = []
        if oculto.get("hoja_oculta"):
            marcas.append(f"hoja «{bloque.get('hoja')}» oculta")
        filas_ocultas = sorted(referidas & set(oculto.get("filas_ocultas") or [])) if referidas else list(oculto.get("filas_ocultas") or [])
        if filas_ocultas:
            marcas.append(f"fila{'s' if len(filas_ocultas) > 1 else ''} {_rangos(filas_ocultas)} oculta{'s' if len(filas_ocultas) > 1 else ''}")
        cols = [letra for c, letra in (oculto.get("columnas_ocultas") or {}).items() if not any(columnas) or c in columnas]
        if cols:
            marcas.append(f"columna{'s' if len(cols) > 1 else ''} {', '.join(cols)} oculta{'s' if len(cols) > 1 else ''}")
        if marcas:
            # Campo para que el informe muestre que el dato sale de contenido no visible.
            ev["visibilidad_origen"] = marcas
    h["descripcion"] = f"[{bloque.get('hoja') or bloque.get('archivo')} · cuadro {bloque['bloque']}] {h['descripcion']}"
    return h


def ejecutar_por_bloques(db, informe, ejecutar_unidad):
    from services import validez

    meta = informe.saneamiento or {}
    cobertura = deepcopy(meta["excel"])
    todo, estados = {}, {}
    for bloque in cobertura["bloques"]:
        datos = informe.datos[bloque["inicio"] : bloque["fin"]]
        unidad = SimpleNamespace(
            id=informe.id,
            datos=datos,
            dominio=informe.dominio,
            periodo=informe.periodo,
            area_id=informe.area_id,
            creado_en=informe.creado_en,
            saneamiento={
                k: v
                for k, v in meta.items()
                if k
                not in {"excel", "formula_hallazgos", "matematico_hallazgos", "origen_excel", "controles_tabulares"}
            },
        )
        unidad.saneamiento["periodos_excel_duplicados"] = bloque.get("periodos_duplicados", [])
        from services import interpretacion_previa as ip
        if meta.get(ip.CAMPO):
            interpretada = ip.unidad(bloque["archivo"], bloque["hoja"], bloque["bloque"], meta[ip.CAMPO])
            unidad.saneamiento["interpretacion_unidad"] = interpretada
            unidad.dominio = ip.dominio(interpretada)
        condicion = validez.evaluar(unidad)
        bloque["condiciones_validez"] = condicion
        if not condicion["puede_concluir"]:
            bloque["controles"] = {
                nombre_lente: {"estado": "no_evaluable", "motivo": "Sin datos agregados o condiciones indispensables pendientes"}
                for nombre_lente in ("ajedrecista", "cronista", "orquesta", "radiologo", "rizoma", "matematico", "conciliador")
            }
        else:
            finales = []
            for lente, nuevos in ejecutar_unidad(
                db, unidad, previos=None if len(cobertura["bloques"]) == 1 and db is not None else []
            ):
                if lente == "_crudos_finales":
                    finales = nuevos
            for h in finales:
                todo.setdefault(h["lente"], []).append(ubicar(h, bloque))
            bloque["controles"] = unidad._condiciones_lentes
        if meta.get(ip.CAMPO):
            bloque["trazabilidad_controles"] = ip.trazabilidad_controles(
                interpretada, bloque["controles"],
                modulo_tabular=bool(meta.get("version_controles_tabulares")),
                opciones=meta.get("opciones_tecnicas") or {})
        for lente, estado in bloque["controles"].items():
            estados.setdefault(lente, []).append(estado)
    for campo, lente in [("formula_hallazgos", "ajedrecista"), ("matematico_hallazgos", "matematico")]:
        todo.setdefault(lente, []).extend({**h, "lente": lente} for h in meta.get(campo, []))
    for h in meta.get("controles_tabulares", []):
        todo.setdefault(h["lente"], []).append(h)
    condiciones = {}
    for lente, items in estados.items():
        n = sum(x["estado"] == "evaluado" for x in items)
        condiciones[lente] = {
            "estado": "evaluado"
            if n
            else ("no_aplicable" if all(x["estado"] == "no_aplicable" for x in items) else "no_evaluable"),
            "motivo": items[0]["motivo"]
            if len(items) == 1
            else f"Evaluable en {n} de {len(items)} cuadros; consultar el detalle por hoja.",
            "bloques_evaluados": n,
            "bloques_totales": len(items),
        }
    for lente, items in estados.items():
        if not condiciones[lente]['bloques_evaluados'] and len(items)>1:
            motivos=list(dict.fromkeys(x.get('motivo','') for x in items))
            condiciones[lente]['motivo']='; '.join(motivos)+f'. No ejecutado en {len(items)} cuadros.'
    for h in meta.get("controles_tabulares", []):
        condiciones[h["lente"]] = {
            **condiciones.get(h["lente"], {}),
            "estado": "evaluado",
            "motivo": "Controles por cuadro y controles tabulares de estructura, escala y conciliación ejecutados.",
        }
    informe._condiciones_lentes = condiciones
    informe._antecedentes_excluidos = []
    informe._condiciones_interdocumentales = (
        {
            "estado": "no_evaluable",
            "motivo": "Los cuadros se revisaron por separado; no se estableció comparabilidad entre hojas, libros ni informes anteriores.",
        }
        if len(cobertura["bloques"]) > 1
        else {"estado": "no_aplicable", "motivo": "Un único cuadro; no hay otros cuadros que conciliar."}
    )
    informe._excel_ejecucion = cobertura
    for lente in ("ajedrecista", "cronista", "orquesta", "radiologo", "rizoma", "matematico", "conciliador"):
        yield lente, todo.get(lente, [])
    finales = [h for hs in todo.values() for h in hs]
    if meta.get("version_controles_tabulares"):
        from services.prioridad_controles import consolidar_contexto

        finales = consolidar_contexto(finales)
    yield "_crudos_finales", finales


def resumen_lineas(cobertura):
    if not cobertura:
        return []
    hojas = cobertura.get("hojas", [])
    incluidas = sum(h.get("filas_preparadas", h.get("filas_datos", 0)) > 0 for h in hojas)
    lineas = []
    if hojas:
        lineas.append(
            f"Excel: {incluidas} de {len(hojas)} hojas con datos incorporados. Los controles se ejecutan por cuadro; incorporar una hoja no significa que todos los controles sean evaluables."
        )
    archivos = cobertura.get("archivos") or []
    if len(archivos) > 1 or any(a.get("estado") == "fallido" for a in archivos):
        leidos = sum(a.get("estado") != "fallido" for a in archivos)
        lineas.append(f"Archivos: {leidos} de {len(archivos)} leídos; cada archivo se revisa como un bloque propio.")
    for a in archivos:
        if a.get("estado") == "fallido":
            lineas.append(f"{a['archivo']}: no se procesó. {a.get('motivo', '')}".strip())
            continue
        if a.get("tipo_detectado") in {"csv", "pdf"}:
            avisos = a.get("avisos", [])
        else:
            avisos = [x for x in a.get("avisos", []) if "la extensión es" in x]
        if len(archivos) == 1 and not avisos:
            continue
        linea = f"{a['archivo']}: lectura como {a.get('tipo_lectura') or a.get('tipo_detectado')}, {a.get('filas', 0)} filas."
        if avisos:
            linea += " " + " ".join(avisos)
        lineas.append(linea)
    for h in hojas:
        linea = f"{h['archivo']} · {h['hoja']}: {h.get('filas_preparadas', h.get('filas_datos', 0))} filas incorporadas, {h['bloques']} cuadros."
        if h.get("motivo"):
            linea += " " + h["motivo"] + "."
        if h.get("filas_estructura"):
            linea += f" {len(h['filas_estructura'])} filas de títulos, notas o encabezados separadas de los datos."
        if h.get("filas_no_interpretadas"):
            linea += " Filas no interpretadas: " + ", ".join(map(str, h["filas_no_interpretadas"])) + "."
        if h.get("no_revisada") and not h.get("advertencias"):
            linea += f" Hoja no revisada: {contar(h['no_revisada']['celdas_numericas'], 'celda numérica', 'celdas numéricas')} sin leer."
        if h.get("advertencias"):
            linea += " " + " ".join(h["advertencias"])
        lineas.append(linea)
    return lineas


def evaluar_carga(informe):
    from services import validez

    bloques = (informe.saneamiento or {}).get("excel", {}).get("bloques", [])
    if not bloques:
        return validez.evaluar(informe)
    resultados = []
    for b in bloques:
        unidad = SimpleNamespace(datos=informe.datos[b["inicio"] : b["fin"]], saneamiento=informe.saneamiento)
        resultados.append((b, validez.evaluar(unidad)))
    salida = {
        "estado": "verificacion_parcial",
        "puede_concluir": any(r["puede_concluir"] for b, r in resultados),
        "esquema": validez.esquema(informe.datos),
        "cambio_esquema": any(r["cambio_esquema"] for b, r in resultados),
    }
    for tipo in ["condiciones_verificadas", "condiciones_no_verificadas", "condiciones_fallidas"]:
        agrupadas = {}
        for b, r in resultados:
            for c in r[tipo]:
                key = (c["condicion"], c.get("motivo"))
                if key not in agrupadas:
                    agrupadas[key] = {**c, "bloques": []}
                agrupadas[key]["bloques"].append(b["bloque"])
        salida[tipo] = list(agrupadas.values())
    if not salida["puede_concluir"]:
        salida["estado"] = "no_evaluable"
    return salida
