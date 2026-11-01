"""Orquesta el análisis completo de un informe ya persistido: lentes disponibles,
triangulación entre modelos y devolución al responsable de decisión. Compartido por el flujo
manual (`POST /informes/{id}/analizar`) y por la carga rápida de archivos."""
from __future__ import annotations

import logging
from collections.abc import Sequence
from typing import Any

from sqlalchemy.orm import Session

import models
import schemas
from enums import Lente
from services import ambito as ambito_svc
from services import casos_revision as casos_revision_svc
from services import controles_solicitados as controles_solicitados_svc
from services import convergencia as convergencia_svc
from services import devolucion as devolucion_svc
from services import guardrails, plural, proveniencia
from services import perfilado as perfilado_svc
from services import triangulacion as triangulacion_svc
from services import triangulacion_validacion as validacion_ia
from services.lentes import ejecutar_lentes_stream
from versiones import (
    ANALYSIS_VERSION,
    APP_VERSION,
    AUTOMATIC_PROTOCOL_VERSION,
    DOCUMENT_VERSION,
    version_codigo,
    version_control,
)


def _es_registro_pdf(fila: dict) -> bool:
    return (
        str(fila.get("tipo_documento") or "").lower() == "pdf"
        or str(fila.get("archivo_origen") or "").lower().endswith(".pdf")
    )


def _nota_limite_pdf(informe: models.Informe) -> str:
    meta = (informe.saneamiento or {}).get("pdf") or {}
    if not isinstance(meta, dict) or not meta.get("truncado_por_paginas"):
        return ""
    analizadas = meta.get("paginas_analizadas") or "-"
    total = meta.get("paginas_documento") or "-"
    limite = meta.get("limite_paginas_pdf") or "-"
    return f" Se analizaron {analizadas} de {plural.contar(total, 'página')}, con límite aplicado de {limite}; no es lectura completa del PDF."


def _orientacion_analisis(informe: models.Informe) -> str:
    saneamiento = informe.saneamiento or {}
    if not isinstance(saneamiento, dict):
        return ""
    opciones = saneamiento.get("opciones_tecnicas")
    opciones = opciones if isinstance(opciones, dict) else {}
    valor = saneamiento.get("orientacion_analisis") or opciones.get("orientacion_analisis")
    return str(valor or "").strip()


def _controles_solicitados(informe: models.Informe) -> dict:
    saneamiento = informe.saneamiento or {}
    if not isinstance(saneamiento, dict):
        return {}
    valor = saneamiento.get("controles_solicitados")
    return valor if isinstance(valor, dict) else {}


logger = logging.getLogger("gobernanza.analisis")


def _estado_lectura(contenido: object, estado_tareas: object, status: object) -> str:
    tareas = estado_tareas if isinstance(estado_tareas, dict) else {}
    estado = (tareas.get("lectura_senales") or {}).get("estado")
    if estado:
        return str(estado)
    if status == "skipped":
        return "omitida"
    if not contenido:
        return "fallida"
    if isinstance(contenido, dict) and not validacion_ia.tiene_lectura(contenido) and (
        validacion_ia.es_lectura_sustituta(contenido) or "hallazgos_exploratorios" in contenido
    ):
        # Registro con exploración del libro pero sin lectura de señales.
        return "fallida"
    return "completada"


def estado_interpretacion_publico(modelo: object, contenido: object, error: object,
                                  estado_tareas: object = None, status: object = None) -> dict:
    """Explica el resultado de un proveedor sin publicar detalles internos.

    La lectura de señales y la exploración del libro son tareas separadas: un
    modelo cuya exploración respondió pero cuya lectura falló no "respondió".
    """
    nombre = str(modelo or "modelo").strip().lower()
    estado = _estado_lectura(contenido, estado_tareas, status)
    if estado == "completada":
        return {"modelo": nombre, "estado": "respondio", "detalle": "Respondió y su lectura fue incorporada."}
    if estado == "omitida":
        return {"modelo": nombre, "estado": "omitido", "detalle": "Modelo suspendido o fuera de servicio; no se realizó la llamada."}
    exploro = isinstance(contenido, dict) and "hallazgos_exploratorios" in contenido
    if estado == "no_solicitada":
        return {"modelo": nombre, "estado": "solo_exploracion" if exploro else "sin_respuesta",
                "detalle": "No había señales para leer; participó sólo de la exploración del libro." if exploro else "No había señales para leer."}
    mensaje = str(error or "").lower()
    if "timeout" in mensaje or "timed out" in mensaje:
        detalle = "No respondió dentro del tiempo disponible."
    elif "token" in mensaje or "truncad" in mensaje:
        detalle = "Su respuesta se cortó por límite de tokens; no acredita una lectura completa."
    elif "429" in mensaje or "resource_exhausted" in mensaje or "rate limit" in mensaje or "quota" in mensaje:
        detalle = "El proveedor rechazó la llamada por límite de uso."
    elif "sustituta" in mensaje:
        detalle = "No completó la lectura de señales."
    elif "json" in mensaje or "objeto" in mensaje or "formato" in mensaje or "esquema" in mensaje:
        detalle = "Respondió, pero no entregó un formato utilizable."
    elif "no configurad" in mensaje:
        detalle = "No estaba disponible en este entorno."
    else:
        detalle = "No pudo completar la lectura."
    if exploro:
        detalle += " Su exploración del libro se registró por separado."
    return {"modelo": nombre, "estado": "sin_respuesta", "detalle": detalle}


def modelos_consultados_publicos(interpretaciones: Sequence[Any], consolidado: dict | None = None, *, manifiesto: dict | None = None) -> list[dict]:
    descartadas = {d.get("modelo"): d for d in (consolidado or {}).get("lecturas_descartadas", [])}
    salida = []
    for item in interpretaciones:
        if isinstance(item, dict):
            publico = estado_interpretacion_publico(item.get("modelo"), item.get("contenido"), item.get("error"),
                                                    item.get("estado_tareas"), item.get("status"))
            if item.get("rol") or item.get("seleccion_modelo"):
                publico["rol"] = item.get("rol") or item.get("seleccion_modelo")
            if publico["estado"] == "respondio" and item.get("modelo") in descartadas:
                publico.update(estado="descartada", detalle="Respondió, pero su lectura se descartó: " + "; ".join(descartadas[item.get("modelo")].get("motivos") or []) + ".")
            salida.append(publico)
        else:
            registro = next((m for m in (manifiesto or {}).get("modelos", []) if m.get("proveedor") == item.modelo), {})
            salida.append(estado_interpretacion_publico(item.modelo, item.contenido, item.error, registro.get("estado_tareas")))
    return salida


CRITERIO_LENTE = {
    Lente.CONCILIADOR.value: "Conciliador — si una señal puede explicarse por criterios de carga, corte, acumulación, clasificación, cobertura o publicación.",
    Lente.AJEDRECISTA.value: "Ajedrecista — estados imposibles o fuera de lugar: totales que no cierran, series acumuladas que retroceden.",
    Lente.CRONISTA.value: "Cronista — cortes, saltos o cambios de régimen dentro de una misma serie temporal.",
    Lente.ORQUESTA.value: "Orquesta — consistencia entre documentos relacionados y magnitudes que deberían conciliar.",
    Lente.RADIOLOGO.value: "Radiólogo — semejanza con antecedentes realmente comparables cuando existe un banco adecuado (calibración actual: dominio presupuestario).",
    Lente.RIZOMA.value: "Rizoma — asociaciones no obvias entre variables, solo como hipótesis exploratorias.",
    Lente.MATEMATICO.value: "Matemático — patrones de cálculo: promedios sin ponderar, valores implícitos, hardcodes.",
}


def _barrido_por_lente(condiciones: dict | None = None) -> str:
    """Solo las lentes que realmente corrieron; las que no aplicaban se nombran aparte."""
    condiciones = condiciones or {}
    corrieron, no_corrieron = [], []
    for lente in Lente:
        if lente.value not in CRITERIO_LENTE:
            continue
        estado = (condiciones.get(lente.value) or {}).get("estado", "evaluado")
        if estado == "evaluado":
            corrieron.append(f"- {CRITERIO_LENTE[lente.value]}")
        else:
            motivo = (condiciones.get(lente.value) or {}).get("motivo") or "no correspondía a este archivo"
            nombre = CRITERIO_LENTE[lente.value].split(" — ")[0]
            no_corrieron.append(f"- {nombre}: {motivo}.")
    texto = "\n\nQué se revisó y no encontró señales comunicables por encima del umbral:\n" + "\n".join(corrieron)
    if no_corrieron:
        texto += "\n\nQué no se revisó en este archivo:\n" + "\n".join(no_corrieron)
    return texto


def _mensaje_sin_hallazgos(informe: models.Informe) -> str:
    datos = informe.datos or []
    barrido = _barrido_por_lente(getattr(informe, "_condiciones_lentes", None))
    orientacion = _orientacion_analisis(informe)
    controles = list(_controles_solicitados(informe).get("controles") or [])
    if controles:
        resultados = "\n".join(
            f"- {control.get('nombre')}: {control.get('resultado')}"
            for control in controles
        )
        cierre_orientado = (
            f"\n\nPregunta declarada: {orientacion}\n"
            "Resultado de los controles específicos:\n"
            f"{resultados}"
        )
    else:
        cierre_orientado = (
            f"\n\nOrientación declarada: {orientacion}\n"
            "El barrido no encontró señales comunicables que respondan directamente a esa orientación. "
            "Esto no la confirma ni la descarta: puede requerir otros datos, comparaciones o criterios de validación."
            if orientacion else ""
        )
    if any(_es_registro_pdf(fila) for fila in datos):
        modos = {str(fila.get("extraccion_pdf") or "").lower() for fila in datos if _es_registro_pdf(fila)}
        if "tabla" in modos:
            filas_tabla = sum(1 for fila in datos if str(fila.get("extraccion_pdf") or "").lower() == "tabla")
            return (
                f"Se extrajeron {plural.contar(filas_tabla, 'fila tabular', 'filas tabulares')} del PDF, pero las lentes no encontraron señales "
                "comunicables por encima del umbral. Esto significa que hubo lectura estructurada, no una revisión documental exhaustiva."
                + _nota_limite_pdf(informe)
                + barrido
                + cierre_orientado
            )
        paginas = sum(1 for fila in datos if _es_registro_pdf(fila))
        return (
            f"Se pudo leer el PDF como texto en {plural.contar(paginas, 'página')}, pero no se extrajeron tablas suficientes para análisis cuantitativo. "
            "Esto no equivale a una revisión documental profunda: para PDFs, esta versión solo registra texto extraído y números detectados "
            "cuando no logra reconstruir tablas."
            + _nota_limite_pdf(informe)
            + barrido
            + cierre_orientado
        )
    return (
        "Se hizo un barrido adicional sobre esta información y no se encontraron señales comunicables con los "
        "controles ejecutados. Esto no garantiza que no haya nada, pero el barrido se hizo con método y quedó registrado."
        + barrido
        + cierre_orientado
    )


NOMBRE_LENTE_EVENTO = {
    Lente.AJEDRECISTA.value: "Ajedrecista",
    Lente.CRONISTA.value: "Cronista",
    Lente.ORQUESTA.value: "Orquesta",
    Lente.RADIOLOGO.value: "Radi\u00f3logo",
    Lente.RIZOMA.value: "Rizoma",
    Lente.MATEMATICO.value: "Matem\u00e1tico",
    Lente.CONCILIADOR.value: "Conciliador",
}


def _en_contexto(flujo, redeclarar):
    """Recorre `flujo` redeclarando el contexto antes de cada paso.

    Una respuesta en streaming reanuda el generador en un hilo con una copia nueva
    del contexto: sin esto, las lentes que corren después del primer evento perderían
    el ámbito del informe y las columnas declaradas como código.
    """
    iterador = iter(flujo)
    while True:
        redeclarar()
        try:
            paso = next(iterador)
        except StopIteration:
            return
        yield paso


def ejecutar_analisis_stream(db: Session, informe: models.Informe, modelos: list[str] | None = None):
    """Version generadora de ejecutar_analisis(): yieldea un evento de
    progreso por cada lente que termina y por cada modelo de IA que responde,
    y al final un evento "final" con el mismo resultado que devuelve la
    version no-streaming. ejecutar_analisis() (abajo) es un envoltorio que
    se queda solo con ese ultimo evento -- misma logica, un solo lugar de
    verdad, comparable con la version anterior hallazgo por hallazgo."""
    from services import casos_revision, columnas_declaradas
    from services import interpretacion_previa as ip
    interpretacion = ip.exigir(informe)
    # Revisión rápida: sin IA. No corren el lector del cuadro ni la lectura de señales con IA.
    rapida = ip.es_rapida(interpretacion)
    if informe.id is not None and getattr(informe, "devolucion", None) is None:
        # Restos de una corrida interrumpida después de guardar los hallazgos
        # del motor y antes de la devolución: nunca se entregaron; se rehacen
        # para no duplicarlos.
        # Las respuestas a señales anteriores quedan (con la copia de su señal); sólo se desvinculan.
        db.query(models.RespuestaSenal).filter(models.RespuestaSenal.informe_id == informe.id).update(
            {models.RespuestaSenal.hallazgo_id: None}, synchronize_session=False)
        db.query(models.Hallazgo).filter(models.Hallazgo.informe_id == informe.id).delete(synchronize_session=False)
    base_revision = casos_revision.buscar_base(db, informe)
    # Lo que la consulta declara como código tiene que tener efecto: se declara antes de que
    # corra cualquier lente, y siempre, aunque no haya nada, para que ninguna corrida herede la
    # restricción de la anterior.
    restriccion_declarada = columnas_declaradas.detectar(_orientacion_analisis(informe), informe.datos or [])
    perfilado_svc.declarar_identificadores(restriccion_declarada["columnas"])
    crudos: list[dict] = []
    from services.excel_cobertura import evaluar_carga
    condicion = evaluar_carga(informe)
    # Un cuadro que la IA no supo clasificar se controla igual: las cuentas no dependen de su tipo.
    flujo = ejecutar_lentes_stream(db, informe) if condicion["puede_concluir"] else []
    ambito_informe = ambito_svc.de_informe(informe)
    if len(interpretacion["cuadros"]) == 1:
        interpretada = interpretacion["cuadros"][0]
        regimen = interpretada["confirmada"]["regimen"]
        ambito_informe = "organizacion" if regimen == "privado" else "publico" if regimen == "publico" else ambito_informe
        restriccion_declarada["columnas"] = list(set(restriccion_declarada["columnas"]) | {
            c for c, rol in interpretada["confirmada"]["roles"].items() if rol == "identificador"})

    def redeclarar() -> None:
        perfilado_svc.declarar_identificadores(restriccion_declarada["columnas"])
        ambito_svc.declarar(ambito_informe)

    for lente, nuevos in _en_contexto(flujo, redeclarar):
        if lente == "_crudos_finales":
            # el conciliador ya no muta los dicts anteriores: esta es la lista
            # completa con las modulaciones (severidad/limites) aplicadas.
            crudos = nuevos
            continue
        crudos.extend(nuevos)
        yield {
            "tipo": "lente",
            "lente": lente,
            "nombre": NOMBRE_LENTE_EVENTO.get(lente, lente),
            "cantidad": len(nuevos),
            # Las variaciones anuales van al comprobante como contexto, no como señales:
            # la pantalla de progreso tiene que contarlas igual.
            "contexto": sum(1 for h in nuevos if casos_revision_svc.es_contexto(h)),
            "condiciones": getattr(informe, "_condiciones_lentes", {}).get(lente),
        }

    controles_solicitados = {"controles": [], "hallazgos": []}
    if condicion["puede_concluir"]:
        bloques_excel = (informe.saneamiento or {}).get("excel", {}).get("bloques", [])
        if bloques_excel:
            from services.excel_cobertura import ubicar
            for bloque in bloques_excel:
                u = ip.unidad(bloque["archivo"], bloque["hoja"], bloque["bloque"], interpretacion)
                if not u.get("plan", {}).get("solicitudes"):
                    continue
                parcial = controles_solicitados_svc.evaluar(informe.datos[bloque['inicio']:bloque['fin']], _orientacion_analisis(informe))
                for c in parcial['controles']:
                    controles_solicitados['controles'].append({**c, 'nombre':f"{bloque.get('hoja') or bloque.get('archivo')} · cuadro {bloque['bloque']}: {c['nombre']}"})
                controles_solicitados['hallazgos'].extend(ubicar(h,bloque) for h in parcial['hallazgos'])
        elif any(u.get("plan", {}).get("solicitudes") for u in interpretacion["cuadros"]):
            controles_solicitados = controles_solicitados_svc.evaluar(informe.datos or [], _orientacion_analisis(informe))
    controles_solicitados["estado"] = controles_solicitados_svc.estado_controles(controles_solicitados["controles"])
    controles_solicitados["orientacion"] = _orientacion_analisis(informe)
    if controles_solicitados["controles"]:
        crudos.extend(controles_solicitados["hallazgos"])
        saneamiento = dict(informe.saneamiento or {})
        saneamiento["controles_solicitados"] = {
            clave: valor
            for clave, valor in controles_solicitados.items()
            if clave != "hallazgos"
        }
        informe.saneamiento = saneamiento
        db.add(informe)
        for control in controles_solicitados["controles"]:
            yield {
                "tipo": "control_solicitado",
                "codigo": control["codigo"],
                "nombre": control["nombre"],
                "estado": control["estado"],
                "detalle": control["resultado"],
            }

    if condicion["puede_concluir"] and not rapida:
        # La lectura ya ocurrió antes de confirmar; ahora sólo se hacen las cuentas.
        from services import lector_mapa
        if (informe.saneamiento or {}).get("grillas_lectura"):
            yield {"tipo": "lector_inicio"}
        try:
            lectura = ip.verificar_relaciones(informe.saneamiento or {})
        except Exception:  # noqa: BLE001 — una falla del lector nunca frena el análisis
            logger.exception("lector del cuadro")
            lectura = lector_mapa.lectura_fallida(informe.saneamiento or {}, "analizar")
            yield {"tipo": "lector_error"}
        if lectura is not None:
            crudos.extend(lectura["senales"])
            saneamiento = dict(informe.saneamiento or {})
            saneamiento["mapa_lectura"] = lectura["estado"]
            informe.saneamiento = saneamiento
            db.add(informe)
            yield {"tipo": "lector", "modo": lectura["estado"]["modo"], "senales": lectura["estado"]["senales"],
                   "preguntas": len(lectura["estado"]["preguntas"]),
                   "resumen": lector_mapa.resumen_lectura(lectura["estado"])}

    origen_excel = (informe.saneamiento or {}).get("origen_excel") or {}
    if origen_excel:
        from openpyxl.utils import column_index_from_string, get_column_letter
        for h in crudos:
            ev = h.get("evidencia") or {}
            columnas = ev.get("columnas") or [ev.get("columna_mayor"), ev.get("columna_menor"), ev.get("columna")]
            letras = [origen_excel['columnas'][c] for c in columnas if c in origen_excel['columnas']]
            filas = [origen_excel['filas'][i] for i in ev.get('filas', []) if isinstance(i, int) and 0 <= i < len(origen_excel['filas'])]
            if letras and filas and not ev.get('celda'):
                indices = sorted(column_index_from_string(c) for c in letras)
                ev['ubicaciones_excel'] = [f"{origen_excel['hoja']}!{get_column_letter(indices[0])}{fila}:{get_column_letter(indices[-1])}{fila}" for fila in filas]
                h['evidencia'] = ev
    crudos = guardrails.normalizar_hallazgos(crudos, informe)
    # Sondeo previo («radiografía»): una regla que aparece en muchos lugares de una hoja, o en la misma celda de
    # varias hojas, no se presenta como señal; queda como contexto y el comienzo del informe lo declara.
    from services import sondeo as sondeo_svc
    abstenciones = sondeo_svc.aplicar(
        crudos, lambda h: not (casos_revision_svc.es_nota(h) or casos_revision_svc.es_contexto(h)))
    saneamiento = dict(informe.saneamiento or {})
    if abstenciones or "sondeo" in saneamiento:
        saneamiento["sondeo"] = dict(version=sondeo_svc.VERSION, abstenciones=abstenciones)
        informe.saneamiento = saneamiento
        db.add(informe)

    hallazgos_orm = []
    for h in crudos:
        obj = models.Hallazgo(
            informe_id=informe.id,
            lente=h["lente"],
            modo=h.get("modo"),
            descripcion=h["descripcion"],
            severidad=h["severidad"],
            evidencia=h.get("evidencia"),
            tipo_evidencia=h.get("tipo_evidencia"),
            alcance=h.get("alcance") or "agregado",
            metodo=h.get("metodo"),
            incertidumbre=h.get("incertidumbre") or "media",
            estado_validacion=h.get("estado_validacion") or "pendiente_decision_humana",
            pregunta_institucional=h.get("pregunta_institucional"),
            resumen_llano=h.get("resumen_llano"),
            limites=h.get("limites"),
        )
        db.add(obj)
        hallazgos_orm.append(obj)
    db.flush()
    # Los hallazgos del motor se conservan aunque la etapa de IA falle.
    db.commit()

    interpretaciones: list[dict] = []
    contexto_semantico = (informe.saneamiento or {}).get("contexto_semantico") or {}
    contexto_informe = {
        "dominio": informe.dominio,
        "ambito": ambito_svc.de_informe(informe),
        "periodo": informe.periodo,
        "orientacion_analisis": _orientacion_analisis(informe),
        "contexto_semantico": contexto_semantico,
        "interpretacion_confirmada": [
            {"archivo": u.get("archivo"), "hoja": u.get("hoja"), "bloque": u.get("bloque"),
             "tipo": u.get("confirmada", {}).get("tipo"), "eje": u.get("confirmada", {}).get("eje"),
             "roles": u.get("confirmada", {}).get("roles", {})}
            for u in interpretacion.get("cuadros", [])
        ],
    }
    etapa_ia: dict[str, Any] = {"estado": "no_requerida", "error": None}
    if rapida:
        etapa_ia = {"estado": "no_solicitada", "error": None, "motivo": "Revisión rápida: sin IA."}
    consolidado = {"coincidencias": [], "divergencias": [], "hallazgos_singulares": []}
    try:
        if not rapida and (crudos or contexto_semantico.get("elementos")):
            etapa_ia["estado"] = "en_curso"
            for it in triangulacion_svc.interpretar_con_modelos_stream(
                contexto_informe, crudos, modelos
            ):
                interpretaciones.append(it)
                publico = estado_interpretacion_publico(
                    it["modelo"], it.get("contenido"), it.get("error"), it.get("estado_tareas"), it.get("status")
                )
                tareas = it.get("estado_tareas") or {}
                yield {
                    "tipo": "modelo",
                    "modelo": it["modelo"],
                    "rol": it.get("rol") or it.get("seleccion_modelo"),
                    "ok": publico["estado"] in {"respondio", "solo_exploracion"},
                    "lectura_senales": (tareas.get("lectura_senales") or {}).get("estado"),
                    "exploracion_libro": (tareas.get("exploracion_libro") or {}).get("estado"),
                    "exploracion_completada": (it.get("estrategia_semantica") or {}).get("evaluada"),
                    "motivos_exploracion": (it.get("estrategia_semantica") or {}).get("motivos",[]),
                    "error": None if publico["estado"] in {"respondio", "solo_exploracion"} else publico["detalle"],
                }
            consolidado = triangulacion_svc.consolidar(crudos, interpretaciones, contexto_informe)
            etapa_ia["estado"] = "completada"
    except Exception as exc:  # noqa: BLE001 - la IA es best-effort; el motor ya quedó guardado
        logger.exception("falló la etapa de IA del informe %s", informe.id)
        etapa_ia = {"estado": "error", "error": f"{type(exc).__name__}: {str(exc)[:300]}"}
        try:
            consolidado = triangulacion_svc.consolidar(crudos, interpretaciones, contexto_informe)
        except Exception:  # noqa: BLE001
            consolidado = {"coincidencias": [], "divergencias": [], "hallazgos_singulares": []}
    for it in interpretaciones:
        db.add(models.InterpretacionModelo(
            informe_id=informe.id,
            modelo=it["modelo"],
            identificador_version=it.get("identificador_version"),
            contenido=it.get("contenido") or {},
            error=it.get("error"),
        ))

    hallazgos_con_id = [
        {**h, "id": obj.id} for h, obj in zip(crudos, hallazgos_orm)
    ]
    convergencia = convergencia_svc.detectar_convergencia(hallazgos_con_id)
    if convergencia:
        yield {"tipo": "convergencia", "grupos": len(convergencia)}

    db.add(models.Triangulacion(
        informe_id=informe.id,
        coincidencias=consolidado["coincidencias"],
        divergencias=consolidado["divergencias"],
        hallazgos_singulares=consolidado["hallazgos_singulares"],
        convergencia_multilente=convergencia,
    ))

    cobertura = guardrails.cobertura_informe(informe, crudos)
    from services import semantica_excel
    try:
        cobertura["exploracion_semantica"] = semantica_excel.resumir(contexto_semantico, interpretaciones)
    except Exception as exc:  # noqa: BLE001
        logger.exception("falló el resumen de la exploración semántica del informe %s", informe.id)
        etapa_ia = {"estado": "error", "error": etapa_ia.get("error") or f"{type(exc).__name__}: {str(exc)[:300]}"}
        cobertura["exploracion_semantica"] = {
            "estado": "no_completada", "titulo_estado": "Exploración no completada", "resultado": "sin_conclusion",
            "hallazgos": [], "modelos": [], "limites": "La exploración con IA no pudo resumirse; los controles del motor no dependen de ella.",
        }
    cobertura["interpretacion_ia"] = {
        "etapa": etapa_ia,
        "linea_estado": consolidado.get("linea_estado", ""),
        "estado_modelos": consolidado.get("estado_modelos", []),
        "senales": consolidado.get("senales", []),
        "cobertura_lectura": consolidado.get("cobertura_lectura"),
        "lecturas_descartadas": consolidado.get("lecturas_descartadas", []),
        "priorizacion_sin_contraste": consolidado.get("priorizacion_sin_contraste", []),
        "respaldo_no_independiente": consolidado.get("respaldo_no_independiente", False),
    }
    cobertura["condiciones_validez"] = condicion
    revision = casos_revision.agrupar(hallazgos_orm)
    cobertura["revision"] = {k:v for k,v in revision.items() if k not in {"casos","contexto"}}
    from services.lectura_temporal import resumen as resumen_temporal
    cobertura["resumen_temporal"] = resumen_temporal(crudos)
    excel = getattr(informe, "_excel_ejecucion", (informe.saneamiento or {}).get("excel"))
    if excel:
        from services import biblioteca_casos, lector_mapa
        indice_experiencia = biblioteca_casos.indice(db, informe.usuario_id)
        fuentes = {f.get("nombre_original"): f.get("sha256_bruto") for f in
                   ((informe.saneamiento or {}).get("proveniencia") or {}).get("fuentes", [])}
        mapa_lector = (informe.saneamiento or {}).get("mapa_lectura") or {}
        for bloque in excel.get("bloques") or []:
            unidad = ip.unidad(bloque["archivo"], bloque["hoja"], bloque["bloque"], interpretacion)
            traza = bloque.get("trazabilidad_controles") or {}
            if traza.get("relaciones", {}).get("seleccionado"):
                if rapida:
                    traza["relaciones"].update(estado="no_solicitado", motivo="Revisión rápida: el lector del cuadro con IA "
                                               "no corrió; las relaciones aritméticas del motor sí se controlaron.")
                elif mapa_lector.get("fallo"):
                    traza["relaciones"].update(estado="no_evaluable", motivo="La verificación de relaciones no se completó.")
                else:
                    huella_cuadro = ip.huella([unidad.get("id"), interpretacion["revision"]])
                    grupos = [g for h in mapa_lector.get("hojas") or []
                              if h.get("huella") == huella_cuadro
                              for g in h.get("grupos") or []]
                    controladas = sum(lector_mapa._grupo_controlado(g, mapa_lector.get("confirmaciones") or {})
                                      for g in grupos)
                    traza["relaciones"].update(
                        estado="ejecutado" if controladas else "pendiente_confirmacion" if grupos else "sin_relaciones_acordadas",
                        motivo=(f"Se contrastaron {controladas} de {len(grupos)} relaciones acordadas."
                                if grupos else "El lector no acordó relaciones para contrastar en este cuadro."),
                        relaciones_acordadas=len(grupos), relaciones_controladas=controladas)
            bloque["trazabilidad_controles"] = traza
            bloque["experiencia_comparable"] = [
                biblioteca_casos.experiencia(indice_experiencia, unidad, codigo, fuentes.get(bloque["archivo"]))
                for codigo in (*ip.LENTES, "relaciones") if (unidad.get("plan") or {}).get(codigo)
            ]
        from services.excel_cobertura import resumen_lineas
        cobertura["excel"] = excel
        cobertura["excel_resumen"] = resumen_lineas(excel)
        # No comunicar controles negativos obtenidos de mezclar cuadros.
        cobertura["controles_negativos"] = {}
        libros = {b.get("archivo") for b in (excel.get("bloques") or []) if b.get("archivo")}
        cruces = ((cobertura.get("registro_revision") or {}).get("cruces") or {})
        comparacion = (
            (cruces.get("texto") or "Se cargó más de una edición: el resultado de la comparación figura en «Qué se leyó y qué se controló».")
            if len(libros) > 1
            else "Sin comparación con ediciones anteriores: si el informe se publica periódicamente, cargar también la edición anterior permite detectar valores históricos que cambiaron."
        )
        cobertura.setdefault("limites_globales", []).append(
            "Excel revisado por cuadros separados, con conciliación entre cuadros de la misma hoja cuando la relación se corrobora en varios años. "
            + comparacion
            + " Un importe reclasificado entre rubros que no altera el total y es chico frente a la variación habitual de las series no es detectable. Consultar cobertura por hoja."
        )
    if not condicion["puede_concluir"]:
        cobertura["estado"] = guardrails.ESTADO_PREGUNTA_NO_EVALUADA
        for item in cobertura["lentes"]:
            item.update(estado="no_evaluable", ejecutado=False)
    else:
        for item in cobertura["lentes"]:
            aplicabilidad = getattr(informe, "_condiciones_lentes", {}).get(item["lente"], {})
            item["ejecutado"] = aplicabilidad.get("estado", "evaluado") == "evaluado"
            if "bloques_totales" in aplicabilidad:
                item.update({k:v for k,v in aplicabilidad.items() if k != "estado"})
            if not item["ejecutado"]:
                item.update(aplicabilidad)
    cobertura["restricciones_declaradas"] = restriccion_declarada
    cobertura["antecedentes_excluidos"] = getattr(informe, "_antecedentes_excluidos", [])
    cobertura["control_interdocumental"] = getattr(informe, "_condiciones_interdocumentales", {})
    if condicion["puede_concluir"] and cobertura["estado"] != guardrails.ESTADO_PREGUNTA_NO_EVALUADA and (any(c["estado"] == "no_evaluable" for c in cobertura["lentes"]) or cobertura["control_interdocumental"].get("estado") == "no_evaluable" or ((excel or {}).get("lectura") or {}).get("estado") in {"parcial", "insuficiente"}):
        cobertura["estado"] = guardrails.ESTADO_COBERTURA_PARCIAL
    manifiesto = {
        "version_aplicacion": APP_VERSION, "version_protocolo": AUTOMATIC_PROTOCOL_VERSION,
        "version_documento_metodologico": DOCUMENT_VERSION,
        "version_motor": ANALYSIS_VERSION, "version_codigo": version_codigo(),
        "fecha": models.now().isoformat(), "dominio": informe.dominio,
        "sha256_procesado": proveniencia.huella(informe.datos or []),
        "proveniencia": (informe.saneamiento or {}).get("proveniencia", {"estado": "no_registrada_en_carga_historica"}),
        "sanitizacion": (informe.saneamiento or {}).get("etapa_sanitizacion", {"estado": "detalle_historico_no_registrado"}),
        "condiciones": condicion,
        "excel": excel,
        "control_interdocumental": cobertura["control_interdocumental"],
        "exploracion_semantica": cobertura["exploracion_semantica"],
        "antecedentes_excluidos": cobertura["antecedentes_excluidos"],
        "controles": [{**item, "version_control": version_control(item["lente"])} for item in cobertura["lentes"]],
        "modelos": [{"proveedor": it["modelo"], "version": it.get("identificador_version"),
                     "respondio": _estado_lectura(it.get("contenido"), it.get("estado_tareas"), it.get("status")) == "completada",
                     "rol": it.get("rol") or it.get("seleccion_modelo"), "estado_tareas": it.get("estado_tareas")} for it in interpretaciones],
        "etapa_ia": etapa_ia,
        "guardrails": [{"hallazgo_id": h["id"], **h["limites"]["guardrails"]} for h in hallazgos_con_id],
        "convergencias": convergencia,
        "naturalezas": [{"hallazgo_id": h["id"], "naturaleza": h["evidencia"]["naturaleza"]} for h in hallazgos_con_id],
    }
    manifiesto.update(informe_id=informe.id, snapshot_senales=casos_revision.snapshot(hallazgos_orm),
        opciones_tecnicas=(informe.saneamiento or {}).get('opciones_tecnicas'))
    for it, item in zip(interpretaciones, manifiesto['modelos']):
        item.update(intentos=it.get('intentos',1), errores_formato=it.get('errores_formato',[]), error=it.get('error'), alcance_ia=it.get('alcance_ia'), status=it.get('status'), diagnostico=it.get('diagnostico'), llamadas=it.get('llamadas',[]), estrategia_semantica=it.get('estrategia_semantica'))
    for it, item in zip(interpretaciones, manifiesto['modelos']):
        item.update(it.get('diagnostico') or {})
    ejecucion = models.EjecucionAnalisis(
        informe_id=informe.id, dominio=informe.dominio, version_aplicacion=APP_VERSION,
        version_protocolo=AUTOMATIC_PROTOCOL_VERSION, version_codigo=version_codigo(),
        manifiesto=manifiesto,
    )
    db.add(ejecucion)
    db.flush()
    manifiesto = {**manifiesto, "ejecucion_id": ejecucion.id}
    ejecucion.manifiesto = manifiesto
    cobertura["manifiesto"] = manifiesto
    cobertura['comparacion_corridas'] = casos_revision.comparar(base_revision,manifiesto) if base_revision else {
        'estado':'sin_base','motivo':'No hay una ejecución previa comparable para esta persona, área, dominio y fuente.'}
    yield {"tipo": "cobertura", "cobertura": cobertura}

    estado = str(cobertura.get("estado") or (
        guardrails.ESTADO_CON_SENALES if crudos else guardrails.ESTADO_SIN_SENALES
    ))
    if crudos:
        mensaje, hay_hallazgos = devolucion_svc.armar_mensaje(crudos, consolidado, informe, cobertura=cobertura)
    else:
        mensaje, hay_hallazgos = _mensaje_sin_hallazgos(informe), False
    if not condicion["puede_concluir"]:
        mensaje = "No evaluable: faltan o fallaron condiciones indispensables. Se requiere información adicional y una nueva ejecución; no se concluye conformidad."
    elif not crudos and cobertura["estado"] == guardrails.ESTADO_COBERTURA_PARCIAL:
        mensaje = "No se detectaron inconsistencias con los controles evaluables. La cobertura es parcial; esto no demuestra la veracidad del informe.\n\n" + "\n".join(
            f"- {NOMBRE_LENTE_EVENTO.get(c['lente'], c['lente'])}: {c['estado'].replace('_', ' ')}. {c.get('motivo', '')}"
            for c in cobertura["lentes"])
    if etapa_ia["estado"] == "error":
        mensaje += "\n\nLa etapa de interpretación con IA no se completó por un error interno. Los hallazgos del motor se conservaron y no dependen de ella."
    if consolidado.get("linea_estado"):
        mensaje += "\n\n" + consolidado["linea_estado"]
    exploratorios = cobertura["exploracion_semantica"]["hallazgos"]
    if exploratorios:
        mensaje += f"\n\nExploración semántica con IA: {plural.contar(len(exploratorios), 'hipótesis exploratoria generada con IA, pendiente', 'hipótesis exploratorias generadas con IA, pendientes')} de revisión humana. Ver evidencia comparativa en la sección de contexto. El resultado mecánico no incluye estas hipótesis."
    devolucion = models.Devolucion(
        informe_id=informe.id,
        mensaje=mensaje,
        hay_hallazgos=hay_hallazgos,
        estado=estado,
        cobertura=cobertura,
    )
    db.add(devolucion)

    db.commit()
    db.refresh(devolucion)

    resultado = schemas.AnalisisOut(
        informe_id=informe.id,
        hallazgos=[schemas.HallazgoOut.model_validate(h) for h in hallazgos_orm],
        coincidencias=consolidado["coincidencias"],
        divergencias=consolidado["divergencias"],
        hallazgos_singulares=consolidado["hallazgos_singulares"],
        modelos_consultados=modelos_consultados_publicos(interpretaciones, consolidado),
        convergencia_multilente=convergencia,
        devolucion=schemas.DevolucionOut.model_validate(devolucion),
    )
    yield {"tipo": "final", "resultado": resultado}


def ejecutar_analisis(
    db: Session, informe: models.Informe, modelos: list[str] | None = None
) -> schemas.AnalisisOut:
    ultimo: dict | None = None
    for evento in ejecutar_analisis_stream(db, informe, modelos):
        ultimo = evento
    if not ultimo or ultimo.get("tipo") != "final":
        raise RuntimeError("el análisis no produjo resultado final")
    resultado = ultimo.get("resultado")
    if not isinstance(resultado, schemas.AnalisisOut):
        raise RuntimeError("resultado final inválido")
    return resultado
