"""Cifras del original que no llegaron al análisis, y cifras analizadas que no están en el original.

Cada forma de leer mal un archivo (una comilla sin cerrar, un cuadro impreso de costado, un signo
que no se reconoce, un encabezado que se toma como dato) termina en lo mismo: cifras que estaban en
el original y no entran al análisis, o cifras analizadas que el original no tiene. Contarlas no
depende de saber por qué falló la lectura, y alcanza para declararlo al comienzo del informe.

Se cuentan las cifras de las filas con dos o más números (las de un cuadro, no las de un título o
una nota) y se dejan afuera los años sueltos, que son rótulos de período.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from typing import Any, Iterable

MINIMO_DECLARAR = 3  # con menos cifras perdidas no se declara (una nota al pie con dos números)
PROPORCION_DECLARAR = 0.02
MAX_CIFRAS = 20000

_ESPACIOS = re.compile(r"[\s   ]")
_FORMA = re.compile(r"^-?\d[\d.,]*$")


def _limpio(texto: str) -> str | None:
    """«−1.234,5 %», «(1.234)», «$ 98,5*» → «-1.234,5»; None si no es una cifra."""
    from services.perfilado import _UNIDAD_PEGADA_RE, quitar_llamada_nota, sin_formula_texto

    t = _ESPACIOS.sub("", sin_formula_texto(str(texto))).replace("−", "-").replace("–", "-")
    t = _UNIDAD_PEGADA_RE.sub("", t)
    t = quitar_llamada_nota(t).strip("$%").strip()
    if t.startswith("(") and t.endswith(")"):
        t = "-" + t[1:-1]
    t = t.rstrip("%").strip()
    if t.startswith("+"):
        t = t[1:]
    from services.perfilado import expandir_cientifica

    signo = "-" if t.startswith("-") else ""
    expandida = expandir_cientifica(t.lstrip("-"))
    if expandida:
        t = signo + expandida
    if not _FORMA.match(t) or t.endswith((".", ",")):
        return None
    return t


def valores(texto: Any) -> set[float]:
    """Las lecturas posibles de una cifra escrita: con coma decimal o con punto decimal."""
    t = _limpio(texto)
    if t is None:
        return set()
    salida = set()
    for convertido in (t.replace(".", "").replace(",", "."), t.replace(",", "")):
        try:
            salida.add(round(float(convertido), 6))
        except ValueError:
            pass
    return salida


def _es_anio(crudo: Any) -> bool:
    """Un año suelto («2024», o el número 2024 en una celda) es un rótulo de período, no una cifra."""
    if isinstance(crudo, (int, float)) and not isinstance(crudo, bool):
        return float(crudo).is_integer() and 1800 <= crudo <= 2200
    return bool(re.fullmatch(r"\s*(?:18|19|20|21)\d{2}\s*", str(crudo)))


def _de_numero(v: Any) -> set[float]:
    if isinstance(v, bool) or not isinstance(v, (int, float)) or v != v:
        return set()
    # Un porcentaje con formato de celda se guarda como 0,125 y puede leerse como 12,5.
    return {round(float(v), 6), round(float(v) * 100, 6)}


def cifras_fila(celdas: Iterable[Any], con_anios: bool = False) -> list[tuple[int, set[float], str]]:
    """Cifras de una fila con dos o más números; sin años sueltos, salvo con_anios (una tabla ya
    leída escribe «1.804» como «1804», que no es un año)."""
    encontradas = []
    for j, c in enumerate(celdas):
        if not con_anios and _es_anio(c):
            continue
        conjunto = _de_numero(c) if not isinstance(c, str) else valores(c)
        if conjunto:
            encontradas.append((j, conjunto, c if isinstance(c, str) else f"{c:g}"))
    return encontradas if len(encontradas) >= 2 else []


def _original_csv(contenido: bytes) -> list[tuple[str, set[float], str]]:
    from services.carga_archivos import _decodificar, _separador_csv, desenvolver_lineas_csv

    texto, _, _ = _decodificar(contenido)
    texto = texto.replace("﻿", "")
    texto, _ = desenvolver_lineas_csv(texto)
    sep = _separador_csv(texto)
    salida = []
    # Línea por línea: una comilla sin cerrar no puede tragarse aquí el resto del archivo. Lo que
    # quedó dentro de esa comilla se vuelve a separar por el separador.
    for n, linea in enumerate(texto.splitlines(), 1):
        campos = []
        for c in next(csv.reader([linea], delimiter=sep), []):
            campos.extend(c.split(sep) if sep in c and c.count('"') == 0 and linea.count('"') % 2 else [c])
        for j, conjunto, crudo in cifras_fila(c.strip().strip('"') for c in campos):
            salida.append((f"fila {n}", conjunto, crudo))
    return salida


def _original_excel(contenido: bytes, hojas_leidas: set[str] | None) -> list[tuple[str, set[float], str]]:
    import datetime as dt
    import warnings

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    salida = []
    try:
        for ws in libro.worksheets:
            if hojas_leidas is not None and ws.title not in hojas_leidas:
                continue
            for r, fila in enumerate(ws.iter_rows(values_only=True), 1):
                celdas = [None if isinstance(c, (dt.date, dt.time, dt.datetime)) else c for c in fila]
                for j, conjunto, crudo in cifras_fila(celdas):
                    salida.append((f"{ws.title}!{get_column_letter(j + 1)}{r}", conjunto, crudo))
                if len(salida) > MAX_CIFRAS:
                    return salida
    finally:
        libro.close()
    return salida


def _incorporadas(filas: Iterable[dict]) -> Counter:
    from services.carga_archivos import META_PDF

    contador: Counter = Counter()
    for f in filas:
        for k, v in f.items():
            if k in META_PDF:
                continue
            for x in (_de_numero(v) if not isinstance(v, str) else valores(v)):
                contador[x] += 1
    return contador


def _sin_par(originales, contador: Counter) -> list[tuple[str, str]]:
    faltan = []
    for donde, conjunto, crudo in originales:
        par = next((x for x in sorted(conjunto) if contador.get(x, 0) > 0), None)
        if par is None:
            faltan.append((donde, crudo))
        else:
            contador[par] -= 1
            # La otra lectura de la misma cifra (0,125 y 12,5) no puede emparejar otra.
            for x in conjunto - {par}:
                if contador.get(x, 0) > 0:
                    contador[x] -= 1
    return faltan


def declarable(faltan: int, total: int) -> bool:
    return faltan >= MINIMO_DECLARAR and faltan >= PROPORCION_DECLARAR * total


def por_archivo(leidos: list[tuple[str, bytes]], datos: list[dict], hojas: list[dict] | None = None) -> list[dict]:
    """Por archivo de planilla o CSV: cuántas cifras del original no llegaron a los datos analizados."""
    from services.carga_archivos import _extension

    leidas_por_archivo: dict[str, set[str]] = {}
    for h in hojas or []:
        if isinstance(h, dict) and h.get("estado") == "incluida":
            leidas_por_archivo.setdefault(h.get("archivo"), set()).add(h.get("hoja"))
    salida = []
    for nombre, contenido in leidos:
        ext = _extension(nombre)
        try:
            if ext in {"csv", "txt"}:
                originales = _original_csv(contenido)
            elif ext in {"xlsx", "xlsm"}:
                originales = _original_excel(contenido, leidas_por_archivo.get(nombre, set()) if hojas else None)
            else:
                continue
        except Exception:  # noqa: BLE001 — el conteo no puede impedir el análisis
            continue
        if not originales:
            continue
        filas = [f for f in datos if isinstance(f, dict) and f.get("archivo_origen") in (nombre, None)]
        faltan = _sin_par(originales, _incorporadas(filas))
        registro = {"archivo": nombre, "en_archivo": len(originales), "sin_incorporar": len(faltan),
                    "ejemplos": [f"{d}: «{c}»" for d, c in faltan[:4]]}
        registro["declarar"] = declarable(len(faltan), len(originales))
        salida.append(registro)
    return salida


# --------------------------------------------------------------------------- PDF con texto

def _sin_rotulos(fichas: list[str]) -> list[str]:
    """«Comuna 7», «9 de Julio», «Calle 25»: un entero corto pegado a una palabra es parte del rótulo."""
    def palabra(k):
        return 0 <= k < len(fichas) and bool(re.search(r"[^\W\d_]{2,}", fichas[k]))

    return [t for k, t in enumerate(fichas)
            if not (re.fullmatch(r"\d{1,3}", t) and (palabra(k - 1) or (k == 0 and palabra(k + 1))))]


def todas_las_cifras(texto: str) -> list[str]:
    """Toda cifra que aparece en el texto, sin filtrar filas ni rótulos: para saber si una cifra
    leída en la tabla existe en el original."""
    return [_limpio(t) for t in re.split(r"\s+", texto or "") if _limpio(t)]


def cifras_texto(texto: str) -> list[str]:
    """Cifras del texto de una página, de las líneas con dos o más números (las de un cuadro)."""
    salida = []
    for linea in (texto or "").splitlines():
        fichas = _sin_rotulos([t for t in re.split(r"[\s ]+", linea) if t])
        palabras = sum(1 for t in fichas if re.search(r"[^\W\d_]{2,}", t))
        if palabras >= 6 and palabras > 2 * sum(1 for t in fichas if valores(t)):
            continue  # un renglón de un párrafo que menciona cifras, no una fila del cuadro
        # «Semáforo en rojo 29.224 …»: las palabras no cuentan; «1.234 (p)» sí.
        for _, conjunto, crudo in cifras_fila(fichas):
            salida.append(_limpio(crudo) or crudo)
    if not salida:
        # Texto con una palabra por línea (así lo devuelven algunos PDF): todas sus cifras.
        salida = [_limpio(crudo) or crudo
                  for _, conjunto, crudo in cifras_fila(_sin_rotulos([t for t in re.split(r"\s+", texto or "") if t]))]
    return salida


def cifras_tabla(encabezado: list, filas: list[list]) -> list[str]:
    salida = []
    for k, f in enumerate([encabezado, *filas]):
        for _, conjunto, crudo in cifras_fila(f, con_anios=k > 0):
            salida.append(_limpio(crudo) or crudo)
    return salida


def _contador_textos(cifras: Iterable[str]) -> Counter:
    contador: Counter = Counter()
    for c in cifras:
        for x in valores(c):
            contador[x] += 1
    return contador


def comparar(originales: list[str], leidas: list[str], texto_completo: list[str] | None = None) -> dict:
    """Cifras del original sin par en lo leído («faltan») y cifras leídas sin par en el original («ajenas»).

    Las ajenas se buscan en todas las cifras del texto (texto_completo), también las de los rótulos.
    """
    faltan = _sin_par([("", valores(c), c) for c in originales if valores(c)], _contador_textos(leidas))
    ajenas = _sin_par([("", valores(c), c) for c in leidas if valores(c)],
                      _contador_textos(texto_completo if texto_completo is not None else originales))
    return {
        "en_original": sum(1 for c in originales if valores(c)),
        "en_lectura": sum(1 for c in leidas if valores(c)),
        "faltan": len(faltan), "ejemplos_faltan": [c for _, c in faltan[:4]],
        "ajenas": len(ajenas), "ejemplos_ajenas": [c for _, c in ajenas[:4]],
    }



def celdas_ajenas(filas: list[list], texto_completo: list[str]) -> list[tuple[int, int]]:
    """(fila, columna) de las cifras de la tabla sin par en el texto del PDF: el mismo emparejamiento que
    `comparar` usa para contar las «ajenas», celda por celda."""
    contador = _contador_textos(texto_completo)
    salida = []
    for i, f in enumerate(filas):
        for j, _, crudo in cifras_fila(f, con_anios=True):
            conjunto = valores(_limpio(crudo) or crudo)
            if not conjunto:
                continue
            par = next((x for x in sorted(conjunto) if contador.get(x, 0) > 0), None)
            if par is None:
                salida.append((i, j))
                continue
            contador[par] -= 1
            for x in conjunto - {par}:
                if contador.get(x, 0) > 0:
                    contador[x] -= 1
    return salida


def aviso_pagina(comp: dict, tablas: int) -> str | None:
    """Lo que la pantalla de revisión dice de la página, en una frase."""
    n = comp["en_original"]
    if not tablas:
        return (f"No se reconoció ningún cuadro en esta página, aunque su texto tiene {n} cifras. Armá la tabla con "
                "«Armar una tabla con el texto leído» y revisala contra el original.") if n >= 6 else None
    partes = []
    if declarable(comp["faltan"], n):
        partes.append(f"la lectura propuesta deja afuera {comp['faltan']} de {n} cifras del texto de la página "
                      f"(por ejemplo «{comp['ejemplos_faltan'][0]}»)")
    if declarable(comp["ajenas"], comp["en_lectura"]):
        partes.append(f"{comp['ajenas']} cifras de la tabla propuesta no están en el texto del PDF "
                      f"(por ejemplo «{comp['ejemplos_ajenas'][0]}»): la lectura puede estar desarmada")
    if not partes:
        return None
    texto = "; ".join(partes)
    return texto[0].upper() + texto[1:] + ". Revisalas contra el original antes de confirmar."


def restar(originales: list[str], otras: list[str]) -> list[str]:
    """Las cifras del original que no son de las otras tablas de la página."""
    return [c for _, c in _sin_par([("", valores(c), c) for c in originales if valores(c)], _contador_textos(otras))]
