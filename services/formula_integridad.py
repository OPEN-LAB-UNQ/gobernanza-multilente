"""Control de integridad de formulas en archivos Excel -- fachada de compatibilidad.

El detalle vive dividido por responsabilidad en services/calculos/: excel_utils
(helpers de celdas), formulas (patrones de formula), integridad (cortes de bloque,
estados explicitos), relaciones (relaciones por nombre de columna) y ponderaciones
(cambios de criterio en promedios ponderados, lente Matematico). Este modulo solo
orquesta y conserva los nombres publicos que ya usaban carga_rapida.py y los tests.
"""
from __future__ import annotations

import io

from openpyxl import load_workbook

from services.calculos.excel_utils import (
    EXCEL_EXTS,
    MAX_HALLAZGOS,
    _extension,
)
from services.calculos.integridad import (
    _analizar_hoja,
)

# re-exportado por compatibilidad: algun test historico lo importa como
# formula_integridad._mapear_componentes_ponderados
from services.calculos.ponderaciones import (
    _analizar_matematico_hoja,
    _mapear_componentes_ponderados,  # noqa: F401
)


def analizar_matematico_archivo(filename: str, contenido: bytes) -> list[dict]:
    if _extension(filename) not in EXCEL_EXTS:
        return []
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=False, read_only=False)
    except Exception:
        return []
    hallazgos: list[dict] = []
    for ws in wb.worksheets:
        from services.interpretacion_previa import permite_hoja
        if not permite_hoja("ponderaciones", filename, ws.title):
            continue
        hallazgos.extend(_analizar_matematico_hoja(ws))
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    wb.close()
    return hallazgos[:MAX_HALLAZGOS]


def analizar_matematico_archivos(archivos: list[tuple[str, bytes]]) -> list[dict]:
    hallazgos: list[dict] = []
    for filename, contenido in archivos:
        hallazgos.extend(analizar_matematico_archivo(filename, contenido))
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    return hallazgos[:MAX_HALLAZGOS]


def analizar_archivo(
    filename: str,
    contenido: bytes,
    detectar_formulas: bool = True,
    inferir_patrones: bool = False,
) -> list[dict]:
    if _extension(filename) not in EXCEL_EXTS:
        return []
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=False, read_only=False)
        wb_valores = load_workbook(io.BytesIO(contenido), data_only=True, read_only=False)
    except Exception:
        return []
    hallazgos: list[dict] = []
    for ws in wb.worksheets:
        from services.interpretacion_previa import permite_hoja
        if not permite_hoja("formulas", filename, ws.title):
            continue
        ws_valores = wb_valores[ws.title] if ws.title in wb_valores.sheetnames else None
        hallazgos.extend(_analizar_hoja(ws, detectar_formulas=detectar_formulas,
            inferir_patrones=inferir_patrones and permite_hoja("patrones", filename, ws.title), ws_valores=ws_valores))
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    wb.close()
    wb_valores.close()
    return hallazgos[:MAX_HALLAZGOS]


def analizar_archivos(
    archivos: list[tuple[str, bytes]],
    detectar_formulas: bool = True,
    inferir_patrones: bool = False,
) -> list[dict]:
    hallazgos: list[dict] = []
    for filename, contenido in archivos:
        hallazgos.extend(analizar_archivo(
            filename,
            contenido,
            detectar_formulas=detectar_formulas,
            inferir_patrones=inferir_patrones,
        ))
        if len(hallazgos) >= MAX_HALLAZGOS:
            break
    return hallazgos[:MAX_HALLAZGOS]
