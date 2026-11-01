"""Lectura de fórmulas simples para controlar valores fijos y totales sin caché.

Sólo se interpretan referencias a celdas (también de otra hoja), rangos dentro de SUM y la
aritmética básica (+ - * /). Cualquier otra función deja la fórmula sin evaluar: nunca se
inventa un resultado. El valor guardado por Excel tiene prioridad; si falta (libro generado
sin recalcular), se calcula desde las celdas de origen.
"""

import re

from openpyxl.utils import column_index_from_string, get_column_letter

from .comun import numero

TOKEN = re.compile(
    r"\s*(?:(?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z0-9_\.ÁÉÍÓÚÑáéíóúñ]+)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"
    r"|(?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)|(?P<fn>[A-Z][A-Z0-9\.]*)\s*\(|(?P<op>[-+*/(),;%]))"
)
CELDA = re.compile(r"^(?:(?P<hoja>'(?:[^']|'')+'|[^!]+)!)?(?P<c1>\$?)(?P<col>[A-Z]{1,3})(?P<r1>\$?)(?P<fila>\d+)$")


class NoEvaluable(Exception):
    pass


def _tokens(texto):
    pos, out = 0, []
    texto = texto.strip()
    while pos < len(texto):
        m = TOKEN.match(texto, pos)
        if not m or m.end() == pos:
            raise NoEvaluable(texto)
        pos = m.end()
        tipo = m.lastgroup
        out.append((tipo, m.group(tipo)))
    return out


def _partir_ref(ref):
    """'Hoja'!$A$1:B2 -> (hoja|None, [(abs_col, col, abs_fila, fila), ...])."""
    hoja = None
    if "!" in ref:
        hoja, ref = ref.rsplit("!", 1)
        hoja = hoja[1:-1].replace("''", "'") if hoja.startswith("'") else hoja
    partes = []
    for p in ref.split(":"):
        m = re.match(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$", p)
        partes.append((bool(m.group(1)), column_index_from_string(m.group(2)), bool(m.group(3)), int(m.group(4))))
    return hoja, partes


def firma(texto, fila, col):
    """Forma relativa de la fórmula (como R1C1): igual en celdas que repiten el mismo cálculo."""
    try:
        toks = _tokens(texto.lstrip("="))
    except NoEvaluable:
        return ("texto", texto)
    out = []
    for tipo, v in toks:
        if tipo == "ref":
            hoja, partes = _partir_ref(v)
            out.append(("ref", hoja, tuple(
                (("C", c) if ac else ("c", c - col), ("R", r) if ar else ("r", r - fila)) for ac, c, ar, r in partes)))
        else:
            out.append((tipo, v.upper() if tipo == "fn" else v))
    return tuple(out)


def trasladar(texto, origen, destino):
    """Copia la fórmula como lo haría Excel al pegarla en otra celda (refs relativas se desplazan)."""
    dr, dc = destino[0] - origen[0], destino[1] - origen[1]
    toks = _tokens(texto.lstrip("="))
    out = []
    for tipo, v in toks:
        if tipo == "ref":
            hoja, partes = _partir_ref(v)
            celdas = []
            for ac, c, ar, r in partes:
                c2, r2 = (c if ac else c + dc), (r if ar else r + dr)
                if c2 < 1 or r2 < 1:
                    raise NoEvaluable(texto)
                celdas.append(f"{'$' if ac else ''}{get_column_letter(c2)}{'$' if ar else ''}{r2}")
            pref = "" if hoja is None else ("'" + hoja.replace("'", "''") + "'!")
            out.append(pref + ":".join(celdas))
        elif tipo == "fn":
            out.append(v + "(")
        else:
            out.append(v)
    return "=" + "".join(out)


class Libro:
    """Valores de un libro: los guardados por Excel y, si faltan, los calculados."""

    def __init__(self, valores, formulas):
        self.valores = valores
        self.formulas = formulas
        self._memo = {}
        self._en_curso = set()

    def formula(self, hoja, fila, col):
        try:
            c = self.formulas[hoja].cell(fila, col)
        except KeyError:
            return None
        return c.value if c.data_type == "f" and isinstance(c.value, str) else None

    def crudo(self, hoja, fila, col):
        try:
            return self.valores[hoja].cell(fila, col).value
        except KeyError:
            return None

    def valor(self, hoja, fila, col):
        clave = (hoja, fila, col)
        if clave in self._memo:
            return self._memo[clave]
        v = numero(self.crudo(hoja, fila, col))
        texto = self.formula(hoja, fila, col)
        if v is None and texto and clave not in self._en_curso and len(self._en_curso) < 60:
            self._en_curso.add(clave)
            try:
                v = self.evaluar(texto, hoja)
            except (NoEvaluable, ZeroDivisionError, RecursionError, IndexError, KeyError, ValueError):
                v = None
            finally:
                self._en_curso.discard(clave)
        self._memo[clave] = v
        return v

    def evaluar(self, texto, hoja):
        toks = _tokens(texto.lstrip("="))
        pos = [0]

        def ver():
            return toks[pos[0]] if pos[0] < len(toks) else (None, None)

        def tomar():
            t = ver()
            pos[0] += 1
            return t

        def rango(ref):
            h, partes = _partir_ref(ref)
            h = h or hoja
            if h not in self.formulas.sheetnames:
                raise NoEvaluable(ref)
            if len(partes) == 1:
                return [self.valor(h, partes[0][3], partes[0][1])]
            (_, c1, _, r1), (_, c2, _, r2) = partes
            if (abs(r2 - r1) + 1) * (abs(c2 - c1) + 1) > 5000:
                raise NoEvaluable(ref)
            return [self.valor(h, r, c) for r in range(min(r1, r2), max(r1, r2) + 1)
                    for c in range(min(c1, c2), max(c1, c2) + 1)]

        def factor():
            tipo, v = tomar()
            if tipo == "op" and v in "+-":
                x = factor()
                return -x if v == "-" else x
            if tipo == "num":
                return float(v)
            if tipo == "ref":
                if ":" in v:
                    raise NoEvaluable(v)
                x = rango(v)[0]
                if x is None:
                    # Celda vacía: Excel la toma como cero en la aritmética.
                    h, partes = _partir_ref(v)
                    crudo = self.crudo(h or hoja, partes[0][3], partes[0][1])
                    if crudo not in (None, "") or self.formula(h or hoja, partes[0][3], partes[0][1]):
                        raise NoEvaluable(v)
                    return 0.0
                return x
            if tipo == "op" and v == "(":
                x = expr()
                if tomar() != ("op", ")"):
                    raise NoEvaluable(texto)
                return x
            if tipo == "fn" and v.upper() == "SUM":
                total = 0.0
                while True:
                    t2, v2 = ver()
                    if t2 == "ref" and ":" in v2:
                        tomar()
                        total += sum(x for x in rango(v2) if x is not None)
                    else:
                        total += expr()
                    t3, v3 = tomar()
                    if (t3, v3) == ("op", ")"):
                        return total
                    if (t3, v3) not in (("op", ","), ("op", ";")):
                        raise NoEvaluable(texto)
            raise NoEvaluable(texto)

        def termino():
            x = factor()
            while ver() in (("op", "*"), ("op", "/"), ("op", "%")):
                _, op = tomar()
                if op == "%":
                    x = x / 100
                    continue
                y = factor()
                x = x * y if op == "*" else x / y
            return x

        def expr():
            x = termino()
            while ver() in (("op", "+"), ("op", "-")):
                _, op = tomar()
                y = termino()
                x = x + y if op == "+" else x - y
            return x

        resultado = expr()
        if pos[0] != len(toks):
            raise NoEvaluable(texto)
        return resultado

    def evaluar_seguro(self, texto, hoja):
        try:
            return self.evaluar(texto, hoja)
        except (NoEvaluable, ZeroDivisionError, RecursionError, IndexError, KeyError, ValueError, TypeError):
            return None
