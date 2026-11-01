"""Totales por fila y por columna que se cruzan: la intersección identifica la celda.

Un cuadro con columna «Total» (suma de las columnas) y fila «Total» (suma de las filas) tiene
dos controles independientes sobre cada componente. Si un componente se altera, no cierra el
total de su fila ni el de su columna, y los dos por la misma diferencia: la celda que está en
la intersección es la ubicación más probable, y los totales son celdas participantes. Cuando
la triangulación no es inequívoca (no hay fila de total, o varias filas fallan por lo mismo)
se señala la fila o la columna completa como rango, no sólo el total.

También señala un valor fuera de escala respecto de su fila y de su columna a la vez (un ×1000
en una celda de un cuadro de rubros por categoría), aunque la serie no tenga vecinos para un
control temporal: el nivel de cada fila y de cada columna se ajusta en logaritmos (pulido de
medianas) y se busca la celda cuyo residuo es una potencia de mil, aislada en su fila y su
columna. Nada depende del nombre del archivo ni del tema: sólo de la estructura del cuadro.

Los totales que el cuadro no rotula también se reconocen por sus cifras: una columna que es la
suma de otras dos a cuatro en casi todas las filas, una fila que es la suma de un tramo de filas
vecinas en casi todas las columnas y, en los cuadros de bloques paralelos («Principal», «Intereses»
debajo de cada sector), el bloque final que suma los bloques de arriba.
"""

import math
import re
from itertools import combinations
from statistics import median

import numpy as np
from openpyxl.utils import get_column_letter

from .comun import es_numero, normalizar
from .matriz import _pulido
from .modelo import es_total
from .transversal import NO_SUMABLE, _h, _tolerancia, _tolerancia_cierre, decimales_efectivos

MINIMO_CIERRE = 0.6  # proporción de filas (o columnas) que cierran para reconocer el total
ESCALA_MINIMA = 2.5  # log10: un factor de ~300 o más
CERCA_POTENCIA = 0.2
# Totales reconocidos por sus cifras (sin la palabra «Total»).
CIERRE_NUMERICO = 0.8  # proporción de filas o columnas que tienen que cerrar
MAXIMO_CANDIDATAS = 14  # columnas entre las que se buscan los sumandos de una columna total
FILAS_FUERTES = 4  # filas con tres o más cifras significativas en las que la columna total cierra
TRAMO_MAXIMO = 60  # filas de un tramo que una fila total puede sumar
FALLO_CERCANO = 0.15  # un desvío mayor ya no es un total que falla: es otra relación
APORTE_MINIMO = 0.01  # una fila que pesa menos del 1 % del total no prueba que el total sea una suma
TOLERANCIA_RELATIVA = 1e-4  # 0,01 %: redondeo de la fuente en cifras grandes
CIERRE_BLOQUES = 0.6  # proporción de columnas en que el bloque total suma los bloques de arriba
REPETICIONES_HIJO = 3  # veces que se repite un rótulo hijo («Principal») para leer bloques paralelos
MAXIMO_BLOQUES = 40  # bloques de arriba entre los que se busca el tramo que suma el total
TRAMO_MINIMO = 3  # filas seguidas en que un total falla por una diferencia pareja: un tramo de la fuente
TRAMO_PAREJO = 3.0  # cociente máximo entre la mayor y la menor diferencia porcentual del tramo
TRAMO_TOPE = 100 * 0.15  # % (FALLO_CERCANO): un desvío mayor no es un tramo de la fuente, se señala fila por fila
TOTAL_APARENTE = re.compile(r"\b(?:total(?:es)?|saldos?|consolidad[oa]s?|agregad[oa]s?|global(?:es)?)\b")


COLUMNA_TIEMPO = re.compile(r"^\s*(?:per[ií]odo|a[ñn]o|anio|mes|fecha|trimestre|semestre)s?\s*$", re.I)


def _numericas(c):
    # «Período» con 2023.11 (año.mes guardado como número) no es una cantidad: no entra en ninguna suma.
    return [x for x in c["columnas"] if x["numerica"] and not x.get("identificador") and not x.get("tiempo")
            and x.get("rol_confirmado") not in {"porcentaje", "indice"} and not COLUMNA_TIEMPO.match(str(x.get("nombre") or ""))]


def _es_col_total(x):
    if es_total(x.get("encabezado") if isinstance(x.get("encabezado"), str) else x.get("nombre")):
        return True
    # Encabezado de varios niveles con el «Total» en un nivel intermedio y la última fila vacía («Sector Bancario» ·
    # «Total Sector Bancario» · vacía): vale el rótulo más bajo de la columna.
    return isinstance(x.get("rotulo_inferior"), str) and es_total(x["rotulo_inferior"])


def _total_columna(c, filas, componentes, total):
    """{fila: diferencia} de la columna total contra la suma de las componentes, si es un total."""
    cierran, fallan = 0, {}
    for r in filas:
        declarado = c["valores"].get((r, total["idx"]))
        partes = [c["valores"].get((r, x["idx"])) for x in componentes]
        if not _comparable(declarado, partes):
            continue
        suma = sum(p for p in partes if p is not None)
        dif = declarado - suma
        if abs(dif) <= max(_tolerancia_cierre([declarado] + partes, len(partes)), _total_redondeado(declarado, partes)):
            cierran += 1
        else:
            fallan[r] = dif
    if cierran < 2 or cierran < MINIMO_CIERRE * (cierran + len(fallan)):
        return None
    return fallan


def _comparable(declarado, partes):
    """Hay un total y dos o más partes con número. Una fila con el total y las partes en cero (años todavía sin datos
    al pie de la serie) cierra con cualquier relación: no prueba que la columna sea la suma de esas partes."""
    return declarado is not None and sum(p is not None for p in partes) >= 2 and bool(declarado or any(partes))


def _total_redondeado(declarado, partes):
    """Un total publicado con menos decimales que sus partes (914.035.694,4 contra partes con dos decimales, o
    11.876.225.248 contra partes con cinco) puede diferir de la suma en una unidad de su último decimal."""
    decs = [decimales_efectivos(p) for p in partes if p]
    propio = decimales_efectivos(declarado)
    return 10.0 ** -propio + 1e-9 if decs and max(decs) > propio else 0.0


def _total_fila(c, grupo, total_fila, columnas):
    """{idx: diferencia} de la fila total contra la suma de su grupo, si es un total."""
    cierran, fallan = 0, {}
    for x in columnas:
        declarado = c["valores"].get((total_fila, x["idx"]))
        partes = [c["valores"].get((r, x["idx"])) for r in grupo]
        if not _comparable(declarado, partes):
            continue
        suma = sum(p for p in partes if p is not None)
        dif = declarado - suma
        if abs(dif) <= _tolerancia_cierre([declarado] + partes, len(partes)):
            cierran += 1
        else:
            fallan[x["idx"]] = dif
    if cierran < 2 or cierran < MINIMO_CIERRE * (cierran + len(fallan)):
        return None
    return fallan


def _grupos(c):
    """(fila total, filas de su grupo) para cada fila «Total» precedida por dos o más filas."""
    salida, inicio = [], 0
    filas = c["filas"]
    for k, r in enumerate(filas):
        if not es_total(c["etiquetas"].get(r)):
            continue
        grupo = filas[inicio:k]
        inicio = k + 1
        if len(grupo) >= 2 and not any(es_total(c["etiquetas"].get(g)) for g in grupo):
            salida.append((r, grupo))
    return salida


def _grupos_arriba(c):
    """(fila total, filas de su grupo) para cada fila «Total» seguida por dos o más filas.

    Muchos cuadros oficiales ponen el total arriba («Total» del país y después las jurisdicciones).
    El grupo llega hasta la próxima fila «Total»; sirve para triangular sólo si cierra con sus cifras."""
    filas = c["filas"]
    totales = [k for k, r in enumerate(filas) if es_total(c["etiquetas"].get(r))]
    return [(filas[k], filas[k + 1:fin]) for k, fin in zip(totales, totales[1:] + [len(filas)]) if fin - k > 2]


def _rango(c, fila, idxs):
    return f"{c['hoja']}!{get_column_letter(min(idxs))}{fila}:{get_column_letter(max(idxs))}{fila}"


def _rango_col(c, idx, filas):
    letra = get_column_letter(idx)
    return f"{c['hoja']}!{letra}{min(filas)}:{letra}{max(filas)}"


def _decimales(v):
    """Decimales con que está escrito el valor (como mucho 4: más allá es ruido de punto flotante)."""
    t = repr(float(v))
    if "e" in t or "." not in t:
        return 0
    return min(4, len(t.split(".")[1].rstrip("0")))


def _tramos_parejos(orden, fallas):
    """Separa las filas en que un total falla en tramos parejos y filas sueltas.

    Un tramo parejo son tres o más filas seguidas del cuadro (`orden`: una que cierra o en cero lo corta; una en
    blanco no es del cuadro y no lo corta, pero el rango de la señal la cubre igual) en
    que el total se aparta de la suma con el mismo signo y por una diferencia porcentual pareja: la mayor, como mucho
    tres veces la menor y no más del 15 %. Así queda una fuente que en un período sumó otra cosa (otra cobertura,
    otro método), y también un mismo error repetido en bloque: por eso la señal del tramo es de prioridad alta y su
    rango cubre cada fila. fallas: {fila: (diferencia con signo, diferencia porcentual)}."""
    posicion = {r: k for k, r in enumerate(orden)}
    seguidas, actual = [], []
    for r in sorted((r for r in fallas if r in posicion), key=posicion.get):
        if actual and (posicion[r] != posicion[actual[-1]] + 1 or (fallas[r][0] > 0) != (fallas[actual[0]][0] > 0)):
            seguidas.append(actual)
            actual = []
        actual.append(r)
    if actual:
        seguidas.append(actual)
    tramos, sueltas = [], [r for r in fallas if r not in posicion]
    for t in seguidas:
        pcts = [fallas[r][1] for r in t]
        if len(t) >= TRAMO_MINIMO and 0 < min(pcts) and max(pcts) <= min(TRAMO_TOPE, TRAMO_PAREJO * min(pcts)):
            tramos.append(t)
        else:
            sueltas.extend(t)
    return tramos, sorted(sueltas)


def _hallazgo_tramo(c, tcol, componentes, tramo, fallas, metodo):
    """Una señal por tramo parejo, de prioridad alta: el total no es la suma en esas filas, con la diferencia de cada
    una."""
    primera, ultima = tramo[0], tramo[-1]
    # La señal se ancla en la fila que más se aparta: si dentro del tramo hay un error, es la más probable.
    ancla = max(tramo, key=lambda r: fallas[r][1])
    pcts = [fallas[r][1] for r in tramo]
    rango = _rango_col(c, tcol["idx"], tramo)
    nombres = ", ".join(f"«{x.get('nombre')}»" for x in componentes)
    lado = "por encima" if fallas[primera][0] > 0 else "por debajo"
    declarado = c["valores"][(ancla, tcol["idx"])]

    def rotulo(r):
        return c["etiquetas"].get(r) or f"fila {r}"

    h = _h(c, ancla, tcol["idx"], "columna_total_no_cierra",
           f"«{tcol.get('nombre')}» es la suma de {nombres} en el resto de las filas, pero no en {len(tramo)} filas "
           f"seguidas, de «{rotulo(primera)}» a «{rotulo(ultima)}» ({rango}): ahí el total queda siempre {lado} de la "
           f"suma, por una diferencia pareja (entre {min(pcts):.3g} % y {max(pcts):.3g} %; la mayor, en «{rotulo(ancla)}»). "
           "Puede ser un tramo en que la fuente suma otra cosa (otra cobertura u otro método) o un mismo error repetido en "
           "esas filas.",
           120, {"total_declarado": declarado, "total_reconstruido": declarado - fallas[ancla][0],
                "diferencia_con_signo": fallas[ancla][0], "diferencia_porcentual": max(pcts),
                "columnas_componentes": [x.get("nombre") for x in componentes],
                "tramo": {"desde": rotulo(primera), "hasta": rotulo(ultima), "filas": len(tramo), "rango": rango,
                          "diferencias_con_signo": {c.ref(r, tcol["idx"]): fallas[r][0] for r in tramo},
                          "diferencia_porcentual_minima": min(pcts), "diferencia_porcentual_maxima": max(pcts)},
                "metodo": metodo + "; tramo parejo de filas seguidas"},
           participantes=[c.ref(r, tcol["idx"]) for r in tramo],
           accion="Buscar en las notas de la fuente qué abarca el total en ese tramo y contrastar esas filas con la "
                  "fuente.",
           lente="conciliador")
    h["evidencia"]["ubicaciones_excel"] = [c.ref(ancla, tcol["idx"]), rango]
    h["evidencia"]["grupo_relacion"] = f"tramo:{rango}"
    return h


def _tolerancia_suma(total, partes):
    """Redondeo admitido entre un total y la suma de sus partes.

    Medio dígito de la resolución MÁS GRUESA de los sumandos por la cantidad de sumandos (un
    componente publicado con un decimal arrastra hasta 0,05 aunque los demás tengan tres), o el
    0,01 % del total, lo que sea mayor. Los ceros no dicen nada de la resolución."""
    decs = [_decimales(p) for p in partes if p]
    dec = min(decs) if decs else 4
    suma = sum(partes)
    return max(0.5 * 10 ** -dec * len(partes), TOLERANCIA_RELATIVA * max(abs(total), abs(suma))) + 1e-9


def _fuerte(v, dec_columna):
    """El valor tiene al menos tres cifras significativas con la resolución de su columna."""
    return bool(v) and abs(v) >= 100 * 10 ** -dec_columna - 1e-12


def _sumables(c):
    return [x for x in _numericas(c) if not NO_SUMABLE.search(str(x.get("nombre") or ""))]


def _columnas_total(c):
    """[(columna total, sumandos, {fila: diferencia})] reconocidas por sus cifras.

    Una columna es total si es la suma de un subconjunto de dos a cuatro columnas numéricas en el
    80 % de las filas comparables y en cuatro o más filas con tres o más cifras significativas. Se
    busca entre catorce columnas como mucho, sin períodos (un período no es la suma de otros)."""
    cands = [x for x in _sumables(c) if not x.get("periodo") and not x.get("total_periodo")]
    if len(cands) < 3 or len(cands) > MAXIMO_CANDIDATAS:
        return []
    filas, m, decs = _matriz(c, cands)
    if not filas:
        return []
    dec_col = decs.max(axis=0)
    salida = []
    for t in range(len(cands)):
        total = m[:, t]
        mejor = None
        otras = [j for j in range(len(cands)) if j != t]
        for k in (2, 3, 4):
            for sub in combinations(otras, k):
                partes = m[:, sub]
                validas = ~np.isnan(total) & ~np.isnan(partes).any(axis=1)
                validas &= (total != 0) | (partes != 0).any(axis=1)
                n = int(validas.sum())
                if n < FILAS_FUERTES:
                    continue
                suma = np.nansum(partes, axis=1)
                d = np.where(decs[:, sub] < 0, 99, decs[:, sub]).min(axis=1)
                d = np.where(d == 99, 4, d)
                tol = np.maximum(0.5 * 10.0 ** -d * k, TOLERANCIA_RELATIVA * np.maximum(np.abs(total), np.abs(suma))) + 1e-9
                cierra = validas & (np.abs(total - suma) <= tol)
                nc = int(cierra.sum())
                if nc < CIERRE_NUMERICO * n:
                    continue
                fuertes = sum(1 for i in np.flatnonzero(cierra) if _fuerte(total[i], dec_col[t]))
                if fuertes < FILAS_FUERTES:
                    continue
                # Cada sumando aporta: una columna en cero en todas las filas que cierran no es parte del total.
                if any(not (partes[cierra, q] != 0).any() for q in range(k)):
                    continue
                clave = (nc, -k)
                if mejor is None or clave > mejor[0]:
                    fallan = {filas[i]: float(total[i] - suma[i]) for i in np.flatnonzero(validas & ~cierra)}
                    mejor = (clave, [cands[j] for j in sub], fallan)
        if mejor:
            salida.append((cands[t], mejor[1], mejor[2]))
    return salida


def totales_por_fila(c, filas_explicadas=frozenset()):
    """Columnas total sin rótulo «Total» (o varias): las filas en que la columna no es la suma de sus sumandos."""
    out = []
    metodo = "columna total reconocida por sus cifras (suma de otras columnas en el resto de las filas)"
    for tcol, partes, fallan in _columnas_total(c):
        pendientes = {}
        for r, dif in fallan.items():
            if r in filas_explicadas:
                continue
            # Una fila de tasas, índices o variaciones («Tasa de actividad», «Var. %») no se suma entre columnas.
            if FILA_TASA.search(normalizar(c["etiquetas"].get(r))):
                continue
            declarado = c["valores"][(r, tcol["idx"])]
            valores = [c["valores"][(r, x["idx"])] for x in partes]
            if abs(dif) <= _tolerancia_suma(declarado, valores):
                continue
            if _diferencia_explicada(c, r, dif, tcol, partes):
                continue
            pendientes[r] = (dif, abs(dif) / max(abs(declarado), abs(declarado - dif), 1e-9) * 100)
        tramos, sueltas = _tramos_parejos(c["filas"], pendientes)
        out += [_hallazgo_tramo(c, tcol, partes, t, pendientes, metodo) for t in tramos]
        for r in sueltas:
            dif = pendientes[r][0]
            declarado = c["valores"][(r, tcol["idx"])]
            rel = abs(dif) / max(abs(declarado), abs(declarado - dif), 1e-9)
            nombres = ", ".join(f"«{x.get('nombre')}»" for x in partes)
            # Un desvío grande ya no es un total que falla por poco: en esa fila puede regir otra relación. Se pregunta.
            lejos = rel > FALLO_CERCANO
            h = _h(c, r, tcol["idx"], "columna_total_no_cierra",
                   f"«{tcol.get('nombre')}» es la suma de {nombres} en el resto de las filas del cuadro; en esta fila "
                   f"declara {es_numero(declarado)} y esas columnas suman {es_numero(declarado - dif)}: diferencia "
                   f"{es_numero(dif)} ({rel * 100:.3g} %)."
                   + (" La diferencia es grande: en esta fila la columna puede no ser la suma de las demás (otra "
                      "relación o una fila de otra naturaleza)." if lejos else ""),
                   120 if TOLERANCIA_RELATIVA <= rel <= FALLO_CERCANO else 70,
                   {"total_declarado": declarado, "total_reconstruido": declarado - dif, "diferencia_con_signo": dif,
                    "diferencia_porcentual": rel * 100, "columnas_componentes": [x.get("nombre") for x in partes],
                    "metodo": metodo},
                   participantes=[c.ref(r, x["idx"]) for x in partes],
                   accion="Revisar el total de la fila y las columnas que suma; la diferencia no indica cuál es incorrecta.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, tcol['idx'])}"
            out.append(h)
    return out


FILA_NO_SUMABLE = re.compile(r"^(?:promedio|media|mediana|tasa|indice|porcentaje|variacion|var\.?)\b|%")
# Rótulo de fila (normalizado) de una tasa, un índice o una variación en cualquier lugar del rótulo.
FILA_TASA = re.compile(r"\btasas?\b|variaci|\bvar\b|indice|porcentaj|%|\bratio\b|promedio")


def _matriz(c, cols):
    """Valores (NaN si falta) y decimales (-1 si es cero o falta) de las filas del cuadro en esas columnas.

    Se calcula una vez por cuadro y juego de columnas (los totales sin rótulo y los bloques la comparten)."""
    clave = tuple(x["idx"] for x in cols)
    cache = c.get("_matrices")
    if cache is None:
        cache = c["_matrices"] = {}
    if clave not in cache:
        cache[clave] = _matriz_nueva(c, cols)
    return cache[clave]


def _matriz_nueva(c, cols):
    filas = list(c["filas"])
    m = np.array([[np.nan if c["valores"].get((r, x["idx"])) is None else c["valores"][(r, x["idx"])]
                   for x in cols] for r in filas], dtype=float).reshape(len(filas), len(cols))
    decs = np.array([[_decimales_publicados(c, r, x, m[i, j])
                      for j, x in enumerate(cols)] for i, r in enumerate(filas)],
                    dtype=float).reshape(len(filas), len(cols))
    return filas, m, decs


def _decimales_publicados(c, fila, columna, valor):
    """La resolución de una cifra es la visible en Excel, si está declarada."""
    if np.isnan(valor) or not valor:
        return -1
    formato = (c.get("formatos_celda") or {}).get((fila, columna["idx"]))
    if isinstance(formato, str) and formato != "General":
        base = formato.split(";")[0].replace('\\.', '')
        decimal = re.search(r"\.(0+|#+|\?+)(?![0#?])", base)
        if decimal:
            return len(decimal.group(1))
        if re.search(r"[0#?]", base) and "." not in base:
            return 0
    return _decimales(valor)


def _tramos(n, i, barrera):
    """Tramos contiguos de 2 a 60 filas justo debajo y justo arriba de la fila i, sin cruzar una barrera."""
    for paso in (1, -1):
        j, tramo = i + paso, []
        while 0 <= j < n and len(tramo) < TRAMO_MAXIMO and not barrera[j]:
            tramo.append(j)
            j += paso
        if len(tramo) >= 2:
            yield paso, tramo


def totales_sin_rotulo(c):
    """Filas que son el total de un tramo de filas vecinas sin decirlo («Año 2022» sobre sus meses,
    «Población con ingresos» debajo de las provincias), reconocidas por sus cifras.

    Para cada fila sin rótulo «Total» se prueban los tramos contiguos de 2 a 60 filas de abajo y de
    arriba que no cruzan un total rotulado. La fila es un total si en dos o más columnas «fuertes»
    (tres cifras significativas con la resolución de la columna) es la suma del tramo, si esas
    columnas son el 80 % de las que cierran o casi cierran (desvío de hasta 15 %) y si superan a las
    que no se parecen (porcentajes, promedios). Se señalan las columnas que casi cierran: un desvío
    de más del redondeo admitido y de hasta 15 %; si el total está bien comprobado (cuatro o más
    columnas que cierran y el 80 % de las comparables), también las que se apartan más, como señal
    media. La tolerancia es la de la resolución de las cifras, no la materialidad de la conciliación
    (0,2 %), que ocultaría un desvío chico."""
    cols = _sumables(c)
    if len(cols) < 2 or len(c["filas"]) < 3:
        return []
    filas, m, decs = _matriz(c, cols)
    n = len(filas)
    presente = ~np.isnan(m)
    m0 = np.where(presente, m, 0.0)
    dec_col = decs.max(axis=0)
    # Una fila sin cifras o un total rotulado cortan el tramo.
    barrera = [not presente[i].any() or es_total(c["etiquetas"].get(r)) for i, r in enumerate(filas)]
    out = []
    for i, r in enumerate(filas):
        etiqueta = normalizar(c["etiquetas"].get(r))
        if barrera[i] or FILA_NO_SUMABLE.search(etiqueta):
            continue
        total = m[i]
        fuertes = np.array([presente[i, j] and _fuerte(total[j], dec_col[j]) for j in range(len(cols))])
        if fuertes.sum() < 2:
            continue
        mejor = None
        for paso, tramo in _tramos(n, i, barrera):
            sumas = np.cumsum(m0[tramo], axis=0)
            completas = np.cumsum(presente[tramo], axis=0)
            d = np.where(decs[tramo] < 0, 99, decs[tramo])
            dmin = np.minimum.accumulate(d, axis=0)
            dmin = np.where(dmin == 99, 4, dmin)
            # Todos los largos del tramo a la vez: fila k-1 = suma de las primeras k filas.
            kk = np.arange(1, len(tramo) + 1)[:, None]
            validas = fuertes[None, :] & (completas == kk)
            dif = total[None, :] - sumas
            escala = np.maximum(np.abs(total)[None, :], np.abs(sumas))
            tol = np.maximum(0.5 * 10.0 ** -dmin * kk, TOLERANCIA_RELATIVA * escala) + 1e-9
            # Sólo es una suma si dos o más filas del tramo aportan (al menos el 1 % del total): igualar a una
            # fila cuando las demás están en cero o casi (saldo = crédito − débito con un débito ínfimo) no prueba nada.
            pesa = np.abs(m0[tramo]) >= APORTE_MINIMO * np.abs(total)[None, :]
            aportan = np.cumsum(pesa, axis=0) >= 2
            dentro = validas & (np.abs(dif) <= tol)
            cierra = dentro & aportan
            rel = np.abs(dif) / np.maximum(escala, 1e-9)
            cerca = validas & ~dentro & (rel <= FALLO_CERCANO)
            lejos = validas & ~dentro & (rel > FALLO_CERCANO)
            nc, ncerca, nlejos = cierra.sum(axis=1), cerca.sum(axis=1), lejos.sum(axis=1)
            ok = (kk[:, 0] >= 2) & (nc >= 2) & (nc >= CIERRE_NUMERICO * (nc + ncerca)) & (nc > nlejos)
            for q in np.flatnonzero(ok):
                clave = (int(nc[q]), -int(ncerca[q]), -int(q))
                if mejor is None or clave > mejor[0]:
                    # Con un total bien comprobado (cuatro o más columnas que cierran, el 80 % de las comparables),
                    # una columna que se aparta más del 15 % también se dice, como señal media.
                    con_lejos = int(nc[q]) >= 4 and nc[q] >= CIERRE_NUMERICO * validas[q].sum()
                    senaladas = cerca[q] | (lejos[q] if con_lejos else False)
                    mejor = (clave, [filas[t] for t in tramo[:q + 1]], senaladas, dif[q], sumas[q], rel[q], int(nc[q]))
        if not mejor:
            continue
        _, grupo, senaladas, dif, suma, rel, nc = mejor
        desde, hasta = min(grupo), max(grupo)
        for j in np.flatnonzero(senaladas):
            x = cols[j]
            declarado = float(total[j])
            lejos_j = rel[j] > FALLO_CERCANO
            h = _h(c, r, x["idx"], "total_sin_rotulo_no_cierra",
                   f"«{c['etiquetas'].get(r) or f'fila {r}'}» es la suma de las filas {desde} a {hasta} en {nc} columnas del "
                   f"cuadro; en esta columna declara {es_numero(declarado)} y esas filas suman {es_numero(float(suma[j]))}: "
                   f"diferencia {es_numero(float(dif[j]))} ({rel[j] * 100:.3g} %)."
                   + (" La diferencia es grande: puede ser un error del total o una columna que no se suma igual que las "
                      "demás." if lejos_j else ""),
                   120 if TOLERANCIA_RELATIVA <= rel[j] <= FALLO_CERCANO else 70,
                   {"total_declarado": declarado, "total_reconstruido": float(suma[j]),
                    "diferencia_con_signo": float(dif[j]), "diferencia_porcentual": float(rel[j] * 100),
                    "filas_componentes": [desde, hasta], "columnas_que_cierran": nc,
                    "metodo": "fila total reconocida por sus cifras (suma de un tramo de filas en el resto de las columnas)"},
                   participantes=[c.ref(f, x["idx"]) for f in grupo][:40],
                   accion="Revisar el total y las filas que suma; la diferencia no indica cuál de ellas es incorrecta.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, x['idx'])}"
            out.append(h)
    return out


DISTRIBUCION_SIN_CATEGORIA = 0.5  # puntos porcentuales que puede faltarle a una distribución que declara 100


def _diferencia_explicada(c, r, dif, total, partes):
    """La diferencia de la fila es exactamente otra columna del cuadro (el total suma también esa columna,
    que en otras filas vale cero o «-»), o una de las partes es un dato no publicado («---», «s», «///»)."""
    if abs((c["valores"].get((r, total["idx"])) or 0) - 100) < 1e-9 and abs(dif) <= DISTRIBUCION_SIN_CATEGORIA:
        return True  # distribución en % que suma 99,8: redondeo o una categoría sin publicar («Ns/Nc»)
    usadas = {total["idx"], *(x["idx"] for x in partes)}
    for x in _numericas(c):
        v = c["valores"].get((r, x["idx"]))
        if x["idx"] not in usadas and v and abs(dif - v) <= _tolerancia_suma(dif, [v]):
            return True
    return dif > 0 and any(c["valores"].get((r, x["idx"])) is None and _marca(c["crudos"].get((r, x["idx"])))
                           for x in partes)


def _marca(raw):
    return isinstance(raw, str) and bool(raw.strip())


def _cruce_panel(c, r, dif_fila, componentes):
    """Si la columna de un componente tampoco cierra contra su fila «Total», y por la misma diferencia que la fila,
    la celda donde se cruzan es la ubicación más probable (la fila y la columna quedan explicadas a la vez)."""
    totales = [t for t in c["filas"] if es_total(c["etiquetas"].get(t))]
    if not totales or r in totales:
        return None
    grupo = [f for f in c["filas"] if f not in totales]
    candidatos = []
    for t in totales:
        for x in componentes:
            declarado = c["valores"].get((t, x["idx"]))
            partes = [c["valores"].get((f, x["idx"])) for f in grupo]
            if declarado is None or sum(p is not None for p in partes) < 2:
                continue
            dif_col = declarado - sum(p for p in partes if p is not None)
            if abs(dif_col - dif_fila) <= max(_tolerancia([dif_fila, dif_col], 2), 0.001 * max(abs(dif_fila), abs(dif_col))):
                candidatos.append((t, x))
    if len(candidatos) != 1 or c["valores"].get((r, candidatos[0][1]["idx"])) is None:
        return None
    t, x = candidatos[0]
    v = c["valores"][(r, x["idx"])]
    ref_col = c.ref(t, x["idx"])
    h = _h(c, r, x["idx"], "componente_no_cierra_fila_y_columna",
           f"{es_numero(v)} hace que no cierren a la vez el total de su fila y el total de su columna ({ref_col}), los dos "
           f"por {es_numero(-dif_fila)}. Con ese valor corregido a {es_numero(v + dif_fila)} cerrarían ambos: la celda en "
           "la intersección es la ubicación más probable.",
           135, {"total_columna": ref_col, "diferencia_fila": dif_fila, "valor_que_cierra": v + dif_fila,
                 "metodo": "triangulación fila × columna de los totales declarados (panel)"},
           participantes=[ref_col], accion="Contrastar la celda con la fuente; los dos totales quedan explicados si se corrige.")
    h["evidencia"]["ubicaciones_excel"] = [c.ref(r, x["idx"]), ref_col]
    h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, x['idx'])}"
    return h


# Componente que se resta en un saldo («Ingresos» y «Egresos» de las operaciones con el FMI, con los egresos en
# positivo): el total es un neto, no una suma.
EGRESO = re.compile(r"\b(?:egresos?|salidas?|debitos?|pagos?)\b")


def _es_saldo_neto(c, componentes, tcol, fallan):
    """El «Total» cierra en todas las filas que fallan si se restan los componentes rotulados como egreso."""
    signos = [-1 if EGRESO.search(normalizar(str(x.get("nombre") or ""))) else 1 for x in componentes]
    if -1 not in signos or 1 not in signos:
        return False
    for r in fallan:
        declarado = c["valores"].get((r, tcol["idx"]))
        valores = [c["valores"].get((r, x["idx"])) for x in componentes]
        if declarado is None or any(v is None for v in valores):
            return False
        neto = sum(s * v for s, v in zip(signos, valores))
        if abs(declarado - neto) > _tolerancia_suma(declarado, valores):
            return False
    return True


def totales_por_panel(c, totales_col, numericas, filas_explicadas=frozenset()):
    """Cada columna «Total» es la suma de las columnas que la siguen hasta la próxima «Total».

    La relación la declara el rótulo y se confirma con las cifras: tiene que cerrar en dos o más filas y en la
    mayoría de las comparables (_total_columna). Las filas en que no cierra se señalan, también la fila «Total»."""
    out = []
    orden = sorted(numericas, key=lambda x: x["idx"])
    for tcol in totales_col:
        siguientes = [x for x in orden if x["idx"] > tcol["idx"]]
        componentes = []
        for x in siguientes:
            if _es_col_total(x) or NO_SUMABLE.search(str(x.get("nombre") or "")):
                break
            componentes.append(x)
        if len(componentes) < 2:
            continue
        fallan = _total_columna(c, c["filas"], componentes, tcol)
        if not fallan or _es_saldo_neto(c, componentes, tcol, fallan):
            continue
        pendientes = {}
        for r, dif in fallan.items():
            if r in filas_explicadas or FILA_TASA.search(normalizar(c["etiquetas"].get(r))):
                continue
            declarado = c["valores"][(r, tcol["idx"])]
            valores = [c["valores"].get((r, x["idx"])) for x in componentes]
            if abs(dif) <= _tolerancia_suma(declarado, [v for v in valores if v is not None]):
                continue
            if _diferencia_explicada(c, r, dif, tcol, componentes):
                continue
            cruce = _cruce_panel(c, r, dif, componentes)
            if cruce is not None:
                out.append(cruce)
                continue
            pendientes[r] = (dif, abs(dif) / max(abs(declarado), abs(declarado - dif), 1e-9) * 100)
        tramos, sueltas = _tramos_parejos(c["filas"], pendientes)
        metodo = "columna «Total» de su panel (las columnas que la siguen hasta la próxima «Total»)"
        out += [_hallazgo_tramo(c, tcol, componentes, t, pendientes, metodo) for t in tramos]
        for r in sueltas:
            dif = pendientes[r][0]
            declarado = c["valores"][(r, tcol["idx"])]
            rel = abs(dif) / max(abs(declarado), abs(declarado - dif), 1e-9)
            nombres = ", ".join(f"«{x.get('nombre')}»" for x in componentes)
            h = _h(c, r, tcol["idx"], "columna_total_no_cierra",
                   f"«{tcol.get('nombre')}» es la suma de {nombres} en el resto de las filas; en «{c['etiquetas'].get(r) or f'la fila {r}'}» "
                   f"declara {es_numero(declarado)} y esas columnas suman {es_numero(declarado - dif)}: diferencia "
                   f"{es_numero(dif)} ({rel * 100:.3g} %).",
                   120, {"total_declarado": declarado, "total_reconstruido": declarado - dif, "diferencia_con_signo": dif,
                         "diferencia_porcentual": rel * 100, "columnas_componentes": [x.get("nombre") for x in componentes],
                         "metodo": metodo},
                   participantes=[c.ref(r, x["idx"]) for x in componentes],
                   accion="Revisar el total de la fila y las columnas que suma; la diferencia no indica cuál es incorrecta.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, tcol['idx'])}"
            out.append(h)
    return out


def cruces(c, filas_explicadas=frozenset()):
    out = []
    numericas = _numericas(c)
    totales_col = [x for x in numericas if _es_col_total(x)]
    if len(totales_col) > 1:
        # Varias columnas «Total» (Estudiantes · Total, Varón, Mujer; Egresados · Total, Varón, Mujer): cada una
        # con su panel. Las que no se confirman así se buscan por sus cifras.
        paneles = totales_por_panel(c, totales_col, numericas, filas_explicadas)
        vistas = {h["evidencia"]["ubicaciones_excel"][0] for h in paneles}
        # Cada fila de un tramo de panel ya está señalada: la búsqueda por cifras no la repite.
        vistas |= {u for h in paneles if h["evidencia"].get("tramo") for u in h["evidencia"]["celdas_participantes"]}
        return paneles + [h for h in totales_por_fila(c, filas_explicadas) if h["evidencia"]["ubicaciones_excel"][0] not in vistas]
    if len(totales_col) != 1:
        # Sin una columna «Total»: las columnas total se reconocen por sus cifras.
        return totales_por_fila(c, filas_explicadas)
    tcol = totales_col[0]
    componentes = [x for x in numericas if x is not tcol and not NO_SUMABLE.search(str(x.get("nombre") or ""))]
    if len(componentes) < 2:
        return out
    grupos = _grupos(c)
    # Sin una fila «Total» debajo de sus filas, el total puede estar arriba de ellas.
    arriba = [] if grupos else _grupos_arriba(c)
    filas_comp = [r for r in c["filas"] if not es_total(c["etiquetas"].get(r))]
    fallan_filas = _total_columna(c, filas_comp, componentes, tcol)
    if fallan_filas is None:
        # Sólo las columnas a la izquierda del total, si hay medidas derivadas a la derecha.
        componentes = [x for x in componentes if x["idx"] < tcol["idx"]]
        if len(componentes) < 2:
            return out
        fallan_filas = _total_columna(c, filas_comp, componentes, tcol)
        if fallan_filas is None:
            return out
    if not fallan_filas:
        return out
    usados, sin_valor, en_su_total = set(), set(), set()
    for total_fila, grupo in grupos + arriba:
        fallan_cols = _total_fila(c, grupo, total_fila, componentes)
        if fallan_cols is None:
            continue
        filas_grupo = {r: d for r, d in fallan_filas.items() if r in grupo}
        for r, dr in filas_grupo.items():
            candidatos = [
                idx for idx, dc in fallan_cols.items()
                if abs(dr - dc) <= max(_tolerancia([dr, dc], 2), 0.001 * max(abs(dr), abs(dc)))
            ]
            if len(candidatos) != 1:
                continue
            idx = candidatos[0]
            mismos = [rr for rr, d in filas_grupo.items() if abs(d - fallan_cols[idx]) <= max(_tolerancia([d], 2), 0.001 * abs(d))]
            if len(mismos) != 1:
                continue  # dos filas fallan por lo mismo en esa columna: no hay una sola intersección
            v = c["valores"].get((r, idx))
            if v is None:
                # Vacía o con una marca («s», «///»): la diferencia puede ser un dato no publicado.
                sin_valor.add(r)
                continue
            esperado = v + dr
            ref_fila, ref_col = c.ref(r, tcol["idx"]), c.ref(total_fila, idx)
            h = _h(c, r, idx, "componente_no_cierra_fila_y_columna",
                   f"{es_numero(v)} hace que no cierren a la vez el total de su fila ({ref_fila}) y el total de su "
                   f"columna ({ref_col}), los dos por {es_numero(-dr)}. Con ese valor corregido a {es_numero(esperado)} "
                   "cerrarían ambos: la celda en la intersección es la ubicación más probable.",
                   135,
                   {"total_fila": ref_fila, "total_columna": ref_col, "diferencia_fila": dr,
                    "diferencia_columna": fallan_cols[idx], "valor_que_cierra": esperado,
                    "metodo": "triangulación fila × columna de los totales declarados"},
                   participantes=[ref_fila, ref_col],
                   accion="Contrastar la celda con la fuente; los dos totales quedan explicados si se corrige.")
            h["evidencia"]["ubicaciones_excel"] = [c.ref(r, idx), ref_fila, ref_col]
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, idx)}"
            out.append(h)
            usados.add(r)
        if fallan_cols is not None and not fallan_cols:
            # La fila total cierra en todas las columnas: si la columna total no cierra en una fila,
            # el error está en el total de esa fila. Con el total abajo ya lo señalan los controles de
            # totales; con el total arriba se señala aquí, nombrando ese total.
            (en_su_total if arriba else usados).update(filas_grupo)
    sin_fila_total = not grupos
    for r, dr in fallan_filas.items():
        if r in usados or r in filas_explicadas or not (sin_fila_total or r in sin_valor):
            continue
        # No hay fila de total debajo, o la de arriba no ubica una sola celda: la diferencia está en
        # algún componente de la fila o en su total. Se señala la fila completa como rango.
        idxs = [x["idx"] for x in componentes] + [tcol["idx"]]
        ref_total = c.ref(r, tcol["idx"])
        puntos, accion, ubicaciones = 100, "Revisar los componentes de la fila y su total contra la fuente.", [_rango(c, r, idxs)]
        if r in sin_valor:
            motivo = ("La celda donde se cruzan los dos totales que no cierran no tiene un número publicado (vacía o "
                      "con una marca): la diferencia puede ser ese dato no publicado. Se señala la fila completa.")
            puntos, accion = 70, "Confirmar con la fuente si la diferencia corresponde al dato no publicado."
        elif r in en_su_total:
            motivo = (f"La fila de total cierra en todas las columnas: lo más probable es que el que no cierre sea el "
                      f"total de esta fila ({ref_total}).")
            accion, ubicaciones = "Revisar el total de la fila contra la fuente.", [ref_total] + ubicaciones
        elif arriba:
            motivo = "La fila de total no alcanza para ubicar una sola celda de la fila: se señala la fila completa."
        else:
            motivo = "Sin fila de total no se puede triangular qué celda de la fila es: se señala la fila completa."
        h = _h(c, r, tcol["idx"], "fila_no_cierra_total",
               f"Los componentes de la fila suman {es_numero(c['valores'].get((r, tcol['idx'])) - dr)} y el total declara "
               f"{es_numero(c['valores'].get((r, tcol['idx'])))} (diferencia {es_numero(dr)}). {motivo}",
               puntos, {"diferencia_con_signo": dr, "rango": _rango(c, r, idxs)},
               participantes=[c.ref(r, i) for i in idxs], accion=accion)
        h["evidencia"]["ubicaciones_excel"] = ubicaciones
        out.append(h)
    return out


def _diferencias(c):
    """{(fila, idx): diferencias de los totales que la contienen} para corroborar un desvío."""
    numericas = _numericas(c)
    totales_col = [x for x in numericas if _es_col_total(x)]
    filas_comp = [r for r in c["filas"] if not es_total(c["etiquetas"].get(r))]
    por_fila, por_col = {}, {}
    componentes = [x for x in numericas if not _es_col_total(x) and not NO_SUMABLE.search(str(x.get("nombre") or ""))]
    if len(totales_col) == 1 and len(componentes) >= 2:
        por_fila = _total_columna(c, filas_comp, componentes, totales_col[0]) or {}
    for total_fila, grupo in _grupos(c):
        for idx, d in (_total_fila(c, grupo, total_fila, componentes) or {}).items():
            for r in grupo:
                por_col[(r, idx)] = d
    return por_fila, por_col


def escala_fila_columna(c):
    """Celda fuera de escala respecto de su fila y de su columna (potencia de mil, aislada).

    En un cuadro de períodos la serie ya tiene su propio control temporal: allí sólo se señala si
    un total declarado no cierra justamente por el desvío de la celda. En un cuadro transversal
    (sin períodos) basta que el desvío sea una potencia de mil aislada en su fila y su columna.
    """
    out = []
    numericas = [x for x in _numericas(c) if not _es_col_total(x) and not NO_SUMABLE.search(str(x.get("nombre") or ""))]
    filas = [r for r in c["filas"] if not es_total(c["etiquetas"].get(r))]
    if len(numericas) < 3 or len(filas) < 3:
        return out
    m = {}
    for i, r in enumerate(filas):
        for j, x in enumerate(numericas):
            v = c["valores"].get((r, x["idx"]))
            if v is not None and v > 0:
                m[(i, j)] = math.log10(v)
    if len(m) < 0.8 * len(filas) * len(numericas):
        return out
    res = _pulido(m, range(len(filas)), range(len(numericas)))
    sigma = max(0.02, 1.4826 * median(abs(v) for v in res.values()))
    con_periodos = any(x.get("periodo") for x in numericas)
    por_fila, por_col = _diferencias(c)
    for (i, j), e in res.items():
        if abs(e) < max(ESCALA_MINIMA, 8 * sigma):
            continue
        potencia = round(abs(e) / 3) * 3
        if potencia < 3 or abs(abs(e) - potencia) > CERCA_POTENCIA:
            continue
        resto_fila = [abs(res[(i, k)]) for k in range(len(numericas)) if k != j and (i, k) in res]
        resto_col = [abs(res[(k, j)]) for k in range(len(filas)) if k != i and (k, j) in res]
        if not resto_fila or not resto_col or max(resto_fila + resto_col) > 0.5:
            continue  # la fila o la columna entera se mueve: no es una celda aislada
        r, x = filas[i], numericas[j]
        v = c["valores"][(r, x["idx"])]
        factor = 10 ** abs(e)
        esperado = 10 ** (m[(i, j)] - e)
        exceso = v - esperado
        corroborada = any(
            d is not None and abs(d + exceso) <= 0.05 * abs(exceso)
            for d in (por_fila.get(r), por_col.get((r, x["idx"])))
        )
        if con_periodos and not corroborada:
            continue
        if not corroborada and abs(abs(e) - potencia) > CERCA_POTENCIA / 2:
            continue
        h = _h(c, r, x["idx"], "valor_fuera_de_escala_fila_columna",
               f"{es_numero(v)} está {factor:,.0f} veces {'por encima' if e > 0 else 'por debajo'} de lo esperable por el nivel "
               f"de su fila y de su columna ({es_numero(esperado)}); el resto de la fila y de la columna sigue la estructura del "
               f"cuadro. El desvío es cercano a {10 ** potencia:,.0f} veces: posible error de escala o de unidad.",
               120,
               {"factor_escala": factor, "valor_esperado_por_estructura": esperado, "residuo_log10": e,
                "dispersion_residuos_log10": sigma, "corroborada_por_totales": corroborada,
                "metodo": "pulido de medianas en logaritmos (nivel de la fila + nivel de la columna)"},
               accion="Verificar la escala y la unidad de la celda contra la fuente.")
        h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, x['idx'])}"
        if corroborada and por_fila.get(r) is not None:
            # El total de la fila no cierra por este desvío: es la misma causa, en la misma señal.
            total = next(t for t in _numericas(c) if _es_col_total(t))
            h["evidencia"]["ubicaciones_excel"].append(c.ref(r, total["idx"]))
            h["evidencia"]["celdas_participantes"].append(c.ref(r, total["idx"]))
            h["evidencia"]["fila_explicada"] = r
        out.append(h)
    return out



STOCK = re.compile(r"stock|existencia|saldo|al cierre|al final|acumulad", re.I)


def _filas_sumables(c):
    """Filas en las que un subtotal de período suma sus períodos: no un promedio, un precio o un stock."""
    return [r for r in c["filas"]
            if not NO_SUMABLE.search(str(c["etiquetas"].get(r) or "")) and not STOCK.search(str(c["etiquetas"].get(r) or ""))]


def subtotales_columna(c):
    """Columnas sin período intercaladas entre los períodos («1.er trim.», «Total 2025») que suman sus períodos.

    Sólo si hay al menos un subtotal intercalado reconocido: un cuadro con una única columna «Total»
    al final ya lo controla cruces(). Devuelve (hallazgos, cierres verificados)."""
    from .periodos import es_auxiliar
    if c.get("eje") != "columnas" or c.get("paneles"):
        return [], []
    nums = sorted(_numericas(c), key=lambda x: x["idx"])
    pers = [x for x in nums if x.get("periodo")]
    if len(pers) < 3:
        return [], []
    otras = [x for x in nums if not x.get("periodo") and x["idx"] > pers[0]["idx"]
             and not es_auxiliar(str(x.get("nombre") or "")) and not NO_SUMABLE.search(str(x.get("nombre") or ""))]
    reconocidas, corte = [], -1
    for x in otras:
        partes = [p for p in pers if p["idx"] < x["idx"] and (_es_col_total(x) or p["idx"] > corte)]
        if len(partes) < 2:
            continue
        fallan = _total_columna(c, _filas_sumables(c), partes, x)
        if fallan is None:
            continue
        reconocidas.append((x, partes, fallan))
        if not _es_col_total(x):
            corte = x["idx"]
    if not any(not _es_col_total(x) for x, _, _ in reconocidas):
        return [], []
    out, cierres = [], []
    for x, partes, fallan in reconocidas:
        for r in _filas_sumables(c):
            d = c["valores"].get((r, x["idx"]))
            if d is None or r in fallan:
                continue
            cierres.append(dict(archivo=c["archivo"], hoja=c["hoja"], cuadro=c["cuadro"], periodo=x.get("nombre"),
                                origen="subtotal_columna", total=c.ref(r, x["idx"]),
                                componentes=[c.ref(r, p["idx"]) for p in partes], diferencia=0))
        for r, dif in fallan.items():
            d = c["valores"][(r, x["idx"])]
            h = _h(c, r, x["idx"], "subtotal_columna_no_cierra",
                   f"«{x.get('nombre')}» declara {es_numero(d)} y sus {len(partes)} columnas de período suman "
                   f"{es_numero(d - dif)}: diferencia {es_numero(dif)}.", 110,
                   {"total_declarado": d, "total_reconstruido": d - dif, "diferencia_con_signo": dif,
                    "columnas_componentes": [p.get("nombre") for p in partes]},
                   participantes=[c.ref(r, p["idx"]) for p in partes],
                   accion="Revisar el subtotal y los períodos que suma; la diferencia no indica cuál es incorrecto.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, x['idx'])}"
            out.append(h)
    return out, cierres


def _cierre_filas(c, cols, total, partes, minimo):
    """(columnas que cierran, {idx: (declarado, suma, diferencia)} de las que no) si la fila total
    es la suma de las filas `partes` en al menos `minimo` de las columnas comparables; si no, None."""
    cierran, fallan = 0, {}
    for x in cols:
        declarado = c["valores"].get((total, x["idx"]))
        vals = [c["valores"].get((r, x["idx"])) for r in partes]
        if declarado is None or any(v is None for v in vals):
            continue
        suma = sum(vals)
        if abs(declarado - suma) <= _tolerancia_suma(declarado, vals):
            cierran += 1
        else:
            fallan[x["idx"]] = (declarado, suma, declarado - suma)
    if cierran < 2 or cierran < minimo * (cierran + len(fallan)):
        return None
    return cierran, fallan


def bloques_paralelos(c):
    """Cuadros de bloques paralelos: cada sector con sus aperturas repetidas («Principal», «Intereses»)
    y un bloque final que es el total de los de arriba.

    Los rótulos hijos que se repiten tres o más veces definen los bloques; cada padre tiene que ser la
    suma de sus hijos en el 80 % de las columnas (las columnas en que no lo es se señalan). Un padre,
    con o sin hijos propios, que es la suma de los bloques siguientes es un nivel intermedio («Otros
    sectores» sobre dos subsectores, una región sobre sus provincias): entra él y no sus bloques, para
    no contarlos dos veces. El último bloque es el total: si cierra contra los bloques de nivel
    superior en el 60 % de las columnas se señalan los desvíos; si sólo cierra con un único tramo
    contiguo de ellos, el total omite los demás; si cierra igual con varios tramos, si hay más de 40
    bloques o si no cierra de ninguna forma, el total aparente queda declarado como no comprobado.
    Devuelve (hallazgos, declaraciones)."""
    cols = _sumables(c)
    filas = list(c["filas"])
    rotulos = {r: normalizar(c["etiquetas"].get(r)) for r in filas}
    conteo = {}
    for r in filas:
        if rotulos[r]:
            conteo[rotulos[r]] = conteo.get(rotulos[r], 0) + 1
    hijos_rot = {t for t, n in conteo.items() if n >= REPETICIONES_HIJO}
    if not hijos_rot or len(cols) < 2:
        return [], []
    # Ítems en orden: (padre, hijos); un padre sin hijos queda con la lista vacía.
    items, k = [], 0
    while k < len(filas):
        r = filas[k]
        if not rotulos[r] or rotulos[r] in hijos_rot:
            k += 1
            continue
        hijos = []
        j = k + 1
        while j < len(filas) and rotulos[filas[j]] in hijos_rot:
            hijos.append(filas[j])
            j += 1
        items.append((r, hijos))
        k = j
    cierres = {}  # padre -> (filas que suma, (columnas que cierran, fallos))
    for r, hs in items:
        if len(hs) >= 2:
            res = _cierre_filas(c, cols, r, hs, CIERRE_NUMERICO)
            if res:
                cierres[r] = (hs, res)
    validos = set(cierres)
    if len(validos) < REPETICIONES_HIJO:
        return [], []
    # Niveles intermedios: un padre (sin hijos, o con sus propias aperturas) que suma los bloques siguientes
    # («Otros sectores» sobre dos subsectores, una región con sus provincias). Entra el padre y no los bloques
    # que suma: si entraran los dos, el total contaría dos veces la región.
    superiores, k = [], 0
    while k < len(items):
        r, hs = items[k]
        if hs and r not in validos:
            k += 1
            continue
        siguientes, salto = [], None
        for q in range(k + 1, len(items) - 1):  # el último bloque (el total) no es parte de un nivel intermedio
            if items[q][0] not in validos:
                break
            siguientes.append(items[q][0])
            if len(siguientes) >= 2:
                res = _cierre_filas(c, cols, r, siguientes, CIERRE_NUMERICO)
                if res:
                    salto = q + 1
                    if not hs:
                        cierres[r] = (list(siguientes), res)
                    break
        if salto:
            superiores.append(r)
            k = salto
        else:
            if hs:
                superiores.append(r)
            k += 1
    out = _padres_que_no_cierran(c, cierres)
    if len(superiores) < 3:
        return out, []
    total, otros = superiores[-1], superiores[:-1]
    rotulo_total = c["etiquetas"].get(total) or f"fila {total}"
    if not (es_total(rotulo_total) or TOTAL_APARENTE.search(rotulos[total])):
        return out, []
    nombres = {r: c["etiquetas"].get(r) or f"fila {r}" for r in otros}
    completo = _cierre_filas(c, cols, total, otros, CIERRE_BLOQUES)
    if completo:
        nc, fallan = completo
        for idx, (declarado, suma, dif) in fallan.items():
            rel = abs(dif) / max(abs(declarado), abs(suma), 1e-9)
            h = _h(c, total, idx, "total_bloques_no_cierra",
                   f"«{rotulo_total}» es la suma de los {len(otros)} bloques de arriba en {nc} columnas; en esta declara "
                   f"{es_numero(declarado)} y los bloques suman {es_numero(suma)}: diferencia {es_numero(dif)} "
                   f"({rel * 100:.3g} %).",
                   120 if TOLERANCIA_RELATIVA <= rel <= FALLO_CERCANO else 70,
                   {"total_declarado": declarado, "total_reconstruido": suma, "diferencia_con_signo": dif,
                    "diferencia_porcentual": rel * 100, "bloques": [nombres[r] for r in otros]},
                   participantes=[c.ref(r, idx) for r in otros],
                   accion="Revisar el total y los bloques que suma; la diferencia no indica cuál es incorrecto.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(total, idx)}"
            out.append(h)
        return _sin_repetir(out), []
    lista = ", ".join(f"«{nombres[r]}»" for r in otros[:8]) + (" …" if len(otros) > 8 else "")
    declaracion = dict(archivo=c["archivo"], hoja=c["hoja"], cuadro=c["cuadro"], concepto=rotulo_total,
                       donde=f"total aparente «{rotulo_total}» (fila {total} de «{c['hoja']}»)",
                       motivo=f"no es la suma de los {len(otros)} bloques de arriba ({lista}) ni de un tramo de ellos: "
                              "el total no se pudo comprobar")
    if len(otros) > MAXIMO_BLOQUES:
        declaracion["motivo"] = (f"tiene {len(otros)} bloques arriba ({lista}), más de los {MAXIMO_BLOQUES} entre los que se "
                                 "busca qué tramo suma: el total no se comprobó")
        return out, [declaracion]
    # ¿Cierra con un tramo contiguo de los bloques? Entonces el total omite los demás.
    candidatos = []
    for a, b, nc, fallan in _cierres_de_tramos(c, cols, total, otros):
        tramo = otros[a:b]
        omitidos = [r for r in otros if r not in tramo]
        peso = sum(abs(c["valores"].get((r, x["idx"])) or 0) for r in omitidos for x in cols)
        if not peso:
            continue  # lo omitido es cero: el total no cambia
        candidatos.append(((nc, len(tramo)), tramo, omitidos, (nc, fallan)))
    if candidatos:
        mejor = max(candidatos, key=lambda x: x[0])
        empatados = [x for x in candidatos if x[0][0] == mejor[0][0] and x[1] != mejor[1]]
        if empatados:
            # Más de un tramo explica el total igual de bien: no se puede afirmar cuáles bloques omite.
            tramos = "; ".join(", ".join(f"«{nombres[r]}»" for r in x[1]) for x in [mejor] + empatados[:2])
            declaracion["motivo"] = (f"cierra igual de bien con más de un tramo de los bloques de arriba ({tramos}): "
                                     "no se puede afirmar qué bloques omite")
            return out, [declaracion]
        _, tramo, omitidos, (nc, fallan) = mejor
        idxs = [x["idx"] for x in cols if c["valores"].get((total, x["idx"])) is not None]
        omitido = {x["idx"]: sum(c["valores"].get((r, x["idx"])) or 0 for r in omitidos) for x in cols}
        # El ancla es la columna donde lo omitido pesa más: ahí se ve la diferencia.
        ancla = max(idxs, key=lambda i: (abs(omitido.get(i) or 0), -i))
        declarado = c["valores"].get((total, ancla))
        h = _h(c, total, ancla, "total_omite_bloques",
               f"«{rotulo_total}» no es la suma de los {len(otros)} bloques de arriba: en {nc} columnas es la suma sólo de "
               + ", ".join(f"«{nombres[r]}»" for r in tramo) + ". El total omite "
               + ", ".join(f"«{nombres[r]}»" for r in omitidos)
               + f" (en esta columna suman {es_numero(omitido[ancla])}; con ellos el total sería "
               f"{es_numero((declarado or 0) + omitido[ancla])}).",
               125,
               {"total_declarado": declarado, "total_reconstruido": (declarado or 0) + omitido[ancla],
                "diferencia_con_signo": -omitido[ancla], "bloques_sumados": [nombres[r] for r in tramo],
                "bloques_omitidos": [nombres[r] for r in omitidos], "columnas_que_cierran_con_el_tramo": nc,
                "rango": _rango(c, total, idxs)},
               participantes=[c.ref(r, ancla) for r in omitidos] + [c.ref(r, ancla) for r in tramo],
               accion="Verificar si el total debe incluir los bloques omitidos; si es así, corregir el total de la fila.",
               lente="conciliador")
        h["evidencia"]["ubicaciones_excel"] = [_rango(c, total, idxs)]
        h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(total, ancla)}"
        return _sin_repetir(out + [h]), []
    return out, [declaracion]


def _sin_repetir(hs):
    """Una señal por celda: la del total contra sus bloques prevalece sobre la del total contra sus aperturas."""
    vistas, salida = set(), []
    for h in reversed(hs):
        clave = (h["evidencia"]["hoja"], h["evidencia"]["celda"])
        if clave not in vistas:
            vistas.add(clave)
            salida.append(h)
    return salida[::-1]


def _padres_que_no_cierran(c, cierres):
    """Fallos de cada padre que es la suma de sus aperturas («Principal», «Intereses») o de los bloques que
    agrupa en el resto de las columnas: la columna en que no cierra es una señal propia, media (la estructura
    se infiere de los rótulos repetidos; en los cuadros reales los desvíos chicos suelen ser redondeos de la fuente)."""
    out = []
    for r, (partes, (nc, fallan)) in cierres.items():
        rotulo = c["etiquetas"].get(r) or f"fila {r}"
        nombres_partes = ", ".join(f"«{c['etiquetas'].get(p) or f'fila {p}'}»" for p in partes[:6])
        for idx, (declarado, suma, dif) in fallan.items():
            rel = abs(dif) / max(abs(declarado), abs(suma), 1e-9)
            h = _h(c, r, idx, "total_bloques_no_cierra",
                   f"El bloque «{rotulo}» es la suma de {nombres_partes} en {nc} columnas; en esta declara "
                   f"{es_numero(declarado)} y esas filas suman {es_numero(suma)}: diferencia {es_numero(dif)} "
                   f"({rel * 100:.3g} %).",
                   70,  # señal media: la estructura del bloque se infiere de rótulos repetidos
                   {"total_declarado": declarado, "total_reconstruido": suma, "diferencia_con_signo": dif,
                    "diferencia_porcentual": rel * 100, "bloque": rotulo,
                    "bloques": [c["etiquetas"].get(p) or f"fila {p}" for p in partes]},
                   participantes=[c.ref(p, idx) for p in partes],
                   accion="Revisar el bloque y las filas que suma; la diferencia no indica cuál es incorrecta.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, idx)}"
            out.append(h)
    return out


def _cierres_de_tramos(c, cols, total, otros):
    """(a, b, columnas que cierran, fallos) de cada tramo contiguo otros[a:b] (sin el completo) cuya suma es la
    fila total en el 80 % de las columnas comparables. Las sumas y la resolución se acumulan al alargar el
    tramo: O(bloques² · columnas) en vez de recalcular cada tramo. Misma tolerancia que _tolerancia_suma."""
    n = len(otros)
    declarados = [c["valores"].get((total, x["idx"])) for x in cols]
    valores = [[c["valores"].get((r, x["idx"])) for x in cols] for r in otros]
    decs = [[_decimales(v) if v else None for v in fila] for fila in valores]
    for a in range(n):
        suma = [0.0] * len(cols)
        falta = [False] * len(cols)
        dmin = [None] * len(cols)
        for b in range(a + 1, n + 1):
            for j, v in enumerate(valores[b - 1]):
                if v is None:
                    falta[j] = True
                    continue
                suma[j] += v
                d = decs[b - 1][j]
                if d is not None:
                    dmin[j] = d if dmin[j] is None else min(dmin[j], d)
            k = b - a
            if k == n:
                continue
            cierran, fallan = 0, {}
            for j, x in enumerate(cols):
                declarado = declarados[j]
                if declarado is None or falta[j]:
                    continue
                dec = 4 if dmin[j] is None else dmin[j]
                tol = max(0.5 * 10 ** -dec * k, TOLERANCIA_RELATIVA * max(abs(declarado), abs(suma[j]))) + 1e-9
                if abs(declarado - suma[j]) <= tol:
                    cierran += 1
                else:
                    fallan[x["idx"]] = (declarado, suma[j], declarado - suma[j])
            if cierran >= 2 and cierran >= CIERRE_NUMERICO * (cierran + len(fallan)):
                yield a, b, cierran, fallan


def analizar(c):
    if c.get("eje") == "filas":
        # Años en filas y variables en columnas (Total · Estatal · Privado): cada fila es un corte en el que la
        # columna «Total» suma su panel. Sólo la relación que declara el rótulo; las demás quedan para las series.
        numericas = _numericas(c)
        return totales_por_panel(c, [x for x in numericas if _es_col_total(x)], numericas)
    if c.get("eje") == "largo":
        return []
    bloques, c["totales_aparentes_sin_comprobar"] = bloques_paralelos(c)
    sin_rotulo = totales_sin_rotulo(c)
    subtotales, c["cierres_columna"] = subtotales_columna(c)
    if subtotales or c["cierres_columna"]:
        return subtotales + escala_fila_columna(c) + bloques + sin_rotulo
    escala = escala_fila_columna(c)
    explicadas = {h["evidencia"]["fila_explicada"] for h in escala if h["evidencia"].get("fila_explicada")}
    return cruces(c, explicadas) + escala + bloques + sin_rotulo
