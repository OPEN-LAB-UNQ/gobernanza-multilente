"""Modelo común de cuadro: ejes reconocidos por contenido antes de controlar.

Un cuadro tiene filas (rubros, provincias, períodos o registros) y columnas (períodos o
atributos). El eje temporal se busca por lo que dicen las celdas, no por el nombre de la
columna: años en el encabezado, meses o trimestres en el encabezado, períodos en una columna
(«2024-01», «Ene», fechas) o un formato largo período × dimensión × valor. Si no hay eje
temporal el cuadro es transversal (por provincia, por rubro, por meta). Los controles se
aplican después sobre esos ejes, sin depender de cómo se llame cada columna.
"""

import re

from openpyxl.utils import column_index_from_string, get_column_letter

from services.excel_estructura import CODIGO_SERIE
from services.perfilado import es_sin_dato
from services.sanitizacion import columna_sensible, redactar_texto_libre

from . import periodos as P
from .comun import normalizar, numero

TOTAL_RE = re.compile(r"^\s*(?:sub)?total(?:es)?\b|^\s*total\s", re.I)
IDENTIFICADOR_RE = re.compile(r"^(?:c[oó]d(?:igo)?\.?|prog(?:rama)?\.?|id|n[°ºo]\.?|nro\.?|jur\.?|sa[f]?\.?|inciso)$", re.I)
NOTA_RE = re.compile(r"^\s*(?:notas?|fuente|aclaraci[oó]n|obs(?:ervaci[oó]n(?:es)?)?|\(\*\)|\*)\b", re.I)
UNIDAD_RE = re.compile(r"\ben\s+(?:millones|miles|pesos|d[oó]lares|porcentaje|%|unidades)", re.I)


def es_total(etiqueta):
    return bool(TOTAL_RE.match(str(etiqueta or "")))


def _texto(v):
    return isinstance(v, str) and bool(v.strip()) and numero(v) is None


def notas_hoja(ws, excluir_filas=()):
    """Textos explicativos de la hoja (notas, fuentes, aclaraciones), saneados."""
    salida = []
    excluir = set(excluir_filas)
    for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
        for c in fila:
            v = c.value
            if c.row in excluir or not isinstance(v, str):
                continue
            t = v.strip()
            if len(t) >= 25 or NOTA_RE.match(t):
                salida.append({"celda": c.coordinate, "fila": c.row, "texto": redactar_texto_libre(t)[:600]})
    return salida


def _encabezado_superior(ws, header, cols):
    """Rótulos del nivel superior (variable) sobre las columnas, con celdas combinadas extendidas."""
    if header <= 1:
        return {}
    fila = header - 1
    combinadas = {}
    for rango in ws.merged_cells.ranges:
        if rango.min_row <= fila <= rango.max_row:
            v = ws.cell(rango.min_row, rango.min_col).value
            for c in range(rango.min_col, rango.max_col + 1):
                combinadas[c] = v
    salida, actual = {}, None
    for idx in sorted(cols):
        v = combinadas.get(idx, ws.cell(fila, idx).value)
        if _texto(v):
            actual = str(v).strip()
        elif v is not None and str(v).strip():
            actual = None
        salida[idx] = actual
    return salida


class Cuadro(dict):
    """Diccionario con ayudas de lectura; se conserva serializable por sus claves."""

    def celda(self, fila, idx):
        return f"{get_column_letter(idx)}{fila}"

    def ref(self, fila, idx):
        return f"{self['hoja']}!{get_column_letter(idx)}{fila}"

    def valor(self, fila, idx):
        return self["valores"].get((fila, idx))

    def columnas_numericas(self, incluir_periodos=True):
        return [c for c in self["columnas"] if c["numerica"] and (incluir_periodos or not c.get("periodo"))]

    def filas_datos(self, sin_totales=True):
        return [r for r in self["filas"] if not (sin_totales and es_total(self["etiquetas"].get(r)))]


def _rotulo_inferior(ws, filas_rotulos, idx):
    """El rótulo más bajo con texto de la columna dentro de las filas del encabezado (no un código de serie)."""
    for r in sorted(filas_rotulos or [], reverse=True):
        v = ws.cell(r, idx).value
        if isinstance(v, str) and v.strip() and not CODIGO_SERIE.match(v.strip()):
            return v.strip()
    return None


def armar(archivo, hoja, cuadro, ws, libro, header, filas, columnas, filas_rotulos=None):
    """columnas: [(idx, nombre)]. Devuelve un Cuadro con ejes reconocidos.

    filas_rotulos: filas del encabezado de varios niveles; de ahí sale `rotulo_inferior`, que sólo se usa para
    reconocer una columna «Total» cuyo rótulo no está en la última fila."""
    valores, crudos, formulas = {}, {}, {}
    cols = []
    for idx, nombre in sorted(columnas):
        letra = get_column_letter(idx)
        nums = 0
        presentes = 0
        for r in filas:
            raw = libro.crudo(hoja, r, idx)
            f = libro.formula(hoja, r, idx)
            crudos[(r, idx)] = raw
            formulas[(r, idx)] = bool(f)
            v = libro.valor(hoja, r, idx)
            valores[(r, idx)] = v
            # «-», «///», «s/d»: marcas de sin dato, no texto. Una columna con muchos guiones sigue siendo numérica
            # (los nuevos inscriptos de una carrera sin ingresantes) y se controla.
            if (raw is not None and str(raw).strip() and not es_sin_dato(raw)) or f:
                presentes += 1
            if v is not None:
                nums += 1
        encabezado = ws.cell(header, idx).value if header else None
        if isinstance(encabezado, str) and CODIGO_SERIE.match(encabezado.strip()):
            # Un código de serie («84.2_CCFSBT_0_T_43») no es el rótulo: vale el nombre, que trae el del cuadro de arriba.
            encabezado = None
        cols.append(dict(idx=idx, letra=letra, nombre=str(nombre), encabezado=encabezado,
                         numerica=presentes > 0 and nums >= max(1, 0.6 * presentes), presentes=presentes,
                         rotulo_inferior=_rotulo_inferior(ws, filas_rotulos, idx)))
    # Columna de rótulos: la primera con texto en la mayoría de las filas.
    textuales = [c for c in cols if not c["numerica"] and sum(_texto(crudos[(r, c["idx"])]) for r in filas) >= 0.5 * len(filas)]
    etiqueta_idx = textuales[0]["idx"] if textuales else None
    # Identificadores numéricos (códigos de programa, de provincia) a la izquierda del rótulo.
    for c in cols:
        if not c["numerica"]:
            continue
        izquierda = etiqueta_idx is not None and c["idx"] < etiqueta_idx
        rotulo_id = IDENTIFICADOR_RE.match(normalizar(c["encabezado"] or c["nombre"]).strip())
        enteros = [valores[(r, c["idx"])] for r in filas if valores[(r, c["idx"])] is not None]
        if (izquierda or rotulo_id) and enteros and all(float(v).is_integer() and 0 <= v < 10 ** 7 for v in enteros):
            c["identificador"] = True
            c["numerica"] = False
    etiquetas, codigos = {}, {}
    for r in filas:
        partes = [str(crudos[(r, c["idx"])]).strip() for c in cols
                  if not c["numerica"] and not c.get("identificador") and _texto(crudos[(r, c["idx"])])
                  and (etiqueta_idx is None or c["idx"] <= etiqueta_idx)]
        etiquetas[r] = redactar_texto_libre(" ".join(partes))[:160] if partes else None
        ids = [crudos[(r, c["idx"])] for c in cols if c.get("identificador") and crudos[(r, c["idx"])] is not None]
        codigos[r] = str(int(numero(ids[0]))) if ids and numero(ids[0]) is not None else None
    sangria = {r: (ws.cell(r, etiqueta_idx).alignment.indent if etiqueta_idx else 0) or 0 for r in filas}
    c = Cuadro(archivo=archivo, hoja=hoja, cuadro=cuadro, header=header, filas=list(filas), columnas=cols,
               valores=valores, crudos=crudos, formulas=formulas, etiquetas=etiquetas, codigos=codigos,
               sangria=sangria, etiqueta_idx=etiqueta_idx, eje=None, frecuencia=None, col_tiempo=None,
               periodos_fila={}, claves=[], paneles={}, notas=[], unidad=None,
               formulas_desconocidas=bool(getattr(libro, "formulas_desconocidas", False)))
    _reconocer_eje(c, ws)
    return c


# Subperíodo en la fila de abajo del encabezado de años: «I»…«IV», «T1», «1er trim.», «Trim. 2», «II trim.».
SUBPERIODO_RE = re.compile(
    r"^(?:(?P<r>i{1,3}|iv)\.?(?:\s*(?:t|trim(?:estre)?)\.?)?|(?:t|q)\s*(?P<t>[1-4])"
    r"|(?P<n>[1-4])\s*(?:°|º|o|er|do|ro|to)?\.?\s*(?:t|trim(?:estre)?)\.?|trim(?:estre)?\.?\s*(?P<m>[1-4]))$"
)
TOTAL_SUBPERIODO_RE = re.compile(r"^(?:sub)?total(?:\s+(?:anual|del?\s+ano|ano|\d{4}))?\.?$")
MINIMO_SUBPERIODOS = 4


def _subperiodo(v):
    """Trimestre (1 a 4) de un rótulo del segundo nivel, «total» o None."""
    if not isinstance(v, str) or not v.strip():
        return None
    t = normalizar(v)
    if TOTAL_SUBPERIODO_RE.match(t):
        return "total"
    m = SUBPERIODO_RE.match(t)
    if not m:
        return None
    if m.group("r"):
        return P.ROMANOS[m.group("r")]
    return int(m.group("t") or m.group("n") or m.group("m"))


def _dos_niveles(c, ws):
    """Encabezado de dos niveles año + subperíodo (balanza de pagos de un instituto de estadística): una fila de años
    («2006» … «2019», «Año 2020») y, debajo, «I», «II», «III», «IV», «Total» (o «1º trimestre» …
    «4º trimestre», «Total»), o los meses («Julio» … «Diciembre», «Subtotal 2025», «Enero» …, como en
    un perfil de vencimientos). El encabezado reconocido puede ser la fila de los años o la de los
    subperíodos: se prueba con la de abajo y, si no, con la de arriba.
    Devuelve True si reconoció el encabezado."""
    header = c["header"]
    if not header:
        return False
    return _dos_niveles_en(c, ws, header, header + 1) or (header > 1 and _dos_niveles_en(c, ws, header - 1, header))


def _dos_niveles_en(c, ws, fila_anio, fila_sub):
    """Encabezado de dos niveles con los años en `fila_anio` y los subperíodos en `fila_sub`.

    El año se arrastra a la derecha desde su rótulo hasta el siguiente; cada columna con trimestre
    (o mes) es un período de esa frecuencia, «Total» o «Subtotal 2025» es el total del año (sin
    período) y un año sin subperíodo es un período anual. Si conviven años sueltos y subperíodos,
    cada frecuencia es su propio panel. Devuelve True si reconoció el encabezado."""
    numericas = sorted((x for x in c["columnas"] if x["numerica"]), key=lambda x: x["idx"])
    if len(numericas) < MINIMO_SUBPERIODOS:
        return False
    sub, frecuencias = {}, set()
    for x in numericas:
        v = ws.cell(fila_sub, x["idx"]).value
        s = _subperiodo(v)
        if s is None and v is not None and str(v).strip():
            sf = _subperiodo_fila(v)  # un mes sin año («Julio»); «Acumulado a junio» conserva su lectura
            if not sf or sf[1] != "mensual" or (P.leer(v) or {}).get("acumulado"):
                return False  # la fila de abajo trae otra cosa (datos, unidades): no es un segundo nivel
            s = sf[0]
            frecuencias.add("mensual")
        elif isinstance(s, int):
            frecuencias.add("trimestral")
        sub[x["idx"]] = s
    if len(frecuencias) != 1 or sum(isinstance(s, int) for s in sub.values()) < MINIMO_SUBPERIODOS:
        return False
    frecuencia = frecuencias.pop()
    if frecuencia == "trimestral" and not any(isinstance(s, int) and s > 2 for s in sub.values()):
        return False  # sólo «I» y «II»: pueden ser semestres (encuesta de hogares), no trimestres
    # Año vigente en cada columna, arrastrado desde su rótulo (también con celdas combinadas).
    anio_col, anio, marca = {}, None, None
    for j in range(numericas[0]["idx"], numericas[-1]["idx"] + 1):
        v = ws.cell(fila_anio, j).value
        p = P.leer(v)
        if p and p["frecuencia"] == "anual":
            anio, marca = p["anio"], p.get("marca")
        elif v is not None and str(v).strip():
            anio, marca = None, None
        anio_col[j] = (anio, marca, p if p and p["frecuencia"] == "anual" else None)
    periodos = {}
    for x in numericas:
        anio, marca, propio = anio_col[x["idx"]]
        s = sub[x["idx"]]
        if s == "total":
            x["total_periodo"] = True
        elif isinstance(s, int) and anio is not None:
            texto = f"{anio}-T{s}" if frecuencia == "trimestral" else f"{anio}-{s:02d}"
            periodos[x["idx"]] = dict(anio=anio, sub=s, frecuencia=frecuencia, marca=marca, acumulado=False,
                                      texto=texto)
        elif s is None and propio is not None:
            periodos[x["idx"]] = dict(propio)
    subperiodos = [p for p in periodos.values() if p["frecuencia"] == frecuencia]
    if len(subperiodos) < MINIMO_SUBPERIODOS:
        return False
    for p in periodos.values():
        p["ordinal"] = P.ordinal(p)
    for f in ("anual", frecuencia):
        ords = [p["ordinal"] for p in periodos.values() if p["frecuencia"] == f]
        if len(ords) != len(set(ords)):
            return False  # un año o subperíodo repetido: no es la secuencia de un solo nivel temporal
    mezcla = len({p["frecuencia"] for p in periodos.values()}) > 1
    for x in numericas:
        p = periodos.get(x["idx"])
        if p:
            x["periodo"] = p
            x["panel"] = None
    if mezcla:
        c["paneles"] = {idx: p["frecuencia"] for idx, p in periodos.items()}
    c.update(eje="columnas", frecuencia=frecuencia, subperiodos=True)
    return True


def _fila_de_unidad(rotulos):
    """La fila tomada como encabezado sólo dice la unidad («En millones de pesos corrientes»)."""
    presentes = [str(v).strip() for v in rotulos if v is not None and str(v).strip()]
    return bool(presentes) and all(UNIDAD_RE.match(v) for v in presentes)


def _reconocer_eje(c, ws):
    numericas = [x for x in c["columnas"] if x["numerica"]]
    # 0. Encabezado de dos niveles: año arriba y trimestre («I» … «IV», «Total») abajo.
    if c["header"] and _dos_niveles(c, ws):
        return
    # 1. Períodos en el encabezado (años, meses, trimestres), en uno o dos niveles.
    if c["header"]:
        rotulos = [x["encabezado"] for x in numericas]
        if c["header"] > 1 and _fila_de_unidad(rotulos):
            # La unidad quedó como encabezado y los años, en la fila de arriba (una jurisdicción, ingresos tributarios).
            arriba = [ws.cell(c["header"] - 1, x["idx"]).value for x in numericas]
            if P.leer_eje(arriba):
                rotulos = arriba
        eje = P.leer_eje(rotulos)
        if eje:
            superior = _encabezado_superior(ws, c["header"], [x["idx"] for x in numericas])
            for x, p in zip(numericas, eje["periodos"]):
                if p:
                    x["periodo"] = p
                    x["panel"] = superior.get(x["idx"])
            con = [x for x in numericas if x.get("periodo")]
            ords = [x["periodo"]["ordinal"] for x in con]
            paneles = {x["panel"] for x in con}
            if len(ords) != len(set(ords)) and len(paneles) > 1 and None not in paneles:
                # Dos niveles: variable × período. Cada panel debe ser una secuencia sin repetidos.
                por_panel = {}
                for x in con:
                    por_panel.setdefault(x["panel"], []).append(x["periodo"]["ordinal"])
                if all(len(v) == len(set(v)) for v in por_panel.values()):
                    c["paneles"] = {x["idx"]: x["panel"] for x in con}
            else:
                for x in con:
                    x["panel"] = None
            c.update(eje="columnas", frecuencia=eje["frecuencia"])
            return
        # Dos ejercicios (balances, estados de resultados): cada uno es un corte comparable.
        corto = P.leer_eje(rotulos, minimo=2)
        if corto:
            for x, p in zip(numericas, corto["periodos"]):
                if p:
                    x["periodo_corto"] = p
    # El año + mes explícito tiene prioridad sobre leer «enero…diciembre» sin su año.
    if _anio_y_subperiodo(c):
        return
    # 2. Períodos en una columna (vertical) o formato largo.
    candidatas = [x for x in c["columnas"] if not x["numerica"] or x.get("identificador")
                  or all(c["valores"][(r, x["idx"])] is None or float(c["valores"][(r, x["idx"])]).is_integer()
                         for r in c["filas"])]
    for x in candidatas:
        vals = [c["crudos"][(r, x["idx"])] for r in c["filas"]]
        eje = P.leer_eje(vals)
        if not eje or sum(1 for p in eje["periodos"] if p) < max(3, 0.8 * len(c["filas"])):
            continue
        if eje["frecuencia"] == "anual" and any(p and p.get("dia") for p in eje["periodos"]):
            continue
        x["numerica"] = False
        x["tiempo"] = True
        c.update(col_tiempo=x["idx"], frecuencia=eje["frecuencia"], eje="filas",
                 periodos_fila={r: p for r, p in zip(c["filas"], eje["periodos"]) if p})
        ords = [p["ordinal"] for p in c["periodos_fila"].values()]
        if len(ords) != len(set(ords)):
            claves = _claves_largo(c, x["idx"])
            if claves:
                c.update(eje="largo", claves=claves)
        return


def _subperiodo_fila(v):
    """(subperíodo, frecuencia) de un rótulo «I»…«IV» o de un mes sin año; None si no lo es."""
    s = _subperiodo(v)
    if isinstance(s, int):
        return s, "trimestral"
    p = P.leer(v) if isinstance(v, str) else None
    if p and p.get("anio") is None and p["frecuencia"] in ("mensual", "trimestral", "semestral"):
        return p["sub"], p["frecuencia"]
    return None


def _anio_y_subperiodo(c):
    """Período compuesto: el año de una columna, arrastrado hacia abajo, y el subperíodo de la columna
    vecina («1993» en A sólo junto a «I»; «II», «III», «IV» debajo en B). Devuelve True si lo reconoce."""
    filas = c["filas"]
    cols = sorted(c["columnas"], key=lambda x: x["idx"])
    for x, y in zip(cols, cols[1:]):
        if y["idx"] != x["idx"] + 1 or (x["numerica"] and not x.get("identificador")) or y["numerica"]:
            continue
        anio, marca, periodos, frecuencias = None, None, {}, set()
        for r in filas:
            v = c["crudos"][(r, x["idx"])]
            p = P.leer(v)
            if p and p["frecuencia"] == "anual":
                anio, marca = p["anio"], p.get("marca")
            elif v is not None and str(v).strip():
                anio = None  # otro texto en la columna del año («Total»): corta el arrastre
            s = _subperiodo_fila(c["crudos"][(r, y["idx"])])
            if anio is not None and s:
                sub, frecuencia = s
                frecuencias.add(frecuencia)
                texto = f"{anio}-T{sub}" if frecuencia == "trimestral" else f"{anio}-{sub:02d}"
                periodos[r] = dict(anio=anio, sub=sub, frecuencia=frecuencia, marca=marca, acumulado=False, texto=texto)
        if len(frecuencias) != 1 or len(periodos) < max(3, 0.8 * len(filas)):
            continue
        if len({p["anio"] for p in periodos.values()}) < 2:
            continue
        for p in periodos.values():
            p["ordinal"] = P.ordinal(p)
        ords = [p["ordinal"] for p in periodos.values()]
        explicito = (normalizar(x["nombre"]) in ("ano", "anio", "year")
                     and normalizar(y["nombre"]) in ("mes", "month", "trimestre", "semestre"))
        if len(ords) != len(set(ords)) and not explicito:
            continue
        x["numerica"] = False
        x["tiempo"] = y["tiempo"] = True
        c.update(col_tiempo=y["idx"], frecuencia=frecuencias.pop(), eje="filas", periodos_fila=periodos)
        return True
    return False


UNIDAD_COLUMNA_RE = re.compile(r"^(?:unidad(?:es)?(?: de medida)?|u\.? ?(?:de )?medida|medida|unit|um)$")


def _claves_largo(c, tiempo):
    """Columnas de texto de pocos valores distintos que, junto al período, identifican cada fila.

    Una columna de unidad no separa series: si una misma combinación (provincia, producto…)
    trae filas en otra unidad, esa fila es la señal —unidad distinta en esa fila— y la serie
    sigue entera, en lugar de partirse y disparar alarmas de período en cada pedazo.
    """
    unidades = [x for x in c["columnas"] if x["idx"] != tiempo and not x["numerica"]
                and UNIDAD_COLUMNA_RE.match(normalizar(x["encabezado"] or x["nombre"]).strip())]
    if len(unidades) == 1:
        u = unidades[0]
        sin_unidad = _elegir_claves(c, tiempo, excluir={u["idx"]})
        if sin_unidad is not None:
            mezcla = _filas_en_otra_unidad(c, sin_unidad, u["idx"])
            if mezcla:
                c["unidad_mezclada"] = dict(idx=u["idx"], nombre=u["nombre"], filas=mezcla)
                return sin_unidad
    return _elegir_claves(c, tiempo, excluir=set(), estricto=False)


def _filas_en_otra_unidad(c, claves, idx):
    """{fila: (unidad de la fila, unidad del resto de su serie)} cuando una serie mezcla unidades."""
    grupos = {}
    for r in c["filas"]:
        if r not in c["periodos_fila"]:
            continue
        clave = tuple(str(c["crudos"][(r, k)]).strip() for k in claves)
        grupos.setdefault(clave, []).append(r)
    salida = {}
    for filas in grupos.values():
        valores = [str(c["crudos"][(r, idx)] or "").strip() for r in filas]
        conteo = {}
        for v in valores:
            conteo[v] = conteo.get(v, 0) + 1
        if len(conteo) < 2:
            continue
        mayoritaria, veces = max(conteo.items(), key=lambda kv: kv[1])
        # Sólo una minoría clara: una serie partida en mitades es otra cosa (cambio de unidad declarado).
        if veces < 3 or veces < 0.75 * len(filas):
            continue
        for r, v in zip(filas, valores):
            if v != mayoritaria and v:
                salida[r] = (v, mayoritaria)
    return salida


def _elegir_claves(c, tiempo, excluir, estricto=True):
    """Claves que identifican cada fila. Con `estricto`, None si no alcanzan a identificarlas."""
    n = len(c["filas"])
    candidatas = []
    for x in c["columnas"]:
        if x["idx"] == tiempo or x["numerica"] or x["idx"] in excluir:
            continue
        vals = [c["crudos"][(r, x["idx"])] for r in c["filas"]]
        distintos = {str(v).strip() for v in vals if v is not None and str(v).strip()}
        if 2 <= len(distintos) <= max(2, n // 2) and sum(v is not None for v in vals) >= 0.9 * n:
            candidatas.append((len(distintos), x["idx"]))
    candidatas.sort()
    elegidas = []
    for _, idx in candidatas:
        elegidas.append(idx)
        combos = [(tuple(str(c["crudos"][(r, k)]).strip() for k in elegidas), c["periodos_fila"].get(r, {}).get("ordinal"))
                  for r in c["filas"] if r in c["periodos_fila"]]
        if len(combos) == len(set(combos)):
            return elegidas
    if estricto:
        return None
    # Las claves no alcanzan para identificar cada fila: el período se repite de verdad.
    return elegidas if elegidas and _mayoria_unica(c, elegidas) else []


def _mayoria_unica(c, claves):
    combos = [(tuple(str(c["crudos"][(r, k)]).strip() for k in claves), c["periodos_fila"].get(r, {}).get("ordinal"))
              for r in c["filas"] if r in c["periodos_fila"]]
    return len(set(combos)) >= 0.9 * len(combos)


def desde_bloque(archivo, ws, libro, b):
    header = max(b["encabezados"])
    columnas = [(column_index_from_string(letra), nombre) for nombre, letra in b["columnas"].items()]
    c = armar(archivo, b["hoja"], b["bloque"], ws, libro, header, b["filas"], columnas,
              filas_rotulos=b.get("encabezados_rotulos") or b["encabezados"])
    return c


def sin_bloque(archivo, ws, libro, numero_cuadro):
    """Cuadro de una hoja que el lector no interpretó: rótulos de texto a la izquierda y cifras a la derecha."""
    hoja = ws.title
    max_col = min(ws.max_column, 60)
    max_row = min(ws.max_row, 2000)
    datos = []
    for r in range(1, max_row + 1):
        celdas = [(j, ws.cell(r, j).value) for j in range(1, max_col + 1)]
        textos = [j for j, v in celdas if _texto(v)]
        cifras = [j for j, v in celdas if libro.valor(hoja, r, j) is not None and not _texto(v)]
        if textos and cifras and min(textos) < min(cifras):
            datos.append((r, min(textos), cifras))
    if len(datos) < 3:
        return None
    etiqueta = max({d[1] for d in datos}, key=lambda j: sum(d[1] == j for d in datos))
    datos = [d for d in datos if d[1] == etiqueta]
    primera, ultima = datos[0][0], datos[-1][0]
    # Encabezado: la fila más cercana arriba del primer dato con dos o más celdas escritas.
    header = None
    for r in range(primera - 1, max(0, primera - 6), -1):
        escritas = [j for j in range(etiqueta, max_col + 1) if ws.cell(r, j).value not in (None, "")]
        if len(escritas) >= 2:
            header = r
            break
    if header is None:
        return None
    usadas = sorted({j for d in datos for j in d[2]})
    filas = [r for r in range(primera, ultima + 1) if any(libro.valor(hoja, r, j) is not None for j in usadas)
             and _texto(ws.cell(r, etiqueta).value)]
    nombres = {}
    for j in [etiqueta] + usadas:
        v = ws.cell(header, j).value
        nombres[j] = str(v).strip() if v not in (None, "") else get_column_letter(j)
    # Nunca incorporar columnas que identifican personas.
    for j in list(usadas):
        if columna_sensible(nombres[j], [ws.cell(r, j).value for r in filas]):
            usadas.remove(j)
    if columna_sensible(nombres[etiqueta], [ws.cell(r, etiqueta).value for r in filas]) or not usadas:
        return None
    return armar(archivo, hoja, f"h{numero_cuadro}", ws, libro, header, filas, [(j, nombres[j]) for j in [etiqueta] + usadas])


_ANIO_TXT = r"((?:18|19|20)\d{2})"
# La unidad: escala con moneda opcional («millones de pesos», «miles», «millones de dólares») o una moneda.
_UNIDAD_TXT = (r"((?:miles de millones|millones|miles|unidades)(?: de (?:pesos|dolares|australes|euros))?"
               r"|australes|pesos|dolares|euros)")
# La unidad tiene que venir inmediatamente después del año («desde 1987: millones», «desde 1987, en
# millones»), sin texto libre en el medio: «Desde 2005 incluye a las provincias, en millones» no es un corte.
_TRAS_ANIO = rf"\s*[:,]?\s*(?:en\s+)?{_UNIDAD_TXT}\b"
_DESDE = re.compile(rf"\b(?:desde|a partir de)\s+(?:el\s+)?(?:(?:ano|ejercicio)\s+)?{_ANIO_TXT}{_TRAS_ANIO}")
_HASTA = re.compile(rf"\bhasta\s+(?:el\s+)?(?:(?:ano|ejercicio)\s+)?{_ANIO_TXT}{_TRAS_ANIO}")
_TRAMO = re.compile(rf"\b{_ANIO_TXT}\s*(?:-|/|a|al)\s*{_ANIO_TXT}{_TRAS_ANIO}")
ESCALA_UNIDAD = {"unidades": 1.0, "miles": 1e3, "millones": 1e6, "miles de millones": 1e9}


def _escala_de(unidad):
    """Factor de la escala de una unidad («millones de pesos» → 1e6); None si sólo nombra una moneda."""
    for nombre in ("miles de millones", "millones", "miles", "unidades"):
        if unidad.startswith(nombre):
            return ESCALA_UNIDAD[nombre]
    return None


def cambios_declarados(texto):
    """[{anio, antes, despues}]: cambios de unidad que declara el texto, con la unidad de cada lado.

    «1980-1986: miles de pesos; Desde 1987: millones de pesos» → 1987, de «miles de pesos» a «millones de
    pesos». «Hasta 1991 en australes; desde 1992 en pesos» → 1992. Una sola unidad nombrada describe el
    cuadro, no un cambio. Si un lado del corte no se nombra junto a un año, se toma la otra unidad del texto
    cuando hay una sola.
    """
    t = normalizar(texto)
    tramos = []  # (desde, hasta, unidad); None = abierto
    tramos += [(int(m.group(1)), None, m.group(2)) for m in _DESDE.finditer(t)]
    tramos += [(None, int(m.group(1)), m.group(2)) for m in _HASTA.finditer(t)]
    tramos += [(int(m.group(1)), int(m.group(2)), m.group(3)) for m in _TRAMO.finditer(t)]
    unidades = {u for *_, u in tramos} | set(re.findall(_UNIDAD_TXT, t))
    if len(unidades) < 2 or not tramos:
        return []
    # Cortes: el comienzo de cada tramo y el año siguiente a un «hasta» (el fin de un tramo cerrado sólo
    # es corte si otro tramo empieza ahí).
    cortes = {d for d, _, _ in tramos if d is not None} | {h + 1 for d, h, _ in tramos if d is None and h is not None}
    primero = min([d for d, _, _ in tramos if d is not None] + [h for _, h, _ in tramos if h is not None])
    salida = []
    for c in sorted(cortes):
        despues = next((u for d, _, u in tramos if d == c), None)
        antes = next((u for _, h, u in tramos if h == c - 1), None)
        if despues is None and antes is not None:
            otras = unidades - {antes}
            despues = next(iter(otras)) if len(otras) == 1 else None
        elif antes is None and despues is not None:
            otras = unidades - {despues}
            antes = next(iter(otras)) if len(otras) == 1 else None
        # El comienzo del primer tramo no es un corte: es el comienzo de la serie.
        if c == primero and not any(h == c - 1 for _, h, _ in tramos if h is not None):
            continue
        if antes is None or despues is None or antes == despues:
            continue
        salida.append({"anio": c, "antes": antes, "despues": despues})
    return salida


def cortes_unidad(texto):
    """Años en que el propio texto declara que cambia la unidad de los valores."""
    return {x["anio"] for x in cambios_declarados(texto)}


def cambios_unidad(ws, header, notas=(), desde=1, hasta=None):
    """[{anio, antes, despues, celda, texto}]: cambios de unidad que declaran el título, las notas o el
    encabezado del cuadro. Sólo se leen las filas desde `desde` (la siguiente al cuadro anterior de la
    hoja) y, para las notas, hasta `hasta` (el comienzo del cuadro siguiente): la nota de otro cuadro de
    la misma hoja no explica éste."""
    textos = {}
    if header:
        for r in range(max(1, desde, header - 12), header + 1):
            for j in range(1, min(ws.max_column, 30) + 1):
                v = ws.cell(r, j).value
                if isinstance(v, str) and v.strip():
                    textos[f"{get_column_letter(j)}{r}"] = v.strip()
    for n in notas or []:
        if n["fila"] >= desde and (hasta is None or n["fila"] < hasta):
            textos.setdefault(n["celda"], n["texto"])
    salida = []
    for celda, texto in textos.items():
        for x in cambios_declarados(texto):
            salida.append({**x, "celda": celda, "texto": redactar_texto_libre(texto)[:300]})
    return salida


def unidad_cuadro(ws, header):
    for r in range(max(1, header - 6), header):
        for j in range(1, min(ws.max_column, 6) + 1):
            v = ws.cell(r, j).value
            if isinstance(v, str) and UNIDAD_RE.search(v):
                return normalizar(v)
    return None
