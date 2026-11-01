"""Controles reproducibles de libros tabulares: sólo filas/columnas saneadas."""

import io
import re
from copy import copy

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from services.excel_estructura import anios_plausibles
from services.sanitizacion import sanear

from . import (
    compensados,
    conciliacion,
    ediciones,
    ejes,
    entre_hojas,
    formulas_fijas,
    identidades,
    jerarquia_codigos,
    lectura,
    matriz,
    per_capita,
    temporal,
    totales_cruzados,
    transversal,
)
from . import modelo as modelo_cuadro
from . import periodos as P
from . import series as motor_series
from .comun import normalizar, numero
from .evaluador import Libro


def _anio_encabezado(v):
    """Año del encabezado; también «2023 (p)», «2024*», «2022 e»: la marca es un metadato, no un error."""
    n = numero(v)
    if n is None and isinstance(v, str):
        p = P.leer(v)
        if p and p["frecuencia"] == "anual" and (p.get("marca") or p.get("unidad")):
            return float(p["anio"])
    return n


def _formatos_texto(ws, filas, letras):
    """Formato de los números guardados como texto, columna por columna, como en la lectura del cuadro.

    La columna decide con sus propios valores; si son ambiguos («475.735»), decide el cuadro; si
    tampoco, el formato local (punto de miles). Sin esto, «475.735» se leía como 475,735 en los
    controles aunque la lectura del cuadro lo hubiera leído como 475.735.
    """
    from services.perfilado import formato_columna, preferencia_formato

    textos = {letra: [ws[f"{letra}{r}"].value for r in filas] for letra in letras}
    preferencia = preferencia_formato(v for vs in textos.values() for v in vs if isinstance(v, str))
    return {letra: formato_columna(vs, preferencia)[0] for letra, vs in textos.items()
            if any(isinstance(v, str) for v in vs)}


def _con_formato(valor, formato):
    from services.perfilado import es_sin_dato, numero_con_formato

    if not isinstance(valor, str) or formato is None or es_sin_dato(valor):
        return valor
    leido = numero_con_formato(valor, formato)
    return valor if leido is None else leido


def _marca(v):
    p = P.leer(v) if isinstance(v, str) else None
    return p.get("marca") if p else None


def _unidad(texto):
    t = normalizar(texto)
    if "%" in t or "porcentaje" in t:
        return t
    if ("peso" in t or "$" in t) and "por " not in t:
        escala = ("miles_de_millones" if "miles de millones" in t else "millones" if "millones" in t
                  else "miles" if "miles" in t else "unidades")
        return "pesos_" + escala + ("_constantes" if "constante" in t else "_corrientes")
    return t


def _libros(leidos):
    """Libros abiertos una sola vez: valores guardados y fórmulas, con evaluación de respaldo."""
    salida = []
    for archivo, contenido in leidos:
        if contenido[:8] == FIRMA_OLE:
            # .xls binario: sus valores guardados, en un libro en memoria, para que tenga los mismos
            # controles (cierres, series) que el mismo cuadro guardado como .xlsx.
            try:
                wb = libro_desde_xls(contenido)
            except Exception:  # noqa: BLE001 — sin libro, el .xls conserva la lectura y los controles por serie
                continue
            libro = Libro(wb, wb)
            # El .xls binario no trae fórmulas legibles: «sin fórmula» no quiere decir «tipeado».
            libro.formulas_desconocidas = True
            salida.append((archivo, libro))
            continue
        # Por contenido: un .xls que por dentro es un libro xlsx (zip) también se controla.
        if not (archivo.lower().endswith((".xlsx", ".xlsm")) or contenido[:2] == b"PK"):
            continue
        valores = load_workbook(io.BytesIO(contenido), data_only=True)
        formulas = load_workbook(io.BytesIO(contenido), data_only=False)
        salida.append((archivo, Libro(valores, formulas)))
    return salida


FIRMA_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def libro_desde_xls(contenido):
    """.xls binario → Workbook de openpyxl con los valores guardados, en las mismas coordenadas que
    lee pandas (fila r, columna c de xlrd → celda r+1, c+1). El formato binario no guarda fórmulas
    legibles: el libro sirve a la vez de valores y de fórmulas."""
    import xlrd
    from openpyxl import Workbook

    try:
        libro = xlrd.open_workbook(file_contents=contenido, formatting_info=True)
    except xlrd.XLRDError:
        libro = xlrd.open_workbook(file_contents=contenido)
    wb = Workbook()
    wb.remove(wb.active)
    for hoja in libro.sheets():
        ws = wb.create_sheet(hoja.name[:31])
        if hoja.visibility:
            ws.sheet_state = "hidden" if hoja.visibility == 1 else "veryHidden"
        for r in range(hoja.nrows):
            for c in range(hoja.ncols):
                celda = hoja.cell(r, c)
                if celda.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    continue
                v = celda.value
                if celda.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, libro.datemode)
                elif celda.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                    v = int(v)
                elif celda.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                destino = ws.cell(r + 1, c + 1, v)
                try:
                    xf = libro.xf_list[celda.xf_index]
                    destino.number_format = libro.format_map[xf.format_key].format_str
                except (AttributeError, IndexError, KeyError):
                    pass  # valores legibles sin metadatos de formato
                if isinstance(v, str) and v.startswith("="):
                    # Un texto que empieza con «=» no es una fórmula: openpyxl lo tomaría como tal.
                    destino.data_type = "s"
    return wb


def _cerrar(libros):
    for _, libro in libros:
        libro.valores.close()
        libro.formulas.close()


def _ventana_cuadro(cobertura, b):
    """Filas de la hoja que pueden traer el título o las notas del cuadro: desde la siguiente al último
    cuadro de arriba (con otro encabezado) hasta el encabezado del cuadro de abajo."""
    header = max(b["encabezados"])
    otros = [x for x in cobertura["bloques"] if (x["archivo"], x["hoja"]) == (b["archivo"], b["hoja"])
             and x.get("filas") and x.get("encabezados") and max(x["encabezados"]) != header]
    arriba = [max(x["filas"]) for x in otros if max(x["encabezados"]) < header]
    abajo = [min(x["encabezados"]) for x in otros if max(x["encabezados"]) > header]
    return (max(arriba) + 1 if arriba else 1), (min(abajo) if abajo else None)


def _modelos(libros, cobertura):
    """Modelo común de cada cuadro leído y de las hojas que el lector no pudo interpretar."""
    salida = {}
    for archivo, libro in libros:
        con_bloque = set()
        for b in cobertura["bloques"]:
            from services.interpretacion_previa import permite
            if not permite("tabular", b["archivo"], b["hoja"], b["bloque"]):
                continue
            if b["archivo"] != archivo or not b.get("filas") or not b.get("columnas"):
                continue
            ws = libro.valores[b["hoja"]]
            try:
                c = modelo_cuadro.desde_bloque(archivo, ws, libro, b)
            except Exception:  # noqa: BLE001 — un cuadro no modelable conserva los controles por serie
                continue
            c["notas"] = modelo_cuadro.notas_hoja(ws, c["filas"])
            desde, hasta = _ventana_cuadro(cobertura, b)
            c["cambios_unidad"] = modelo_cuadro.cambios_unidad(ws, c["header"], c["notas"], desde, hasta)
            c["unidad"] = modelo_cuadro.unidad_cuadro(ws, c["header"])
            c["formatos_celda"] = {(r, x["idx"]): ws.cell(r, x["idx"]).number_format
                                   for r in c["filas"] for x in c["columnas"]}
            _aplicar_interpretacion(c)
            salida[(archivo, b["hoja"], b["bloque"])] = c
            con_bloque.add(b["hoja"])
        for n, ws in enumerate(libro.valores.worksheets, 1):
            from services.interpretacion_previa import actual
            if actual() is not None:
                # Un cuadro omitido por el inventario no recibe controles por una vía auxiliar.
                continue
            if ws.title in con_bloque or ws.sheet_state != "visible":
                continue
            try:
                c = modelo_cuadro.sin_bloque(archivo, ws, libro, n)
            except Exception:  # noqa: BLE001
                c = None
            if c:
                c["notas"] = modelo_cuadro.notas_hoja(ws, c["filas"])
                c["cambios_unidad"] = modelo_cuadro.cambios_unidad(ws, c["header"], c["notas"])
                c["unidad"] = modelo_cuadro.unidad_cuadro(ws, c["header"])
                salida[(archivo, ws.title, c["cuadro"])] = c
    return salida


def _aplicar_interpretacion(c):
    from services import interpretacion_previa as ip
    u = ip.unidad(c["archivo"], c["hoja"], c["cuadro"])
    if u is None:
        return
    i = u.get("confirmada", {})
    roles = i.get("roles", {})
    c["eje_confirmado"] = i.get("eje")
    for col in c["columnas"]:
        rol = roles.get(col["nombre"])
        col["rol_confirmado"] = rol
        if rol in {"identificador", "categoria", "fecha_operacion", "periodo_serie", "otro"}:
            col["numerica"] = False
            col["identificador"] = rol == "identificador"
        elif rol in {"importe", "cantidad", "porcentaje", "indice"}:
            col["identificador"] = False
            col["tiempo"] = False
    if i.get("eje") in {"operaciones", "ninguno"}:
        c.update(eje=None, frecuencia=None, col_tiempo=None, periodos_fila={}, paneles={})
    elif i.get("eje") == "serie_filas" and c.get("eje") == "columnas":
        c.update(eje=None, frecuencia=None, paneles={})
    elif i.get("eje") == "serie_columnas" and c.get("eje") in {"filas", "largo"}:
        c.update(eje=None, frecuencia=None, col_tiempo=None, periodos_fila={}, paneles={})


def _cuadros(leidos, datos, cobertura, libros=None, modelos=None):
    salida = []
    headers = []
    vistos = set()
    propios = libros is None
    libros = _libros(leidos) if libros is None else libros
    modelos = _modelos(libros, cobertura) if modelos is None else modelos
    try:
        for archivo, libro in libros:
            valores, formulas = libro.valores, libro.formulas
            desconocidas = bool(getattr(libro, "formulas_desconocidas", False))
            for b in cobertura["bloques"]:
                from services.interpretacion_previa import permite
                if not permite("series", b["archivo"], b["hoja"], b["bloque"]):
                    continue
                if b["archivo"] != archivo or not b.get("filas") or not b.get("columnas"):
                    continue
                ws = valores[b["hoja"]]
                wf = formulas[b["hoja"]]
                from services.estructura_revision import contexto_cuadro
                contexto = contexto_cuadro(ws, b)
                header = max(b["encabezados"])
                modelo = modelos.get((archivo, b["hoja"], b["bloque"])) or {}
                if modelo.get("eje_confirmado") == "serie_filas":
                    continue  # los años repetidos por panel no son un eje de columnas
                if modelo.get("subperiodos"):
                    # Año + trimestre en dos filas de encabezado: lo lee sólo el modelo común (un panel por
                    # frecuencia). La lectura anual tomaría «Año 2020» como un año y los trimestres como columnas sueltas.
                    continue
                # Sólo columnas temporales del cuadro reconocido; conservar
                # incluso año duplicado o no entero en el encabezado de origen.
                columnas = []
                for nombre, letra in b["columnas"].items():
                    v = ws[f"{letra}{header}"].value
                    n = _anio_encabezado(v)
                    if n is not None and 1800 <= n <= 2200:
                        columnas.append((column_index_from_string(letra), letra, int(n) if n.is_integer() else n))
                columnas.sort()
                # Importes que caen entre 1800 y 2200 (1920,5 · 1934,2 · …) no son un encabezado de años.
                if columnas and not anios_plausibles([float(x[2]) for x in columnas]):
                    columnas = []
                # Tres períodos alcanzan para leer el cuadro; cada regla exige después su propio mínimo.
                if len(columnas) < 3:
                    if modelo.get("eje") == "largo" or (modelo.get("eje") == "filas" and modelo.get("frecuencia") != "anual"):
                        continue  # formato largo y períodos no anuales: se leen con el modelo común
                    verticales, errores = _vertical(ws, wf, b, datos, archivo)
                    for serie in verticales:
                        serie["ambito"] = contexto
                        serie["formulas_desconocidas"] = desconocidas
                    salida.extend(verticales)
                    headers.extend(errores)
                    continue
                # Conservar también un encabezado corrupto dentro del tramo temporal.
                extremo_min = min(x[0] for x in columnas)
                extremo_max = max(x[0] for x in columnas)
                columnas = []
                for _, letra in b["columnas"].items():
                    col = column_index_from_string(letra)
                    if extremo_min <= col <= extremo_max:
                        raw = ws[f"{letra}{header}"].value
                        if P.es_llamada(raw):
                            continue  # «(2)» entre los años: columna de notas, no el período −2
                        n = _anio_encabezado(raw)
                        if n is None and P.es_auxiliar(raw):
                            continue  # «CV (%)» entre los años: no es un período mal escrito (se declara aparte)
                        columnas.append((col, letra, int(n) if n is not None and n.is_integer() else raw))
                columnas.sort()
                # Encabezado de dos niveles (variable × período): cada variable es su propio panel.
                paneles_modelo = dict(modelo.get("paneles") or {})
                # Columnas del tramo que el modelo no asignó a una variable (por ejemplo, sin números en este
                # bloque): toman el rótulo de arriba, combinado sobre sus años. Si no, el panel queda
                # incompleto y «2014» se lee repetido en una sola secuencia.
                faltan = [x[0] for x in columnas if x[0] not in paneles_modelo]
                if paneles_modelo and faltan:
                    superior = modelo_cuadro._encabezado_superior(ws, header, faltan)
                    paneles_modelo.update({k: v for k, v in superior.items() if v})
                grupos_panel = {}
                for x in columnas:
                    grupos_panel.setdefault(paneles_modelo.get(x[0]), []).append(x)
                if len(grupos_panel) < 2 or None in grupos_panel:
                    grupos_panel = {None: columnas}
                for panel, cols_panel in grupos_panel.items():
                    series_panel, errores_panel = _horizontal(
                        archivo, b, ws, wf, datos, header, cols_panel, contexto, panel, vistos)
                    for serie in series_panel:
                        serie["formulas_desconocidas"] = desconocidas
                    salida.extend(series_panel)
                    headers.extend(errores_panel)
    finally:
        if propios:
            _cerrar(libros)
    return salida, headers


def _horizontal(archivo, b, ws, wf, datos, header, columnas, contexto, panel, vistos):
    salida, headers = [], []
    if not columnas:
        return salida, headers
    periodos = [x[2] for x in columnas]
    key = (archivo, b["hoja"], header, tuple(x[0] for x in columnas))
    s_header = dict(
        archivo=archivo,
        hoja=b["hoja"],
        cuadro=b["bloque"],
        etiqueta="Encabezado temporal" + (f" · {panel}" if panel else ""),
        periodos=periodos,
        celdas=[f"{letra}{header}" for _, letra, _ in columnas],
        valores=periodos,
    )
    errores, validos = temporal.validar(s_header)
    if key not in vistos:
        vistos.add(key)
        headers.extend(errores)
    marcas = [_marca(ws[f"{letra}{header}"].value) for _, letra, _ in columnas]
    etiqueta_col = min(column_index_from_string(letra) for letra in b["columnas"].values())
    # La unidad explícita más cercana evita mezclar importes y porcentajes.
    unidades = [
        str(ws.cell(r, etiqueta_col).value)
        for r in range(1, header)
        if re.match(
            r"^en\s+(?:millones|miles|pesos|%|d[oó]lares)", str(ws.cell(r, etiqueta_col).value or ""), re.I
        )
    ]
    unidad = _unidad(unidades[-1]) if unidades else None
    formatos = _formatos_texto(ws, b["filas"], [letra for _, letra, _ in columnas])
    for local, row in enumerate(b["filas"]):
        saneada = datos[b["inicio"] + local]
        etiquetas = [
            str(v)
            for k, v in saneada.items()
            if k in b["columnas"]
            and column_index_from_string(b["columnas"][k]) < columnas[0][0]
            and v is not None
        ]
        if not etiquetas:
            continue
        label = etiquetas[0] + (f" · {panel}" if panel else "")
        # Sanear textos recuperados del original: nunca reintroducir datos personales.
        raws = [ws[f"{letra}{row}"].value for _, letra, _ in columnas]
        strings = [{"valor": v} for v in raws]
        saneados, _ = sanear(strings)
        raws = [_con_formato(x.get("valor"), formatos.get(letra)) for x, (_, letra, _) in zip(saneados, columnas)]
        nivel = (
            wf.cell(row, etiqueta_col).alignment.indent
            + len(str(wf.cell(row, etiqueta_col).value or ""))
            - len(str(wf.cell(row, etiqueta_col).value or "").lstrip())
        )
        salida.append(
            dict(
                archivo=archivo,
                hoja=b["hoja"],
                cuadro=b["bloque"],
                fila=row,
                encabezado=header,
                orientacion="horizontal",
                columna_etiqueta=etiqueta_col,
                etiqueta=label,
                # Un «Total» en otra columna de rótulos (bloques apilados por indicador): el motivo de los cierres lo dice.
                total_en_otra_columna=any(normalizar(e).split(" ")[0] in ("total", "subtotal") for e in etiquetas[1:]),
                ambito=contexto,
                nivel=nivel,
                unidad=_unidad(label) if normalizar(label).startswith("en ") else unidad,
                periodos=periodos,
                marcas=marcas if any(marcas) else None,
                panel=panel,
                periodos_validos=validos,
                celdas=[f"{letra}{row}" for _, letra, _ in columnas],
                valores=raws,
                formulas=[wf[f"{letra}{row}"].data_type == "f" for _, letra, _ in columnas],
                periodo_con_datos=[
                    any(numero(ws.cell(r, col).value) is not None for r in b["filas"] if r != row)
                    or sum(numero(v) is not None for v in raws) >= 4
                    for col, _, _ in columnas
                ],
            )
        )
    return salida, headers


FILA_RESUMEN = re.compile(r"^(?:total|subtotal|promedio|media|acumulado|suma|totales)\b")


def _vertical(ws, wf, b, datos, archivo):
    candidatas = [(c, letra) for c, letra in b["columnas"].items() if normalizar(c) in ("ano", "anio", "periodo", "year")]
    if len(candidatas) != 1:
        return [], []
    nombre, letra = candidatas[0]
    periodos = []
    locales = []
    for local, row in enumerate(b["filas"]):
        v = ws[f"{letra}{row}"].value
        # «Total», «Promedio» o «Acumulado» al pie de un cuadro con años en filas resumen los
        # períodos: no son un período (sus cierres los controla la coherencia interna).
        if isinstance(v, str) and FILA_RESUMEN.match(normalizar(v)):
            continue
        # «2025 (p)», «2025*», «2025 est.»: el año con su marca de provisorio sigue siendo el período.
        n = _anio_encabezado(v)
        periodos.append(int(n) if n is not None and n.is_integer() else v)
        locales.append(local)
    filas = [b["filas"][k] for k in locales]
    if sum(numero(v) is not None and 1800 <= numero(v) <= 2200 for v in periodos) < 3:
        return [], []
    # Si el año se repite en la mayoría de las filas (país × año, datos diarios con su año), la
    # columna clasifica filas: no es una serie anual y sus repeticiones no son años duplicados.
    anios = [numero(v) for v in periodos if numero(v) is not None]
    if len(set(anios)) * 2 <= len(anios):
        return [], []
    # Años con sus meses debajo en la misma columna: mezcla títulos de año (a veces vacíos) y meses; no se
    # arma una serie que los compare punto a punto (cada fila de año sería una «celda vacía»).
    formato = temporal.formato_oficial(periodos)[0]
    if formato == "anios_con_subperiodos":
        return [], []
    if formato == "bloques_por_frecuencia":
        # Bloques apilados de distinta frecuencia (años, trimestres, meses): cada bloque es su propia serie, con sus
        # períodos, su validación y sus controles de valores; un trimestre no se compara con un año.
        ss, errores = [], []
        for k, (desde, hasta) in enumerate(temporal.bloques_de_frecuencia(periodos), 1):
            parte, errores_parte = _series_verticales(ws, wf, b, datos, archivo, nombre, letra, periodos[desde:hasta],
                                                      filas[desde:hasta], locales[desde:hasta], f"{b['bloque']}.{k}")
            ss += parte
            errores += errores_parte
        return ss, errores
    return _series_verticales(ws, wf, b, datos, archivo, nombre, letra, periodos, filas, locales, b["bloque"])


def _series_verticales(ws, wf, b, datos, archivo, nombre, letra, periodos, filas, locales, cuadro):
    """Una serie por columna de valores de un cuadro con los períodos en la columna `letra` (o de un bloque de él)."""
    base = dict(archivo=archivo, hoja=b["hoja"], cuadro=cuadro, periodos=periodos, nivel=0, unidad=None)
    errors, validos = temporal.validar(
        {**base, "etiqueta": nombre, "valores": periodos, "celdas": [f"{letra}{r}" for r in filas]}
    )
    ss = []
    for c, letra_col in b["columnas"].items():
        if c == nombre:
            continue
        valores = [datos[b["inicio"] + i].get(c) for i in locales]
        # Valor crudo sólo dentro de columnas ya saneadas; recuperar NA sin perder su tipo.
        raws, _ = sanear([{"valor": ws[f"{letra_col}{r}"].value} for r in filas])
        formato = _formatos_texto(ws, filas, [letra_col]).get(letra_col)
        vs = [_con_formato(x.get("valor"), formato) for x in raws]
        if sum(numero(v) is not None for v in valores) < 3:
            continue
        ss.append(
            {
                **base,
                "etiqueta": c,
                "orientacion": "vertical",
                "columna": column_index_from_string(letra_col),
                "encabezado": max(b["encabezados"]),
                "valores": vs,
                "periodos_validos": validos,
                "celdas": [f"{letra_col}{r}" for r in filas],
                "formulas": [wf[f"{letra_col}{r}"].data_type == "f" for r in filas],
            }
        )
    # Un período que abarcan otras series del cuadro sin número en ninguna no tiene datos (la fuente no lo publicó,
    # como un año sin encuesta): un vacío ahí no es un faltante de la serie.
    for s, marcas in zip(ss, ejes.periodos_con_datos([[numero(v) is not None for v in s["valores"]] for s in ss])):
        s["periodo_con_datos"] = marcas
    return ss, errors


def _todas_series(leidos, datos, cobertura, libros, modelos):
    """Series anuales del lector histórico más las del modelo común (meses, trimestres, formato largo)."""
    ss, headers = _cuadros(leidos, datos, cobertura, libros, modelos)
    cubiertos = {(s["archivo"], s["hoja"], s["cuadro"]) for s in ss}
    cubiertos |= {(h["evidencia"]["archivo_origen"], h["evidencia"]["hoja"], h["evidencia"]["bloque_excel"]) for h in headers}
    from services.estructura_revision import contexto_cuadro
    for clave, c in modelos.items():
        from services.interpretacion_previa import permite
        if not permite("series", c["archivo"], c["hoja"], c["cuadro"]):
            continue
        if c.get("eje_confirmado") == "serie_filas" and c.get("eje") is None:
            mixtas, avisos = _series_mixtas(c)
            ss.extend(mixtas)
            headers.extend(avisos)
            if mixtas:
                continue
        if clave in cubiertos or c["eje"] not in ("columnas", "filas", "largo"):
            continue
        if c["eje"] == "columnas" and c["frecuencia"] == "anual" and not c.get("paneles"):
            # Años en el encabezado que el lector histórico no tomó (rótulos como «Ejercicio 2025»).
            if sum(1 for x in c["columnas"] if x.get("periodo")) < 3:
                continue
        acumuladas, totales = _columnas_derivadas(c)
        nuevas = ejes.series(c, acumuladas, totales)
        b = next((b for b in cobertura["bloques"] if (b["archivo"], b["hoja"], b["bloque"]) == clave), None)
        ambito = None
        if b is not None:
            try:
                ambito = contexto_cuadro(_libro(libros, c["archivo"]).valores[c["hoja"]], b)
            except Exception:  # noqa: BLE001
                ambito = None
        for s in nuevas:
            s["ambito"] = ambito
            s["formulas_desconocidas"] = bool(c.get("formulas_desconocidas"))
        ss.extend(nuevas)
        headers.extend(ejes.validar_cuadro(c))
    return ss, headers


def _series_mixtas(c):
    """Separa años y meses bajo cada año cuando la persona confirmó períodos en filas."""
    tiempo = [x for x in c["columnas"] if x.get("rol_confirmado") == "periodo_serie"]
    if len(tiempo) != 1:
        return [], []
    idx = tiempo[0]["idx"]
    anual, mensual = {}, {}
    anio = None
    for r in c["filas"]:
        p = P.leer(c["crudos"].get((r, idx)))
        if p is None:
            continue
        if p["frecuencia"] == "anual":
            anio = p["anio"]
            p["ordinal"] = P.ordinal(p)
            anual[r] = p
        elif p["frecuencia"] == "mensual" and (p.get("anio") or anio):
            p["anio"] = p.get("anio") or anio
            p["ordinal"] = P.ordinal(p)
            mensual[r] = p
    if len(anual) < 3 or len(mensual) < 3:
        return _series_con_anio_en_columnas(c, idx)
    series, errores = [], []
    for frecuencia, periodos in (("anual", anual), ("mensual", mensual)):
        sub = copy(c)
        sub.update(eje="filas", frecuencia=frecuencia, col_tiempo=idx,
                   filas=list(periodos), periodos_fila=periodos, claves=[])
        series.extend(ejes.series(sub))
        errores.extend(ejes.validar_cuadro(sub))
    return series, errores


def _series_con_anio_en_columnas(c, tiempo):
    """Mes en una columna y ejercicios repetidos por variable en otras columnas."""
    meses = {}
    for r in c["filas"]:
        p = P.leer(c["crudos"].get((r, tiempo)))
        if p and p["frecuencia"] == "mensual" and p.get("anio") is None:
            meses[r] = p
    if len(meses) < 3:
        return [], []
    series, errores, anios_validados = [], [], set()
    for x in c["columnas"]:
        if not x["numerica"] or x.get("rol_confirmado") not in {"importe", "cantidad", "indice"}:
            continue
        m = re.search(r"(?<!\d)((?:18|19|20|21)\d{2})(?!\d)", str(x.get("encabezado") or x["nombre"]))
        if not m:
            continue
        anio = int(m.group(1))
        periodos = {r: {**p, "anio": anio, "ordinal": anio * 12 + p["sub"] - 1} for r, p in meses.items()}
        sub = copy(c)
        sub.update(eje="filas", frecuencia="mensual", col_tiempo=tiempo,
                   filas=list(periodos), periodos_fila=periodos, claves=[],
                   columnas=[next(t for t in c["columnas"] if t["idx"] == tiempo), x])
        series.extend(ejes.series(sub))
        if anio not in anios_validados:
            errores.extend(ejes.validar_cuadro(sub))
            anios_validados.add(anio)
    return series, errores


def _libro(libros, archivo):
    return next(lb for a, lb in libros if a == archivo)


def _columnas_derivadas(c):
    """Columnas que son acumulado de otra (se reinicia al cambiar de año) o total de las demás."""
    if c["eje"] not in ("filas", "largo"):
        return frozenset(), frozenset()
    filas = sorted(c["periodos_fila"], key=lambda r: c["periodos_fila"][r]["ordinal"])
    cols = [x for x in c["columnas"] if x["numerica"] and not x.get("tiempo")
            and x.get("rol_confirmado") not in {"porcentaje", "indice"}]
    acumuladas, totales = set(), set()
    for x in cols:
        for y in cols:
            if x is y:
                continue
            buenas = total = 0
            for k, r in enumerate(filas):
                vx, vy = c["valores"][(r, x["idx"])], c["valores"][(r, y["idx"])]
                if vx is None or vy is None:
                    continue
                p = c["periodos_fila"][r]
                previo = filas[k - 1] if k else None
                nuevo_anio = previo is None or c["periodos_fila"][previo]["anio"] != p["anio"]
                base = 0 if nuevo_anio else c["valores"][(previo, x["idx"])]
                if base is None:
                    continue
                total += 1
                buenas += abs(vx - (base + vy)) <= 0.001 * max(abs(vx), 1)
            if total >= 4 and buenas >= 0.75 * total:
                acumuladas.add(x["idx"])
    for x in cols:
        if x["idx"] in acumuladas:
            continue
        otras = [y for y in cols if y is not x and y["idx"] not in acumuladas]
        if len(otras) < 2:
            continue
        buenas = total = 0
        for r in filas:
            vx = c["valores"][(r, x["idx"])]
            partes = [c["valores"][(r, y["idx"])] for y in otras]
            if vx is None or any(v is None for v in partes):
                continue
            total += 1
            buenas += abs(vx - sum(partes)) <= 0.001 * max(abs(vx), 1)
        if total >= 4 and buenas >= 0.75 * total:
            totales.add(x["idx"])
    return frozenset(acumuladas), frozenset(totales)


def cobertura_lectura(leidos, datos, cobertura):
    """Proporción del archivo que pudo interpretarse como series por período."""
    libros = _libros(leidos)
    try:
        ss, _ = _todas_series(leidos, datos, cobertura, libros, _modelos(libros, cobertura))
    finally:
        _cerrar(libros)
    return lectura.resumen(leidos, ss)


SIN_DATO_CONVENCION = 2  # series distintas de la hoja con la misma marca para tomarla como convención del cuadro
SIN_DATO_CONVENCION_SERIE = 3  # o celdas con la marca dentro de una misma serie («///» en el primer año de una variación)


def _sin_dato_por_convencion(ss):
    """Marcas de dato no disponible («s/i», «s/d», «///», «...») que la hoja usa en dos o más
    series, o tres o más veces en una misma serie.

    Usada así, la marca es la convención del cuadro y la lectura ya la declara como dato no
    disponible: no es un cambio de tipo que haya que revisar. Una o dos marcas en una sola serie
    siguen siendo señal: pueden ser datos faltantes en esa serie."""
    from services.perfilado import SIN_DATO

    celdas = {}
    for s in ss:
        for v in s.get("valores") or []:
            if isinstance(v, str) and v.strip().lower() in SIN_DATO:
                celdas.setdefault((s["archivo"], s["hoja"], v.strip().lower()), []).append(id(s))
    por_hoja = {}
    for (archivo, hoja, marca), usos in celdas.items():
        if len(set(usos)) >= SIN_DATO_CONVENCION or len(usos) >= SIN_DATO_CONVENCION_SERIE:
            por_hoja.setdefault((archivo, hoja), set()).add(marca)
    for s in ss:
        s["sin_dato_convencion"] = tuple(sorted(por_hoja.get((s["archivo"], s["hoja"]), ())))


def analizar(leidos, datos, cobertura):
    from services import interpretacion_previa as ip
    from services.estructura_revision import enriquecer, preparar, totales_separados
    def elegidas(series, control):
        return [s for s in series if ip.permite(control, s["archivo"], s["hoja"], s["cuadro"])]
    def relaciones_compatibles(series):
        if ip.actual() is None:
            return [series]
        grupos = {}
        for s in series:
            u = ip.unidad(s["archivo"], s["hoja"], s["cuadro"])
            i = u.get("confirmada", {})
            firma = (i.get("tipo"), i.get("institucion"), i.get("regimen"))
            grupos.setdefault(firma, []).append(s)
        return list(grupos.values())
    libros = _libros(leidos)
    try:
        modelos = _modelos(libros, cobertura)
        ss, headers = _todas_series(leidos, datos, cobertura, libros, modelos)
        ss = elegidas(ss, "tabular")
        series_temporales = elegidas(ss, "series")
        cobertura["estructura_revision"] = preparar(ss)
        motor_series.anotar_volatilidad_cuadro(series_temporales)
        _sin_dato_por_convencion(ss)
        crudos = [h for s in series_temporales if not s.get("acumulada") for h in motor_series.analizar(s)]
        _agrupar_totales(series_temporales, crudos)
        en_matriz, escalas = _matrices(series_temporales, crudos)
        celdas_escala = {u for h in escalas for u in h["evidencia"]["celdas_participantes"]}
        crudos = [h for h in crudos if h["evidencia"]["ubicaciones_excel"][0] not in celdas_escala]
        puntos = motor_series.contexto_compartido(series_temporales, crudos) + en_matriz + escalas
        relaciones = []
        # La revisión rápida no cruza ediciones: el informe lo declara al comienzo.
        rapida = ip.es_rapida(ip.actual())
        for compatibles in relaciones_compatibles(elegidas(ss, "relaciones")):
            relaciones += (conciliacion.analizar(compatibles) + per_capita.analizar(compatibles)
                           + compensados.analizar(compatibles) + ([] if rapida else ediciones.analizar(compatibles))
                           + identidades.analizar(compatibles) + entre_hojas.analizar(compatibles))
        propios = []
        cuadros = _unir_fragmentos(modelos)
        for c in cuadros.values():
            propios += transversal.analizar(c, _libro(libros, c["archivo"]))
            propios += totales_cruzados.analizar(c)
            extra = c.get("cierres_columna") or []
            if extra:
                er = cobertura["estructura_revision"]
                er["cierres_verificados"] = er["cierres_verificados"] + extra
                er["cierres_columna"] = int(er.get("cierres_columna") or 0) + len(extra)
            aparentes = c.get("totales_aparentes_sin_comprobar") or []
            if aparentes:
                # Un bloque final que parece el total y no suma los bloques de ninguna forma: no se calla.
                er = cobertura["estructura_revision"]
                er["totales_aparentes_sin_comprobar"] = (er.get("totales_aparentes_sin_comprobar") or []) + aparentes
        er = cobertura["estructura_revision"]
        # Celdas que la conciliación por series ya señala: la jerarquía por código no las repite.
        ya_conciliadas = {u for h in relaciones for u in (h["evidencia"].get("ubicaciones_excel") or [])[:1]}
        for c in cuadros.values():
            # Jerarquía por código (1 / 1.1 / 1.1.1): cada padre es el total de sus hijos inmediatos.
            nuevos, cierres_codigo, sin_comprobar = jerarquia_codigos.analizar(c, ya_conciliadas)
            relaciones += nuevos
            contados = {x.get("total") for x in er["cierres_verificados"]}
            cierres_codigo = [x for x in cierres_codigo if x["total"] not in contados]
            if cierres_codigo:
                er["cierres_verificados"] = er["cierres_verificados"] + cierres_codigo
                er["cantidad_cierres"] = int(er.get("cantidad_cierres") or 0) + len(cierres_codigo)
                er["motivo_sin_cierres"] = None
            if sin_comprobar:
                er["jerarquia_sin_comprobar"] = (er.get("jerarquia_sin_comprobar") or []) + sin_comprobar
        separados = totales_separados(cobertura.get("bloques") or [], ss, er.get("cierres_verificados") or [])
        if separados:
            er["totales_separados"] = separados
        for archivo, libro in libros:
            hojas = None if ip.actual() is None else {ws.title for ws in libro.formulas.worksheets
                if ip.permite_hoja("formulas", archivo, ws.title)}
            propios += formulas_fijas.analizar(archivo, libro, hojas=hojas)
        propios += _unidades_encabezado(cobertura, modelos)
        from . import rubro_escala

        propios += rubro_escala.analizar(series_temporales)
        # Una fila total sin rótulo que ya es el total de una jerarquía (sangría o código) se controla
        # por la conciliación, con su regla y sus componentes: la señal por cifras no la reemplaza.
        conciliadas = {u for h in relaciones for u in (h["evidencia"].get("ubicaciones_excel") or [])[:1]}
        propios = [h for h in propios if not (h["evidencia"].get("regla") == "total_sin_rotulo_no_cierra"
                                              and (h["evidencia"].get("ubicaciones_excel") or [None])[0] in conciliadas)]
        salida = consolidar(headers + puntos + _sin_repetir(propios, headers + puntos), relaciones)
        _agrupar_misma_celda(salida)
        _contexto_declarado(salida, modelos)
        _cambio_unidad_declarado(salida, modelos)
        _cierre_exacto(salida, er.get("cierres_verificados") or [], _hojas_cronograma(libros))
        return enriquecer(salida, ss)
    finally:
        _cerrar(libros)


def _escala(texto):
    """Escala declarada en un rótulo de unidad: la magnitud, no el objeto medido."""
    t = normalizar(texto)
    if "%" in t or "porcentaje" in t:
        return "porcentaje"
    if "miles de millones" in t:
        return "miles de millones"
    if re.search(r"\bmillon(?:es)?\b", t):
        return "millones"
    if re.search(r"\bmiles\b", t):
        return "miles"
    return None


def _unidades_encabezado(cobertura, modelos):
    """Una columna cuyo rótulo declara una unidad distinta de la del resto del cuadro."""
    salida = []
    for b in cobertura.get("bloques") or []:
        # Comparar unidades exige el mismo indicador a través de períodos.
        # En tablas verticales, «Ventas» y «Variación» son variables distintas.
        anotadas = {c: u for c, u in (b.get("unidades_columna") or {}).items() if P.leer(c)}
        if not anotadas or not b.get("encabezados"):
            continue
        header = max(b["encabezados"])
        escalas = {c: _escala(u["unidad"]) for c, u in anotadas.items()}
        modelo = modelos.get((b["archivo"], b["hoja"], b["bloque"])) or {}
        del_cuadro = _escala(modelo.get("unidad") or "")
        for c, u in anotadas.items():
            propia = escalas[c]
            if not propia:
                continue
            otras = [e for k, e in escalas.items() if k != c and e]
            referencia_cuadro = None
            if otras and len(set(otras)) == 1 and otras[0] != propia:
                referencia_cuadro = f"las demás columnas declaran {otras[0]}"
            elif not otras and del_cuadro and del_cuadro != propia:
                referencia_cuadro = f"el cuadro declara {del_cuadro} («{(modelo.get('unidad') or '')[:80]}»)"
            if not referencia_cuadro:
                continue
            celda = f"{u['letra']}{header}"
            ref = f"{b['hoja']}!{celda}"
            accion = "Confirmar la unidad de esa columna antes de comparar sus valores con los demás períodos."
            salida.append(dict(
                lente="ajedrecista",
                modo="control_tabular",
                severidad="media",
                descripcion=(f"{ref} · {u['rotulo']}: la columna declara su unidad en el encabezado ({u['unidad']}) "
                             f"y {referencia_cuadro}. Es una señal de unidad mezclada en esa columna. {accion}"),
                pregunta_institucional=accion,
                metodo="Control tabular reproducible; unidad leída del rótulo de la columna",
                evidencia=dict(
                    archivo_origen=b["archivo"], hoja=b["hoja"], bloque_excel=b["bloque"], celda=celda,
                    indicador=c, periodo=None, valor_observado=u["rotulo"], regla="unidad_columna_distinta",
                    naturaleza="control", puntaje_prioridad=60,
                    factores_prioridad=[{"criterio": "unidad_columna_distinta", "puntos": 60}],
                    ubicaciones_excel=[ref], celdas_participantes=[ref], accion_revision=accion,
                    unidad_columna=u["unidad"],
                ),
            ))
    return salida


def _unir_fragmentos(modelos):
    """Un estado contable partido en secciones («ACTIVO CORRIENTE», «PASIVO…») es un solo cuadro:
    los cortes transversales de la misma hoja con las mismas columnas se controlan juntos."""
    grupos = {}
    for clave, c in modelos.items():
        if c["eje"] is not None:
            grupos[clave] = [c]
            continue
        from services import interpretacion_previa as ip
        u = ip.unidad(c["archivo"], c["hoja"], c["cuadro"])
        firma_semantica = ip.huella(u.get("confirmada")) if u is not None else None
        firma = (c["archivo"], c["hoja"], tuple(x["letra"] for x in c["columnas"] if x["numerica"]), firma_semantica)
        grupos.setdefault(firma, []).append(c)
    salida = {}
    for clave, cs in grupos.items():
        if len(cs) == 1:
            salida[(cs[0]["archivo"], cs[0]["hoja"], cs[0]["cuadro"])] = cs[0]
            continue
        base = modelo_cuadro.Cuadro(cs[0])
        for k in ("valores", "crudos", "formulas", "formatos_celda", "etiquetas", "codigos", "sangria"):
            base[k] = {kk: vv for c in cs for kk, vv in c.get(k, {}).items()}
        base["filas"] = sorted({r for c in cs for r in c["filas"]})
        base["cuadro"] = cs[0]["cuadro"]
        salida[(base["archivo"], base["hoja"], base["cuadro"])] = base
    return salida


def _matrices(ss, crudos):
    """Valores aislados y cambios de escala comunes, cuadro por cuadro."""
    senaladas = {h["evidencia"]["ubicaciones_excel"][0] for h in crudos if h["evidencia"]["naturaleza"] == "control"}
    grupos = {}
    for s in ss:
        grupos.setdefault((s["archivo"], s["hoja"], s["cuadro"], s.get("orientacion"), s.get("panel")), []).append(s)
    aislados, escalas = [], []
    for series_cuadro in grupos.values():
        a, e = matriz.analizar(series_cuadro, senaladas)
        aislados += a
        escalas += e
    return aislados, escalas


def _agrupar_totales(ss, crudos):
    """El salto de un total que acompaña al de uno de sus componentes es la misma causa."""
    from .modelo import es_total
    por_serie = {(s["archivo"], s["hoja"], s["cuadro"], s["etiqueta"]): s for s in ss}
    picos = [h for h in crudos if h["evidencia"]["regla"] in ("pico_aislado", "escala_extremo", "signo_aislado")]
    for h in picos:
        e = h["evidencia"]
        s = por_serie.get((e["archivo_origen"], e["hoja"], e["bloque_excel"], e["indicador"]))
        if not s or not (s.get("total") or es_total(s["etiqueta"])):
            continue
        for otro in picos:
            o = otro["evidencia"]
            so = por_serie.get((o["archivo_origen"], o["hoja"], o["bloque_excel"], o["indicador"]))
            if otro is h or not so or so.get("total") or es_total(so["etiqueta"]):
                continue
            if (o["archivo_origen"], o["hoja"], o["bloque_excel"], str(o["periodo"])) != (e["archivo_origen"], e["hoja"], e["bloque_excel"], str(e["periodo"])):
                continue
            grupo = o.get("grupo_relacion") or f"celda:{o['ubicaciones_excel'][0]}"
            o["grupo_relacion"] = e["grupo_relacion"] = grupo
            e["total_acompana_a"] = o["ubicaciones_excel"][0]
            break


def _sin_repetir(propios, previos):
    """Una misma regla sobre la misma celda se informa una vez."""
    vistos = {(h["evidencia"]["hoja"], h["evidencia"]["celda"], h["evidencia"]["regla"]) for h in previos}
    out = []
    for h in propios:
        k = (h["evidencia"]["hoja"], h["evidencia"]["celda"], h["evidencia"]["regla"])
        if k not in vistos:
            vistos.add(k)
            out.append(h)
    return out


def _agrupar_misma_celda(hallazgos):
    """Señales de reglas distintas sobre la misma celda son una misma causa: un único caso."""
    por_celda = {}
    for h in hallazgos:
        e = h["evidencia"]
        if e.get("naturaleza") != "control" or not e.get("celda"):
            continue
        por_celda.setdefault((e["archivo_origen"], e["hoja"], e["celda"]), []).append(h)
    for (archivo, hoja, celda), hs in por_celda.items():
        if len(hs) < 2:
            continue
        grupo = next((h["evidencia"]["grupo_relacion"] for h in hs if h["evidencia"].get("grupo_relacion")), f"celda:{hoja}!{celda}")
        periodo = next((h["evidencia"].get("periodo") for h in hs if h["evidencia"].get("periodo") is not None), None)
        for h in hs:
            h["evidencia"]["grupo_relacion"] = grupo
            if h["evidencia"].get("periodo") is None:
                h["evidencia"]["periodo"] = periodo


MESES_NOTA = {1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio", 7: "julio", 8: "agosto",
              9: "septiembre|setiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"}


def _contexto_declarado(hallazgos, modelos):
    """Una nota del propio cuadro que documenta el movimiento lo vuelve contexto declarado.

    La nota debe nombrar el rubro (por su rótulo o su código) y el período de la celda; así una
    desafectación documentada «en marzo, el programa 22» no se lee como un signo erróneo.
    """
    import re as _re
    por_hoja = {}
    for c in modelos.values():
        por_hoja.setdefault((c["archivo"], c["hoja"]), []).append(c)
    for h in hallazgos:
        e = h["evidencia"]
        if e.get("naturaleza") != "control" or not e.get("celda"):
            continue
        cuadros = por_hoja.get((e.get("archivo_origen"), e.get("hoja"))) or []
        notas = [n for c in cuadros for n in c.get("notas") or []]
        if not notas:
            continue
        m = _re.match(r"^([A-Z]+)(\d+)$", e["celda"])
        if not m:
            continue
        fila = int(m.group(2))
        col = column_index_from_string(m.group(1))
        cuadro = next((c for c in cuadros if fila in c["filas"]), None)
        if not cuadro:
            continue
        rotulo = normalizar(cuadro["etiquetas"].get(fila) or "")
        codigo = cuadro["codigos"].get(fila)
        x = next((x for x in cuadro["columnas"] if x["idx"] == col), {})
        p = x.get("periodo") or cuadro["periodos_fila"].get(fila)
        for n in notas:
            t = normalizar(n["texto"])
            fila_ok = (len(rotulo) >= 6 and rotulo in t) or (
                codigo and _re.search(rf"\b(?:programa|prog\.?|codigo|cod\.?|partida|rubro)\s*{codigo}\b", t))
            if not fila_ok:
                continue
            # Sólo movimientos de una celda en un período: la nota debe nombrar ese período.
            if not p:
                continue
            if p["frecuencia"] == "mensual":
                periodo_ok = bool(_re.search(rf"\b(?:{MESES_NOTA[p['sub']]})\b", t))
            else:
                periodo_ok = bool(p.get("anio") and not p.get("anio_inferido") and str(p["anio"]) in t)
            if not periodo_ok:
                continue
            e.update(naturaleza="contexto", contexto_declarado={"celda": n["celda"], "texto": n["texto"]},
                     puntaje_prioridad=20)
            e.setdefault("factores_prioridad", []).append({"criterio": "contexto_declarado_en_nota", "puntos": 0})
            h["severidad"] = "media"
            h["descripcion"] += f" Contexto declarado en la nota {n['celda']}: «{n['texto'][:200]}»."
            break


REGLAS_ESCALA = {"cambio_escala_sostenido", "periodo_en_otra_escala", "salto_escala_final"}
TOLERANCIA_FACTOR_LOG = 0.5  # el salto observado mezcla el cambio de unidad con la evolución del año


def _cambio_unidad_declarado(hallazgos, modelos):
    """Un cambio de escala en el año en que el propio cuadro declara que cambia la unidad
    («1980-1986: miles de pesos; Desde 1987: millones de pesos») no es un error: queda como
    contexto, con la nota que lo explica. La señal no se borra; deja de ser una alarma.

    La nota tiene que ser del mismo cuadro (no de otro de la misma hoja) y explicar el salto: el
    sentido (de miles a millones los valores bajan) y el factor (×1.000 entre miles y millones) deben
    coincidir. Un salto de ×10, o en el sentido contrario, sigue siendo una señal.
    """
    import math

    por_cuadro = {(c["archivo"], c["hoja"], c["cuadro"]): c.get("cambios_unidad") or [] for c in modelos.values()}
    for h in hallazgos:
        e = h["evidencia"]
        if e.get("regla") not in REGLAS_ESCALA or e.get("naturaleza") == "contexto":
            continue
        cambios = por_cuadro.get((e.get("archivo_origen"), e.get("hoja"), e.get("bloque_excel"))) or []
        if not cambios or not e.get("factor_escala") or not e.get("sentido"):
            continue
        m = re.match(r"^\s*((?:18|19|20)\d{2})", str(e.get("periodo") if e.get("periodo") is not None else ""))
        if not m:
            continue
        anio = int(m.group(1))
        observado = math.log10(e["factor_escala"]) * (1 if e["sentido"] > 0 else -1)
        nota = None
        for x in cambios:
            antes, despues = modelo_cuadro._escala_de(x["antes"]), modelo_cuadro._escala_de(x["despues"])
            if x["anio"] != anio or not antes or not despues or antes == despues:
                continue
            esperado = math.log10(antes / despues)  # de miles a millones: -3 (los valores bajan)
            if abs(observado - esperado) <= TOLERANCIA_FACTOR_LOG:
                nota = x
                break
        if not nota:
            continue
        e.update(naturaleza="contexto", contexto_declarado={"celda": nota["celda"], "texto": nota["texto"]},
                 cambio_unidad_declarado=nota["anio"], puntaje_prioridad=20)
        e.setdefault("factores_prioridad", []).append({"criterio": "cambio_de_unidad_declarado", "puntos": 0})
        h["severidad"] = "media"
        h["descripcion"] += (f" El cuadro declara el cambio de unidad en {nota['anio']}, de {nota['antes']} a "
                             f"{nota['despues']} ({nota['celda']}: «{nota['texto'][:160]}»), y el salto coincide en "
                             "sentido y factor: queda como contexto; para comparar períodos hay que llevar los valores a "
                             "una misma unidad.")


# Señales de patrón de una serie (pico, salto, valor repetido, signo) y controles aritméticos que no cierran.
PATRON_SERIE = {"pico_aislado", "salto_escala_final", "escala_extremo", "valor_repetido", "signo_aislado",
                "valor_identico_otra_fila"}
ARITMETICAS = {"conciliacion_aritmetica", "total_columna_no_cierra", "subtotal_columna_no_cierra",
               "total_bloques_no_cierra", "total_sin_rotulo_no_cierra"}
CIERRE_EXACTO = 1e-6  # diferencia del cierre, relativa al valor de la celda
# Título de la hoja que declara un cronograma de pagos: los montos van y vienen entre períodos por construcción.
CRONOGRAMA = re.compile(r"perfil\s+(?:anual\s+|mensual\s+|trimestral\s+)?de\s+vencimientos"
                        r"|cronograma\s+de\s+(?:pagos|vencimientos|amortizacion|servicios)")


def _hojas_cronograma(libros, filas=12):
    """(archivo, hoja) cuyas primeras filas declaran un perfil o cronograma de vencimientos."""
    out = set()
    for archivo, libro in libros:
        for ws in libro.valores.worksheets:
            texto = " ".join(v for fila in ws.iter_rows(min_row=1, max_row=filas, values_only=True)
                             for v in fila if isinstance(v, str))
            if CRONOGRAMA.search(normalizar(texto)):
                out.add((archivo, ws.title))
    return out


def _cierre_exacto(hallazgos, cierres, cronogramas):
    """En un perfil de vencimientos, un pico, salto o valor repetido en una celda que cierra exacto con su total
    y sus componentes, y que no participa de ninguna suma que no cierra, es coherente con el cuadro: los pagos van
    y vienen entre períodos, y un error de escala o de dígito en esa celda rompería la suma. Se anota el cierre.
    Fuera de un cronograma no: un error de carga que se propaga a subtotales y total recalculados también cierra."""
    por_celda = {}
    for c in cierres:
        for ref in [c.get("total"), *(c.get("componentes") or [])]:
            por_celda.setdefault((c.get("archivo"), ref), []).append(c)
    rotas = {(h["evidencia"].get("archivo_origen"), u) for h in hallazgos
             if h["evidencia"].get("regla") in ARITMETICAS for u in h["evidencia"].get("celdas_participantes") or []}
    for h in hallazgos:
        e = h["evidencia"]
        if e.get("regla") not in PATRON_SERIE or (e.get("archivo_origen"), e.get("hoja")) not in cronogramas:
            continue
        clave = (e.get("archivo_origen"), f"{e.get('hoja')}!{e.get('celda')}")
        valor = numero(e.get("valor_observado"))
        if (valor is None or clave in rotas
                or any(x.get("regla") in ARITMETICAS for x in e.get("conciliaciones") or [])):
            continue
        exactos = [c for c in por_celda.get(clave, [])
                   if abs(c.get("diferencia") or 0) <= CIERRE_EXACTO * max(abs(valor), 1e-9)]
        if e.get("regla") == "valor_identico_otra_fila":
            # Idéntico a otra fila porque una es el total y la otra su único componente distinto de cero
            # («CORTO PLAZO» = «En moneda nacional» cuando no hay deuda en moneda extranjera).
            otra = e.get("celda_igual")
            exactos = [c for c in exactos if otra == c.get("total") or otra in (c.get("componentes") or [])]
        if exactos:
            e["cierre_exacto"] = {"total": exactos[0]["total"], "componentes": exactos[0]["componentes"]}


def consolidar(puntos, relaciones):
    # Una conciliación corrobora una celda ya señalada, sin repetirla como otra alarma.
    for h in sorted(relaciones, key=lambda x: len(x["evidencia"]["celdas_participantes"])):
        ev = h["evidencia"]
        relacionados = [
            p
            for p in puntos
            if p["evidencia"]["archivo_origen"] == ev["archivo_origen"]
            and p["evidencia"]["naturaleza"] == "control"
            and p["evidencia"]["ubicaciones_excel"][0] in ev["celdas_participantes"]
        ]
        if relacionados:
            for p in relacionados:
                e = p["evidencia"]
                independiente = ev.get("regla") in {"conciliacion_aritmetica", "poblacion_implicita"}
                if independiente and not e.get("corroboracion_independiente"):
                    e["corroboracion_independiente"] = True
                    e["puntaje_prioridad"] += 20
                    e["factores_prioridad"].append({"criterio": "conciliacion_independiente", "puntos": 20})
                e.setdefault("conciliaciones", []).append(ev)
                if independiente:
                    p["severidad"] = "alta"
        else:
            # Evitar repetir una igualdad en sentidos opuestos.
            clave = (ev["archivo_origen"], tuple(sorted(ev["celdas_participantes"])))
            anterior = next(
                (
                    p
                    for p in puntos
                    if (p["evidencia"]["archivo_origen"], tuple(sorted(p["evidencia"]["celdas_participantes"])))
                    == clave
                ),
                None,
            )
            if not anterior:
                puntos.append(h)
    puntos = [
        p for p in puntos if not p["evidencia"].get("requiere_conciliacion") or p["evidencia"].get("conciliaciones")
    ]
    return sorted(
        puntos,
        key=lambda h: (
            h["evidencia"]["naturaleza"] != "control",
            -h["evidencia"]["puntaje_prioridad"],
            h["evidencia"]["hoja"],
            h["evidencia"]["celda"],
        ),
    )
