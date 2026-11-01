"""Controles de cuadros sin eje temporal (por provincia, por rubro, por meta) y de atributos.

Se aplican sobre columnas que no son períodos: porcentajes que deben sumar 100, partes que
deben sumar el total de la columna, relaciones de orden que el rótulo declara (indigencia ≤
pobreza), identidades contables reconocidas por sus rótulos (activo = pasivo + patrimonio neto;
resultado bruto = ventas + costo), signos que rompen la convención del propio cuadro, valores
fuera de la escala de su columna y valores idénticos copiados entre filas o entre ejercicios.
Cada regla exige que la estructura se confirme en el propio cuadro antes de afirmar un desvío.
"""

import math
import re
from fractions import Fraction
from statistics import median

from openpyxl.utils import get_column_letter

from . import periodos as P
from .comun import es_numero, mad, normalizar
from .modelo import es_total

NO_SUMABLE = re.compile(
    r"%|porcentaj|tasa|promedio|medi[oa]\b|mediana|[ií]ndice|ratio|variaci|\bvar\.|\bvar\b|crecim|participaci|desv[ií]o|"
    r"precio|unitari|per\s*c[aá]pita|por\s+(?:niñ|habitante|persona|hogar)|ejecuci[oó]n|cumplimiento|avance|"
    r"(?:\(|\b)veces\b|liquidez|endeudamiento|solvencia|rentabilidad|"
    # Límites del intervalo de confianza («LI», «LS») y coeficiente de variación: acompañan a una estimación, no se suman.
    r"\bl[is]\b|l[ií]mite|intervalo|\bc\.?\s?v\b|coeficiente de variaci|error est[aá]ndar|"
    r"distribuci[oó]n\s+(?:relativa|porcentual)",
    re.I,
)
# Columnas de variación o diferencia («Var. ia.», «Variación %», «Diferencia»): sus valores cruzan el cero y cambian de orden de
# magnitud por naturaleza; no tienen una escala propia que controlar. Un porcentaje o una tasa de nivel
# («Desocupación (%)») sí la tiene: se excluye sólo si la columna cruza el cero (es una variación).
SIN_ESCALA = re.compile(r"variaci|\bvar\b|\bdiferencias?\b", re.I)
PORCENTAJE_O_TASA = re.compile(r"%|tasa|porcentaj", re.I)
FILA_PORCENTUAL = re.compile(r"%|porcentaj", re.I)
# Separación respecto del valor más cercano por debajo de la cual, sin un factor cercano a una potencia
# de diez, un valor fuera de escala queda como señal media.
SEPARACION_CLARA = 10
CERCA_POTENCIA_DIEZ = 0.35
ADMITE_SIGNO = re.compile(
    r"variaci|diferencia|saldo|resultado|excedente|super[aá]vit|d[eé]ficit|neto|otros|ajuste|desaf|reversi|"
    r"ganancia|p[eé]rdida|balance|margen|tenencia|recpam|financier",
    re.I,
)
EGRESO = re.compile(r"\bcosto|\bgasto|egreso|amortizaci|depreciaci|devoluci|reintegro|previsi[oó]n(?:es)? del", re.I)
ORDEN = [
    (re.compile(r"indigencia", re.I), re.compile(r"pobreza", re.I), "indigencia_supera_pobreza",
     "La indigencia es un subconjunto de la pobreza: su tasa no puede superar a la de pobreza."),
    (re.compile(r"desocupa", re.I), re.compile(r"actividad", re.I), "desocupacion_supera_actividad",
     "La tasa de desocupación se mide sobre la población activa y no puede superar a la tasa de actividad."),
]


def _sig(v):
    return len(re.sub(r"[^0-9]", "", f"{abs(v):.10g}").strip("0")) or 1


# Promedios habituales en las series oficiales: de semestres, trimestres, cuatrimestres, bimestres y meses.
DIVISORES_PROMEDIO = (1, 2, 3, 4, 6, 12)


def _promedio_de(v):
    """(k, d) si el valor, escrito con más de cuatro decimales, es el promedio de k valores con d decimales o menos
    (73,9666… = 221,9 / 3; 0,09075 = 0,363 / 4) o la misma cifra con ruido de punto flotante (k = 1); None si no
    (una cifra calculada, sin una resolución de origen que se pueda leer)."""
    t = repr(float(v))
    if not v or "e" in t or "." not in t or len(t.split(".")[1].rstrip("0")) <= 4:
        return None

    def encaja(p, d):
        return abs(p - round(p, d)) <= 8 * math.ulp(p)  # pocas ulp: sólo el redondeo de la cuenta

    # La misma cifra con ruido de punto flotante (100,00000000000001; 22027,600000000002).
    for d in range(4):
        if encaja(float(v), d):
            return 1, d
    for k in DIVISORES_PROMEDIO[1:]:
        for d in range(4):
            if encaja(float(v) * k, d):
                # Un promedio de enteros no es una resolución gruesa (sumar uno a un entero lo deja entero). Límite: un
                # promedio de 6 o de 12 enteros encaja antes con k = 3 (7/6 = 3,5 / 3) y se lee con uno o dos decimales;
                # mirarlo con todos los k confundiría promedios de tasas (0,25 / 3 = 1/12) con promedios de enteros.
                return (k, d) if d else None
    return None


def decimales_efectivos(v):
    """Decimales con que está escrito el valor (como mucho 4: más allá es ruido de punto flotante o una cifra
    calculada). Un promedio de valores escritos con pocos decimales tiene la resolución de sus sumandos."""
    t = repr(float(v))
    if "e" in t or "." not in t:
        return 0
    promedio = _promedio_de(v)
    return promedio[1] if promedio else min(4, len(t.split(".")[1].rstrip("0")))


def _sig_efectivas(v):
    """Cifras significativas con la resolución efectiva: el promedio de tres valores con un decimal (73,9666…) tiene
    las cifras de su suma (221,9), no diez."""
    promedio = _promedio_de(v)
    if promedio is None:
        return _sig(v)
    k, d = promedio
    return _sig(round(float(v) * k, d))


def periodo_de(x):
    p = x.get("periodo") or x.get("periodo_corto")
    if not p:
        return None
    return p["anio"] if p["frecuencia"] == "anual" else P.etiqueta(p)


def _h(c, fila, idx, regla, descripcion, puntos, calculo=None, participantes=None, accion=None, periodo=None,
       lente="ajedrecista"):
    celda = c.celda(fila, idx)
    ref = f"{c['hoja']}!{celda}"
    accion = accion or "Contrastar la celda con la fuente y documentar la corrección o la explicación."
    etiqueta = c["etiquetas"].get(fila) or f"fila {fila}"
    col = next((x for x in c["columnas"] if x["idx"] == idx), {})
    ev = dict(
        archivo_origen=c["archivo"], hoja=c["hoja"], bloque_excel=c["cuadro"], celda=celda,
        indicador=etiqueta, columna=col.get("nombre"), periodo=periodo if periodo is not None else periodo_de(col),
        valor_observado=c["valores"].get((fila, idx)), regla=regla, naturaleza="control",
        puntaje_prioridad=puntos, factores_prioridad=[{"criterio": regla, "puntos": puntos}],
        ubicaciones_excel=[ref], celdas_participantes=list(dict.fromkeys([ref] + (participantes or []))),
        accion_revision=accion, **(calculo or {}),
    )
    return dict(
        lente=lente, modo="control_tabular",
        severidad="alta" if puntos >= 80 else "media",
        descripcion=f"{ref} · {etiqueta} · {col.get('nombre')}: {descripcion} {accion}",
        pregunta_institucional=accion,
        metodo="Control tabular reproducible sobre el cuadro reconocido; cálculo documentado en evidencia",
        evidencia=ev,
    )


def _columnas_valor(c):
    """Columnas numéricas que se leen como corte transversal (no períodos de una serie)."""
    if c["eje"] in ("columnas",):
        # «Total» del año en un encabezado año + trimestre: es el total de sus períodos, no un corte.
        return [x for x in c["columnas"] if x["numerica"] and not x.get("periodo") and not x.get("total_periodo")]
    if c["eje"] in ("filas", "largo"):
        return []
    return [x for x in c["columnas"] if x["numerica"]]


TOTAL_DE_PERIODOS = re.compile(r"^total(?:es)?(?:[\s_]*\d+)?(?:\s+(?:anual|del?\s+ano|ano))?\.?$")


def _columnas_corte(c):
    """Columnas sobre las que corren las reglas «por columna» (signo, escala): las de valor, sin los períodos.

    Si el cuadro tiene tres o más columnas cuyo encabezado es un período («2019», «2019 4º trimestre»,
    «2024T4*»), aunque el eje no se haya podido leer entero (años sueltos y trimestres en una misma
    fila), esas columnas son una serie: se controlan por fila, no como un corte entre rubros distintos.
    Lo mismo el «Total» del año intercalado entre ellas."""
    cols = _columnas_valor(c)
    numericas = [x for x in c["columnas"] if x["numerica"]]
    periodos = {x["idx"] for x in numericas if x.get("periodo") or P.leer(x.get("encabezado"))}
    if len(periodos) < 3:
        return cols
    return [x for x in cols if x["idx"] not in periodos
            and not TOTAL_DE_PERIODOS.match(normalizar(x.get("encabezado") if x.get("encabezado") is not None else ""))]


def _tolerancia(valores, n=1):
    dec = 0
    for v in valores:
        if v is None:
            continue
        t = repr(float(v))
        if "e" not in t and "." in t:
            dec = max(dec, min(4, len(t.split(".")[1].rstrip("0"))))
    return 0.5 * 10 ** -dec * max(1, n) + 1e-9


def _tolerancia_cierre(valores, n=1):
    """La de _tolerancia, con la resolución efectiva: para decidir si un total cierra con sus partes cuando son
    promedios de valores escritos con pocos decimales (tasas anuales de trimestres con tres decimales)."""
    dec = max((decimales_efectivos(v) for v in valores if v is not None), default=0)
    return 0.5 * 10 ** -dec * max(1, n) + 1e-9


def signos(c):
    out = []
    filas = c.filas_datos()
    cols = _columnas_corte(c)
    for x in cols:
        if ADMITE_SIGNO.search(x["nombre"]):
            continue
        vals = [(r, c["valores"][(r, x["idx"])]) for r in filas]
        vals = [(r, v) for r, v in vals if v]
        if len(vals) >= 5:
            neg = [r for r, v in vals if v < 0]
            pos = [r for r, v in vals if v > 0]
            for raros, comun in ((neg, "positivos"), (pos, "negativos")):
                if len(raros) == 1 and len(vals) - 1 >= 4:
                    r = raros[0]
                    if ADMITE_SIGNO.search(c["etiquetas"].get(r) or ""):
                        continue
                    v = c["valores"][(r, x["idx"])]
                    if abs(v) < 1e-6 * median(abs(u) for _, u in vals):
                        continue  # ruido de punto flotante (−0,0000014 publicado como −0,00), no un signo cambiado
                    out.append(_h(c, r, x["idx"], "signo_opuesto_columna",
                                  f"Único valor con signo distinto ({es_numero(v)}) en una columna con {len(vals) - 1} valores {comun}.",
                                  120, {"valores_con_signo_comun": len(vals) - 1},
                                  accion="Confirmar si el signo es correcto o si corresponde a una reversión documentada."))
        # Convención de egresos: costos y gastos con el mismo signo en todo el cuadro.
        egresos = [(r, c["valores"][(r, x["idx"])]) for r in filas
                   if EGRESO.search(c["etiquetas"].get(r) or "") and not ADMITE_SIGNO.search(c["etiquetas"].get(r) or "")
                   and not re.search(r"ingreso", c["etiquetas"].get(r) or "", re.I) and c["valores"][(r, x["idx"])]]
        if len(egresos) >= 3:
            neg = [r for r, v in egresos if v < 0]
            pos = [r for r, v in egresos if v > 0]
            for raros, convencion in ((pos, "negativo"), (neg, "positivo")):
                if len(raros) == 1 and len(egresos) - 1 >= 2:
                    r = raros[0]
                    v = c["valores"][(r, x["idx"])]
                    out.append(_h(c, r, x["idx"], "signo_egreso_fuera_de_convencion",
                                  f"«{c['etiquetas'].get(r)}» vale {es_numero(v)}, pero los demás costos y gastos del cuadro "
                                  f"({len(egresos) - 1}) se registran en {convencion}.",
                                  115, {"egresos_con_convencion": len(egresos) - 1},
                                  accion="Verificar el signo: un egreso cargado con el signo contrario altera los resultados que lo suman."))
    return _unicos(out)


def _unicos(hs):
    vistos, out = set(), []
    for h in hs:
        k = (h["evidencia"]["hoja"], h["evidencia"]["celda"])
        if k not in vistos:
            vistos.add(k)
            out.append(h)
    return out


def orden(c):
    out = []
    cols = [x for x in c["columnas"] if x["numerica"]]
    for menor_re, mayor_re, regla, lectura in ORDEN:
        menores = [x for x in cols if menor_re.search(x["nombre"])]
        mayores = [x for x in cols if mayor_re.search(x["nombre"]) and not menor_re.search(x["nombre"])]
        if len(menores) != 1 or len(mayores) != 1:
            continue
        a, b = menores[0], mayores[0]
        filas = [r for r in c["filas"] if c["valores"][(r, a["idx"])] is not None and c["valores"][(r, b["idx"])] is not None]
        malas = [r for r in filas if c["valores"][(r, a["idx"])] > c["valores"][(r, b["idx"])] + 1e-9]
        if len(filas) < 3 or not malas or len(malas) > max(1, 0.2 * len(filas)):
            continue
        for r in malas:
            va, vb = c["valores"][(r, a["idx"])], c["valores"][(r, b["idx"])]
            out.append(_h(c, r, a["idx"], regla,
                          f"{a['nombre']} = {es_numero(va)} supera a {b['nombre']} = {es_numero(vb)}. {lectura} "
                          f"En las otras {len(filas) - len(malas)} filas la relación se cumple.",
                          120, {"columna_mayor_esperada": b["nombre"], "valor_mayor_esperado": vb},
                          participantes=[c.ref(r, a["idx"]), c.ref(r, b["idx"])],
                          accion="Verificar si los valores de ambas columnas quedaron intercambiados o mal cargados."))
    return out


def porcentajes(c):
    """Partes porcentuales de cada fila que deben sumar 100 (estructura confirmada en el cuadro)."""
    cols = [x for x in _columnas_valor(c) if "%" in x["nombre"] or "porcentaje" in normalizar(x["nombre"])]
    if len(cols) < 3:
        return []
    total = [x for x in cols if es_total(x["nombre"]) or normalizar(x["nombre"]).startswith("total")]
    candidatas = [x for x in cols if x not in total and not NO_SUMABLE.search(x["nombre"].replace("%", "").replace("porcentaje", ""))]
    # Las partes son columnas contiguas cuya suma da 100 en la mayoría de las filas (la apertura
    # por edad, por ejemplo); otras tasas del mismo cuadro no participan.
    mejor = None
    for i in range(len(candidatas)):
        for j in range(i + 2, len(candidatas) + 1):
            partes = candidatas[i:j]
            filas = [r for r in c["filas"] if all(c["valores"][(r, x["idx"])] is not None for x in partes)]
            if len(filas) < 4:
                continue
            sumas = {r: sum(c["valores"][(r, x["idx"])] for x in partes) for r in filas}
            tol = _tolerancia([c["valores"][(r, x["idx"])] for r in filas for x in partes], len(partes)) + 0.05
            cierran = [r for r in filas if abs(sumas[r] - 100) <= tol]
            if len(cierran) >= 0.7 * len(filas) and (not mejor or len(cierran) > len(mejor[3])):
                mejor = (partes, filas, sumas, cierran)
    if not mejor:
        return []
    partes, filas, sumas, cierran = mejor
    out = []
    for r in filas:
        if r in cierran:
            continue
        dif = sumas[r] - 100
        ancla = total[0]["idx"] if total else partes[0]["idx"]
        participantes = [c.ref(r, x["idx"]) for x in partes] + ([c.ref(r, total[0]["idx"])] if total else [])
        declarado = c["valores"].get((r, total[0]["idx"])) if total else None
        out.append(_h(c, r, ancla, "porcentajes_no_suman_100",
                      f"Las partes porcentuales de la fila suman {es_numero(sumas[r])} (diferencia {es_numero(dif)} puntos)"
                      + (f"; la columna total muestra {es_numero(declarado)}" if declarado is not None else "")
                      + f". En {len(cierran)} de {len(filas)} filas las partes suman 100.",
                      120 if abs(dif) >= 0.5 else 70,
                      {"suma_partes": sumas[r], "total_declarado": declarado, "filas_que_cierran": len(cierran)},
                      participantes=participantes,
                      accion="Revisar qué parte de la fila quedó mal cargada; el total no indica cuál de ellas es."))
        out[-1]["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, ancla)}"
        out[-1]["evidencia"]["periodo"] = None
    return out


FILA_SUMA_MINIMA = 3  # columnas distintas con valores que cierran en la fila para tratarla como una suma
DIFERENCIA_MAXIMA_FILA_SUMA = 0.5


def totales_columna(c):
    """Fila «Total» que debe ser la suma de las filas de su grupo, columna por columna.

    Las columnas de períodos de una serie quedan para la conciliación por series; acá se
    controlan los atributos (total anual, crédito, población) y los cuadros transversales.
    Una diferencia mayor que el 15 % se toma como otra cosa que una suma, salvo que la fila
    sea claramente una suma (cierran al menos tres columnas distintas con valores y al menos
    la mitad de las comparadas) y el total supere a la suma de sus partes, todas presentes y
    no negativas, como deja un dígito corrido hacia abajo en un mes. Un promedio, un saldo o
    un conteo de personas únicas declaran menos que la suma y no entran; tampoco una parte
    reservada («s») o vacía, que el total publicado incluye.
    """
    out = []
    cols = [x for x in _columnas_valor(c) if not NO_SUMABLE.search(x["nombre"])]
    if not cols:
        return []
    inicio = 0
    filas = c["filas"]
    totales = [k for k, r in enumerate(filas) if es_total(c["etiquetas"].get(r))]
    for k, r in enumerate(filas):
        if not es_total(c["etiquetas"].get(r)):
            continue
        grupo = filas[inicio:k]
        inicio = k + 1
        if not grupo and k == totales[0]:
            # «Total» arriba (el país y después las jurisdicciones): su grupo son las filas que siguen hasta la
            # próxima «Total». Se confirma igual que abajo: otra columna tiene que cerrar.
            fin = next((j for j in totales if j > k), len(filas))
            grupo = filas[k + 1:fin]
        if len(grupo) < 2:
            continue
        cierran, fallan, lejanas, comparadas = [], [], [], 0
        for x in cols:
            declarado = c["valores"][(r, x["idx"])]
            partes = [c["valores"][(f, x["idx"])] for f in grupo]
            if declarado is None or sum(p is not None for p in partes) < 2:
                continue
            comparadas += 1
            suma = sum(p for p in partes if p is not None)
            tol = _tolerancia([declarado] + partes, len(partes))
            dif = declarado - suma
            if abs(dif) <= tol:
                cierran.append(x)
            elif abs(dif) <= 0.15 * max(abs(suma), abs(declarado)):
                fallan.append((x, declarado, suma, dif, partes, tol, False))
            elif (dif > 0 and abs(dif) <= DIFERENCIA_MAXIMA_FILA_SUMA * max(abs(suma), abs(declarado))
                  and all(p is not None and p >= 0 for p in partes)):
                lejanas.append((x, declarado, suma, dif, partes, tol, True))
        # Columnas distintas con valores: las de ceros o copiadas de otra no prueban que la fila sea una suma.
        con_valores = {(c["valores"][(r, x["idx"])], tuple(c["valores"][(f, x["idx"])] for f in grupo))
                       for x in cierran if c["valores"][(r, x["idx"])]}
        if len(con_valores) >= FILA_SUMA_MINIMA and 2 * len(con_valores) >= comparadas:
            fallan += lejanas
        # La fila es una suma si otra columna cierra, o si es la única columna sumable y casi cierra.
        if not fallan or not (cierran or (len(cols) == 1 and len(fallan) == 1)):
            continue
        for x, declarado, suma, dif, partes, tol, lejana in fallan:
            rel = abs(dif) / max(abs(suma), abs(declarado), 1e-9)
            participantes = [c.ref(r, x["idx"])] + [c.ref(f, x["idx"]) for f in grupo][:40]
            h = _h(c, r, x["idx"], "total_columna_no_cierra",
                   f"«{c['etiquetas'].get(r)}» declara {es_numero(declarado)} y las {len(grupo)} filas de su grupo suman "
                   f"{es_numero(suma)}: diferencia {es_numero(dif)} ({rel * 100:.3g} %), mayor que el redondeo admitido "
                   f"({es_numero(tol, 4)})."
                   + (f" En la misma fila cierran {', '.join(y['nombre'] for y in cierran[:4])}." if cierran else "")
                   + (f" La diferencia pasa el 15 %, pero la fila es una suma: cierran {len(con_valores)} columnas y el "
                      "total supera la suma de sus partes sin que falte ninguna." if lejana else ""),
                   120 if rel >= 0.0001 else 70,
                   {"total_declarado": declarado, "total_reconstruido": suma, "diferencia_con_signo": dif,
                    "diferencia_porcentual": rel * 100, "tolerancia_redondeo": tol,
                    "columnas_que_cierran": [y["nombre"] for y in cierran],
                    "fila_suma_por_otras_columnas": lejana},
                   participantes=participantes,
                   accion="Revisar el total y sus componentes; la diferencia no identifica por sí sola cuál de ellos es incorrecto.",
                   lente="conciliador")
            h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(r, x['idx'])}"
            out.append(h)
    return out


def _fila(c, patron, excluir=None):
    cand = [r for r in c["filas"] if re.search(patron, normalizar(c["etiquetas"].get(r)))
            and not (excluir and re.search(excluir, normalizar(c["etiquetas"].get(r))))]
    return cand[0] if len(cand) == 1 else None


def identidades(c):
    """Identidades contables reconocidas por rótulo, columna por columna."""
    out = []
    activo = _fila(c, r"^total (?:del )?activo$")
    pasivo = _fila(c, r"^total (?:del )?pasivo$")
    pn = _fila(c, r"^total (?:del )?patrimonio neto$")
    pasivo_pn = _fila(c, r"^total (?:del )?pasivo y (?:del )?patrimonio neto$")
    ventas = _fila(c, r"^ventas|^ingresos por (?:ventas|servicios)", r"costo")
    costo = _fila(c, r"^costo de (?:las|los)? ?(?:ventas|servicios|mercader)")
    bruto = _fila(c, r"^resultado bruto|^utilidad bruta|^margen bruto")
    reglas = []
    if activo and pasivo and pn:
        reglas.append(("activo_distinto_pasivo_mas_pn", activo, [pasivo, pn], "Total del activo = Total del pasivo + Total del patrimonio neto"))
    if pasivo_pn and pasivo and pn:
        reglas.append(("pasivo_mas_pn_no_cierra", pasivo_pn, [pasivo, pn], "Total del pasivo y patrimonio neto = pasivo + patrimonio neto"))
    if activo and pasivo_pn and not (pasivo and pn):
        reglas.append(("activo_distinto_pasivo_mas_pn", activo, [pasivo_pn], "Total del activo = Total del pasivo y patrimonio neto"))
    if bruto and ventas and costo:
        reglas.append(("resultado_bruto_no_cierra", bruto, [ventas, costo], "Resultado bruto = Ventas + Costo (con el costo en negativo)"))
    for x in [x for x in c["columnas"] if x["numerica"]
              and x.get("rol_confirmado") not in {"porcentaje", "indice"}]:
        for regla, total, partes, identidad in reglas:
            declarado = c["valores"][(total, x["idx"])]
            vals = [c["valores"][(p, x["idx"])] for p in partes]
            if declarado is None or any(v is None for v in vals):
                continue
            suma = sum(vals)
            if regla == "resultado_bruto_no_cierra" and vals[1] > 0:
                suma = vals[0] - vals[1]  # costo presentado en positivo
            tol = _tolerancia([declarado] + vals, len(vals))
            dif = declarado - suma
            if abs(dif) <= tol:
                continue
            rel = abs(dif) / max(abs(declarado), abs(suma), 1e-9)
            h = _h(c, total, x["idx"], regla, lente="conciliador", descripcion=
                   f"{identidad}: declarado {es_numero(declarado)}, calculado {es_numero(suma)}; diferencia {es_numero(dif)} ({rel * 100:.3g} %).",
                   puntos=125 if rel >= 0.0001 else 70,
                   calculo={"identidad": identidad, "total_declarado": declarado, "total_reconstruido": suma,
                    "diferencia_con_signo": dif, "diferencia_porcentual": rel * 100},
                   participantes=[c.ref(p, x["idx"]) for p in partes],
                   accion="Identificar qué rubro rompe la identidad; la diferencia no indica cuál es.")
            out.append(h)
    return out


def _complemento_a_100(c, r, x, columnas):
    """La fila es una distribución porcentual (trae un 100) y el valor tiene su complemento en la misma fila."""
    v = c["valores"][(r, x["idx"])]
    otros = [c["valores"].get((r, y["idx"])) for y in columnas if y["idx"] != x["idx"]]
    otros = [w for w in otros if isinstance(w, (int, float))]
    return any(abs(w - 100) <= 0.05 for w in otros) and any(abs(v + w - 100) <= 0.5 for w in otros)


def escala(c, libro=None):
    """Valores fuera de la escala de su columna: posible ×10, decimal corrido o unidad distinta."""
    out = []
    filas = c.filas_datos()
    # «Como % del total de servicios» entre los importes de cada rubro: otra magnitud, no un valor fuera de escala.
    porcentuales = [r for r in filas if FILA_PORCENTUAL.search(str(c["etiquetas"].get(r) or ""))]
    if porcentuales and len(porcentuales) < len(filas) / 2:
        filas = [r for r in filas if r not in porcentuales]
    columnas = _columnas_corte(c)
    for x in columnas:
        nombre = str(x.get("nombre") or "")
        if SIN_ESCALA.search(nombre):
            continue  # variación: no tiene una escala propia
        if re.match(r"^\s*(?:sub)?total(?:es)?\b", nombre, re.I):
            continue  # columna «Total» de cada fila: mezcla total general, subtotales y partidas por construcción
        vals = [(r, c["valores"][(r, x["idx"])]) for r in filas]
        vals = [(r, v) for r, v in vals if v]
        if len(vals) < 6:
            continue
        negativos = sum(v < 0 for _, v in vals)
        if PORCENTAJE_O_TASA.search(nombre) and min(negativos, len(vals) - negativos) >= max(2, 0.1 * len(vals)):
            continue  # porcentaje o tasa que cruza el cero: es una variación
        signo = 1 if sum(v > 0 for _, v in vals) >= sum(v < 0 for _, v in vals) else -1
        vals = [(r, v) for r, v in vals if v * signo > 0]
        if len(vals) < 6:
            continue
        # Distribución porcentual (la columna suma 100): una participación chica es natural, no un error de
        # escala. Un valor anormalmente grande sí se sigue marcando.
        distribucion = signo > 0 and abs(sum(v for _, v in vals) - 100) <= 1.5
        logs = {r: math.log10(abs(v)) for r, v in vals}
        med = median(logs.values())
        dispersion = max(mad(list(logs.values())) * 1.4826, 0.03)
        orden = sorted(logs.values())
        for r, lv in logs.items():
            z = (lv - med) / dispersion
            if abs(z) < 3.5 or (z < 0 and (distribucion or _complemento_a_100(c, r, x, columnas))):
                continue
            if sum(1 for o in orden if abs(o - lv) < 1e-9) >= 3:
                continue  # el mismo valor en tres o más filas es una convención (p. ej. 0,00001), no un error de escala
            otros = [o for o in orden if o != lv]
            hueco = (lv - max(o for o in otros if o < lv)) if z > 0 and any(o < lv for o in otros) else \
                    (min(o for o in otros if o > lv) - lv) if z < 0 and any(o > lv for o in otros) else 0
            # Muy lejos de la mediana y separado del resto, o separado del resto por dos órdenes de magnitud.
            if hueco < math.log10(4) or (abs(z) < 6 and hueco < 2):
                continue
            if _fila_proporcionada(c, filas, r, x, columnas):
                continue
            v = c["valores"][(r, x["idx"])]
            factor = 10 ** abs(lv - med)
            potencia = round(math.log10(factor))
            participantes = [c.ref(r, x["idx"])] + _origenes(c, libro, r, x["idx"])
            texto = (f"{es_numero(v)} está {factor:,.1f} veces {'por encima' if z > 0 else 'por debajo'} de la mediana de la "
                     f"columna ({es_numero(10 ** med * signo)}) y separado del valor más cercano por un factor de {10 ** hueco:,.1f}.")
            cerca_potencia = potencia >= 1 and abs(math.log10(factor) - potencia) < CERCA_POTENCIA_DIEZ
            if cerca_potencia:
                texto += " La distancia es cercana a una potencia de diez: posible error de escala, de decimal o de unidad (miles/pesos)."
            puntos = 110 + min(15, round(abs(z)))
            if 10 ** hueco < SEPARACION_CLARA and not cerca_potencia:
                # Poco separado del resto y sin un factor de escala reconocible: un valor chico o grande de
                # la columna, no necesariamente un error de unidad. Queda como señal media.
                puntos = 70
            h = _h(c, r, x["idx"], "valor_fuera_de_escala_columna", texto, puntos,
                   {"mediana_columna": 10 ** med * signo, "desviacion_robusta_log": z, "factor_escala": factor,
                    "separacion_valor_mas_cercano": 10 ** hueco, "valores_comparados": len(vals)},
                   participantes=participantes,
                   accion="Verificar la escala y la unidad de la celda contra la fuente; si es una fórmula, revisar sus celdas de origen.")
            h["evidencia"]["grupo_relacion"] = f"fila_escala:{c['hoja']}:{r}"
            if puntos < 80:
                # La prioridad del caso no la sube el factor: sin separación clara ni potencia de diez, la rareza es media.
                h["evidencia"]["rareza_media_sin_separacion"] = True
            out.append(h)
    return out


def _fila_proporcionada(c, filas, r, x, columnas):
    """Una fila entera más chica (o más grande) que las demás es un rubro de otro tamaño, no una
    celda fuera de escala: si en esa fila la razón entre esta columna y las otras columnas de
    valores es la habitual del cuadro, la celda acompaña a su fila y no se señala."""
    v = c["valores"][(r, x["idx"])]
    coherentes = comparadas = 0
    for y in columnas:
        if y is x:
            continue
        w = c["valores"][(r, y["idx"])]
        if not w or v * w <= 0:
            continue
        razones = []
        for f in filas:
            a, b = c["valores"][(f, x["idx"])], c["valores"][(f, y["idx"])]
            if f != r and a and b and a * b > 0:
                razones.append(math.log10(a / b))
        if len(razones) < 4:
            continue
        comparadas += 1
        propia = math.log10(v / w)
        dispersion = max(mad(razones) * 1.4826, 0.05)
        if abs(propia - median(razones)) <= max(3.5 * dispersion, math.log10(3)):
            coherentes += 1
    return comparadas >= 1 and coherentes * 2 > comparadas


def _periodos_cortos(c):
    """Cuadros de dos ejercicios (balances): cada ejercicio también es un corte transversal."""
    return [x for x in c["columnas"] if x["numerica"] and x.get("periodo_corto")]


def _origenes(c, libro, fila, idx):
    if libro is None or not c["formulas"].get((fila, idx)):
        return []
    from .evaluador import _partir_ref, _tokens
    texto = libro.formula(c["hoja"], fila, idx)
    refs = []
    try:
        for tipo, v in _tokens(texto.lstrip("=")):
            if tipo == "ref" and ":" not in v:
                hoja, partes = _partir_ref(v)
                refs.append(f"{hoja or c['hoja']}!{get_column_letter(partes[0][1])}{partes[0][3]}")
    except Exception:  # noqa: BLE001 — una fórmula no interpretable no aporta celdas de origen
        return []
    return refs[:6]


def repetidos(c):
    """Valores idénticos de muchas cifras en dos filas distintas o en dos ejercicios de la misma fila."""
    out = []
    if c.get("formulas_desconocidas"):
        # Libro sin fórmulas legibles (.xls binario): no se sabe qué celda es tipeada y qué es calculada,
        # y «valor idéntico tipeado» es justamente esa distinción.
        return out
    filas = c.filas_datos()
    cols = [x for x in c["columnas"] if x["numerica"]]
    for x in cols:
        if NO_SUMABLE.search(x["nombre"]):
            continue
        # Columna sin encabezado (auxiliar de un gráfico): sus valores se repiten a propósito.
        nombre = str(x["nombre"] or "").strip()
        letra = re.sub(r"\d", "", c.celda(1, x["idx"]))
        if not nombre or re.fullmatch(r"columna_\d+", nombre) or nombre == letra:
            continue
        por_valor = {}
        for r in filas:
            v = c["valores"][(r, x["idx"])]
            if v and _sig_efectivas(v) >= 7 and not c["formulas"][(r, x["idx"])]:
                por_valor.setdefault(v, []).append(r)
        if _coincidencias_esperables(c, filas, x):
            continue
        columna = [c["valores"][(r, x["idx"])] for r in filas]
        for v, rs in por_valor.items():
            if len(rs) != 2:
                continue
            a, b = rs
            # Filas duplicadas completas son otra regla; acá sólo un valor coincidente.
            iguales = sum(c["valores"][(a, y["idx"])] == c["valores"][(b, y["idx"])] for y in cols)
            if iguales > 1:
                continue
            if _coincidencias_por_unidad(v, columna) > COINCIDENCIA_ESPERABLE:
                continue  # serie empalmada: con su unidad común, una coincidencia casual no es rara
            out.append(_h(c, b, x["idx"], "valor_identico_otra_fila",
                          f"{es_numero(v)} aparece idéntico en «{c['etiquetas'].get(a)}» ({c.celda(a, x['idx'])}) y en "
                          f"«{c['etiquetas'].get(b)}» ({c.celda(b, x['idx'])}); con {_sig_efectivas(v)} cifras significativas la "
                          "coincidencia casual es improbable.",
                          100, {"celda_igual": c.ref(a, x["idx"]), "cifras_significativas": _sig_efectivas(v)},
                          participantes=[c.ref(a, x["idx"]), c.ref(b, x["idx"])],
                          accion="Verificar si el valor de una de las dos filas quedó copiado de la otra."))
    # Mismo valor en dos ejercicios de una fila mientras el resto del cuadro cambia.
    periodos = [x for x in cols if x.get("periodo_corto")]
    if len(periodos) == 2:
        a, b = periodos
        cambian = [r for r in filas if c["valores"][(r, a["idx"])] not in (None, 0) and c["valores"][(r, a["idx"])] != c["valores"][(r, b["idx"])]]
        for r in filas:
            va, vb = c["valores"][(r, a["idx"])], c["valores"][(r, b["idx"])]
            if va and va == vb and _sig_efectivas(va) >= 7 and len(cambian) >= 5 and not c["formulas"][(r, a["idx"])]:
                reciente = a if a["periodo_corto"]["ordinal"] > b["periodo_corto"]["ordinal"] else b
                previo = b if reciente is a else a
                out.append(_h(c, r, reciente["idx"], "valor_identico_otro_ejercicio",
                              f"{es_numero(va)} es idéntico en {reciente['nombre']} y {previo['nombre']}, mientras "
                              f"{len(cambian)} rubros del cuadro cambian entre ejercicios; con {_sig_efectivas(va)} cifras la coincidencia casual es improbable.",
                              100, {"celda_igual": c.ref(r, previo["idx"]), "cifras_significativas": _sig_efectivas(va)},
                              participantes=[c.ref(r, reciente["idx"]), c.ref(r, previo["idx"])],
                              accion="Verificar si el valor del ejercicio quedó copiado del ejercicio anterior."))
    return out


# Serie empalmada: un índice viejo de dos decimales (hasta 200,00) por un coeficiente. Sus valores se escriben con
# diez cifras pero son múltiplos de una misma unidad y se repiten como los del índice viejo.
DENOMINADOR_UNIDAD = 200000  # índices viejos de dos decimales hasta 2.000,00
COINCIDENCIAS_DE_UNIDAD = 10  # valores de la columna con la misma unidad (y al menos el 10 % de ella)
MAXIMO_MULTIPLO = 10 ** 7  # más allá, la unidad inferida es demasiado chica para tomarla como tal
COINCIDENCIA_ESPERABLE = 0.05  # coincidencias casuales esperadas a partir de las cuales una igualdad no es rara


def _coincidencias_por_unidad(v, valores):
    """Coincidencias casuales esperables en la columna si v es múltiplo de la misma unidad que buena parte de ella
    (una serie empalmada: un índice viejo de dos decimales por un coeficiente); 0 si no lo es.

    Comparte la unidad un valor w cuyo cociente con v es de enteros con denominador chico (n ≤ 200.000; en valores
    continuos casi ninguno lo es); hacen falta diez o más, y el 10 % de la columna. La unidad es v dividido por el
    mínimo común múltiplo de los enteros que le tocan a v en esos cocientes. Las coincidencias esperadas se cuentan
    sólo entre los valores que comparten la unidad (n(n-1)/2 pares, su rango y la unidad, como en
    _coincidencias_esperables): un tramo empalmado corto dentro de una columna continua, o una columna corta, sigue
    dando una coincidencia improbable."""
    vals = [w for w in valores if w]
    otros = [w for w in vals if w != v]
    if not v or len(otros) < COINCIDENCIAS_DE_UNIDAD:
        return 0.0
    comunes, multiplo, unidad_comun = 0, 1, [abs(v)]
    for w in otros:
        menor = abs(v) <= abs(w)
        q = abs(v) / abs(w) if menor else abs(w) / abs(v)
        f = Fraction(q).limit_denominator(DENOMINADOR_UNIDAD)
        if abs(float(f) - q) <= 1e-13:
            comunes += 1
            unidad_comun.append(abs(w))
            propio = f.numerator if menor else f.denominator
            multiplo = multiplo * propio // math.gcd(multiplo, propio)
            if multiplo > MAXIMO_MULTIPLO:
                return 0.0
    if comunes < max(COINCIDENCIAS_DE_UNIDAD, 0.1 * len(otros)):
        return 0.0
    unidad = abs(v) / multiplo
    n, rango = len(unidad_comun), max(unidad_comun) - min(unidad_comun)
    return n * (n - 1) / 2 * unidad / max(rango, unidad)


def _coincidencias_esperables(c, filas, x):
    """En columnas largas dos valores iguales pueden ser azar: con n filas hay n(n-1)/2 pares y la
    chance de que alguno coincida crece con el cuadrado de n. Se compara con la resolución de los
    valores (sus decimales) y el rango de la columna; si una coincidencia casual no es improbable,
    la igualdad no se señala."""
    vals = [c["valores"][(r, x["idx"])] for r in filas]
    vals = [v for v in vals if v]
    n = len(vals)
    if n < 50:
        return False
    resolucion = _tolerancia(vals) * 2
    rango = max(vals) - min(vals)
    esperadas = n * (n - 1) / 2 * resolucion / max(rango, resolucion)
    return esperadas > COINCIDENCIA_ESPERABLE


SALDO_INICIAL = re.compile(r"^(?:saldos?|stocks?|existencias?|posicion(?:es)?)\s+(?:al\s+)?(?:inicial(?:es)?|inicio|comienzo|de apertura)\b")
SALDO_FINAL = re.compile(r"^(?:saldos?|stocks?|existencias?|posicion(?:es)?)\s+(?:al\s+)?(?:final(?:es)?|cierre|de cierre)\b")


def _nivel(c, r):
    """Sangría del rótulo: la del formato de la celda y los espacios con que empieza el texto."""
    crudo = c["crudos"].get((r, c["etiqueta_idx"])) if c.get("etiqueta_idx") else None
    espacios = len(crudo) - len(crudo.lstrip()) if isinstance(crudo, str) else 0
    return (c.get("sangria") or {}).get(r, 0) * 4 + espacios


def _tolerancia_resolucion(total, partes):
    """Medio dígito de la resolución más gruesa de los sumandos por su cantidad, o el 0,01 % del total."""
    from .totales_cruzados import _tolerancia_suma
    return _tolerancia_suma(total, partes)


def saldos(c):
    """Bloques «Saldo inicial … Saldo final»: el saldo final es el inicial más los flujos del primer nivel.

    Los flujos son las filas entre los dos saldos con la sangría mínima («Transacciones», «Revalorizaciones»,
    «Otras variaciones»); sus aperturas con sangría («Variaciones de precio») no se suman otra vez. La
    identidad tiene que cerrar en el 80 % de las columnas para afirmar un desvío en las demás."""
    filas = c["filas"]
    rot = {r: normalizar(c["etiquetas"].get(r)) for r in filas}
    iniciales = [r for r in filas if SALDO_INICIAL.match(rot[r])]
    finales = [r for r in filas if SALDO_FINAL.match(rot[r])]
    if len(iniciales) != 1 or len(finales) != 1 or iniciales[0] >= finales[0]:
        return []
    ini, fin = iniciales[0], finales[0]
    medio = [r for r in filas if ini < r < fin and rot[r]]
    if not medio:
        return []
    minimo = min(_nivel(c, r) for r in medio)
    flujos = [r for r in medio if _nivel(c, r) == minimo]
    cols = [x for x in c["columnas"] if x["numerica"]
            and x.get("rol_confirmado") not in {"porcentaje", "indice"}
            and not NO_SUMABLE.search(str(x.get("nombre") or ""))]
    cierran, fallan = 0, []
    for x in cols:
        final = c["valores"].get((fin, x["idx"]))
        partes = [c["valores"].get((r, x["idx"])) for r in [ini] + flujos]
        if final is None or any(v is None for v in partes):
            continue
        suma = sum(partes)
        dif = final - suma
        if abs(dif) <= _tolerancia_resolucion(final, partes):
            cierran += 1
        else:
            fallan.append((x, final, suma, dif))
    if cierran < 2 or cierran < 0.8 * (cierran + len(fallan)):
        return []
    out = []
    nombres = ", ".join(f"«{c['etiquetas'].get(r)}»" for r in flujos)
    for x, final, suma, dif in fallan:
        rel = abs(dif) / max(abs(final), abs(suma), 1e-9)
        h = _h(c, fin, x["idx"], "saldo_final_no_cierra",
               f"Saldo final = saldo inicial + {nombres}: declarado {es_numero(final)}, calculado {es_numero(suma)}; "
               f"diferencia {es_numero(dif)} ({rel * 100:.3g} %). En otras {cierran} columnas del cuadro la identidad cierra.",
               125 if rel >= 0.0001 else 70,
               {"identidad": "Saldo final = Saldo inicial + flujos", "total_declarado": final, "total_reconstruido": suma,
                "valor_declarado": final, "valor_calculado": suma, "diferencia_con_signo": dif,
                "diferencia_porcentual": rel * 100, "columnas_que_cierran": cierran},
               participantes=[c.ref(r, x["idx"]) for r in [ini] + flujos],
               accion="Identificar qué flujo o qué saldo rompe la identidad; la diferencia no indica cuál es.",
               lente="conciliador")
        h["evidencia"]["grupo_relacion"] = f"celda:{c.ref(fin, x['idx'])}"
        out.append(h)
    return out


def _subperiodo_rotulo(v):
    from .modelo import _subperiodo
    if _subperiodo(v) not in (None, "total"):
        return True
    p = P.leer(v) if isinstance(v, str) else None
    return bool(p and p["frecuencia"] in ("mensual", "trimestral", "semestral", "bimestral"))


def filas_repetidas(c):
    """Dos filas seguidas iguales en todas las columnas: el período repite el anterior.

    Sólo en un cuadro de períodos en filas (o formato largo), o si los dos rótulos son subperíodos
    («II», «III», «Marzo»): dos rubros distintos con las mismas cifras en un estado (dos resultados
    iguales) son otra cosa. Se exigen tres columnas comparadas y tres valores de cuatro cifras, y que la
    repetición sea aislada: tres o más filas iguales seguidas (un dato trimestral o anual repetido en sus
    meses) o repeticiones en más de dos lugares del cuadro son la forma de publicarlo, no un error."""
    cols = [x for x in c["columnas"] if x["numerica"] and not x.get("tiempo")]
    filas = c["filas"]
    periodos = c.get("eje") in ("filas", "largo")
    iguales = []  # (k, fila, fila siguiente, columnas comparadas): la fila k + 1 repite la k
    for k, (a, b) in enumerate(zip(filas, filas[1:])):
        if not periodos and not (_subperiodo_rotulo(c["etiquetas"].get(a)) and _subperiodo_rotulo(c["etiquetas"].get(b))):
            continue
        if periodos and c.get("periodos_fila") and a not in c["periodos_fila"]:
            continue  # la fila de arriba no es un período: la categoría («BONOS DE CONSOLIDACIÓN») sobre su único título
        pares = [(c["valores"].get((a, x["idx"])), c["valores"].get((b, x["idx"]))) for x in cols]
        pares = [(u, v) for u, v in pares if u is not None and v is not None]
        if len(pares) < 3 or any(u != v for u, v in pares):
            continue
        if sum(1 for u, _ in pares if u and _sig(u) >= 4) < 3:
            continue
        iguales.append((k, a, b, len(pares)))
    posiciones = {k for k, *_ in iguales}
    if sum(1 for k in posiciones if k - 1 not in posiciones) > 2:
        return []
    out = []
    for k, a, b, comparadas in iguales:
        if k - 1 in posiciones or k + 1 in posiciones:
            continue
        idxs = [x["idx"] for x in cols if c["valores"].get((b, x["idx"])) is not None]
        rango = f"{c['hoja']}!{get_column_letter(min(idxs))}{b}:{get_column_letter(max(idxs))}{b}"
        pa = c.get("periodos_fila", {}).get(a, {}).get("texto") or c["etiquetas"].get(a)
        pb = c.get("periodos_fila", {}).get(b, {}).get("texto") or c["etiquetas"].get(b)
        h = _h(c, b, idxs[0], "fila_repetida_periodo_anterior",
               f"La fila de {pb} repite exactamente la de {pa} en las {comparadas} columnas, con cifras de cuatro o más "
               "dígitos: el período parece copiado del anterior o la serie quedó corrida.",
               110, {"fila_igual": a, "columnas_iguales": comparadas, "rango": rango},
               participantes=[c.ref(a, i) for i in idxs] + [c.ref(b, i) for i in idxs],
               accion="Verificar si el período repite los valores del anterior o si la serie quedó corrida un período.",
               periodo=pb)
        h["evidencia"]["ubicaciones_excel"] = [rango]
        h["evidencia"]["grupo_relacion"] = f"fila_repetida:{c['hoja']}:{b}"
        out.append(h)
    return out


ENTEROS_MINIMOS = 8  # valores enteros de la columna que la muestran como de conteos


def decimal_aislado(c):
    """Un valor con decimales en una columna de conteos: todos los demás son enteros (personas, establecimientos,
    delitos) y éste no (51,47). Es el rastro típico de un valor tipeado o pegado de otra columna."""
    out = []
    for x in _columnas_valor(c):
        if NO_SUMABLE.search(x["nombre"]) or x.get("periodo"):
            continue
        vals = [(r, c["valores"].get((r, x["idx"]))) for r in c["filas"]]
        vals = [(r, v) for r, v in vals if v is not None]
        # 57,99999999999999 es un 58 con ruido de punto flotante; los ceros no muestran una columna de conteos.
        enteros = [r for r, v in vals if abs(v - round(v)) < 1e-6 and v != 0]
        decimales = [(r, v) for r, v in vals if abs(v - round(v)) >= 1e-6]
        if len(decimales) != 1 or len(enteros) < ENTEROS_MINIMOS:
            continue
        r, v = decimales[0]
        if c["formulas"].get((r, x["idx"])):
            continue  # un promedio o cociente calculado puede dar decimales
        out.append(_h(c, r, x["idx"], "decimal_en_columna_de_enteros",
                      f"«{c['etiquetas'].get(r) or 'fila ' + str(r)}» tiene {es_numero(v)} en «{x['nombre']}», una columna en la "
                      f"que los otros {len(enteros)} valores son enteros. Un conteo no lleva decimales.",
                      120, {"enteros_en_columna": len(enteros)},
                      accion="Contrastar el valor con la fuente: puede estar tipeado o pegado de otra columna."))
    return out


def porcentajes_como_cantidades(c):
    """Una unidad confirmada se contrasta con una relación repetida del original.

    No impone un techo de 100 a todas las tasas: exige partes no negativas que
    cierren repetidamente contra un total de cantidades, en un corte categórico.
    Genera una sola pregunta de unidad por cuadro, no errores por cada celda.
    """
    if c.get("eje_confirmado") != "ninguno":
        return []
    cols = _columnas_valor(c)
    partes = [x for x in cols if x.get("rol_confirmado") == "porcentaje"
              and not re.search(r"variaci|crecim|cumplimiento|ejecuci|\bvar\b|diferencia", x["nombre"], re.I)]
    totales = [x for x in cols if x.get("rol_confirmado") == "cantidad" and es_total(x["nombre"])]
    if len(partes) < 2 or len(totales) != 1:
        return []
    total = totales[0]
    filas, compatibles = [], []
    for r in c.filas_datos():
        valores = [c["valores"].get((r, x["idx"])) for x in partes]
        cantidad = c["valores"].get((r, total["idx"]))
        if cantidad is None or any(v is None or v < 0 for v in valores):
            continue
        filas.append(r)
        if cantidad > 100 and max(valores) > 100 and abs(sum(valores) - cantidad) <= _tolerancia(valores, len(partes)):
            compatibles.append(r)
    if len(compatibles) < 4 or len(compatibles) < 0.8 * len(filas):
        return []
    r = compatibles[0]
    ancla = next(x for x in partes if c["valores"][(r, x["idx"])] > 100)
    valores = [c["valores"][(r, x["idx"])] for x in partes]
    cantidad = c["valores"][(r, total["idx"])]
    participantes = [c.ref(r, x["idx"]) for x in [total] + partes]
    h = _h(c, r, ancla["idx"], "porcentajes_como_cantidades",
           f"En {len(compatibles)} de {len(filas)} filas, las columnas confirmadas como porcentaje suman "
           f"la cantidad de «{total['nombre']}» y contienen valores mayores que 100. "
           f"Por ejemplo, las partes suman {es_numero(sum(valores))} y el total de cantidades es {es_numero(cantidad)}. "
           "La aritmética cierra; la unidad declarada requiere revisión.",
           70, {"suma_partes": sum(valores), "total_declarado": cantidad,
                "filas_compatibles": compatibles, "filas_comprobadas": len(filas),
                "columnas_porcentaje": [x["nombre"] for x in partes],
                "roles_confirmados": {x["nombre"]: x["rol_confirmado"] for x in [total] + partes}},
           participantes=participantes,
           accion="Confirmar si son cantidades o porcentajes y corregir el encabezado o la interpretación; no cambiar las cifras sin contrastar la fuente.")
    return [h]


def analizar(c, libro=None):
    return (signos(c) + orden(c) + porcentajes(c) + totales_columna(c) + identidades(c) + saldos(c) + escala(c, libro)
            + repetidos(c) + filas_repetidas(c) + decimal_aislado(c) + porcentajes_como_cantidades(c))
