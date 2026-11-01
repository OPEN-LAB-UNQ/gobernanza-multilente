"""Cuánto del archivo pudo leerse como serie temporal antes de aplicar ningún control.

Un control sólo puede fallar sobre lo que el motor logró interpretar. Esta medida se calcula
aparte de los controles: recorre el libro y cuenta los cuadros que parecen tablas por período
(una fila de encabezado con años y filas de rubros debajo), y los compara con los que quedaron
efectivamente disponibles como series. La diferencia se informa con su motivo: sin ella, un
archivo ilegible y un archivo correcto se ven igual.
"""

import io
import re

from openpyxl import load_workbook

from services.excel_estructura import anio_rotulo, encabezados_aislados, presente
from services.plural import contar

from .comun import numero
from .periodos import es_auxiliar

ENCABEZADO = re.compile(r"^(indicador|categor|clase|concepto|rubro|partida)", re.I)
MINIMO_COLUMNAS = 3
UMBRAL_PARCIAL = 0.8
UMBRAL_INSUFICIENTE = 0.5


def _anios(fila):
    return [(j, numero(v)) for j, v in enumerate(fila[1:], 1) if numero(v) is not None and 1800 <= numero(v) <= 2200]


def _cuadros_aparentes(contenido):
    """Inventario por coordenadas; no depende de cuántas filas devolvió el lector."""
    from openpyxl.utils import get_column_letter

    from services.excel_estructura import encabezados_apilados, filas_encabezado_anual, paneles_anuales

    from .comun import normalizar
    libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    salida = []
    try:
        for ws in libro.worksheets:
            filas = list(ws.iter_rows(values_only=True))
            ancho = max((len(f) for f in filas), default=0)
            filas = [tuple(f) + (None,) * (ancho - len(f)) for f in filas]
            headers = filas_encabezado_anual(filas)
            # Un cuadro apilado con su propio encabezado (trimestres, otra serie) cierra el anterior.
            cortes = sorted(set(headers) | set(encabezados_apilados(filas, set(headers))))
            for n, i in enumerate(headers):
                fin = next((k for k in cortes if k > i), len(filas))
                cuerpo = filas[i + 1:fin]
                aisladas = set(encabezados_aislados(filas[i], cuerpo))
                for panel in paneles_anuales(filas[i], cuerpo):
                    anios = [(j, anio_rotulo(filas[i][j])) for j in panel
                             if j not in aisladas and anio_rotulo(filas[i][j]) is not None]
                    if not anios:
                        continue
                    etiquetas = [j for j in panel if j < anios[0][0] and presente(filas[i][j])]
                    if not etiquetas:
                        continue
                    etiqueta = etiquetas[0]
                    # Las columnas auxiliares («CV (%)» después de cada año) no son parte de la serie.
                    auxiliares = {j for j in panel if es_auxiliar(filas[i][j])}
                    unidades = []
                    for k in range(i + 1, fin):
                        celdas = [f"{get_column_letter(j + 1)}{k + 1}" for j in panel
                                  if anios[0][0] <= j <= anios[-1][0] and presente(filas[k][j]) and j not in auxiliares]
                        if not celdas:
                            continue
                        unidades.append({'id': ('horizontal', k + 1, etiqueta + 1), 'celdas': celdas,
                                         'sin_rotulo': not presente(filas[k][etiqueta])})
                    # Un encabezado sin ningún valor guardado debajo no es una serie que leer
                    # (p. ej. todas fórmulas sin resultado guardado: eso se declara por hoja).
                    if not unidades:
                        continue
                    salida.append({'hoja': ws.title, 'fila': i + 1, 'series': len(unidades),
                                   'unidades': unidades, 'periodos': len(anios),
                                   'motivo': f'el cuadro tiene {len(anios)} períodos y se necesitan {MINIMO_COLUMNAS}' if len(anios) < MINIMO_COLUMNAS else None})
            if headers:
                continue
            # Año + mes en columnas vecinas: inventario independiente de todas las celdas
            # mensuales, incluso cuando los rótulos de variables ocupan varias filas.
            mensual = False
            from . import periodos as P
            for i, fila in enumerate(filas):
                pares = [j for j in range(len(fila) - 1)
                         if normalizar(fila[j]) in ('ano', 'anio', 'year')
                         and normalizar(fila[j + 1]) in ('mes', 'month')]
                if len(pares) != 1:
                    continue
                eje = pares[0]
                anio, rows = None, []
                for k in range(i + 1, len(filas)):
                    v = filas[k][eje]
                    p = P.leer(v)
                    if p and p['frecuencia'] == 'anual':
                        anio = p['anio']
                    elif presente(v):
                        anio = None
                    mes = P.leer(filas[k][eje + 1])
                    if anio is not None and mes and mes['frecuencia'] == 'mensual' and mes.get('anio') is None:
                        rows.append(k)
                if len(rows) < 3:
                    continue
                unidades = []
                for j in range(ancho):
                    if j in (eje, eje + 1) or not any(numero(filas[k][j]) is not None for k in rows):
                        continue
                    rotulo = any(isinstance(filas[k][j], str) and presente(filas[k][j]) for k in range(i, rows[0]))
                    celdas = [f"{get_column_letter(j + 1)}{k + 1}" for k in rows if presente(filas[k][j])]
                    unidades.append({'id': ('vertical', i + 1, j + 1), 'celdas': celdas, 'sin_rotulo': not rotulo})
                salida.append({'hoja': ws.title, 'fila': i + 1, 'series': len(unidades), 'unidades': unidades,
                               'periodos': len(rows), 'motivo': None})
                mensual = True
                break
            if mensual:
                continue
            # Tablas verticales con eje anual explícito: una serie por columna de valores.
            for i, fila in enumerate(filas):
                ejes = [j for j, v in enumerate(fila) if normalizar(v) in ('ano','anio','periodo','year')]
                if len(ejes) != 1:
                    continue
                eje = ejes[0]
                rows = [k for k in range(i + 1, len(filas)) if numero(filas[k][eje]) is not None and 1800 <= numero(filas[k][eje]) <= 2200]
                if not rows:
                    continue
                unidades = []
                for j, nombre in enumerate(fila):
                    if j == eje or not presente(nombre) or not any(numero(filas[k][j]) is not None for k in rows):
                        continue
                    celdas = [f"{get_column_letter(j + 1)}{k + 1}" for k in rows if presente(filas[k][j])]
                    unidades.append({'id': ('vertical', i + 1, j + 1), 'celdas': celdas, 'sin_rotulo': False})
                salida.append({'hoja': ws.title, 'fila': i + 1, 'series': len(unidades), 'unidades': unidades,
                               'periodos': len(rows), 'motivo': f'el cuadro tiene {len(rows)} períodos y se necesitan {MINIMO_COLUMNAS}' if len(rows) < MINIMO_COLUMNAS else None})
                break
        return salida
    finally:
        libro.close()


def resumen(leidos, series):
    """Una fila leída nunca compensa otra omitida, aunque el total coincida."""
    por_archivo = []
    for archivo, contenido in leidos:
        if not (archivo.lower().endswith((".xlsx", ".xlsm")) or contenido[:2] == b"PK"):
            continue
        aparentes = _cuadros_aparentes(contenido)
        if not aparentes:
            continue
        propias = [s for s in series if s['archivo'] == archivo]
        disponibles = {}
        for s in propias:
            if s.get('orientacion') == 'vertical':
                clave = (s['hoja'], 'vertical', s['encabezado'], s['columna'])
            else:
                clave = (s['hoja'], 'horizontal', s.get('fila'), s.get('columna_etiqueta'))
            disponibles.setdefault(clave, set()).update(s['celdas'])
        esperadas = {}
        origen = {}
        for c in aparentes:
            for unidad in c['unidades']:
                clave = (c['hoja'], *unidad['id'])
                esperadas[clave] = set(unidad['celdas'])
                origen[clave] = (c, unidad)
        # Un encabezado multinivel puede localizarse por su primera o última fila. En
        # vertical se exige la misma columna y todas las coordenadas, no el número de encabezado.
        verticales = {}
        for k, celdas in disponibles.items():
            if k[1] == 'vertical':
                verticales.setdefault((k[0], k[3]), set()).update(celdas)
        incorporadas = {k for k, celdas in esperadas.items()
                        if celdas <= (verticales.get((k[0], k[3]), set()) if k[1] == 'vertical'
                                      else disponibles.get(k, set()))}
        motivos = {}
        for k in esperadas.keys() - incorporadas:
            c, unidad = origen[k]
            motivo = ('la serie no tiene rótulo identificable' if unidad['sin_rotulo'] else c['motivo']) or 'no se incorporó íntegramente la serie a los controles temporales'
            ubicacion = f"{c['hoja']} ({', '.join(unidad['celdas'][:2])})"
            motivos.setdefault(motivo, []).append(ubicacion)
        total = len(esperadas)
        leidas = len(incorporadas)
        por_archivo.append({'archivo': archivo, 'series_leidas': leidas, 'series_aparentes': total,
                            'hojas_leidas': len({k[0] for k in incorporadas}),
                            'hojas_aparentes': len({k[0] for k in esperadas}),
                            'proporcion': leidas / total if total else 0.0,
                            'motivos': {k: sorted(v)[:6] for k, v in motivos.items()}})
    if not por_archivo:
        return {}
    leidas = sum(a['series_leidas'] for a in por_archivo)
    total = sum(a['series_aparentes'] for a in por_archivo)
    proporcion = leidas / total if total else 0.0
    estado = 'completa' if total and leidas == total else 'parcial' if proporcion >= UMBRAL_INSUFICIENTE else 'insuficiente'
    return {'estado': estado, 'proporcion': proporcion, 'series_leidas': leidas,
            'series_aparentes': total, 'archivos': por_archivo}


def porcentaje(datos):
    total = datos.get('series_aparentes', 0)
    leidas = datos.get('series_leidas', 0)
    proporcion = min(1.0, leidas / total) if total else 0
    return f"{min(99.9, proporcion * 100):.1f}".rstrip("0").rstrip(".") if proporcion < 1 else '100'


def lineas(datos):
    """Texto para el comprobante: primero cuánto se leyó, después por qué no se leyó el resto."""
    if not datos:
        return []
    total = datos.get("series_aparentes", 0)
    leidas = datos.get("series_leidas", 0)
    if not total:
        return ["**Series temporales: cobertura no calculable.** No se identificó un conjunto de "
                "series por período para calcular este porcentaje. Esto no significa que el archivo "
                "no se haya leído; consultá la cobertura de los demás controles y de la exploración IA."]
    pct = porcentaje(datos)
    cabecera = (f"**Cobertura de lectura: {pct} % de las series temporales.** Se reconocieron {leidas} de las "
                f"{total} series por período. Este porcentaje no mide la lectura total del archivo.")
    if leidas < total:
        cabecera += (" Los controles que dependen de estas series sólo alcanzan la parte reconocida. "
                     "Otros controles o la exploración IA pueden tener un alcance diferente. En este alcance, "
                     "la ausencia de observaciones no indica que el archivo esté bien.")
    salida = [cabecera]
    for archivo in datos["archivos"]:
        if archivo["proporcion"] >= UMBRAL_PARCIAL and not archivo["motivos"]:
            continue
        salida.append(
            f"{archivo['archivo']}: {archivo['series_leidas']} de {archivo['series_aparentes']} series, "
            f"{archivo['hojas_leidas']} de {archivo['hojas_aparentes']} hojas."
        )
        for motivo, donde in archivo["motivos"].items():
            salida.append(f"Serie no reconocida porque {motivo}: {', '.join(donde)}.")
    return salida


FILAS_MINIMAS_PDF = 20


def resumen_pdf(meta):
    """Cuánto de un PDF pudo convertirse en tablas por período.

    Los controles trabajan sobre series: un documento del que sólo se extrae texto, o unas pocas
    filas sueltas, no queda revisado aunque el análisis termine sin observaciones.
    """
    if not isinstance(meta, dict) or not meta.get("cantidad_archivos"):
        return {}
    paginas = meta.get("paginas_documento")
    analizadas = meta.get("paginas_analizadas")
    filas = meta.get("filas_pdf") or 0
    # Una fila cuyas celdas traen dos o más números juntos no es una fila de tabla: el cuadro no se
    # pudo separar en columnas. Contarla como revisable sería declarar una cobertura inexistente.
    utiles = meta.get("filas_pdf_utiles")
    utiles = filas if utiles is None else int(utiles)
    cuadros = int(meta.get("cuadros_pdf") or 0)
    sin_separar = int(meta.get("cuadros_pdf_sin_separar") or 0)
    sin_rotulos = int(meta.get("cuadros_pdf_sin_rotulos") or 0)
    con_tablas = any(str(m).startswith("tabla") for m in (meta.get("modos_extraccion") or []))
    sin_texto = int(meta.get("paginas_sin_texto") or 0)
    if not con_tablas or utiles < FILAS_MINIMAS_PDF:
        estado = "insuficiente"
    elif (meta.get("truncado_por_paginas") or (paginas and analizadas and analizadas < paginas) or sin_separar
          or sin_rotulos or sin_texto):
        estado = "parcial"
    else:
        estado = "completa"
    return {"estado": estado, "paginas_documento": paginas, "paginas_analizadas": analizadas,
            "filas": filas, "filas_utiles": utiles, "cuadros": cuadros, "cuadros_sin_separar": sin_separar,
            "cuadros_sin_rotulos": sin_rotulos, "paginas_sin_texto": sin_texto,
            "con_tablas": con_tablas, "truncado": bool(meta.get("truncado_por_paginas")),
            "archivos": meta.get("cantidad_archivos")}


def lineas_pdf(datos):
    if not datos:
        return []
    sin_texto = int(datos.get("paginas_sin_texto") or 0)
    leidas = (datos.get("paginas_analizadas") or 0) - sin_texto
    paginas = f"{leidas} de {datos['paginas_documento']} páginas" if datos.get("paginas_documento") else "las páginas disponibles"
    base = f"**PDF: se leyeron {paginas} y se extrajeron {datos['filas']} filas de datos.**"
    if sin_texto:
        base += (f" {contar(sin_texto, 'página no tiene', 'páginas no tienen')} texto extraíble (probablemente "
                 f"escaneada{'s' if sin_texto > 1 else ''}): su contenido no se leyó ni se controló.")
    sin_separar = datos.get("cuadros_sin_separar") or 0
    sin_rotulos = datos.get("cuadros_sin_rotulos") or 0
    if sin_separar:
        base += (
            f" {sin_separar} de los {datos.get('cuadros')} cuadros no pudieron separarse en columnas "
            "—sus celdas traen varios números juntos, como pasa con los encabezados de dos niveles—."
        )
    if sin_rotulos:
        base += (
            f" En {sin_rotulos} de los {datos.get('cuadros')} cuadros las columnas quedaron separadas pero sin "
            "rótulo confiable: el encabezado se reparte entre columnas vecinas y no se puede decir de qué es "
            "cada número."
        )
    if sin_separar or sin_rotulos:
        base += f" Quedan {datos.get('filas_utiles')} filas revisables."
    if datos["estado"] == "insuficiente":
        motivo = (
            "El contenido se leyó como texto por página, no como tablas por período."
            if not datos["con_tablas"]
            else f"Con {datos.get('filas_utiles')} filas revisables no hay series suficientes para aplicar los controles."
        )
        return [f"{base} {motivo} Los controles trabajan sobre series por período, así que la ausencia de "
                "observaciones no indica que el documento esté bien: casi no se lo pudo revisar. "
                "Si el documento tiene un anexo de datos, conviene cargar ese archivo."]
    if datos["estado"] == "parcial":
        return [f"{base} El documento quedó recortado o no se recorrió entero: lo que no se leyó no fue revisado."]
    return [base]
