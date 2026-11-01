"""Valores fijos donde el cuadro calcula: control básico, siempre activo.

Si una columna (o una fila) repite el mismo cálculo —«=SUM(C5:N5)», «=H5/D5*100», una
referencia a otra hoja— y una celda de esa misma línea tiene un número tipeado, se copia la
fórmula vecina a esa celda (como lo haría Excel) y se compara el resultado con el valor fijo.
Si difieren más que el redondeo, el valor quedó desactualizado: prioridad alta, con la celda.
Si coinciden, no se informa acá (queda para la revisión técnica de fórmulas). La tolerancia
combina una parte absoluta (medio último decimal por término) y una relativa mínima, para que
una diferencia de 25 millones sobre 70.600 millones no pase por redondeo.

En las planillas oficiales es habitual que los años viejos estén tipeados y los recientes
calculados, o que la fórmula trasladada no aplique a esa celda. Por eso la prioridad alta exige
más evidencia: tres o más fórmulas hermanas en la línea, con el valor fijo entre ellas o después
de ellas (el total al pie), nunca antes de que empiece el tramo calculado (los períodos viejos
tipeados), y un resultado plausible de la fórmula trasladada (ni cerca de cero ni una diferencia
de casi el 100 %). Si falta algo de eso, la señal queda como media. Tampoco alcanza para la
prioridad alta una columna de enlaces a celdas de la misma columna («=+C25» en C23, «=+C28» en C26; lo
mismo en una fila): copiar un dato no es un cálculo, y trasladar el enlace sólo apunta a otra celda de
la columna, no al valor que correspondía (los cuadros de deuda encadenan así filas que no guardan
relación entre sí). Un enlace a otra columna de la misma fila («=+B6» en toda la columna C) sí se
traslada con sentido y conserva la prioridad alta.
"""

import re
from collections import Counter, defaultdict

from openpyxl.utils import column_index_from_string, get_column_letter

from .comun import es_numero
from .evaluador import NoEvaluable, firma, trasladar

HERMANAS_MINIMAS = 3  # fórmulas iguales en la línea para afirmar con prioridad alta
DIFERENCIA_IMPLAUSIBLE = 0.9  # la fórmula trasladada da ~0 o algo sin relación con el valor: no aplica
# Enlace suelto a otra celda de la misma hoja («=+C25», «=$D$30»): copia un dato, no calcula nada.
ENLACE = re.compile(r"^=\s*\+?\s*\$?(?P<col>[A-Z]{1,3})\$?(?P<fila>\d+)\s*$", re.I)


def _enlace_en_la_linea(formula, origen, motivo):
    """La fórmula de la línea sólo enlaza otra celda de esa misma columna (o fila): trasladada, apunta a
    otra celda de la línea, no a un cálculo."""
    m = ENLACE.match(formula)
    if not m:
        return False
    if motivo.startswith("la columna"):
        return column_index_from_string(m.group("col").upper()) == origen[1]
    if motivo.startswith("la fila"):
        return int(m.group("fila")) == origen[0]
    return False


def _decimales(v):
    t = repr(float(v))
    if "e" in t or "." not in t:
        return 0
    return min(6, len(t.split(".")[1].rstrip("0")))


def _hallazgo(archivo, hoja, fila, col, regla, texto, puntos, calculo, participantes):
    celda = f"{get_column_letter(col)}{fila}"
    ref = f"{hoja}!{celda}"
    accion = ("Restituir la fórmula o documentar por qué la celda se cargó a mano; verificar el valor contra su origen."
              if puntos >= 80 else "Restituir la fórmula para que la celda se actualice con sus componentes.")
    return dict(
        lente="ajedrecista", modo="control_tabular", severidad="alta" if puntos >= 80 else "media",
        descripcion=f"{ref}: {texto} {accion}", pregunta_institucional=accion,
        metodo="Comparación del valor fijo con la fórmula de su misma línea, trasladada a la celda",
        evidencia=dict(
            archivo_origen=archivo, hoja=hoja, bloque_excel=None, celda=celda, indicador=None, periodo=None,
            valor_observado=calculo.get("valor_fijo"), regla=regla, naturaleza="control", puntaje_prioridad=puntos,
            factores_prioridad=[{"criterio": regla, "puntos": puntos}], ubicaciones_excel=[ref],
            celdas_participantes=list(dict.fromkeys([ref] + participantes)), accion_revision=accion,
            grupo_relacion=f"celda:{ref}", **calculo,
        ),
    )


def analizar(archivo, libro, hojas=None):
    out = []
    for ws in libro.formulas.worksheets:
        if hojas is not None and ws.title not in hojas:
            continue
        hoja = ws.title
        formulas = {}
        fijos = {}
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 3000), max_col=min(ws.max_column, 80)):
            for c in fila:
                if c.data_type == "f" and isinstance(c.value, str):
                    formulas[(c.row, c.column)] = c.value
                elif isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    fijos[(c.row, c.column)] = float(c.value)
        if not formulas or not fijos:
            continue
        firmas = {k: firma(v, *k) for k, v in formulas.items()}
        candidatos = defaultdict(list)  # celda fija -> [(celda_origen, motivo)]
        # 1. Líneas que repiten un mismo cálculo (columna o fila).
        for eje in (0, 1):
            lineas = defaultdict(list)
            for k in formulas:
                lineas[k[1 - eje]].append(k)
            for _, ks in lineas.items():
                cuenta = Counter(firmas[k] for k in ks)
                for f, n in cuenta.items():
                    if n < 2:
                        continue
                    del_patron = sorted(k for k in ks if firmas[k] == f)
                    pos = [k[eje] for k in del_patron]
                    lo, hi = min(pos), max(pos)
                    if n < 3 and hi - lo + 1 > n + 1:
                        continue
                    posiciones = set(pos)
                    for p in range(lo - 1, hi + 2):
                        destino = (p, del_patron[0][1]) if eje == 0 else (del_patron[0][0], p)
                        # Sólo dentro de un tramo contiguo del cálculo: una celda vecina con la misma
                        # fórmula y al menos otra cerca. Totales de bloques separados no forman un tramo.
                        entre = p - 1 in posiciones and p + 1 in posiciones
                        tras_tramo = (p - 1 in posiciones and p - 2 in posiciones) or (p + 1 in posiciones and p + 2 in posiciones)
                        if not (entre or tras_tramo):
                            continue
                        if destino in fijos:
                            origen = min(del_patron, key=lambda k: abs(k[eje] - p))
                            # Con hermanas antes (dentro del tramo o al pie): no es un período viejo tipeado antes de que
                            # empiecen las fórmulas.
                            candidatos[destino].append((origen, f"la {'columna' if eje == 0 else 'fila'} repite el mismo cálculo en {n} celdas",
                                                        n >= HERMANAS_MINIMAS and p > lo))
        # 2. Columnas paralelas: en otras filas ambas columnas calculan lo mismo, acá una quedó fija.
        por_fila = defaultdict(dict)
        for (r, col), f in firmas.items():
            por_fila[r][col] = f
        columnas_con_formula = {col for (_, col) in firmas}
        # Sólo en cuadros chicos de pocas columnas calculadas (estados contables por ejercicio): en
        # cuadros jerárquicos anchos cada columna puede tener su propia estructura.
        minimo_apoyo = 1 if len(columnas_con_formula) <= 3 else None
        for (r, col) in (fijos if minimo_apoyo else ()):
            otras = por_fila.get(r) or {}
            for col2 in otras:
                apoyo = sum(1 for r2, fs in por_fila.items() if r2 != r and col in fs and col2 in fs and fs[col] == fs[col2])
                if apoyo >= minimo_apoyo:
                    candidatos[(r, col)].append(((r, col2), f"en {apoyo} filas las columnas {get_column_letter(col)} y {get_column_letter(col2)} tienen el mismo cálculo",
                                                 apoyo >= HERMANAS_MINIMAS))
                    break
        for (r, col), origenes in candidatos.items():
            fijo = fijos[(r, col)]
            mejor = None
            for origen, motivo, respaldo in origenes:
                from services import interpretacion_previa as ip
                if ip.actual() is not None:
                    rol_destino = ip.rol_celda(archivo, hoja, r, col)
                    rol_origen = ip.rol_celda(archivo, hoja, origen[0], origen[1])
                    if rol_destino not in {"importe", "cantidad", "porcentaje", "indice"} or rol_destino != rol_origen:
                        continue
                try:
                    texto = trasladar(formulas[origen], origen, (r, col))
                except NoEvaluable:
                    continue
                if f"{get_column_letter(col)}{r}" in texto.replace("$", ""):
                    continue  # la fórmula trasladada se referiría a sí misma
                calculado = libro.evaluar_seguro(texto, hoja)
                if calculado is None:
                    continue
                mejor = (origen, motivo, texto, calculado, respaldo)
                break
            if not mejor:
                continue
            origen, motivo, texto, calculado, respaldo = mejor
            terminos = max(1, texto.count(":") * 4 + texto.count("+") + texto.count("-") + 1)
            tol = max(0.5 * 10 ** -_decimales(fijo), 0.005) * terminos + 1e-9 * max(abs(calculado), abs(fijo))
            dif = fijo - calculado
            ref_origen = f"{hoja}!{get_column_letter(origen[1])}{origen[0]}"
            calculo = {"valor_fijo": fijo, "valor_segun_formula": calculado, "diferencia_con_signo": dif,
                       "formula_de_referencia": formulas[origen], "celda_formula_de_referencia": ref_origen,
                       "formula_trasladada": texto, "tolerancia_redondeo": tol, "motivo_patron": motivo}
            if abs(dif) > tol:
                rel = abs(dif) / max(abs(calculado), abs(fijo), 1e-9)
                calculo["diferencia_porcentual"] = rel * 100
                plausible = rel < DIFERENCIA_IMPLAUSIBLE
                enlace = _enlace_en_la_linea(formulas[origen], origen, motivo)
                calculo.update(formulas_hermanas_a_ambos_lados=respaldo, resultado_plausible=plausible,
                               formula_es_enlace=enlace)
                texto_fijo = (f"valor fijo {es_numero(fijo)} donde {motivo}; con la fórmula de {ref_origen} trasladada ({texto}) "
                              f"daría {es_numero(calculado)}. Diferencia {es_numero(dif)} ({rel * 100:.3g} %), mayor que el "
                              f"redondeo admitido ({es_numero(tol, 4)})")
                if respaldo and plausible and not enlace:
                    texto_fijo += ": el valor quedó desactualizado respecto de sus componentes."
                elif enlace:
                    texto_fijo += (". La fórmula de la línea sólo copia otra celda de la misma línea (un enlace, no un "
                                   "cálculo): trasladada apunta a una celda vecina, que no tiene por qué ser el valor que correspondía. "
                                   "Conviene confirmar si la celda debía enlazarse y a cuál.")
                else:
                    # Años viejos tipeados y recientes calculados, o una fórmula que no aplica a esa celda: se pregunta.
                    texto_fijo += (". " + ("La fórmula trasladada da un resultado sin relación con el valor (cerca de cero o "
                                           "con una diferencia de casi el 100 %): puede no aplicar a esta celda. "
                                           if not plausible else
                                           "El valor está antes de que empiece el tramo calculado (o hay pocas fórmulas "
                                           "iguales): puede ser un dato histórico cargado a mano a propósito. ")
                                   + "Conviene confirmar si la celda debía calcularse.")
                fuerte = abs(dif) > 100 * tol or rel >= 0.001
                out.append(_hallazgo(
                    archivo, hoja, r, col, "valor_fijo_distinto_de_formula", texto_fijo,
                    130 if fuerte and respaldo and plausible and not enlace else 75, calculo, [ref_origen]))
            # Un valor fijo que hoy coincide con su fórmula es un riesgo de mantenimiento, no un error:
            # queda para la revisión técnica de fórmulas (opción avanzada), no para el control básico.
    return out
