"""Controles por celda para tablas de CSV y de PDF con texto.

Cada tabla leída de un CSV o de un PDF se vuelca a un libro en memoria (una hoja por tabla) y
pasa por el mismo lector y los mismos controles que un Excel. Las ubicaciones vuelven en
términos del archivo de origen:

- CSV: la hoja lleva el nombre del archivo y la celda es la que mostraría una planilla al abrir
  el CSV (fila N = línea N del archivo contando el encabezado; columna A = primer campo). La
  evidencia agrega «ubicacion_legible»: «fila N («rótulo»), columna X («encabezado»)».
- PDF: una hoja por cuadro extraído («p3 cuadro 1»); la evidencia agrega página, cuadro, fila del
  cuadro y columna en «ubicacion_legible», y la descripción usa esa ubicación.

Nada depende del nombre del archivo ni del tema: sólo de la estructura leída.
"""

from __future__ import annotations

import csv
import io
import re

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

PROHIBIDOS_HOJA = re.compile(r"[\\/?*\[\]:]")
EXTRACCIONES_TABLA = {"tabla", "tabla_coordenadas", "tabla_texto"}
MAX_FILAS = 5000
MAX_COLUMNAS = 200


def _titulo(texto, usados):
    base = PROHIBIDOS_HOJA.sub("_", str(texto)).strip("' ") or "Tabla"
    base = base[:31]
    titulo, n = base, 2
    while titulo.lower() in usados:
        sufijo = f" ({n})"
        titulo = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(titulo.lower())
    return titulo


def rotulo_original(nombre, anteriores):
    """El rótulo tal como figura en el documento: «2022_2» es el segundo «2022» (el sufijo sólo
    distingue columnas repetidas al leer) y se muestra como «2022»."""
    m = re.fullmatch(r"(.+)_(\d+)", nombre)
    if m and m.group(1) in {str(a) for a in anteriores}:
        return m.group(1)
    return nombre


def _anio_rotulo(v):
    if isinstance(v, str) and re.fullmatch(r"\s*(?:18|19|20|21)\d{2}(?:\.0)?\s*", v):
        return int(float(v))
    # La revisión de lectura distingue un rótulo repetido como «2024 · 2024»: el período es 2024
    # y así el control de encabezado lo señala como año duplicado.
    m = isinstance(v, str) and re.fullmatch(r"\s*((?:18|19|20|21)\d{2}) · \1\s*", v)
    return int(m.group(1)) if m else v


def _columnas_numericas(grilla, desde):
    """Por columna: valores convertidos con un único criterio numérico (como el lector de CSV)."""
    from services.perfilado import interpretar_columna, preferencia_formato

    ancho = max((len(f) for f in grilla), default=0)
    todos = [c for f in grilla[desde:] for c in f if isinstance(c, str) and c.strip()]
    preferencia = preferencia_formato(todos)
    salida = [list(f) + [None] * (ancho - len(f)) for f in grilla]
    for j in range(ancho):
        valores = [salida[i][j] if isinstance(salida[i][j], str) and salida[i][j].strip() else None
                   for i in range(desde, len(salida))]
        nuevos, informe = interpretar_columna(valores, preferencia)
        if not informe["numerica"]:
            continue
        for k, v in enumerate(nuevos):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                salida[desde + k][j] = v
    return salida


def _sensibles(encabezado, filas):
    """Columnas que la sanitización excluiría: no entran al libro en memoria."""
    from services import sanitizacion

    out = set()
    for j, nombre in enumerate(encabezado):
        valores = [f[j] if j < len(f) else None for f in filas]
        if sanitizacion.columna_sensible(str(nombre or f"columna_{j + 1}"), valores):
            out.add(j)
    return out


def _texto(v):
    from services.sanitizacion import redactar_texto_libre

    return redactar_texto_libre(v) if isinstance(v, str) else v


def es_lista_de_registros(filas, excluir=()):
    """Una lista de registros (transacciones, beneficiarios, movimientos) no es un cuadro de series:
    sus fechas se repiten y se desordenan dentro de cada categoría. Un cuadro en formato largo,
    aun con un período duplicado o faltante, tiene a lo sumo uno o dos tropiezos por categoría."""
    from .periodos import leer, ordinal

    if len(filas) < 8:
        return False
    ancho = max(len(f) for f in filas)
    columnas = [[f[j] if j < len(f) else None for f in filas] for j in range(ancho)]
    tiempo = None
    for j, col in enumerate(columnas):
        presentes = [v for v in col if v not in (None, "")]
        leidos = [leer(v) for v in presentes if isinstance(v, str)]
        if presentes and len([p for p in leidos if p and p.get("anio")]) >= 0.8 * len(presentes):
            tiempo = j
            break
    if tiempo is None:
        return False
    # Categorías: columnas de texto que se repiten (un nombre o un código por fila identifica
    # registros, no categorías; las columnas sensibles tampoco cuentan).
    texto = [j for j, col in enumerate(columnas) if j != tiempo and j not in excluir
             and sum(isinstance(v, str) and v.strip() != "" for v in col) >= 0.8 * len(col)
             and len(set(col)) <= 0.5 * len(col)]
    grupos = {}
    for f in filas:
        p = leer(f[tiempo]) if tiempo < len(f) and isinstance(f[tiempo], str) else None
        if not p or not p.get("anio"):
            continue
        clave = tuple(f[j] if j < len(f) else None for j in texto)
        grupos.setdefault(clave, []).append((ordinal(p), p.get("dia")))
    tropiezos = []
    for serie in grupos.values():
        if len(serie) < 4:
            continue
        repetidos = len(serie) - len(set(serie))
        retrocesos = sum(1 for a, b in zip(serie, serie[1:]) if b < a)
        tropiezos.append((repetidos + retrocesos) / len(serie))
    if not tropiezos:
        return False
    tropiezos.sort()
    return tropiezos[len(tropiezos) // 2] > 0.2


# --------------------------------------------------------------------------- CSV

def _grilla_csv(contenido):
    from services.carga_archivos import _decodificar, _separador_csv

    texto, _, _ = _decodificar(contenido)
    texto = texto.replace("﻿", "")
    from services.carga_archivos import _comentarios_y_corte_csv, desenvolver_lineas_csv

    texto, _ = desenvolver_lineas_csv(texto)
    sep = _separador_csv(texto)

    # Sin comentarios «#» ni la última línea de un archivo cortado: no son registros (el lector los declara).
    texto, _, _ = _comentarios_y_corte_csv(texto, sep)
    filas = []
    for fila in csv.reader(io.StringIO(texto), delimiter=sep):
        filas.append([c.strip() for c in fila][:MAX_COLUMNAS])
        if len(filas) >= MAX_FILAS:
            break
    while filas and not any(filas[-1]):
        filas.pop()
    return filas


def _hoja_csv(archivo, contenido, usados):
    grilla = _grilla_csv(contenido)
    if len(grilla) < 3:
        return None
    # Primera fila con dos o más campos: el encabezado; lo de arriba son títulos.
    cabecera = next((i for i, f in enumerate(grilla) if sum(1 for c in f if c) >= 2), None)
    if cabecera is None:
        return None
    # Una fila con más campos que el encabezado no entra corrida: si lo que sobra está vacío se
    # recorta; si trae datos, la fila queda afuera (el lector ya la declaró).
    ancho = len(grilla[cabecera])
    grilla = grilla[: cabecera + 1] + [
        [] if len(f) > ancho and any(c for c in f[ancho:]) else f[:ancho] for f in grilla[cabecera + 1:]
    ]
    rotulos = [c.strip() for c in grilla[cabecera]]
    grilla = grilla[: cabecera + 1] + [
        [] if sum(1 for c in f if c) >= 2 and [c.strip() for c in f[: len(rotulos)]] == rotulos else f
        for f in grilla[cabecera + 1:]
    ]
    sensibles = _sensibles(grilla[cabecera], grilla[cabecera + 1:])
    if es_lista_de_registros(grilla[cabecera + 1:], sensibles):
        return None
    valores = _columnas_numericas(grilla, cabecera + 1)
    titulo = _titulo(archivo, usados)
    celdas = []
    for i, fila in enumerate(valores):
        for j, v in enumerate(fila):
            if j in sensibles or v is None or v == "":
                continue
            if i == cabecera:
                v = _anio_rotulo(v)
            celdas.append((i + 1, j + 1, _texto(v)))
    return dict(titulo=titulo, celdas=celdas, tipo="csv", archivo=archivo, filas={}, encabezado=cabecera + 1)


# --------------------------------------------------------------------------- PDF

def _tablas_pdf(archivo, filas, lectura):
    """Agrupa las filas leídas del PDF por cuadro; un cuadro que sigue en la página siguiente
    sin repetir el encabezado (declarado por el lector) es el mismo cuadro."""
    from services.carga_archivos import META_PDF, fila_sin_separar

    heredados = {(h.get("pagina"), h.get("tabla")) for h in (lectura or {}).get("encabezados_heredados") or []}
    grupos, orden = {}, []
    previo = None
    for f in filas:
        if f.get("archivo_origen") != archivo or f.get("extraccion_pdf") not in EXTRACCIONES_TABLA:
            continue
        clave = (f.get("extraccion_pdf"), f.get("pagina_pdf"), f.get("tabla_pdf"))
        columnas = tuple(k for k in f if k not in META_PDF)
        if clave not in grupos:
            if previo is not None and (clave[1], clave[2]) in heredados and grupos[previo]["columnas"] == columnas:
                grupos[clave] = grupos[previo]
            else:
                grupos[clave] = {"columnas": columnas, "filas": [], "pagina": clave[1], "tabla": clave[2]}
                orden.append(clave)
        grupos[clave]["filas"].append(f)
        previo = clave
    salida = []
    for clave in orden:
        g = grupos[clave]
        if len(g["filas"]) < 3 or len(g["columnas"]) < 2:
            continue
        if sum(fila_sin_separar(f) for f in g["filas"]) > 0.3 * len(g["filas"]):
            continue  # columnas pegadas: no hay celdas confiables que controlar
        salida.append(g)
    return salida


def _celda_pdf(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    return str(v)


def _hoja_pdf(archivo, tabla, usados):
    columnas = list(tabla["columnas"])
    encabezado = [None if str(c).startswith("columna_") else rotulo_original(str(c), columnas[:k])
                  for k, c in enumerate(columnas)]
    # Los números ya decididos por el lector (p. ej. una columna en inglés) entran como números:
    # volver a texto «1.234» los expondría a otra lectura del separador.
    grilla = [encabezado] + [[_celda_pdf(f.get(c)) for c in columnas] for f in tabla["filas"]]
    sensibles = _sensibles(columnas, grilla[1:])
    if es_lista_de_registros(grilla[1:], sensibles):
        return None
    valores = _columnas_numericas(grilla, 1)
    titulo = _titulo(f"p{tabla['pagina']} cuadro {tabla['tabla']}", usados)
    celdas = []
    for i, fila in enumerate(valores):
        for j, v in enumerate(fila):
            if j in sensibles or v is None or v == "":
                continue
            celdas.append((i + 1, j + 1, _texto(_anio_rotulo(v) if i == 0 else v)))
    origen = {}
    for i, f in enumerate(tabla["filas"], start=2):
        origen[i] = {"pagina": f.get("pagina_pdf"), "tabla": f.get("tabla_pdf"), "fila": f.get("fila_tabla_pdf"),
                     "pagina_impresa": f.get("pagina_pdf_impresa")}
    return dict(titulo=titulo, celdas=celdas, tipo="pdf", archivo=archivo, filas=origen,
                pagina=tabla["pagina"], tabla=tabla["tabla"], encabezado=1)


# --------------------------------------------------------------------------- libro y controles

def _libro(hojas):
    wb = Workbook()
    wb.remove(wb.active)
    for h in hojas:
        ws = wb.create_sheet(h["titulo"])
        for r, c, v in h["celdas"]:
            ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def hojas_virtuales(leidos, datos, lectura_archivos):
    """Por archivo CSV o PDF leído: las hojas que lo representan en el libro en memoria."""
    from services.carga_archivos import tipo_por_contenido

    estado = {a.get("archivo"): a for a in lectura_archivos or []}
    salida = []
    for archivo, contenido in leidos:
        info = estado.get(archivo) or {}
        if info.get("estado") in {"fallido", "sin_datos"}:
            continue
        try:
            tipo = tipo_por_contenido(archivo, contenido)
        except Exception:  # noqa: BLE001
            continue
        usados: set[str] = set()
        if tipo == "csv":
            h = _hoja_csv(archivo, contenido, usados)
            hojas = [h] if h else []
        elif tipo == "pdf":
            hojas = [h for t in _tablas_pdf(archivo, datos, info) if (h := _hoja_pdf(archivo, t, usados))]
        else:
            continue
        hojas = [h for h in hojas if h["celdas"]]
        if hojas:
            salida.append((archivo, hojas))
    return salida


def analizar(leidos, datos, lectura_archivos, normalizar, sanear):
    """(hallazgos, resumen): controles por celda de cada tabla de CSV o PDF, ubicados en el origen."""
    from services import controles_excel
    from services.excel_cobertura import leer_libro, preparar

    hallazgos, resumen = [], []
    for archivo, hojas in hojas_virtuales(leidos, datos, lectura_archivos):
        try:
            contenido = _libro(hojas)
            filas, meta = leer_libro(contenido, archivo)
            cobertura = {"version": 1, "hojas": meta["hojas"], "bloques": meta["bloques"]}
            if not cobertura["bloques"]:
                continue
            saneados, reporte, _ = preparar(filas, cobertura, normalizar, sanear)
            propios = controles_excel.analizar([(archivo, contenido)], saneados, reporte["excel"])
        except Exception:  # noqa: BLE001 — una tabla que no se puede controlar no frena la carga
            import logging

            logging.getLogger(__name__).exception("controles por celda no aplicados a %s", archivo)
            continue
        por_titulo = {h["titulo"]: h for h in hojas}
        for h in propios:
            ubicar(h, por_titulo)
        hallazgos.extend(propios)
        _declarar_columnas_no_controladas(archivo, meta["hojas"], lectura_archivos)
        estructura = reporte["excel"].get("estructura_revision") or {}
        resumen.append({
            "archivo": archivo,
            "tipo": hojas[0]["tipo"],
            "tablas": len(hojas),
            "cuadros_controlados": len(reporte["excel"]["bloques"]),
            "hallazgos": len(propios),
            # Cuántos cierres se comprobaron y, si ninguno, por qué: el informe lo declara al comienzo.
            "estructura": {k: estructura.get(k) for k in ("cantidad_cierres", "relaciones_candidatas", "motivo_sin_cierres",
                                                         "cierres_sin_comprobar", "cierres_columna", "cierres_verificados",
                                                         "totales_separados", "jerarquia_sin_comprobar",
                                                         "totales_aparentes_sin_comprobar")},
        })
    return hallazgos, resumen



def _declarar_columnas_no_controladas(archivo, hojas, lectura_archivos):
    """Columnas con números que quedaron fuera de los controles (un rótulo que no se reconoce como período,
    columnas auxiliares): en un Excel el informe las toma de la hoja; en un CSV o un PDF, de la lectura del
    archivo, para que el comienzo del informe las diga."""
    lectura = next((a for a in lectura_archivos or [] if a.get("archivo") == archivo), None)
    if lectura is None:
        return
    varias = len(hojas) > 1
    pendientes = lectura.setdefault("pendientes_lectura", [])
    for h in hojas:
        for c in h.get("columnas_no_controladas") or []:
            if not isinstance(c, dict):
                continue
            donde = (f"«{h.get('hoja')}», " if varias else "") + f"columna «{c.get('columna')}» ({c.get('celda') or c.get('letra')})"
            item = dict(tipo="no_controlado", donde=donde, motivo=c.get("motivo") or "columna numérica fuera de los controles")
            if item not in pendientes:
                pendientes.append(item)


def _patron(titulos):
    nombres = "|".join(re.escape(t) for t in sorted(titulos, key=len, reverse=True))
    return re.compile(rf"(?P<hoja>{nombres})!(?P<c1>[A-Z]{{1,3}})(?P<r1>\d+)(?::(?P<c2>[A-Z]{{1,3}})(?P<r2>\d+))?")


def _rotulo(hoja, fila):
    """Primer texto de la fila: el rótulo que la identifica."""
    for r, c, v in hoja["celdas"]:
        if r == fila and isinstance(v, str) and v.strip():
            return v.strip()[:60]
    return None


def _encabezado(hoja, fila, col):
    """Texto de la columna en la fila de encabezado de la tabla (o la más cercana por encima)."""
    propia = hoja.get("encabezado")
    if propia and propia < fila:
        valor = next((v for r, c, v in hoja["celdas"] if r == propia and c == col), None)
        if valor is not None:
            return str(valor)[:40]
    candidatos = {}
    for r, c, v in hoja["celdas"]:
        if c == col and r < fila and v is not None:
            candidatos[r] = v
    for r in sorted(candidatos, reverse=True):
        v = candidatos[r]
        if isinstance(v, str) and not re.fullmatch(r"[-+()\d.,%\s]+", v) or isinstance(v, int) and 1800 <= v <= 2200:
            return str(v)[:40]
    return None


def legible(hoja, letra, fila):
    col = column_index_from_string(letra)
    rotulo = _rotulo(hoja, fila)
    encabezado = _encabezado(hoja, fila, col)
    if hoja["tipo"] == "csv":
        texto = f"fila {fila}" + (f" («{rotulo}»)" if rotulo else "") + f", columna {letra}"
        return texto + (f" («{encabezado}»)" if encabezado else "")
    o = hoja["filas"].get(fila) or {}
    pagina = o.get("pagina") or hoja.get("pagina")
    texto = f"página {pagina}"
    if o.get("pagina_impresa") and o["pagina_impresa"] != pagina:
        texto += f" (impresa {o['pagina_impresa']})"
    texto += f", cuadro {o.get('tabla') or hoja.get('tabla')}"
    if fila == 1:
        valor = next((v for r, c, v in hoja["celdas"] if r == 1 and c == col), None)
        return texto + ", encabezado" + (f", columna «{str(valor)[:40]}»" if valor is not None else f", columna {col}")
    else:
        texto += f", fila {o.get('fila') or fila - 1}" + (f" («{rotulo}»)" if rotulo else "")
    return texto + (f", columna «{encabezado}»" if encabezado else f", columna {col}")


def ubicar(h, por_titulo):
    """Traduce las referencias de la hoja en memoria a la ubicación en el archivo de origen."""
    ev = h.get("evidencia") or {}

    def traducir(m):
        hoja = por_titulo.get(m.group("hoja").strip())
        if not hoja:
            return m.group(0)
        prefijo = hoja["archivo"] + " · " if hoja["tipo"] == "pdf" else m.group("hoja").strip() + " · "
        uno = legible(hoja, m.group("c1"), int(m.group("r1")))
        if m.group("c2") and (m.group("c2"), m.group("r2")) != (m.group("c1"), m.group("r1")):
            return prefijo + uno + " a " + legible(hoja, m.group("c2"), int(m.group("r2")))
        return prefijo + uno

    hoja = por_titulo.get(ev.get("hoja"))
    if hoja:
        ev["origen_tabla"] = hoja["tipo"]
        celda = re.match(r"^([A-Z]{1,3})(\d+)$", str(ev.get("celda") or ""))
        if celda:
            ev["ubicacion_legible"] = (hoja["archivo"] + " · " if hoja["tipo"] == "pdf" else ev["hoja"] + " · ") + legible(
                hoja, celda.group(1), int(celda.group(2)))
            if hoja["tipo"] == "pdf":
                o = hoja["filas"].get(int(celda.group(2))) or {}
                ev["pagina_pdf"] = o.get("pagina") or hoja.get("pagina")
                ev["tabla_pdf"] = o.get("tabla") or hoja.get("tabla")
        if hoja["tipo"] == "csv":
            ev["referencia_csv"] = "Celda como la muestra una planilla al abrir el CSV: fila = línea del archivo, A = primer campo."
    if h.get("descripcion"):
        h["descripcion"] = _patron(por_titulo).sub(traducir, h["descripcion"])
    return h
