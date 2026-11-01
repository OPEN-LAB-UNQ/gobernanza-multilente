"""Variantes de presentación: misma aritmética, lectura y detección esperadas."""

import argparse
import io
import json
import math
import random
import sys
from pathlib import Path

from openpyxl import Workbook

VARIANTES = ("base", "desplazado", "rotulo", "fila_vacia", "laterales", "apilados", "aislados", "hojas")


def libro(semilla, variante, alterado=False):
    rng = random.Random(semilla)
    w = Workbook()
    ws = w.active
    ws.title = "Cuadro"
    col = rng.randint(2, 5) if variante == "desplazado" else 1
    fila = rng.randint(2, 5) if variante == "desplazado" else 1
    anio = rng.randint(1990, 2020)
    valores = [rng.randint(20, 100), rng.randint(110, 200)]
    crecimiento = rng.uniform(1.04, 1.15)
    rotulo = "Descripción" if variante == "rotulo" else "Concepto"

    def escribir(hoja, row, column):
        for j, x in enumerate([rotulo] + list(range(anio, anio + 6)), column):
            hoja.cell(row, j, x)
        for k, (nombre, base) in enumerate(
            [("Salud", valores[0]), ("Educación", valores[1]), ("Total", sum(valores))], 1
        ):
            hoja.cell(row + k, column, nombre)
            for j in range(6):
                hoja.cell(row + k, column + j + 1, base * crecimiento**j)

    escribir(ws, fila, col)
    series = 3
    if variante == "laterales":
        escribir(ws, fila, col + 9)
        series = 6
    if variante == "apilados":
        escribir(ws, fila + 7, col)
        series = 6
    if variante == "hojas":
        otra = w.create_sheet("Otra")
        escribir(otra, fila, col)
        w._sheets.reverse()
        series = 6
    if variante == "aislados":
        ws.cell(fila, col + 17, 1980)
        ws.cell(fila, col + 20, 1981)
    target = ws.cell(fila + 3, col + 3)
    diferencia = target.value * 0.35 if alterado else 0
    if alterado:
        target.value *= 1.35
    if variante == "fila_vacia":
        ws.insert_rows(fila + 2)
    celda = ws.cell(fila + 4 if variante == "fila_vacia" else fila + 3, col + 3).coordinate
    data = io.BytesIO()
    w.save(data)
    return data.getvalue(), {
        "series": series,
        "celda": celda,
        "hoja": "Cuadro",
        "alterado": alterado,
        "diferencia": diferencia,
    }


def evaluar(contenido, esperado):
    from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion

    leidos = [("sintetico.xlsx", contenido)]
    d, cv = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, cv, lambda x: (x, {}), sanitizacion.sanear)
    cov = controles_excel.cobertura_lectura(leidos, d, r["excel"])
    hs = controles_excel.analizar(leidos, d, r["excel"])
    errores = []
    if (
        cov.get("series_leidas") != esperado["series"]
        or cov.get("series_aparentes") != esperado["series"]
        or cov.get("estado") != "completa"
    ):
        errores.append("Cobertura o cantidad de series distinta de la verdad terreno")
    if not esperado["alterado"] and hs:
        errores.append("Alertas sobre un libro aritméticamente correcto")
    # La conciliación puede ser primaria o adjunta a otra señal de la misma celda.
    evidencias = [ev for h in hs for ev in [h["evidencia"], *h["evidencia"].get("conciliaciones", [])]]
    if esperado["alterado"] and not any(
        ev.get("regla") == "conciliacion_aritmetica"
        and ev.get("hoja") == esperado["hoja"]
        and ev.get("celda") == esperado["celda"]
        and math.isclose(ev.get("diferencia_absoluta", 0), esperado["diferencia"], rel_tol=1e-9, abs_tol=1e-7)
        for ev in evidencias
    ):
        errores.append("No se detectó la conciliación alterada en su celda")
    return {"errores": errores, "cobertura": cov, "hallazgos": hs}


def ejecutar(inicio, semillas, salida):
    from versiones import ANALYSIS_VERSION, version_codigo

    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=False)
    casos = []
    fallos = []
    for semilla in range(inicio, inicio + semillas):
        for variante in VARIANTES:
            for alterado in (False, True):
                data, verdad = libro(semilla, variante, alterado)
                try:
                    resultado = evaluar(data, verdad)
                except Exception as e:
                    resultado = {"errores": [f"{type(e).__name__}: {e}"]}
                nombre = f"{semilla}-{variante}-{int(alterado)}"
                casos.append({"caso": nombre, "verdad": verdad, "errores": resultado["errores"]})
                if resultado["errores"]:
                    fallos.append(nombre)
                    p = salida / "fallos" / nombre
                    p.mkdir(parents=True)
                    (p / "archivo.xlsx").write_bytes(data)
                    (p / "resultado.json").write_text(
                        json.dumps({"verdad": verdad, **resultado}, ensure_ascii=False, indent=2)
                    )
    resumen = {
        "motor": ANALYSIS_VERSION,
        "huella": version_codigo(),
        "inicio": inicio,
        "semillas": semillas,
        "variantes": list(VARIANTES),
        "casos": len(casos),
        "fallos": len(fallos),
        "casos_fallidos": fallos,
        "alcance": "Datos sintéticos con verdad conocida. No estima precisión en cualquier informe real ni llama a IA.",
    }
    (salida / "resumen.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2))
    (salida / "casos.json").write_text(json.dumps(casos, ensure_ascii=False, indent=2))
    print(json.dumps(resumen, ensure_ascii=False), flush=True)
    return 1 if fallos else 0


def main():
    # Mismo aislamiento que el ensayo de controles: no cargar .env ni la base real.
    import os
    import shutil
    import subprocess
    import tempfile

    from verificar_calidad import CARPETAS, ROOT

    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--inicio", type=int, default=0)
    p.add_argument("--semillas", type=int, default=4)
    p.add_argument("--salida", type=Path, required=True)
    args = p.parse_args()
    if args.inicio < 0 or not 1 <= args.semillas <= 100:
        p.error("inicio >=0; semillas entre 1 y 100")
    destino = args.salida.resolve()
    if destino.exists():
        p.error("Elegir un directorio nuevo para conservar las campañas anteriores")
    env = {k: os.environ[k] for k in ("PATH", "LANG", "LC_ALL", "SYSTEMROOT") if k in os.environ}
    env.update(DATABASE_URL="sqlite://", SECRET_KEY="test-only", GOBERNANZA_PASSWORD="test-only")
    with tempfile.TemporaryDirectory(prefix="gobernanza-estructuras-") as tmp:
        stage = Path(tmp)
        for nombre in CARPETAS:
            if (ROOT / nombre).exists():
                shutil.copytree(
                    ROOT / nombre, stage / nombre, ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.log")
                )
        for patron in ("*.py", "requirements*.txt", "pyproject.toml"):
            for archivo in ROOT.glob(patron):
                shutil.copy2(archivo, stage / archivo.name)
        r = subprocess.run(
            [
                sys.executable,
                "-c",
                "from scripts.ensayo_estructuras import ejecutar; import sys; sys.exit(ejecutar(int(sys.argv[1]),int(sys.argv[2]),sys.argv[3]))",
                str(args.inicio),
                str(args.semillas),
                str(destino),
            ],
            cwd=stage,
            env=env,
        )
    return r.returncode


if __name__ == "__main__":
    sys.exit(main())
