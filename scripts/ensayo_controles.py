#!/usr/bin/env python3
"""Laboratorio reproducible de errores sembrados; sin base, IA ni datos de usuarios."""

import argparse
import hashlib
import io
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

VERSION = 1


def guardar_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def libro(semilla, vertical=False):
    from openpyxl import Workbook

    rng = random.Random(semilla)
    w = Workbook()
    s = w.active
    s.title = "Serie " + str(rng.randrange(100000))
    fila, col = rng.randint(1, 6), rng.randint(1, 4)
    anio = rng.randint(1990, 2005)
    escala = rng.choice([1, 10, 1000])
    valores = []
    for _ in range(3):
        base = rng.uniform(200, 800) * escala
        crecimiento = rng.uniform(0.002, 0.025)
        valores.append([round(base * (1 + crecimiento) ** j * rng.uniform(0.999, 1.001), 4) for j in range(24)])
    valores.append([sum(v[j] for v in valores) for j in range(24)])
    etiquetas = ["Rubro A", "Rubro B", "Rubro C", "Total"]
    celdas, periodos = [], []
    if vertical:
        s.cell(fila, col, "Año")
        for i, etiqueta in enumerate(etiquetas):
            s.cell(fila, col + i + 1, etiqueta)
        for j in range(24):
            periodos.append(s.cell(fila + j + 1, col, anio + j).coordinate)
        for i, serie in enumerate(valores):
            celdas.append([s.cell(fila + j + 1, col + i + 1, v).coordinate for j, v in enumerate(serie)])
    else:
        s.cell(fila, col, "Concepto")
        for j in range(24):
            periodos.append(s.cell(fila, col + j + 1, anio + j).coordinate)
        for i, serie in enumerate(valores):
            s.cell(fila + i + 1, col, etiquetas[i])
            celdas.append([s.cell(fila + i + 1, col + j + 1, v).coordinate for j, v in enumerate(serie)])
    return w, {"hoja": s.title, "celdas": celdas, "periodos": periodos, "indice": rng.randint(3, 6)}


def bytes_libro(w):
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def casos(semilla, vertical):
    """La verdad se escribe al mutar, sin consultar reglas o respuestas del motor."""
    from openpyxl import load_workbook

    original, mapa = libro(semilla, vertical)
    base = bytes_libro(original)
    specs = [
        ("vacio", "vacio_interior"),
        ("texto", "cambio_tipo"),
        ("signo", "signo_aislado"),
        ("por_diez", "pico_aislado"),
        ("entre_diez", "pico_aislado"),
        ("duplicado", "anio_duplicado"),
        ("anio_fraccionario", "anio_no_entero"),
        ("conciliacion", "conciliacion_aritmetica"),
    ]

    def mutar(w, tipo, regla, i=0, j=None):
        s = w.active
        j = mapa["indice"] if j is None else j
        celda = mapa["periodos"][j] if tipo in ("duplicado", "anio_fraccionario") else mapa["celdas"][i][j]
        anterior = s[celda].value
        if tipo == "vacio":
            nuevo = None
        elif tipo == "texto":
            nuevo = "s/d"
        elif tipo == "signo":
            nuevo = -anterior
        elif tipo == "por_diez":
            nuevo = anterior * 10
        elif tipo == "entre_diez":
            nuevo = anterior / 10
        elif tipo == "duplicado":
            nuevo = s[mapa["periodos"][j - 1]].value
        elif tipo == "anio_fraccionario":
            nuevo = anterior + 0.5
        else:
            nuevo = anterior * 1.25
        s[celda] = nuevo
        return dict(
            tipo=tipo,
            regla=regla,
            hoja=s.title,
            celda=celda,
            original=anterior,
            alterado=nuevo,
            localizacion="participante" if tipo == "conciliacion" else "directa",
            diferencia=abs(nuevo - anterior) if tipo == "conciliacion" else None,
        )

    yield "original", base, base, []
    for tipo, regla in specs:
        w = load_workbook(io.BytesIO(base))
        verdad = [mutar(w, tipo, regla)]
        yield tipo, base, bytes_libro(w), verdad
    w = load_workbook(io.BytesIO(base))
    combinadas = [
        ("duplicado", "anio_duplicado", 0, 23),
        ("por_diez", "pico_aislado", 0, 3),
        ("signo", "signo_aislado", 1, 7),
        ("vacio", "vacio_interior", 1, 10),
        ("texto", "cambio_tipo", 2, 12),
        ("entre_diez", "pico_aislado", 2, 15),
        ("conciliacion", "conciliacion_aritmetica", 1, 19),
    ]
    verdad = [mutar(w, *args) for args in combinadas]
    yield "siete_combinados", base, bytes_libro(w), verdad
    for tipo in (
        "negativos_validos",
        "ceros_validos",
        "bordes_vacios",
        "numeros_textuales",
        "salto_compartido",
        "formula_sin_cache",
    ):
        w = load_workbook(io.BytesIO(base))
        s = w.active
        for i, cs in enumerate(mapa["celdas"]):
            for j, c in enumerate(cs):
                if tipo == "negativos_validos":
                    s[c] = -s[c].value
                elif tipo == "ceros_validos":
                    s[c] = 0
                elif tipo == "bordes_vacios" and j in (0, 23):
                    s[c] = None
                elif tipo == "numeros_textuales":
                    s[c] = str(s[c].value).replace(".", ",")
                elif tipo == "salto_compartido" and j == mapa["indice"]:
                    s[c] = s[c].value * 10
                elif tipo == "formula_sin_cache" and i == 0 and j == mapa["indice"]:
                    # Fórmula sin valor guardado que da el mismo valor: el caso sigue siendo válido.
                    s[c] = f"={s[c].value}+0"
        yield tipo, base, bytes_libro(w), []


def analizar(data):
    from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion

    leidos = [("ensayo.xlsx", data)]
    filas, c = carga_archivos.parse_archivos_con_cobertura(leidos)
    filas, meta, _ = excel_cobertura.preparar(filas, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar(leidos, filas, meta["excel"]), {
        "filas": len(filas),
        "hojas": len(meta["excel"]["hojas"]),
        "cuadros": len(meta["excel"]["bloques"]),
    }


def evaluar(hallazgos, verdad):
    controles = [h for h in hallazgos if h["evidencia"].get("naturaleza") == "control"]
    recuperados = []
    for d in verdad:
        matches = []
        for n, h in enumerate(controles, 1):
            ev = h["evidencia"]
            evidencias = [ev] + ev.get("conciliaciones", [])
            for e in evidencias:
                if e.get("regla") != d["regla"]:
                    continue
                if d["localizacion"] == "directa":
                    ok = e.get("hoja") == d["hoja"] and e.get("celda") == d["celda"]
                else:
                    ok = d["hoja"] + "!" + d["celda"] in e.get("celdas_participantes", []) and abs(
                        e.get("diferencia_absoluta", 0) - d["diferencia"]
                    ) <= max(1e-6, d["diferencia"] * 1e-9)
                if ok:
                    matches.append(n)
        recuperados.append({**d, "detectado": bool(matches), "puesto": min(matches) if matches else None})
    errores = []
    if not verdad and controles:
        errores.append("Observaciones de control en un caso válido conocido")
    if any(not x["detectado"] for x in recuperados):
        errores.append("Mutación no recuperada")
    if any(x["puesto"] and x["puesto"] > 20 for x in recuperados):
        errores.append("Mutación fuera de los primeros 20 controles")
    return recuperados, errores, len(controles)


def ejecutar(inicio, semillas, salida):
    from versiones import ANALYSIS_VERSION, version_codigo

    salida = Path(salida)
    salida.mkdir(parents=True, exist_ok=False)
    resultados = []
    tipos = Counter()
    detectados = Counter()
    for seed in range(inicio, inicio + semillas):
        for vertical in (False, True):
            for nombre, base, data, verdad in casos(seed, vertical):
                clave = f"{seed:06d}-" + ("vertical-" if vertical else "horizontal-") + nombre
                hs = []
                cobertura = {}
                recuperados = []
                errores = []
                n = 0
                try:
                    hs, cobertura = analizar(data)
                    recuperados, errores, n = evaluar(hs, verdad)
                    if cobertura["hojas"] != 1 or cobertura["cuadros"] != 1:
                        errores.append("Cobertura inesperada")
                except Exception as exc:
                    errores = [f"{type(exc).__name__}: {exc}"]
                for d in verdad:
                    tipos[d["tipo"]] += 1
                for d in recuperados:
                    if d["detectado"]:
                        detectados[d["tipo"]] += 1
                row = dict(
                    caso=clave,
                    semilla=seed,
                    orientacion="vertical" if vertical else "horizontal",
                    familia=nombre,
                    sha256=hashlib.sha256(data).hexdigest(),
                    esperados=len(verdad),
                    detectados=sum(d["detectado"] for d in recuperados),
                    controles=n,
                    cobertura=cobertura,
                    errores=errores,
                )
                resultados.append(row)
                if errores:
                    caso = salida / "fallos" / clave
                    caso.mkdir(parents=True)
                    (caso / "original.xlsx").write_bytes(base)
                    (caso / "variante.xlsx").write_bytes(data)
                    guardar_json(
                        caso / "evidencia.json",
                        dict(resumen=row, verdad=verdad, recuperacion=recuperados, hallazgos=hs),
                    )
    fallos = sum(bool(r["errores"]) for r in resultados)
    resumen = dict(
        version_generador=VERSION,
        huella_generador=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        motor=ANALYSIS_VERSION,
        huella_motor=version_codigo(),
        inicio=inicio,
        semillas=semillas,
        casos=len(resultados),
        casos_validos=sum(r["esperados"] == 0 for r in resultados),
        casos_validos_con_alertas=sum(r["esperados"] == 0 and r["controles"] > 0 for r in resultados),
        mutaciones=sum(tipos.values()),
        detectadas=sum(detectados.values()),
        fallos=fallos,
        por_tipo={k: dict(esperados=v, detectados=detectados[k]) for k, v in sorted(tipos.items())},
        alcance="Datos sintéticos conocidos; no estima precisión en informes reales. No ejecuta proveedores IA ni modifica el motor.",
    )
    guardar_json(salida / "resumen.json", resumen)
    guardar_json(salida / "casos.json", resultados)
    (salida / "resumen.md").write_text(
        "# Ensayo automático de controles\n\n"
        + f"Motor: {ANALYSIS_VERSION}. Semillas: {inicio}–{inicio + semillas - 1}.\n\n"
        + f"{len(resultados)} casos; {sum(detectados.values())}/{sum(tipos.values())} mutaciones recuperadas; {fallos} casos con fallas.\n\n"
        + "| Tipo | Esperados | Detectados |\n| --- | ---: | ---: |\n"
        + "".join(f"| {k} | {v} | {detectados[k]} |\n" for k, v in sorted(tipos.items()))
        + "\nLos casos válidos miden falsas alertas sólo en este corpus. Las conciliaciones identifican una relación rota, no necesariamente la celda causal.\n\nCada falla conserva original, variante, verdad terreno y respuesta del motor en `fallos/`.\n",
        encoding="utf-8",
    )
    print(json.dumps(resumen, ensure_ascii=False), flush=True)
    return 1 if fallos else 0


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--inicio", type=int, default=0)
    p.add_argument("--semillas", type=int, default=32)
    p.add_argument("--salida", type=Path, required=True, help="Directorio nuevo; nunca sobrescribe otra campaña")
    args = p.parse_args()
    if args.inicio < 0 or not 1 <= args.semillas <= 10000:
        p.error("Usar inicio >=0 y entre 1 y 10000 semillas")
    destino = args.salida.resolve()
    if destino.exists():
        p.error("La salida ya existe; elegir otro directorio")
    # Exactamente la misma barrera de aislamiento del control de calidad.
    from verificar_calidad import CARPETAS, ROOT

    env = {k: os.environ[k] for k in ("PATH", "LANG", "LC_ALL", "SYSTEMROOT") if k in os.environ}
    env.update(DATABASE_URL="sqlite://", SECRET_KEY="test-only", GOBERNANZA_PASSWORD="test-only")
    with tempfile.TemporaryDirectory(prefix="gobernanza-mutaciones-") as tmp:
        stage = Path(tmp)
        for nombre in CARPETAS:
            if (ROOT / nombre).exists():
                shutil.copytree(
                    ROOT / nombre, stage / nombre, ignore=shutil.ignore_patterns("__pycache__", "*.pyc", "*.log")
                )
        for patron in ("*.py", "requirements*.txt", "pyproject.toml"):
            for archivo in ROOT.glob(patron):
                shutil.copy2(archivo, stage / archivo.name)
        cmd = [
            sys.executable,
            "-c",
            "from scripts.ensayo_controles import ejecutar; import sys; sys.exit(ejecutar(int(sys.argv[1]),int(sys.argv[2]),sys.argv[3]))",
            str(args.inicio),
            str(args.semillas),
            str(destino),
        ]
        return subprocess.run(cmd, cwd=stage, env=env).returncode


if __name__ == "__main__":
    raise SystemExit(main())
