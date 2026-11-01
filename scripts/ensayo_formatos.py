"""Casos adicionales de representación, ausencia de datos y fórmulas."""

import io
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook

from scripts.ensayo_estructuras import evaluar
from services import carga_archivos, controles_excel, excel_cobertura, formula_integridad, sanitizacion

VARIANTES = (
    "encabezados_texto",
    "numeros_es",
    "moneda_formato",
    "titulo_combinado",
    "hoja_oculta",
    "total_primero",
    "vacio",
    "texto",
    "formula",
)


def construir(seed, variante, alterado):
    rng = random.Random(seed)
    w = Workbook()
    s = w.active
    s.title = "Cuadro"
    base = rng.randint(1100, 1700)
    totalfila = 2 if variante == "total_primero" else 4
    rows = (
        [("Total", 3), ("Salud", 1), ("Educación", 2)]
        if variante == "total_primero"
        else [("Salud", 1), ("Educación", 2), ("Total", 3)]
    )
    s.append(["Concepto"] + list(range(2018, 2024)))
    for nombre, mul in rows:
        s.append([nombre] + [(base + j * 13) * mul for j in range(6)])
    truth = {"series": 3, "celda": f"D{totalfila}", "hoja": "Cuadro", "alterado": alterado, "diferencia": 0}
    if variante == "formula":
        for j in range(2, 8):
            col = s.cell(1, j).column_letter
            s.cell(4, j, f"=SUM({col}2:{col}3)")
        if alterado:
            s["D4"] = "=SUM(C2:C3)"
    elif variante in ("vacio", "texto"):
        if alterado:
            s["D2"] = None if variante == "vacio" else "error de carga"
        truth["celda"] = "D2"
    elif alterado:
        cell = s[f"D{totalfila}"]
        original = cell.value
        cell.value = original + 1000
        truth["diferencia"] = 1000
    if variante == "encabezados_texto":
        for c in s[1][1:]:
            c.value = " " + str(c.value) + " "
    if variante == "numeros_es":
        for row in s.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.value = f"{c.value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if variante == "moneda_formato":
        for row in s.iter_rows(min_row=2, min_col=2):
            for c in row:
                c.number_format = '"$" #,##0.00'
    if variante == "titulo_combinado":
        s.insert_rows(1, 2)
        s.merge_cells("A1:G1")
        s["A1"] = "Ejecución presupuestaria · en pesos"
        truth["celda"] = "D6"
    if variante == "hoja_oculta":
        s.sheet_state = "hidden"
        w.create_sheet("Portada")["A1"] = "El cuadro está en la hoja Cuadro."
    data = io.BytesIO()
    w.save(data)
    return data.getvalue(), truth


def comprobar(raw, truth, variante):
    if variante == "formula":
        hs = formula_integridad.analizar_archivo("sintetico.xlsx", raw)
        errores = []
        if not truth["alterado"] and hs:
            errores.append("Alertas en fórmulas coherentes")
        if truth["alterado"] and not any(
            h["evidencia"].get("celda") == "D4" and h["evidencia"].get("regla") == "formula_fuera_de_patron" for h in hs
        ):
            errores.append("No detectó fórmula desplazada en D4")
        return {"errores": errores, "hallazgos": hs}
    if variante not in ("vacio", "texto") or not truth["alterado"]:
        return evaluar(raw, truth)
    leidos = [("sintetico.xlsx", raw)]
    d, cv = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, cv, lambda x: (x, {}), sanitizacion.sanear)
    hs = controles_excel.analizar(leidos, d, r["excel"])
    return {
        "errores": []
        if any(h["evidencia"].get("celda") == truth["celda"] for h in hs)
        else ["No comunicó la celda alterada"],
        "hallazgos": hs,
    }


def ejecutar(salida):
    out = Path(salida)
    out.mkdir(exist_ok=False)
    casos = []
    for seed in (731, 1193, 2053):
        for variante in VARIANTES:
            for alterado in (False, True):
                raw, truth = construir(seed, variante, alterado)
                try:
                    res = comprobar(raw, truth, variante)
                except Exception as exc:
                    res = {"errores": [type(exc).__name__ + ": " + str(exc)]}
                nombre = f"{seed}-{variante}-{int(alterado)}"
                casos.append({"caso": nombre, "errores": res["errores"]})
                if res["errores"]:
                    (out / (nombre + ".xlsx")).write_bytes(raw)
                    (out / (nombre + ".json")).write_text(
                        json.dumps({"verdad": truth, **res}, ensure_ascii=False, indent=2)
                    )
    resumen = {"casos": len(casos), "fallos": sum(bool(c["errores"]) for c in casos), "detalle": casos}
    (out / "resumen.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2))
    print(json.dumps(resumen, ensure_ascii=False))
    return resumen["fallos"] > 0


if __name__ == "__main__":
    sys.exit(ejecutar(sys.argv[1]))
