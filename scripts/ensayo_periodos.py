"""Pares sintéticos con períodos y columnas heterogéneos."""

import io
import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook

from scripts.ensayo_relaciones import evaluar

VARIANTES = (
    "inverso",
    "mezclado",
    "columna_vacia",
    "columna_texto",
    "componente_ausente",
    "periodo_cero",
    "magnitud_grande",
    "encabezado_textual",
)


def construir(seed, variante, alterado):
    rng = random.Random(seed)
    base = rng.randint(1100, 1700)
    years = list(range(2018, 2024))
    if variante == "inverso":
        years.reverse()
    if variante == "mezclado":
        rng.shuffle(years)
    w = Workbook()
    s = w.active
    s.title = "Cuadro"
    s.append(["Concepto"] + [str(y) if variante == "encabezado_textual" else y for y in years])
    scale = 1000000 if variante == "magnitud_grande" else 1
    for label, mult in [("Atención", 1), ("Prevención", 2), ("Total", 3)]:
        s.append([label] + [(base + (y - 2018) * 11) * mult * scale for y in years])
    if variante == "componente_ausente":
        s["G2"] = None
    if variante == "periodo_cero":
        for row in (2, 3, 4):
            s.cell(row, 7, 0)
    target = "C4"
    if variante in ("columna_vacia", "columna_texto"):
        s.insert_cols(3)
        target = "D4"
        if variante == "columna_texto":
            for row, value in enumerate(("Responsable", "Equipo A", "Equipo B", "Dirección"), 1):
                s.cell(row, 3, value)
    delta = 1000 * scale
    if alterado:
        s[target] = s[target].value + delta
    truth = {"hoja": "Cuadro", "celda": target, "alterado": alterado, "diferencia_esperada": delta}
    out = io.BytesIO()
    w.save(out)
    return out.getvalue(), truth


def ejecutar(salida):
    out = Path(salida)
    out.mkdir(parents=True, exist_ok=False)
    resultados = []
    for seed in (4523, 6079, 8291):
        for variante in VARIANTES:
            for alterado in (False, True):
                raw, truth = construir(seed, variante, alterado)
                r = evaluar(raw, truth)
                name = f"{seed}-{variante}-{int(alterado)}"
                resultados.append({"caso": name, "errores": r["errores"]})
                if r["errores"]:
                    (out / (name + ".xlsx")).write_bytes(raw)
                    (out / (name + ".json")).write_text(
                        json.dumps({"verdad": truth, **r}, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
    resumen = {"casos": len(resultados), "fallos": sum(bool(x["errores"]) for x in resultados), "detalle": resultados}
    (out / "resumen.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(resumen, ensure_ascii=False))
    return bool(resumen["fallos"])


if __name__ == "__main__":
    sys.exit(ejecutar(sys.argv[1]))
