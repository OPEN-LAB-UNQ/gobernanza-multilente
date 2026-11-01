"""Pendientes 0.9.4: lectura declarada (formato numérico, cuadros apilados, períodos provisorios,
páginas escaneadas, CSV con campos de más, PDF dañado) y señales calibradas (volatilidad,
triangulación de totales, escala por fila y columna).

Todos los archivos se construyen en el test, con temas variados; nada depende del nombre del
archivo ni del tema, sólo de la estructura y del contenido.
"""

import io
import random

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import models
from services import carga_archivos, casos_revision, controles_excel, excel_cobertura, perfilado
from services.controles_excel import periodos as P
from services.controles_excel import tablas_planas
from services.pdf_controles import _numero_pdf
from services.sanitizacion import sanear
from tests.flujo_confirmado import cargar_confirmado

pytest_plugins = ["tests.test_ejemplos_zip"]

ANIOS = [2019, 2020, 2021, 2022, 2023, 2024]


# Constructores -------------------------------------------------------------------------------

def _xlsx(filas, hoja="Datos"):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila, 1):
            if v is not None:
                ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _serie(rnd, base, n=len(ANIOS), g=(0.97, 1.07)):
    out = []
    for _ in range(n):
        base *= rnd.uniform(*g)
        out.append(round(base))
    return out


def _cuadro(titulo, rotulo, nombres, semilla, base=800, anios=ANIOS, total=True):
    rnd = random.Random(semilla)
    filas = [[titulo], [rotulo] + list(anios)]
    datos = [[n] + _serie(rnd, base * rnd.uniform(0.5, 2), len(anios)) for n in nombres]
    filas += datos
    if total:
        filas.append(["Total"] + [sum(d[j] for d in datos) for j in range(1, len(anios) + 1)])
    return filas


def _en(v, dec=False):
    return f"{v:,.1f}" if dec else f"{v:,}"


def _es(v):
    return f"{v:,}".replace(",", ".")


def _pdf(titulo, encabezado, filas, escaneada=False):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    _, alto = A4
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, alto - 60, titulo)
    c.setFont("Helvetica", 9)
    xs = [50] + [200 + i * 75 for i in range(len(encabezado) - 1)]
    y = alto - 100
    for fila in [encabezado] + filas:
        for i, v in enumerate(fila):
            if i == 0:
                c.drawString(xs[i], y, str(v))
            else:
                c.drawRightString(xs[i] + 60, y, str(v))
        y -= 15
    c.showPage()
    if escaneada:
        from PIL import Image, ImageDraw
        from reportlab.lib.utils import ImageReader

        imagen = Image.new("RGB", (600, 800), "white")
        dibujo = ImageDraw.Draw(imagen)
        for k in range(15):
            dibujo.text((40, 40 + 40 * k), f"Fila {k}   {1200 + 7 * k}   {3400 + 3 * k}", fill="black")
        datos = io.BytesIO()
        imagen.save(datos, "PNG")
        datos.seek(0)
        c.drawImage(ImageReader(datos), 40, 40, 500, 700)
        c.showPage()
    c.save()
    return buf.getvalue()


def _tabla_pdf(nombres, semilla, base, fmt):
    rnd = random.Random(semilla)
    datos = [[n] + _serie(rnd, base * rnd.uniform(0.5, 2)) for n in nombres]
    total = ["Total"] + [sum(d[j] for d in datos) for j in range(1, len(ANIOS) + 1)]
    return datos, [[d[0]] + [fmt(v) for v in d[1:]] for d in datos + [total]]


def _controles_xlsx(nombre, contenido):
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    saneados, reporte, _ = excel_cobertura.preparar(datos, cobertura, lambda f: (f, {}), sanear)
    return controles_excel.analizar([(nombre, contenido)], saneados, reporte["excel"])


def _controles_planos(nombre, contenido):
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([(nombre, contenido)])
    hs, _ = tablas_planas.analizar([(nombre, contenido)], datos, lectura, lambda f: (f, {}), sanear)
    return hs


def _altas(hs):
    return [h for h in hs if h["severidad"] == "alta" and h["evidencia"].get("naturaleza") == "control"]


def _registro(datos, cobertura=None, lectura=None):
    from services.reportes.registro import construir

    saneamiento = {"excel": cobertura or {}, "lectura_archivos": lectura or []}
    informe = models.Informe(id=1, datos=datos, saneamiento=saneamiento, periodo="2026-09")
    return construir(informe, [], {})


# 1. Formato numérico de PDF en inglés y en castellano -----------------------------------------

def test_numero_pdf_lee_formatos_inequivocos():
    assert _numero_pdf("2,078,000.0") == 2078000.0
    assert _numero_pdf("1,234.5") == 1234.5
    assert _numero_pdf("1.234,5") == 1234.5
    assert _numero_pdf("-1,234,567") == -1234567.0
    assert _numero_pdf("12,5") == 12.5


def test_formato_columna_enteros_cortos_y_un_separador_de_tres_cifras():
    assert perfilado.formato_columna(["984", "1,054", "2,310"]) == ("en", "columna")
    assert perfilado.formato_columna(["984", "1.054", "2.310"]) == ("es", "columna")
    # Sin enteros cortos que lo confirmen, «31,821» sigue sin decidirse por la columna.
    assert perfilado.formato_columna(["31,821", "45,002"])[1] == "predeterminado_ambiguo"


def test_pdf_en_ingles_con_valores_inequivocos_se_lee_en_miles():
    nombres = ["North", "South", "East", "West", "Center"]
    datos, filas = _tabla_pdf(nombres, 3, 400000, lambda v: _en(v, dec=True))
    contenido = _pdf("Table 1. Exports by region (USD thousand)", ["Region"] + [str(a) for a in ANIOS], filas)
    leidas, _, lectura = carga_archivos.parse_archivos_con_lectura([("exports.pdf", contenido)])
    norte = next(f for f in leidas if f.get("Region") == "North")
    assert norte["2019"] == pytest.approx(datos[0][1])
    assert not lectura[0].get("formato_ambiguo")
    assert _altas(_controles_planos("exports.pdf", contenido)) == []


def test_pdf_en_ingles_ambiguo_se_declara_y_un_valor_inequivoco_desempata():
    nombres = ["Agriculture", "Mining", "Industry", "Trade", "Transport", "Finance"]
    _, filas = _tabla_pdf(nombres, 7, 40000, _en)
    contenido = _pdf("Table 3. Gross value added (thousand USD)", ["Sector"] + [str(a) for a in ANIOS], filas)
    leidas, _, lectura = carga_archivos.parse_archivos_con_lectura([("gva.pdf", contenido)])
    assert lectura[0]["formato_ambiguo"]
    reg = _registro(leidas, lectura=lectura)
    assert any("formato numérico ambiguo" in p["motivo"] for p in reg["pendientes"])
    assert reg["estado"] != "completo"
    # Un valor de la misma tabla con dos comas decide el formato de todo el cuadro.
    filas[0][1] = _en(1234567)
    contenido = _pdf("Table 3. Gross value added (thousand USD)", ["Sector"] + [str(a) for a in ANIOS], filas)
    leidas, _, lectura = carga_archivos.parse_archivos_con_lectura([("gva.pdf", contenido)])
    assert not lectura[0].get("formato_ambiguo")
    mining = next(f for f in leidas if f.get("Sector") == "Mining")
    assert mining["2020"] > 1000


def test_pdf_en_castellano_sigue_leyendo_punto_de_miles():
    datos, filas = _tabla_pdf(["Norte", "Sur", "Centro", "Cuyo", "Litoral"], 5, 40000, _es)
    contenido = _pdf("Cuadro 2. Exportaciones por región (miles de USD)", ["Región"] + [str(a) for a in ANIOS], filas)
    leidas, _, lectura = carga_archivos.parse_archivos_con_lectura([("expo.pdf", contenido)])
    assert not lectura[0].get("formato_ambiguo")
    assert _altas(_controles_planos("expo.pdf", contenido)) == []
    assert perfilado.numero(next(f for f in leidas if f.get("Región") == "Sur")["2021"]) == datos[1][3]


def test_csv_con_coma_ambigua_avisa_que_la_leyo_como_decimal():
    texto = "Concepto;Valor\nA;31,821\nB;45,002\nC;12,250\n".encode()
    _, _, lectura = carga_archivos.parse_archivos_con_lectura([("v.csv", texto)])
    assert any("coma" in a and "decimal" in a for a in lectura[0]["avisos"])


# 2. Cuadros apilados y salvaguarda de cobertura -------------------------------------------------

def test_segundo_cuadro_sin_fila_vacia_se_separa():
    primero = _cuadro("Cuadro 1. Hogares por región", "Región", ["NOA", "NEA", "Cuyo", "Centro"], 1)
    segundo = _cuadro("x", "Tipo", ["Urbano", "Rural", "Disperso"], 2)[1:]
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("hogares.xlsx", _xlsx(primero + segundo))])
    assert len(cobertura["bloques"]) == 2
    assert cobertura["bloques"][1]["columnas"]["Tipo"] == "A"


@pytest.mark.parametrize("anios", [[2022, 2023, 2024], ANIOS])
def test_segundo_cuadro_con_titulo_pegado_y_otras_columnas_se_separa(anios):
    primero = _cuadro("Cuadro 1. Escuelas por región", "Región", ["NOA", "NEA", "Cuyo", "Centro"], 3)
    segundo = _cuadro("Cuadro 2. Docentes por nivel", "Nivel", ["Inicial", "Primario", "Secundario"], 4, anios=anios)
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("escuelas.xlsx", _xlsx(primero + segundo))])
    assert [list(b["columnas"])[0] for b in cobertura["bloques"]] == ["Región", "Nivel"]


def test_bloque_numerico_fuera_de_los_cuadros_se_declara():
    resumen = [["Resumen"], ["Indicador 1", 410, 512], ["Indicador 2", 300, 377], ["Indicador 3", 211, 245], [None]]
    cuadro = _cuadro("Cuadro 1. Pasajeros por línea", "Línea", ["A", "B", "C", "D"], 5)
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([("pasajeros.xlsx", _xlsx(resumen + cuadro))])
    hoja = cobertura["hojas"][0]
    assert hoja["fuera_de_cuadros"] == [{"filas": "2-4", "columnas": "B-C", "valores": 6}]
    reg = _registro(datos, cobertura)
    assert reg["estado"] == "parcial"
    assert any("filas 2-4" in p["donde"] and "6 valores numéricos" in p["motivo"] for p in reg["pendientes"])


def test_cuadro_completo_no_declara_nada_fuera():
    cuadro = _cuadro("Cuadro 1. Pasajeros por línea", "Línea", ["A", "B", "C", "D"], 5)
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("pasajeros.xlsx", _xlsx(cuadro))])
    assert "fuera_de_cuadros" not in cobertura["hojas"][0]


def test_encabezado_de_periodos_corrido_una_columna_se_declara():
    primero = _cuadro("Cuadro 1. Camas por región", "Región", ["NOA", "NEA", "Cuyo"], 6)
    rnd = random.Random(8)
    segundo = [[None], ["Cuadro 2. Camas por sector"], ["Sector", None] + ANIOS]
    for n in ["Público", "Privado"]:
        segundo.append([n] + _serie(rnd, 300) + [None])
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("camas.xlsx", _xlsx(primero + segundo))])
    declaradas = cobertura["hojas"][0]["columnas_no_controladas"]
    assert any("corrido una columna" in c["motivo"] and c["letra"] == "B" for c in declaradas)


# 3. Períodos provisorios o estimados -----------------------------------------------------------

@pytest.mark.parametrize("rotulo", ["2025 (p)", "2025p", "2025*", "2025 prov.", "2025 (e)", "2025 est.", "2025 prel.", "2025 (1)"])
def test_marcas_de_periodo_provisorio(rotulo):
    p = P.leer(rotulo)
    assert p and p["anio"] == 2025 and p["frecuencia"] == "anual" and p["marca"]


@pytest.mark.parametrize("rotulo", ["2025 (p)", "2025*", "2025 est."])
def test_columna_provisoria_entra_en_los_controles(rotulo):
    rnd = random.Random(12)
    filas = [["Toneladas por puerto"], ["Puerto"] + ANIOS + [rotulo]]
    for n in ["Rosario", "Bahía Blanca", "Quequén", "San Lorenzo"]:
        filas.append([n] + _serie(rnd, rnd.uniform(500, 2000), 7, g=(0.98, 1.05)))
    filas[3][7] *= 1000
    hs = _controles_xlsx("puertos.xlsx", _xlsx(filas))
    assert any(h["evidencia"]["celda"] == "H4" for h in _altas(hs))


def test_anio_provisorio_en_columna_de_periodos_no_es_error_y_se_controla():
    rnd = random.Random(14)
    filas = [["Turistas por región (miles)"], ["Año", "Norte", "Sur", "Centro"]]
    bases = [300.0, 500.0, 800.0]
    for anio in [2018, 2019, 2020, 2021, 2022, 2023, 2024, "2025 (p)"]:
        bases = [b * rnd.uniform(0.98, 1.05) for b in bases]
        filas.append([anio] + [round(b, 1) for b in bases])
    hs = _controles_xlsx("turismo.xlsx", _xlsx(filas))
    assert not [h for h in hs if h["evidencia"]["regla"] == "anio_no_entero"]
    filas[-1][2] = round(filas[-1][2] * 1000, 1)
    hs = _controles_xlsx("turismo.xlsx", _xlsx(filas))
    assert any(h["evidencia"]["celda"] == "C10" for h in _altas(hs))


def test_periodo_observado_incluye_el_anio_provisorio():
    from services.reportes.markdown import _periodo_observado

    datos = [{"Rubro": "A", "2023": 1.0, "2024": 2.0, "2025 (p)": 3.0}]
    assert _periodo_observado(datos).startswith("2023 a 2025")
    assert _periodo_observado([{"Año": "2025 (p)", "x": 1}, {"Año": 2024, "x": 2}]) == "2024 a 2025"


# 4. Página escaneada dentro de un PDF con texto -----------------------------------------------

def test_pagina_escaneada_cuenta_como_no_leida():
    _, filas = _tabla_pdf(["Inicial", "Primario", "Secundario", "Superior", "Adultos"], 9, 4000, _es)
    contenido = _pdf("Cuadro 5. Matrícula por nivel", ["Nivel"] + [str(a) for a in ANIOS], filas, escaneada=True)
    leidas, _, lectura = carga_archivos.parse_archivos_con_lectura([("matricula.pdf", contenido)])
    assert lectura[0]["paginas_sin_texto"] == [2]
    reg = _registro(leidas, lectura=lectura)
    assert any(p["tipo"] == "no_leido" and p["donde"] == "página 2" and "escaneada" in p["motivo"] for p in reg["pendientes"])
    from services.controles_excel.lectura import lineas_pdf, resumen_pdf

    texto = lineas_pdf(resumen_pdf({"cantidad_archivos": 1, "paginas_documento": 2, "paginas_analizadas": 2,
                                    "paginas_sin_texto": 1, "filas_pdf": 30, "modos_extraccion": ["tabla_coordenadas"]}))[0]
    assert "se leyeron 1 de 2 páginas" in texto and "escaneada" in texto


# 5. CSV con filas de más campos que el encabezado ----------------------------------------------

def _csv_provincias(extra=None, vacio_final=False):
    rnd = random.Random(21)
    lineas = ["Provincia,2019,2020,2021,2022,2023,2024,Total"]
    for k, p in enumerate(["Salta", "Jujuy", "Tucumán", "Catamarca", "La Rioja", "Formosa"]):
        v = _serie(rnd, rnd.uniform(1000, 3000))
        campos = [p] + [str(x) for x in v] + [str(sum(v))]
        if extra == k:
            campos = campos[:3] + ["123"] + campos[3:]
        linea = ",".join(campos) + ("," if vacio_final and k == 2 else "")
        lineas.append(linea)
    return ("\n".join(lineas) + "\n").encode()


def test_csv_fila_con_campo_de_mas_no_entra_corrida():
    contenido = _csv_provincias(extra=3)
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([("prov.csv", contenido)])
    assert "Catamarca" not in [f["Provincia"] for f in datos]
    assert lectura[0]["filas_descartadas"] == [5]
    assert any("fila 5: 9 campos" in a and "(8)" in a for a in lectura[0]["avisos"])
    hs = _controles_planos("prov.csv", contenido)
    assert not [h for h in hs if str(h["evidencia"].get("celda", "")).endswith("5")]
    reg = _registro(datos, lectura=lectura)
    assert any(p["donde"] == "fila 5" and "más campos" in p["motivo"] for p in reg["pendientes"])


def test_csv_campo_vacio_de_mas_al_final_se_repara():
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([("prov.csv", _csv_provincias(vacio_final=True))])
    assert len(datos) == 6 and not lectura[0].get("filas_descartadas")


# 6. PDF dañado -----------------------------------------------------------------------------------

def _pdf_danado():
    _, filas = _tabla_pdf(["A", "B", "C"], 1, 1000, _es)
    return _pdf("Cuadro 1. Prueba", ["Rubro"] + [str(a) for a in ANIOS], filas)[:500] + b"\x00roto" * 40


def test_pdf_danado_solo_nombra_el_archivo_y_el_motivo():
    with pytest.raises(carga_archivos.ArchivoNoSoportado) as error:
        carga_archivos.parse_archivos_con_lectura([("anexo.pdf", _pdf_danado())])
    assert "anexo.pdf" in str(error.value) and "PDF" in str(error.value)


def test_pdf_danado_en_carga_mixta_se_declara_y_el_resto_se_analiza():
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura(
        [("anexo.pdf", _pdf_danado()), ("prov.csv", _csv_provincias())])
    assert len(datos) == 6
    fallido = next(x for x in lectura if x["archivo"] == "anexo.pdf")
    assert fallido["estado"] == "fallido" and "anexo.pdf" in fallido["motivo"]


def test_pdf_danado_en_la_carga_rapida_dice_cual(flujo):  # noqa: F811
    # Desde 0.9.7 un PDF no se analiza directo: la carga lo deriva a la revisión de lectura, con su nombre.
    client, _, _ = flujo
    r = cargar_confirmado(client, "/api/analisis-rapido", data={"periodo": "2026-09", "notificar": "false"},
                    files={"archivos": ("anexo.pdf", _pdf_danado(), "application/pdf")})
    detalle = r.json()["detail"]
    assert r.status_code == 422 and detalle["codigo"] == "requiere_revision_lectura" and detalle["archivos"] == ["anexo.pdf"]


# 7. Volatilidad: sin falsas alarmas en series irregulares correctas ----------------------------

def _mensual(semilla, error=None, uniforme=True):
    rnd = random.Random(semilla)
    filas = [["Lluvia mensual por estación (mm)"], ["Mes", "Estación A", "Estación B", "Estación C", "Estación D"]]
    bases = [5000.0, 3000.0, 800.0, 400.0]
    for i in range(36):
        if uniforme:
            fila = [round(rnd.uniform(5, 200), 1) for _ in range(4)]
        else:
            bases = [b * rnd.uniform(0.98, 1.04) for b in bases]
            fila = [round(b, 1) for b in bases]
        filas.append([f"{2021 + i // 12}-{i % 12 + 1:02d}"] + fila)
    if error:
        filas[18][2] = round(filas[18][2] * error, 1)
    return _xlsx(filas)


@pytest.mark.parametrize("semilla", [31, 32, 33, 34, 35, 36])
def test_series_uniformes_aleatorias_sin_alta(semilla):
    # Saltos de pico, escala o compensación: ninguno alto. (Un valor repetido exacto en dos
    # meses seguidos es otra regla —valor copiado— y puede darse por azar con un decimal.)
    altas = [h for h in _altas(_controles_xlsx("lluvia.xlsx", _mensual(semilla)))
             if h["evidencia"]["regla"] != "valor_repetido"]
    assert altas == []


def test_series_volatiles_siguen_detectando_errores_grandes():
    hs = _controles_xlsx("lluvia.xlsx", _mensual(31, error=1000))
    assert any(h["evidencia"]["celda"] == "C19" for h in _altas(hs))
    hs = _controles_xlsx("lluvia.xlsx", _mensual(31, error=-1))
    assert any(h["evidencia"]["celda"] == "C19" for h in _altas(hs))


@pytest.mark.parametrize("factor", [10, 1000])
def test_serie_estable_con_error_de_escala_sigue_alta(factor):
    hs = _controles_xlsx("recaudacion.xlsx", _mensual(77, error=factor, uniforme=False))
    assert any(h["evidencia"]["celda"] == "C19" and h["evidencia"]["regla"] == "pico_aislado" for h in _altas(hs))


def test_csv_anual_aleatorio_sin_alta():
    rnd = random.Random(5)
    lineas = ["Provincia,2019,2020,2021,2022,2023,2024"]
    for p in ["Salta", "Jujuy", "Tucumán", "Catamarca", "La Rioja", "Santiago", "Formosa", "Chaco"]:
        lineas.append(p + "," + ",".join(str(rnd.randint(1000, 9000)) for _ in range(6)))
    assert _altas(_controles_planos("p.csv", ("\n".join(lineas) + "\n").encode())) == []


# 8 y 9. Totales por fila y columna; escala respecto de fila y columna ---------------------------

NIVELES = ["Primaria", "Secundaria", "Superior", "Adultos", "Especial"]
PROVINCIAS = ["Salta", "Jujuy", "Tucumán", "Catamarca", "La Rioja", "Santiago", "Formosa", "Chaco"]


def _matriz(error=None, total_fila=True, columnas=NIVELES):
    rnd = random.Random(5)
    niveles = [rnd.uniform(0.3, 3) for _ in columnas]
    orig = [[round(1200 * f * n * rnd.uniform(0.9, 1.1)) for n in niveles] for f in [rnd.uniform(0.2, 4) for _ in PROVINCIAS]]
    vals = [list(r) for r in orig]
    if error:
        i, j, cambio = error
        vals[i][j] = cambio(vals[i][j])
    filas = [["Matrícula por nivel y provincia"], ["Provincia"] + list(columnas) + ["Total"]]
    filas += [[p] + v + [sum(o)] for p, v, o in zip(PROVINCIAS, vals, orig)]
    if total_fila:
        tot = [sum(o[j] for o in orig) for j in range(len(columnas))]
        filas.append(["Total"] + tot + [sum(tot)])
    return filas


@pytest.mark.parametrize("columnas", [NIVELES, ANIOS])
def test_componente_alterado_se_ubica_en_la_interseccion(columnas):
    filas = _matriz((3, 2, lambda v: v + 97), columnas=columnas)
    hs = _controles_xlsx("matricula.xlsx", _xlsx(filas))
    tri = [h for h in hs if h["evidencia"]["regla"] == "componente_no_cierra_fila_y_columna"]
    assert len(tri) == 1 and tri[0]["evidencia"]["celda"] == "D6"
    total = get_column_letter(len(columnas) + 2)
    assert {f"Datos!{total}6", "Datos!D11"} <= set(tri[0]["evidencia"]["celdas_participantes"])
    casos = casos_revision.agrupar(hs)["casos"]
    assert casos[0]["primary_finding"]["evidencia"]["celda"] == "D6"


def test_componente_alterado_en_csv_y_pdf_se_ubica_en_la_celda():
    filas = _matriz((3, 2, lambda v: v + 97))[1:]
    csv = ("\n".join(",".join(str(x) for x in f) for f in filas) + "\n").encode()
    hs = _controles_planos("matricula.csv", csv)
    assert any(h["evidencia"]["regla"] == "componente_no_cierra_fila_y_columna" and h["evidencia"]["celda"] == "D5"
               for h in _altas(hs))
    pdf = _pdf("Cuadro 4. Matrícula por nivel y provincia", filas[0], [[f[0]] + [_es(x) for x in f[1:]] for f in filas[1:]])
    hs = _controles_planos("matricula.pdf", pdf)
    assert any(h["evidencia"]["regla"] == "componente_no_cierra_fila_y_columna" and h["evidencia"]["celda"] == "D5"
               for h in _altas(hs))


def test_sin_fila_total_se_senala_la_fila_completa():
    filas = _matriz((3, 2, lambda v: v + 97), total_fila=False)
    hs = _controles_xlsx("matricula.xlsx", _xlsx(filas))
    fila = [h for h in hs if h["evidencia"]["regla"] == "fila_no_cierra_total"]
    assert fila and fila[0]["evidencia"]["ubicaciones_excel"] == ["Datos!B6:G6"]


def test_escala_por_mil_en_una_celda_se_senala_en_la_celda_sin_fila_total():
    filas = _matriz((5, 3, lambda v: v * 1000), total_fila=False)
    hs = _controles_xlsx("matricula.xlsx", _xlsx(filas))
    escala = [h for h in _altas(hs) if h["evidencia"]["regla"] == "valor_fuera_de_escala_fila_columna"]
    assert escala and escala[0]["evidencia"]["celda"] == "E8"


@pytest.mark.parametrize("formato", ["xlsx", "csv", "pdf"])
def test_matriz_correcta_sin_alta(formato):
    filas = _matriz()
    if formato == "xlsx":
        hs = _controles_xlsx("m.xlsx", _xlsx(filas))
    elif formato == "csv":
        hs = _controles_planos("m.csv", ("\n".join(",".join(str(x) for x in f) for f in filas[1:]) + "\n").encode())
    else:
        hs = _controles_planos("m.pdf", _pdf("Cuadro 4. Matrícula", filas[1], [[f[0]] + [_es(x) for x in f[1:]] for f in filas[2:]]))
    assert _altas(hs) == []
