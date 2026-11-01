import io

import pytest
from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion
from services.controles_excel import lectura
from services.prioridad_controles import consolidar_contexto, portada


def analizar_libro(w):
    b = io.BytesIO()
    w.save(b)
    return analizar_bytes(b.getvalue())


def analizar_bytes(b, nombre="libro.xlsx"):
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, b)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, b)], d, r["excel"])


def caso(valores, anios=None, vertical=False):
    w = Workbook()
    s = w.active
    s.title = "Cuadro de prueba"
    anios = anios or list(range(2000, 2000 + len(valores)))
    if vertical:
        s.append(["Año", "Volumen"])
        for a, v in zip(anios, valores):
            s.append([a, v])
    else:
        s.append(["Concepto", *anios])
        s.append(["Volumen", *valores])
    return w


@pytest.mark.parametrize("vertical", [False, True])
@pytest.mark.parametrize(
    "valor,regla",
    [
        (None, "vacio_interior"),
        ("s/d", "cambio_tipo"),
        ("N/D", "cambio_tipo"),
        ("-", "cambio_tipo"),
        ("NA", "cambio_tipo"),
        ("otro", "cambio_tipo"),
        (-130, "signo_aislado"),
        (1300, "pico_aislado"),
        (13, "pico_aislado"),
    ],
)
def test_estructuras_y_escalas_en_ambas_orientaciones(valor, regla, vertical):
    w = caso([100, 110, 120, valor, 140, 150, 160], vertical=vertical)
    hs = analizar_libro(w)
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["regla"] == regla
    assert ev["valor_observado"] == valor
    assert ev["valor_anterior"] == 120 and ev["valor_posterior"] == 140
    assert ev["puntaje_prioridad"] >= 80


@pytest.mark.parametrize(
    "anios,regla",
    [
        ([2000, 2001, 2002, 2002, 2004, 2005], "anio_duplicado"),
        ([2000, 2001, 2003, 2002, 2004, 2005], "anios_no_crecientes"),
        ([2000, 2001, 2002, 2003.5, 2004, 2005], "anio_no_entero"),
        ([2000, 2001, 2002, 2004, 2005, 2006], "discontinuidad_anual"),
    ],
)
def test_periodos_invalidos_explicitos(anios, regla):
    hs = analizar_libro(caso([100, 110, 120, 130, 140, 150], anios))
    assert regla in {h["evidencia"]["regla"] for h in hs}
    assert not any(h["evidencia"]["regla"] == "pico_aislado" for h in hs)


@pytest.mark.parametrize(
    "valores",
    [
        [None, 100, 110, 120, 130, 140],
        [100, 110, 120, 130, 140, None],
        [100, 110, 120, 130, 140, 150],
        [-10, -9, -8, -7, -6, -5],
        [0, 0, 0, 0, 0, 0],
    ],
)
def test_no_confundir_bordes_tendencia_y_negativos_habituales(valores):
    assert analizar_libro(caso(valores)) == []


def test_jerarquia_sin_doble_conteo_y_redondeo():
    w = Workbook()
    s = w.active
    s.append(["Concepto", 2000, 2001, 2002, 2003, 2004, 2005])
    for fila in [["Rama", 30] * 1, ["  Parte A", 10], ["  Parte B", 20], ["Otra rama", 40], ["Total", 70.01]]:
        s.append([fila[0]] + [fila[1]] * 6)
    assert analizar_libro(w) == []
    celda = s.cell(3, 4)
    celda.value = 12.5
    hs = analizar_libro(w)
    assert hs and any(h["evidencia"]["regla"] == "conciliacion_aritmetica" for h in hs)
    assert any(f"{s.title}!{celda.coordinate}" in h["evidencia"]["celdas_participantes"] for h in hs)
    ev = hs[0]["evidencia"]
    assert ev["diferencia_absoluta"] == 2.5


def test_movimiento_compartido_no_es_error_puntual():
    w = caso([100, 110, 120, 1300, 140, 150, 160])
    s = w.active
    for etiqueta in ("Segunda", "Tercera", "Cuarta"):
        s.append([etiqueta, 100, 110, 120, 1300, 140, 150, 160])
    hs = analizar_libro(w)
    assert len(hs) == 4 and all(h["evidencia"]["naturaleza"] == "contexto" for h in hs)
    assert portada(hs) == []
    grupos = consolidar_contexto(hs)
    assert len(grupos) == 1 and len(grupos[0]["evidencia"]["detalle_contexto"]) == 4


def test_importe_absoluto_no_aumenta_prioridad():
    base = [100, 110, 120, 1300, 140, 150, 160]
    a = analizar_libro(caso(base))
    b = analizar_libro(caso([v * 10000 for v in base]))
    assert a[0]["evidencia"]["puntaje_prioridad"] == b[0]["evidencia"]["puntaje_prioridad"]


def test_formula_sin_cache_no_es_vacio_eliminado():
    w = caso([100, 110, 120, "=120+10", 140, 150, 160])
    assert analizar_libro(w) == []


def test_numero_textual_convertible_no_es_cambio_de_tipo():
    assert analizar_libro(caso([100, 110, 120, "130,00", 140, 150, 160])) == []


def test_fila_de_titulo_y_separacion_no_generan_faltantes():
    w = caso([100, 110, 120, 130, 140, 150])
    s = w.active
    s.append([])
    s.append(["Nota: datos agregados"])
    s.append([])
    assert analizar_libro(w) == []


def test_no_reintroduce_columna_personal_excluida():
    w = Workbook()
    s = w.active
    s.append(["Año", "DNI", "Cantidad"])
    for j in range(7):
        s.append([2000 + j, 12345678 + j, 100 + j])
    hs = analizar_libro(w)
    assert hs == []


def test_extremo_moderado_requiere_conciliacion_independiente():
    valores = [100, 110, 120, 130, 140, 800]
    assert analizar_libro(caso(valores)) == []
    w = caso(valores)
    s = w.active
    s.append(["Componente B", 20, 21, 22, 23, 24, 25])
    s.append(["Total", 120, 131, 142, 153, 164, 175])
    hs = analizar_libro(w)
    assert any(h["evidencia"]["regla"] == "escala_extremo" and h["evidencia"]["conciliaciones"] for h in hs)


def _cuadro_corto(total_2021):
    w = Workbook()
    s = w.active
    s.append(["Categoría", 2019, 2020, 2021, 2022])
    s.append(["Educación", 100.0, 120.0, 150.0, 190.0])
    s.append(["Salud", 50.0, 60.0, 70.0, 90.0])
    s.append(["Otros", 10.0, 12.0, 15.0, 19.0])
    s.append(["Total", 160.0, 192.0, total_2021, 299.0])
    return w, s


def test_serie_de_cuatro_periodos_detecta_total_alterado():
    w, s = _cuadro_corto(235.0 * 1.07)
    hs = [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "conciliacion_aritmetica"]
    assert hs and hs[0]["evidencia"]["celda"] == "D5"


def test_serie_de_cuatro_periodos_sana_no_genera_conciliacion():
    w, s = _cuadro_corto(235.0)
    assert not [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "conciliacion_aritmetica"]


def test_serie_de_cuatro_con_dos_desvios_no_afirma_relacion():
    w, s = _cuadro_corto(235.0 * 1.07)
    s["C5"] = 192.0 * 1.3
    assert not [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "conciliacion_aritmetica"]


def _indicadores(millones, por_nino):
    w = Workbook()
    s = w.active
    s.append(["Indicador", 2016, 2017, 2018, 2019, 2020, 2021])
    s.append(["En millones de $", *millones])
    s.append(["En $ por niño/a", *por_nino])
    return w, s


def _con_poblacion(hs):
    """Hallazgos propios de la regla o corroborados por ella (no se duplica la alarma)."""
    return [
        h
        for h in hs
        if h["evidencia"]["regla"] == "poblacion_implicita"
        or any(c["regla"] == "poblacion_implicita" for c in h["evidencia"].get("conciliaciones", []))
    ]


MILLONES = [100.0, 130.0, 170.0, 250.0, 330.0, 500.0]
POBLACION = [1_000_000, 1_010_000, 1_020_000, 1_030_000, 1_040_000, 1_050_000]


def test_per_capita_dividido_por_diez_se_detecta_en_la_celda_correcta():
    por_nino = [m * 1e6 / p for m, p in zip(MILLONES, POBLACION)]
    por_nino[3] /= 10
    w, s = _indicadores(MILLONES, por_nino)
    hs = _con_poblacion(analizar_libro(w))
    assert len(hs) == 1 and hs[0]["evidencia"]["celda"] == "E3"
    relacionadas = [hs[0]["evidencia"]] + hs[0]["evidencia"].get("conciliaciones", [])
    assert any(f"{s.title}!E2" in r["celdas_participantes"] for r in relacionadas)


def test_per_capita_del_ultimo_anio_se_detecta():
    por_nino = [m * 1e6 / p for m, p in zip(MILLONES, POBLACION)]
    por_nino[-1] /= 10
    w, s = _indicadores(MILLONES, por_nino)
    hs = _con_poblacion(analizar_libro(w))
    assert hs and hs[0]["evidencia"]["celda"] == "G3"


def test_per_capita_coherente_no_genera_observacion():
    por_nino = [m * 1e6 / p for m, p in zip(MILLONES, POBLACION)]
    w, s = _indicadores(MILLONES, por_nino)
    assert not _con_poblacion(analizar_libro(w))


def test_poblacion_inestable_no_permite_afirmar_la_relacion():
    poblacion = [1_000_000, 1_400_000, 900_000, 1_500_000, 1_000_000, 1_300_000]
    por_nino = [m * 1e6 / p for m, p in zip(MILLONES, poblacion)]
    w, s = _indicadores(MILLONES, por_nino)
    assert not _con_poblacion(analizar_libro(w))


@pytest.mark.parametrize(
    ("valores", "regla"),
    [([100, "s/d", 120, 130], "cambio_tipo"), ([100, None, 120, 130], "vacio_interior"), ([100, -110, 120, 130], "signo_aislado")],
)
def test_series_de_cuatro_periodos_conservan_controles_de_estructura(valores, regla):
    w = caso(valores, [2019, 2020, 2021, 2022])
    w.active.append(["Otro rubro", 50, 55, 60, 65])  # el período tiene datos en el cuadro
    hs = analizar_libro(w)
    assert regla in {h["evidencia"]["regla"] for h in hs}


@pytest.mark.parametrize("valores", [[100, 110, 120, 130], [100, 110, 120, None], [100, 1100, 120, 130]])
def test_series_de_cuatro_sin_error_de_estructura_no_se_marcan(valores):
    hs = analizar_libro(caso(valores, [2019, 2020, 2021, 2022]))
    assert not {h["evidencia"]["regla"] for h in hs} & {"cambio_tipo", "vacio_interior", "signo_aislado"}


@pytest.mark.parametrize(("valores", "hay"), [
    ([201, 297, 504, 9751, 2449], True),  # x10 entre vecinos que crecen fuerte (inflación)
    ([433, 545, 888, 174, 4244], True),  # /10 en el mismo contexto
    ([201, 297, 504, 975, 2449], False),  # la serie original: crecimiento alto, sin error
    ([43, 49, 76, 208, 307], False),
    ([5, 40, 2, 0.2, 12], False),  # serie chica y volátil: no se afirma error
])
def test_pico_aislado_con_inflacion_alta(valores, hay):
    w = caso(valores, [2019, 2020, 2021, 2022, 2023])
    w.active.append(["Otro rubro", 50, 55, 60, 65, 70])
    hs = [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "pico_aislado" and h["evidencia"]["celda"] == "E2"]
    assert bool(hs) == hay


@pytest.mark.parametrize(("valores", "hay"), [
    ([1048, 1539, 2031, 3029, 3251, 7272, 12838, 210793], True),  # x10 al final con crecimiento previo
    ([207, 258, 382, 1200, 31667], True),  # x10 al final con inflación alta
    ([1048, 1539, 2031, 3029, 3251, 7272, 12838, 21079], False),  # serie original
])
def test_extremo_final_con_crecimiento_previo(valores, hay):
    w = caso(valores, list(range(2023 - len(valores) + 1, 2024)))
    w.active.append(["Otro rubro", *[50 + i for i in range(len(valores))]])
    hs = [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "escala_extremo"]
    assert bool(hs) == hay


CORRIDA = [56.12, 41.11, 54.36, 95.24, 95.24, 136.52, 222.88, 316.61, 318.95, 398.09, 594.89, 879.33]
ORIGINAL = [56.12, 41.11, 54.36, 95.24, 136.52, 222.88, 316.61, 318.95, 398.09, 594.89, 879.33, 687.66]


def _repetidos(valores):
    return [h for h in analizar_libro(caso(valores)) if h["evidencia"]["regla"] == "valor_repetido"]


def test_valor_repetido_en_dos_periodos_seguidos_se_senala():
    hs = _repetidos(CORRIDA)
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["celda"] == "F2" and ev["periodo"] == 2004
    assert any(c.endswith("!E2") for c in ev["celdas_participantes"])  # también el período anterior


def test_serie_sin_repeticiones_no_se_marca():
    assert not _repetidos(ORIGINAL)


@pytest.mark.parametrize("valores", [
    [0.072, 0.076, 0.070, 0.069, 0.065, 0.065, 0.064],  # porcentajes redondeados: coinciden por azar
    [66.9, 66.9, 71.2, 74.8, 69.3, 72.1, 75.4],  # tres dígitos significativos: no alcanza
    [100, 200, 300, 400, 400, 500, 700, 800],  # valores redondos, sin resolución
    [56.12, 56.12, 54.36, 95.24, 95.24, 136.52, 222.88, 316.61],  # dos repeticiones: no se afirma
    [1050.1, 1050.2, 1050.2, 1050.4, 1050.1, 1050.6, 1050.2],  # serie casi quieta
])
def test_repeticiones_que_no_alcanzan_para_afirmar_nada(valores):
    assert not _repetidos(valores)


PROVINCIAS = ["Buenos Aires", "Catamarca", "Chaco", "Chubut", "Córdoba", "Corrientes", "Entre Ríos",
              "Formosa", "Jujuy", "La Pampa"]
ANIOS_P = [2016, 2017, 2018, 2019, 2020, 2021]


def _pais(ajuste=None, ruido=1.0):
    """Un libro con diez hojas provinciales de la misma estructura y distinta escala."""
    w = Workbook()
    w.remove(w.active)
    for n, provincia in enumerate(PROVINCIAS):
        escala = 1 + n * 0.7  # cada provincia tiene su tamaño
        s = w.create_sheet(provincia)
        s.append(["En millones de $"])
        s.append(["Concepto", *ANIOS_P])
        filas = {
            "Salud": [400, 520, 660, 900, 1180, 1500],
            "Educación": [300, 380, 500, 640, 810, 1030],
            "Ciencia y técnica": [12 * (ruido**n), 15 * (ruido**n), 19 * (ruido**n), 24 * (ruido**n),
                                  31 * (ruido**n), 40 * (ruido**n)],
        }
        if ajuste:
            filas = ajuste(provincia, filas)
        s.append(["Total", *[sum(x) * escala for x in zip(*filas.values())]])
        for nombre, vs in filas.items():
            s.append([nombre, *[v * escala for v in vs]])
    return w


ENTRE_HOJAS = ("escala_entre_hojas", "rubros_intercambiados", "referencia_entre_hojas")


def _entre_hojas(w):
    return [h for h in analizar_libro(w) if h["evidencia"]["regla"] in ENTRE_HOJAS]


def test_un_rubro_en_otra_unidad_se_ve_contra_las_demas_hojas():
    def dividir(provincia, filas):
        if provincia == "Chaco":
            filas["Educación"] = [v / 1000 for v in filas["Educación"]]
        return filas

    hs = _entre_hojas(_pais(dividir))
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["regla"] == "escala_entre_hojas" and ev["hoja"] == "Chaco" and ev["indicador"] == "Educación"
    assert ev["factor_desvio"] > 500 and ev["hojas_comparadas"] == 9


def test_dos_rubros_intercambiados_se_nombran_juntos():
    def intercambiar(provincia, filas):
        if provincia == "Jujuy":
            filas["Salud"], filas["Ciencia y técnica"] = filas["Ciencia y técnica"], filas["Salud"]
        return filas

    hs = _entre_hojas(_pais(intercambiar))
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["regla"] == "rubros_intercambiados" and ev["hoja"] == "Jujuy"
    assert {ev["indicador"], ev["serie_intercambiada"]} == {"Salud", "Ciencia y técnica"}
    celdas = " ".join(ev["celdas_participantes"])
    assert "Jujuy!B4" in celdas and "Jujuy!B6" in celdas  # las dos filas quedan señaladas


def _pais_con_per_capita(ajuste=None):
    """Igual que _pais, con una fila por persona: comparable en su propio valor entre hojas."""
    w = _pais()
    for n, provincia in enumerate(PROVINCIAS):
        s = w[provincia]
        valores = [1000 + 40 * n + 90 * i for i in range(len(ANIOS_P))]
        if ajuste:
            valores = ajuste(provincia, valores)
        s.append(["En $ por niño/a", *valores])
    return w


def test_un_valor_por_persona_fuera_de_escala_se_ve_contra_las_demas_hojas():
    def dividir(provincia, valores):
        return [v / 1000 for v in valores] if provincia == "Formosa" else valores

    hs = _entre_hojas(_pais_con_per_capita(dividir))
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["hoja"] == "Formosa" and ev["indicador"] == "En $ por niño/a"
    assert ev["serie_referencia"] is None  # se compara con el valor de las otras hojas, sin referencia


def test_valores_por_persona_parecidos_entre_hojas_no_generan_observacion():
    assert not _entre_hojas(_pais_con_per_capita())


def test_si_falla_la_fila_de_referencia_se_senala_esa_y_no_cada_rubro():
    w = _pais()
    s = w["Chaco"]
    for fila in s.iter_rows(min_row=3, max_row=3):  # la fila de totales de la hoja
        for celda in fila[1:]:
            celda.value = celda.value / 1000
    hs = _entre_hojas(w)
    assert len(hs) == 1
    ev = hs[0]["evidencia"]
    assert ev["regla"] == "referencia_entre_hojas" and ev["hoja"] == "Chaco"
    assert ev["rubros_afectados"] == ev["rubros_comparados"] == 3


def test_hojas_de_la_misma_estructura_sin_error_no_generan_observaciones():
    assert not _entre_hojas(_pais())


def test_un_rubro_que_ya_varia_mucho_entre_hojas_no_permite_afirmar_nada():
    def dividir(provincia, filas):
        if provincia == "Chaco":
            filas["Ciencia y técnica"] = [v / 1000 for v in filas["Ciencia y técnica"]]
        return filas

    # Ciencia y técnica va de 12 a 12 * 2,2^9 entre provincias: sin patrón, no se afirma nada.
    assert not _entre_hojas(_pais(dividir, ruido=2.2))


def _cuadro_rubros(salud, educacion, total=None):
    w = Workbook()
    s = w.active
    anios = [2016, 2017, 2018, 2019, 2020, 2021]
    s.append(["Categoría", *anios])
    otros = [300, 390, 500, 700, 950, 1300]
    s.append(["Salud", *salud])
    s.append(["Educación", *educacion])
    s.append(["Otros", *otros])
    s.append(["Total", *(total or [a + b + c for a, b, c in zip(salud, educacion, otros)])])
    return w


SALUD = [900, 1170, 1500, 2100, 2850, 3900]
EDUCACION = [4800, 6240, 8000, 11200, 15200, 20800]


def _compensados(w):
    return [
        h
        for h in analizar_libro(w)
        if h["evidencia"]["regla"] == "movimiento_compensado"
        or any(c.get("regla") == "movimiento_compensado" for c in h["evidencia"].get("conciliaciones", []))
    ]


def test_monto_pasado_de_un_rubro_a_otro_se_detecta():
    salud, educacion = SALUD[:], EDUCACION[:]
    salud[3] += 1260  # +60 % en Salud
    educacion[3] -= 1260  # sale de Educación; el total no cambia
    hs = _compensados(_cuadro_rubros(salud, educacion, [a + b + c for a, b, c in zip(SALUD, EDUCACION, [300, 390, 500, 700, 950, 1300])]))
    assert hs
    ev = hs[0]["evidencia"]
    assert {"Sheet!E2", "Sheet!E3"} <= set(ev["celdas_participantes"]) or any(
        {"Sheet!E2", "Sheet!E3"} <= set(c["celdas_participantes"]) for c in ev.get("conciliaciones", [])
    )


def test_series_regulares_no_generan_movimiento_compensado():
    assert not _compensados(_cuadro_rubros(SALUD, EDUCACION))


def test_salto_de_ambos_rubros_en_el_mismo_sentido_no_es_compensado():
    salud, educacion = SALUD[:], EDUCACION[:]
    salud[3] *= 1.6
    educacion[3] *= 1.3
    assert not _compensados(_cuadro_rubros(salud, educacion))


def _edicion(anios, filas):
    w = Workbook()
    s = w.active
    s.title = "Buenos Aires"
    s.append(["Categoría", *anios])
    for nombre, valores in filas:
        s.append([nombre, *valores])
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()


def _analizar_ediciones(anterior, nueva):
    leidos = [("edicion_2022.xlsx", anterior), ("edicion_2023.xlsx", nueva)]
    d, c = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar(leidos, d, r["excel"])


def _ediciones(hs):
    return [
        h
        for h in hs
        if h["evidencia"]["regla"] == "cambio_entre_ediciones"
        or any(c.get("regla") == "cambio_entre_ediciones" for c in h["evidencia"].get("conciliaciones", []))
    ]


SALUD_E = [100.0, 130.0, 170.0, 220.0, 290.0, 380.0]
EDUC_E = [500.0, 640.0, 820.0, 1050.0, 1350.0, 1700.0]


def test_valor_historico_cambiado_entre_ediciones_se_detecta():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    salud = SALUD_E + [500.0]
    salud[2] = 155.0  # 2019 cambia sólo en la edición nueva (-9 %)
    nueva = _edicion(range(2017, 2024), [("Salud", salud), ("Educación", EDUC_E + [2200.0])])
    hs = _ediciones(_analizar_ediciones(anterior, nueva))
    assert hs
    lugares = [hs[0]["evidencia"]] + hs[0]["evidencia"].get("conciliaciones", [])
    assert any("Buenos Aires!D2" in x["celdas_participantes"] for x in lugares)


def test_ediciones_sin_cambios_y_ultimo_anio_provisorio_no_se_marcan():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    salud = SALUD_E[:-1] + [310.0, 500.0]  # 2022 era provisorio en la edición anterior
    nueva = _edicion(range(2017, 2024), [("Salud", salud), ("Educación", EDUC_E + [2200.0])])
    assert not _ediciones(_analizar_ediciones(anterior, nueva))


MOVIDAS = {"serie_de_otro_rubro_entre_ediciones", "series_intercambiadas_entre_ediciones"}


def _movidas(hs):
    return [
        h
        for h in hs
        if h["evidencia"]["regla"] in MOVIDAS
        or any(c.get("regla") in MOVIDAS for c in h["evidencia"].get("conciliaciones", []))
    ]


def _regla(hs):
    ev = hs[0]["evidencia"]
    return ev["regla"] if ev["regla"] in MOVIDAS else next(c["regla"] for c in ev["conciliaciones"] if c["regla"] in MOVIDAS)


def test_dos_rubros_intercambiados_entre_ediciones_se_nombran_juntos():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    nueva = _edicion(range(2017, 2024), [("Salud", EDUC_E + [2200.0]), ("Educación", SALUD_E + [500.0])])
    hs = _movidas(_analizar_ediciones(anterior, nueva))
    assert hs and _regla(hs) == "series_intercambiadas_entre_ediciones"
    celdas = " ".join(c for h in hs for c in h["evidencia"]["celdas_participantes"])
    assert "Buenos Aires!B2" in celdas and "Buenos Aires!B3" in celdas


def test_una_fila_que_toma_la_serie_de_otra_se_senala():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    nueva = _edicion(range(2017, 2024), [("Salud", EDUC_E + [2200.0]), ("Educación", EDUC_E + [2200.0])])
    hs = _movidas(_analizar_ediciones(anterior, nueva))
    assert hs and _regla(hs) == "serie_de_otro_rubro_entre_ediciones"


def test_ediciones_sin_filas_movidas_no_generan_observacion():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    nueva = _edicion(range(2017, 2024), [("Salud", SALUD_E + [500.0]), ("Educación", EDUC_E + [2200.0])])
    assert not _movidas(_analizar_ediciones(anterior, nueva))


def test_una_revision_de_la_serie_no_es_una_fila_movida():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    revisada = [v * 1.1 for v in SALUD_E] + [500.0]
    nueva = _edicion(range(2017, 2024), [("Salud", revisada), ("Educación", EDUC_E + [2200.0])])
    assert not _movidas(_analizar_ediciones(anterior, nueva))


def test_revision_de_toda_la_serie_no_es_un_valor_aislado():
    anterior = _edicion(range(2017, 2023), [("Salud", SALUD_E), ("Educación", EDUC_E)])
    revisada = [v * 1.1 for v in SALUD_E] + [500.0]
    nueva = _edicion(range(2017, 2024), [("Salud", revisada), ("Educación", EDUC_E + [2200.0])])
    assert not _ediciones(_analizar_ediciones(anterior, nueva))


@pytest.mark.parametrize(("desvio", "hay"), [(1.0006, False), (1.01, True)])
def test_conciliacion_ignora_diferencias_de_redondeo(desvio, hay):
    w = Workbook()
    s = w.active
    anios = [2016, 2017, 2018, 2019, 2020, 2021]
    s.append(["Categoría", *anios])
    s.append(["Salud", *[100.0 * (1.2**i) for i in range(6)]])
    s.append(["Educación", *[500.0 * (1.2**i) for i in range(6)]])
    s.append(["Total", *[600.0 * (1.2**i) for i in range(6)]])
    s["D4"] = s["D4"].value * desvio
    hs = [h for h in analizar_libro(w) if h["evidencia"]["regla"] == "conciliacion_aritmetica"]
    assert bool(hs) == hay


def _dos_cuadros(obras_clases, obras_categorias):
    w = Workbook()
    s = w.active
    anios = [2016, 2017, 2018, 2019, 2020, 2021]
    s.append(["Clase", *anios])
    s.append(["GEPI", *[100.0 * (1.3**i) for i in range(6)]])
    s.append(["OO.SS", *obras_clases])
    s.append(["Total", *[(100.0 * (1.3**i)) + o for i, o in enumerate(obras_clases)]])
    s.append([None])
    s.append(["Categoría", *anios])
    s.append(["Educación", *[100.0 * (1.3**i) for i in range(6)]])
    s.append(["Obras sociales", *obras_categorias])
    s.append(["Total", *[(100.0 * (1.3**i)) + o for i, o in enumerate(obras_categorias)]])
    return w


OBRAS = [50.0, 66.0, 88.0, 120.0, 160.0, 210.0]


def _identidades(w):
    return [
        h
        for h in analizar_libro(w)
        if h["evidencia"]["regla"] == "identidad_entre_cuadros"
        or any(c.get("regla") == "identidad_entre_cuadros" for c in h["evidencia"].get("conciliaciones", []))
    ]


def test_misma_serie_en_dos_cuadros_que_deja_de_coincidir():
    desfasada = OBRAS[:]
    desfasada[3] = 150.0
    hs = _identidades(_dos_cuadros(OBRAS, desfasada))
    assert hs and hs[0]["evidencia"]["periodo"] == 2019


def test_series_iguales_en_dos_cuadros_no_generan_observacion():
    assert not _identidades(_dos_cuadros(OBRAS, OBRAS[:]))


def test_series_distintas_en_dos_cuadros_no_se_toman_como_identidad():
    otra = [o * 1.4 for o in OBRAS]
    otra[3] = OBRAS[3]
    assert not _identidades(_dos_cuadros(OBRAS, otra))


def _cobertura(w):
    b = io.BytesIO()
    w.save(b)
    leidos = [("libro.xlsx", b.getvalue())]
    d, c = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.cobertura_lectura(leidos, d, r["excel"])


def test_cobertura_de_lectura_completa_no_agrega_advertencias():
    datos = _cobertura(caso([100, 110, 120, 130, 140, 150]))
    assert datos["estado"] == "completa" and datos["proporcion"] == 1.0
    assert len(controles_excel.lectura.lineas(datos)) == 1


def test_cobertura_de_lectura_insuficiente_explica_el_motivo():
    w = Workbook()
    s = w.active
    s.append(["Categoría", 2020, 2021])  # dos períodos: no alcanza ni para leer la serie
    for nombre in ("Salud", "Educación", "Otros"):
        s.append([nombre, 10, 20])
    datos = _cobertura(w)
    assert datos["estado"] == "insuficiente" and datos["proporcion"] == 0.0
    texto = " ".join(controles_excel.lectura.lineas(datos))
    assert "Cobertura de lectura: 0 %" in texto and "se necesitan 3" in texto
    assert "la ausencia de observaciones no indica que el archivo esté bien" in texto


@pytest.mark.parametrize(("meta", "estado", "frase"), [
    ({"cantidad_archivos": 1, "paginas_documento": 58, "paginas_analizadas": 58, "filas_pdf": 14,
      "modos_extraccion": ["tabla"]}, "insuficiente", "no hay series suficientes"),
    ({"cantidad_archivos": 1, "paginas_documento": 30, "paginas_analizadas": 30, "filas_pdf": 400,
      "modos_extraccion": ["texto_por_pagina"]}, "insuficiente", "se leyó como texto por página"),
    ({"cantidad_archivos": 1, "paginas_documento": 80, "paginas_analizadas": 40, "filas_pdf": 300,
      "modos_extraccion": ["tabla"], "truncado_por_paginas": True}, "parcial", "no fue revisado"),
    ({"cantidad_archivos": 1, "paginas_documento": 20, "paginas_analizadas": 20, "filas_pdf": 300,
      "modos_extraccion": ["tabla"]}, "completa", "se extrajeron 300 filas de datos"),
])
def test_cobertura_de_lectura_de_pdf(meta, estado, frase):
    datos = controles_excel.lectura.resumen_pdf(meta)
    assert datos["estado"] == estado
    assert frase in " ".join(controles_excel.lectura.lineas_pdf(datos))


def test_pdf_sin_metadatos_no_declara_cobertura():
    assert controles_excel.lectura.resumen_pdf({}) == {}
    assert controles_excel.lectura.lineas_pdf({}) == []


def test_lectura_de_pdf_reconoce_tablas_maquetadas_sin_lineas():
    from services.carga_archivos import _filas_texto_alineado_pdf

    texto = "\n".join([
        "TABLA 14. GASTO POR AÑO",
        "Concepto        2018      2019      2020",
        "En millones   252.954   365.035   531.635",
        "Por niño/a     56.316    81.517   119.139",
        "% del PIB         1,7       1,7       1,9",
    ])
    filas = [f for f in _filas_texto_alineado_pdf(texto) if f]
    assert len(filas) == 4 and filas[0][0] == "Concepto" and filas[1][0] == "En millones"


def test_lectura_de_pdf_ignora_los_ejes_de_un_grafico():
    from services.carga_archivos import _filas_texto_alineado_pdf

    texto = "\n".join(["GRÁFICO 5", "0  1  2  3  4", "5  6  7  8  9"])
    assert [f for f in _filas_texto_alineado_pdf(texto) if f] == []


def _meta_pdf(**extra):
    base = {"cantidad_archivos": 1, "paginas_documento": 58, "paginas_analizadas": 58,
            "modos_extraccion": ["tabla"], "filas_pdf": 73}
    return {**base, **extra}


def test_un_cuadro_que_no_se_pudo_separar_en_columnas_no_cuenta_como_leido():
    """Caso real: 5 de 7 cuadros del libro de gasto 2001-2020 quedan con varios números por celda."""
    r = lectura.resumen_pdf(_meta_pdf(filas_pdf_utiles=9, cuadros_pdf=7, cuadros_pdf_sin_separar=5))
    assert r["estado"] == "insuficiente" and r["filas_utiles"] == 9
    texto = " ".join(lectura.lineas_pdf(r))
    assert "5 de los 7 cuadros no pudieron separarse en columnas" in texto
    assert "9 filas revisables" in texto
    assert "la ausencia de observaciones no indica que el documento esté bien" in texto


def test_cuadros_bien_separados_conservan_la_cobertura_completa():
    r = lectura.resumen_pdf(_meta_pdf(filas_pdf_utiles=73, cuadros_pdf=7, cuadros_pdf_sin_separar=0))
    assert r["estado"] == "completa"
    assert "no pudieron separarse" not in " ".join(lectura.lineas_pdf(r))


def test_un_solo_cuadro_sin_separar_deja_la_lectura_parcial():
    r = lectura.resumen_pdf(_meta_pdf(filas_pdf=90, filas_pdf_utiles=70, cuadros_pdf=6, cuadros_pdf_sin_separar=1))
    assert r["estado"] == "parcial"
    assert "1 de los 6 cuadros" in " ".join(lectura.lineas_pdf(r))


@pytest.mark.parametrize(("celda", "varios"), [
    ("994 120 1.115", True),
    ("39.626,2", False),
    ("2001-2020", False),
    ("Condiciones de vida", False),
    ("", False),
])
def test_reconoce_las_celdas_con_varios_numeros(celda, varios):
    assert carga_archivos.celda_con_varios_numeros(celda) is varios


def test_un_cuadro_separado_pero_sin_rotulos_tampoco_queda_revisado():
    """Caso real: reconstruir columnas por coordenadas separa las celdas pero parte el encabezado."""
    r = lectura.resumen_pdf(_meta_pdf(filas_pdf=277, filas_pdf_utiles=0, cuadros_pdf=7,
                                      cuadros_pdf_sin_separar=0, cuadros_pdf_sin_rotulos=7,
                                      modos_extraccion=["tabla_coordenadas"]))
    assert r["estado"] == "insuficiente"
    texto = " ".join(lectura.lineas_pdf(r))
    assert "sin rótulo confiable" in texto and "7 de los 7 cuadros" in texto
    assert "277 filas de datos" in texto and "0 filas revisables" in texto


@pytest.mark.parametrize(("columnas", "dudoso"), [
    (["Nivel de gobierno", "Ayuda directa", "Educación", "Salud"], False),
    (["columna_1", "2016", "columna_4", "2018", "columna_6"], True),
    (["Impuesto al Valor Agregado", "34.758", "39,4", "99.956"], True),
    (["Concepto", "Sociales", "de vida", "alimentación"], True),
])
def test_reconoce_los_rotulos_de_columna_poco_confiables(columnas, dudoso):
    fila = {c: "1" for c in columnas}
    fila["extraccion_pdf"] = "tabla_coordenadas"
    assert carga_archivos.rotulos_poco_confiables(fila) is dudoso


def test_encabezado_de_dos_niveles_compone_el_nombre_de_la_columna():
    """«Primera infancia» arriba y «M $ / %» abajo son una columna cada uno, no un rótulo partido."""
    def palabra(texto, x0, x1, top):
        return {"text": texto, "x0": x0, "x1": x1, "top": top}

    nivel1 = [palabra("Primera", 100, 130, 10), palabra("infancia", 132, 160, 10),
              palabra("Niñez", 200, 225, 10), palabra("y", 227, 232, 10), palabra("adolescencia", 234, 280, 10)]
    nivel2 = [palabra("M", 100, 110, 20), palabra("$", 112, 118, 20), palabra("%", 150, 158, 20),
              palabra("M", 200, 210, 20), palabra("$", 212, 218, 20), palabra("%", 250, 258, 20)]
    franjas = [(100, 118), (150, 158), (200, 218), (250, 258)]
    assert carga_archivos._rotulos_de_columna_pdf([nivel1, nivel2], franjas) == [
        "Primera infancia · M $", "Primera infancia · %", "Niñez y adolescencia · M $", "Niñez y adolescencia · %",
    ]


def test_un_rotulo_cortado_en_dos_lineas_no_se_lee_como_dos_niveles():
    def palabra(texto, x0, x1, top):
        return {"text": texto, "x0": x0, "x1": x1, "top": top}

    filas = [[palabra("Condiciones", 100, 150, 10)], [palabra("de", 100, 110, 20), palabra("vida", 112, 130, 20)]]
    assert carga_archivos._rotulos_de_columna_pdf(filas, [(100, 150)]) == ["Condiciones de vida"]


def test_el_titulo_del_cuadro_no_se_toma_como_encabezado():
    def palabra(texto, x0, x1, top=10):
        return {"text": texto, "x0": x0, "x1": x1, "top": top}

    titulo = [palabra("TABLA", 40, 70), palabra("17.", 72, 85), palabra("GASTO", 87, 105)]
    ancho = [palabra("Categoría", 40, 300)]
    assert carga_archivos._es_titulo_pdf(titulo, 300) is True
    assert carga_archivos._es_titulo_pdf(ancho, 300) is True
    assert carga_archivos._es_titulo_pdf([palabra("Año", 40, 60), palabra("Total", 200, 230)], 300) is False


def test_una_columna_sin_rotulo_que_es_el_porcentaje_de_otra_se_nombra_por_lo_que_hace():
    """Si B es el porcentaje de A sobre un total, 100·A/B da ese total y es el mismo en cada fila."""
    montos = [557.0, 4915.0, 1085.0, 1397.0, 9871.0, 2247.0]
    total = sum(montos)
    registros = [
        {"Rubro": f"Rubro {i}", "2017": monto, "columna_3": round(100 * monto / total, 1),
         "extraccion_pdf": "tabla_coordenadas", "pagina_pdf": 44, "tabla_pdf": 1,
         "fila_tabla_pdf": i, "archivo_origen": "informe.pdf"}
        for i, monto in enumerate(montos, start=1)
    ]
    carga_archivos._nombrar_columnas_por_relacion_pdf(registros)
    columnas = [k for k in registros[0] if k not in carga_archivos.META_PDF]
    assert columnas == ["Rubro", "2017", "2017 · %"]  # el orden de las columnas se conserva


def test_una_columna_sin_rotulo_que_no_es_un_porcentaje_no_se_nombra():
    registros = [
        {"Rubro": f"Rubro {i}", "2017": 100.0 * i, "columna_3": 7.0 * i + 3,
         "extraccion_pdf": "tabla_coordenadas", "pagina_pdf": 44, "tabla_pdf": 1,
         "fila_tabla_pdf": i, "archivo_origen": "informe.pdf"}
        for i in range(1, 7)
    ]
    carga_archivos._nombrar_columnas_por_relacion_pdf(registros)
    assert "columna_3" in registros[0]


def test_un_ano_es_un_rotulo_de_columna_valido():
    fila = {"columna_1": "Rubro", "2016": "1", "2017": "2", "2018": "3", "2019": "4",
            "extraccion_pdf": "tabla_coordenadas"}
    assert carga_archivos.rotulos_poco_confiables(fila) is False
    fila_mala = {"columna_1": "Rubro", "34.758": "1", "39,4": "2", "99.956": "3",
                 "extraccion_pdf": "tabla_coordenadas"}
    assert carga_archivos.rotulos_poco_confiables(fila_mala) is True


@pytest.mark.parametrize(("valores", "esperado"), [
    (["1.738", "9.765", "16.911", "123.726"], ["1738", "9765", "16911", "123726"]),
    (["0,6", "5,2", "1,1", "3,5"], ["0,6", "5,2", "1,1", "3,5"]),
    (["1.234,56", "2.000,10", "3.500,00"], ["1234,56", "2000,10", "3500,00"]),
    (["1.5", "2.25", "3.75"], ["1.5", "2.25", "3.75"]),
])
def test_el_punto_de_miles_de_un_cuadro_de_pdf_no_se_lee_como_decimal(valores, esperado):
    """«1.738» en estos cuadros es mil setecientos treinta y ocho, no uno con setecientos treinta y ocho."""
    registros = [
        {"Rubro": f"R{i}", "Monto": valor, "extraccion_pdf": "tabla_coordenadas", "pagina_pdf": 54,
         "tabla_pdf": 1, "fila_tabla_pdf": i, "archivo_origen": "informe.pdf"}
        for i, valor in enumerate(valores, start=1)
    ]
    carga_archivos._normalizar_miles_pdf(registros)
    assert [r["Monto"] for r in registros] == esperado


def test_una_columna_con_pocos_valores_no_se_toca():
    registros = [{"Monto": "1.738", "extraccion_pdf": "tabla_coordenadas", "pagina_pdf": 1,
                  "tabla_pdf": 1, "fila_tabla_pdf": 1, "archivo_origen": "i.pdf"}]
    carga_archivos._normalizar_miles_pdf(registros)
    assert registros[0]["Monto"] == "1.738"


def test_un_csv_con_punto_de_miles_no_divide_los_importes_por_mil():
    """«9.765» en un CSV argentino es nueve mil setecientos sesenta y cinco."""
    contenido = "rubro;monto;porcentaje\nEducación;9.765;6,1\nSalud;1.781;1,2\nTotal;16.911;100,0\n".encode()
    filas = carga_archivos.parse_archivo("gasto.csv", contenido)
    assert [f["monto"] for f in filas] == [9765, 1781, 16911]
    assert [f["porcentaje"] for f in filas] == [6.1, 1.2, 100.0]


def test_un_csv_exportado_en_ingles_conserva_su_punto_decimal():
    contenido = b"rubro,amount\nEducation,1.5\nHealth,2.25\nTotal,3.75\n"
    filas = carga_archivos.parse_archivo("gasto.csv", contenido)
    assert [f["amount"] for f in filas] == [1.5, 2.25, 3.75]


def test_una_columna_de_texto_con_puntos_no_cambia_los_numeros():
    contenido = "codigo;monto\n8.123.456;1500\n8.123.457;2500\n8.123.458;3500\n".encode()
    filas = carga_archivos.parse_archivo("codigos.csv", contenido)
    assert [f["monto"] for f in filas] == [1500, 2500, 3500]
