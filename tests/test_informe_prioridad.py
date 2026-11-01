"""Prioridad única (rareza, impacto, evidencia) y agrupación por causa común.

Caso P1 de la evaluación del 2026-10-03: un error de carga en Gasto!G6 se propaga a subtotales,
total y a la hoja de variaciones; antes quedaba como «contexto» y 9 de 10 señales caían en la hoja
derivada.
"""
import datetime
import io
from types import SimpleNamespace

from openpyxl import Workbook

import models
from services import carga_archivos, casos_revision, controles_excel, excel_cobertura, guardrails, sanitizacion
from services import lentes as lentes_pkg
from services.excel_cobertura import evaluar_carga

ANIOS = list(range(2017, 2025))
BASE = {"Salud": 4100.0, "Educación": 3650.0, "Promoción social": 5200.0, "Defensa": 1450.0, "Seguridad interior": 1730.0}


def libro_p1():
    vals = {r: {a: round(v * 1.30 ** i, 1) for i, a in enumerate(ANIOS)} for r, v in BASE.items()}
    vals["Salud"][2020] = round(vals["Salud"][2020] * 1.55, 1)
    vals["Educación"][2022] = round(vals["Educación"][2022] * 1.9, 1)  # error que se propaga
    vals["Defensa"][2024] = round(vals["Defensa"][2024] * 9.6, 1)
    social = {a: round(vals["Salud"][a] + vals["Educación"][a] + vals["Promoción social"][a], 1) for a in ANIOS}
    servicios = {a: round(vals["Defensa"][a] + vals["Seguridad interior"][a], 1) for a in ANIOS}
    total = {a: round(social[a] + servicios[a], 1) for a in ANIOS}
    filas = [("Salud", vals["Salud"]), ("Educación", vals["Educación"]), ("Promoción social", vals["Promoción social"]),
             ("Subtotal gasto social", social), ("Defensa", vals["Defensa"]), ("Seguridad interior", vals["Seguridad interior"]),
             ("Subtotal servicios", servicios), ("Total", total)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Gasto"
    ws["A1"] = "Gasto por finalidad - Administración Nacional"
    ws["A2"] = "En millones de pesos"
    ws.cell(4, 1, "Concepto")
    for j, a in enumerate(ANIOS):
        ws.cell(4, 2 + j, a)
    for r, (et, v) in enumerate(filas, 5):
        ws.cell(r, 1, et)
        for j, a in enumerate(ANIOS):
            ws.cell(r, 2 + j, v[a])
    ws2 = wb.create_sheet("Variacion")
    ws2["A1"] = "Variación interanual del gasto por finalidad"
    ws2["A2"] = "En %"
    ws2.cell(4, 1, "Concepto")
    for j, a in enumerate(ANIOS[1:]):
        ws2.cell(4, 2 + j, a)
    for r, (et, v) in enumerate(filas, 5):
        ws2.cell(r, 1, et)
        for j, a in enumerate(ANIOS[1:]):
            ws2.cell(r, 2 + j, round((v[a] / v[a - 1] - 1) * 100, 1))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def correr(nombre, contenido):
    """El mismo recorrido que la carga rápida y el análisis, sin base de datos."""
    leidos = [(nombre, contenido)]
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura(leidos)
    saneados, reporte, _ = excel_cobertura.preparar(datos, cobertura, lambda f: (f, {}), sanitizacion.sanear)
    reporte["controles_tabulares"] = controles_excel.analizar(leidos, saneados, reporte["excel"])
    reporte["excel"]["lectura"] = controles_excel.cobertura_lectura(leidos, saneados, reporte["excel"])
    reporte["version_controles_tabulares"] = 2
    reporte["estructura_revision"] = reporte["excel"].pop("estructura_revision", {})
    reporte.update(archivos_origen=[nombre], opciones_tecnicas={}, ambito="publico")
    informe = models.Informe(id=None, area_id=-1, periodo="2026-10", dominio="general", modo="intraarea", usuario_id=7,
                             creado_por="test", datos=saneados, saneamiento=reporte, creado_en=datetime.datetime(2026, 10, 3))
    assert evaluar_carga(informe)["puede_concluir"]
    crudos = []
    for lente, nuevos in excel_cobertura.ejecutar_por_bloques(None, informe, lentes_pkg._ejecutar_unidad):
        if lente == "_crudos_finales":
            crudos = nuevos
            continue
    crudos = guardrails.normalizar_hallazgos(crudos, informe)
    return informe, [{**h, "id": i} for i, h in enumerate(crudos)]


def celdas(caso):
    return set(caso["referencias"])


def test_p1_la_causa_raiz_encabeza_el_caso_y_la_hoja_derivada_queda_subordinada():
    _, hs = correr("gasto_finalidad_prioridades.xlsx", libro_p1())
    r = casos_revision.agrupar(hs)
    educacion = next(c for c in r["casos"] if "Gasto!G6" in celdas(c))
    principal = educacion["primary_finding"]["evidencia"]
    # La señal principal está en la hoja fuente, no en la tabla de variaciones.
    assert principal["hoja"] == "Gasto" or any(u.startswith("Gasto!") for u in principal.get("ubicaciones_excel", []))
    # Se agrupan reglas distintas de dos hojas del mismo libro y se conservan todas las celdas.
    reglas = {(x.get("evidencia") or {}).get("regla") for x in educacion["related_findings"]}
    assert {"signo_aislado", "pico_aislado"} <= reglas
    assert {"Variacion!G6", "Variacion!G8", "Variacion!G12"} <= celdas(educacion)
    assert educacion["prioridad"] == "alta"
    # Ya no hay diez casos sueltos: la Educación de 2022 es un solo caso, separado de Salud y Defensa.
    assert len(r["casos"]) <= 4
    assert not any("Gasto!G6" in celdas(c) for c in r["casos"] if c is not educacion)


def test_p1_tabla_derivada_marcada_y_subordinada():
    _, hs = correr("gasto_finalidad_prioridades.xlsx", libro_p1())
    derivadas = [h for h in hs if (h["evidencia"].get("tabla_derivada") or {}).get("subordinada")]
    assert derivadas and all(h["evidencia"].get("hoja") == "Variacion" for h in derivadas)
    assert all(casos_revision.ejes(h)["impacto"]["nivel"] == "no_aplica" for h in derivadas)
    assert all(casos_revision.prioridad(h) != "alta" for h in derivadas)


def test_impacto_se_calcula_sobre_el_total_del_cuadro_para_todas_las_senales():
    _, hs = correr("gasto_finalidad_prioridades.xlsx", libro_p1())
    defensa = next(h for h in hs if h["evidencia"].get("hoja") == "Gasto" and h["evidencia"].get("celda") == "I9")
    x = casos_revision.ejes(defensa)
    assert x["impacto"]["nivel"] == "alta" and x["impacto"]["porcentaje"] > 30
    # Evidencia débil (último año, sin confirmación bilateral) pero impacto alto: desde 0.10.13 sigue en ALTA;
    # no poder confirmarlo con un vecino posterior no baja por sí solo la prioridad de lo que pesa en el total.
    assert x["evidencia"]["nivel"] == "patron_debil" and x["prioridad"] == "alta" and x["motivo_tope"] is None


def test_una_sola_politica_para_ejecutivo_y_completo():
    from services.reporte_ejecutivo import lineas
    from services.reportes.markdown import _lineas_priorizadas

    base, hs = correr("gasto_finalidad_prioridades.xlsx", libro_p1())
    informe = SimpleNamespace(id=1, datos=base.datos, saneamiento=base.saneamiento, periodo=base.periodo,
                              creado_en=base.creado_en, devolucion=SimpleNamespace(cobertura={}))
    hallazgos = [SimpleNamespace(**{"tipo_evidencia": None, **h}) for h in hs]
    r = casos_revision.agrupar(hallazgos)
    ejecutivo = "\n".join(lineas(informe, hallazgos))
    completo = "\n".join(_lineas_priorizadas(hallazgos))
    for n, c in enumerate(r["casos"], 1):
        assert f"{c['prioridad'].upper()} · caso {n}" in completo
    # 0.9.6: el ejecutivo presenta consultas (casos agrupados por causa); cada una lleva la prioridad
    # de la misma política (la mayor de sus casos), en el detalle o en la tabla del anexo.
    from services.reportes.consultas import consultas

    lista = consultas(r)
    assert sum(len(c["casos"]) for c in lista) == len(r["casos"])
    for n, c in enumerate(lista, 1):
        assert c["prioridad"] == max((x["prioridad"] for x in c["casos"]), key=casos_revision.NIVELES.get)
        assert f"### {n}. {c['prioridad'].upper()} · " in ejecutivo or f"| {n} | {c['prioridad']} |" in ejecutivo
    assert "puntos:" not in completo  # el puntaje del control ya no se presenta como prioridad
    # Los tres ejes se conservan en el completo; el ejecutivo no los repite en cada caso.
    assert "Rareza:" not in ejecutivo and "Rareza:" in completo and "Impacto:" in completo and "Evidencia:" in completo


def test_ejes_explicitos_y_coherentes():
    h = dict(lente="ajedrecista", modo="control_tabular", severidad="alta", descripcion="x",
             evidencia=dict(regla="conciliacion_aritmetica", hoja="H", celda="B5", periodo=2020, diferencia_porcentual=0.05,
                            impacto=dict(porcentaje=0.03, celda_total="H!B9")))
    x = casos_revision.ejes(h)
    assert x["evidencia"]["nivel"] == "exacta" and x["impacto"]["nivel"] == "baja"
    assert x["prioridad"] == "media"  # error seguro pero chico: no compite con uno material
    h["evidencia"]["impacto"]["porcentaje"] = 4.0
    assert casos_revision.ejes(h)["prioridad"] == "alta"
    assert "Rareza:" in casos_revision.texto_ejes(x) and "Impacto: baja" in casos_revision.texto_ejes(x)
