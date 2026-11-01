"""Una suma escrita con rango (=+U69+SUM(U72:U80)) es la misma que la escrita celda por celda.

Caso real: libro de gasto, «Según categorías»!U81 aparecía en «No cierra» como fórmula fuera de patrón.
"""
from openpyxl import Workbook

from services.calculos.integridad import _analizar_hoja, _suma_relativa


def _hoja(ultima):
    w = Workbook()
    ws = w.active
    for r in range(1, 12):
        ws.cell(r, 1, f"Rubro {r}")
        for c in range(2, 7):
            ws.cell(r, c, 10 * r + c)
    for c in range(2, 6):
        letra = "ABCDEFG"[c - 1]
        ws.cell(12, c, "=+" + "+".join(f"{letra}{r}" for r in (1, *range(3, 12))))
    ws.cell(12, 1, "Total")
    ws.cell(12, 6, ultima)
    return ws


def test_suma_equivalente_no_es_formula_fuera_de_patron():
    hs = _analizar_hoja(_hoja("=+F1+SUM(F3:F11)"))
    assert not [h for h in hs if h["evidencia"]["regla"] == "formula_fuera_de_patron"]


def test_rango_incompleto_sigue_marcado():
    hs = _analizar_hoja(_hoja("=+F1+SUM(F3:F10)"))
    assert [h for h in hs if h["evidencia"]["regla"] == "formula_fuera_de_patron" and h["evidencia"]["celda"] == "F12"]


def test_suma_relativa_solo_para_sumas():
    assert _suma_relativa("=A1*2", 2, 1) is None
    assert _suma_relativa("=SUM(A1:A3)", 4, 1) == _suma_relativa("=A1+A2+A3", 4, 1)
