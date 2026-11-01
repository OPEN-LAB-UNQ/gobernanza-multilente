"""Regresiones de la revisión de ocho ZIP. Sin proveedores externos."""

import os
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

import models
from services.calculos.contraste_formula import calcular
from services.reportes.markdown import _texto_corto
from tests.test_ejemplos_zip import ejecutar_par

pytest_plugins = ["tests.test_ejemplos_zip"]


@pytest.mark.parametrize("caso", range(1, 9))
def test_evidencia_de_los_ocho_pares(flujo, caso):
    client, factory, _ = flujo
    ids = ejecutar_par(client, caso)
    respuesta = client.post(f"/api/ejemplos/{caso}/zip", json=ids)
    assert respuesta.status_code == 200
    out = os.environ.get("REVISION_ZIPS_DIR")
    if out:
        Path(out).mkdir(parents=True, exist_ok=True)
        (Path(out) / f"ejemplos-caso-{caso}-corregido.zip").write_bytes(respuesta.content)
    with ZipFile(BytesIO(respuesta.content)) as z:
        for name in (z.namelist()[1], z.namelist()[3]):
            reader = PdfReader(BytesIO(z.read(name)))
            texto = "\n".join(p.extract_text() for p in reader.pages)
            assert "La pregunta declarada fue evaluada parcialmente" not in texto
            assert "Sin evidencia porque" not in texto
            assert "Período observado en los datos" in texto
            assert len(reader.pages[-1].extract_text().split()) > 90
        observado = " ".join(
            " ".join(p.extract_text() for p in PdfReader(BytesIO(z.read(z.namelist()[3]))).pages).split()
        )
        if caso == 4:
            assert "310" in observado and "340" in observado and "295" in observado
            assert "Acumulados!C5" in observado and "Acumulados!D5" in observado
            assert "C5:C5" not in observado and "D5:D5" not in observado
        if caso == 2:
            assert "Servicios!C5:E5" in observado
        if caso == 5:
            assert "REGULAR" in observado and "5,7" in observado
            assert "El método cambia, pero el resultado de este registro coincide" in observado
            assert "Lecturas relacionadas" in observado
            assert "Qué llamó la atención: En la hoja 'Grupos_sinteticos', H3" in observado
            primera = observado.split("1. Señal de prioridad Alta", 1)[1].split("2. Señal de prioridad Alta", 1)[0]
            assert "Grupos_sinteticos!H3" in primera
        if caso == 6:
            assert "57" in observado and "Cierre de la serie" in observado
            assert observado.index("En 2026-03") < observado.index("salto temporal simultáneo")
            assert "Tramos relacionados: señales 1, 2" in observado
            assert "no deben contarse como casos independientes" in observado
            assert "p.p.." not in observado
            assert "puntos porcentuales" in observado
        if caso == 7:
            assert "Cociente analizado" in observado and "Lecturas relacionadas" in observado
            assert "señales 1, 2, 3 corresponden a una misma transición temporal" in observado
            assert "no deben contarse automáticamente como problemas independientes" in observado.lower()
            assert "Radiólogo: no evaluable (No hay preinformes con comparabilidad verificada)" in observado
        if caso == 8:
            assert "0,9963" in observado and "0,9721" in observado
            assert "Señales exploratorias separadas" in observado
            assert "Se encontraron asociaciones exploratorias para contextualizar" in observado
            assert "no son observaciones mecánicas" in observado
            assert "sin observaciones mecánicas; 1 señal exploratoria informada aparte" in observado.lower()
            assert ".;" not in observado
    with factory() as db:
        bueno = db.get(models.Informe, ids["sin_senal"])
        malo = db.get(models.Informe, ids["con_senal"])
        assert not bueno.hallazgos
        if caso == 4:
            assert len(malo.hallazgos) == 2
            assert all(h.lente == "ajedrecista" for h in malo.hallazgos)
            assert all((h.limites or {}).get("contexto_tecnico") for h in malo.hallazgos)
            assert len({h.severidad for h in malo.hallazgos}) == 1


def test_pares_v2_solo_cambian_celdas_previstas():
    import json

    root = Path(__file__).resolve().parents[1] / "static/samples"
    catalogo = json.loads((root / "manifest.json").read_text())
    esperados = {
        1: {"E4"},
        4: {"C5", "D5"},
        5: {"G4", "H3"},
        6: {"D4", "D5", "D6", "E4", "E5", "E6"},
        8: {f"D{i}" for i in range(2, 10)},
    }
    for caso, permitidas in esperados.items():
        par = {
            d["columna"]: load_workbook(root / Path(d["archivo"]).name, data_only=False).worksheets[0]
            for d in catalogo
            if d["orden"] == caso
        }
        a, b = par["bueno"], par["malo"]
        assert a.title == b.title and a.max_row == b.max_row and a.max_column == b.max_column
        distintas = {c.coordinate for row in a for c in row if c.value != b[c.coordinate].value}
        assert distintas == permitidas


def test_coordenadas_no_se_inventan_si_cambia_orden():
    from services import carga_archivos

    root = Path(__file__).resolve().parents[1] / "static/samples"
    archivo = ("caso-4-v2-con-senal.xlsx", (root / "caso-4-v2-con-senal.xlsx").read_bytes())
    datos = carga_archivos.parse_archivos([archivo])
    assert carga_archivos.origen_excel_verificado([archivo], datos)["filas"] == [2, 3, 4, 5, 6]
    assert carga_archivos.origen_excel_verificado([archivo], list(reversed(datos))) == {}


def test_formula_segura_rechaza_codigo_y_ciclos():
    ws = Workbook().active
    with pytest.raises(ValueError):
        calcular(ws, '=__import__("os").system("id")')
    ws["A1"] = "=A1+1"
    with pytest.raises(ValueError):
        calcular(ws, "=A1")
    ws["B1"] = 5.7
    ws["C1"] = 0.7
    assert calcular(ws, '=IF(AND(B1>=7,C1>=0.75),"PROMOCIONA",IF(AND(B1>=4,C1>=0.6),"REGULAR","LIBRE"))') == "REGULAR"


def test_explicacion_completa_sin_truncar():
    texto = "Explicación completa. " * 80
    assert not _texto_corto(texto, 40).endswith("…")
    assert _texto_corto(texto, 40).endswith("completa.")
