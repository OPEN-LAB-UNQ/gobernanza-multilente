"""Casos de la decimocuarta medición (16/10/2026)."""

import io
import struct
import zipfile
from types import SimpleNamespace

import pdfplumber
import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, TextStringObject
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from services import carga_archivos, controles_excel, ocr_local, perfilado
from services.controles_excel import conciliacion
from services.controles_excel import periodos as P
from services.controles_excel import series as motor_series
from services.controles_excel.comun import normalizar
from services.estructura_revision import preparar as preparar_estructura
from services.lentes import cronista
from services.lentes.util import columna_etiqueta_fila
from services.reportes.markdown import _periodos_por_contenido

ANIOS = list(range(2017, 2026))


def _xlsx(filas, formato=None, hojas=1):
    wb = Workbook()
    wb.remove(wb.active)
    for k in range(hojas):
        ws = wb.create_sheet(f"Hoja{k + 1}")
        for f in filas:
            ws.append(list(f))
        if formato:
            for fila in ws.iter_rows(min_row=2, min_col=2):
                for c in fila:
                    c.number_format = formato
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cuadro_anual():
    regiones = [("Norte", 1000), ("Sur", 800), ("Centro", 1500)]
    filas = [["Región", *ANIOS]]
    for k, (n, b) in enumerate(regiones):
        filas.append([n, *(b + 7 * i + 3 * k for i in range(len(ANIOS)))])
    filas.append(["Total", *(sum(f[j] for f in filas[1:]) for j in range(1, len(ANIOS) + 1))])
    return filas


def _avisos_hojas(cobertura):
    return [a for h in cobertura["hojas"] for a in h.get("avisos_lectura") or []]


# L01 · formato de porcentaje -------------------------------------------------------------------


def test_formato_porcentaje_se_avisa_sin_cambiar_valores():
    filas = [["Rubro", "2024", "2025"], ["Norte", 0.304, 0.31], ["Sur", 0.696, 0.69], ["Total", 1.0, 1.0]]
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([("p.xlsx", _xlsx(filas, "0.0%"))])
    [aviso] = [a for a in _avisos_hojas(cobertura) if "porcentaje" in a]
    assert "6 celdas" in aviso and "fracción (30,4 % → 0,304)" in aviso and "B2" in aviso
    assert cobertura["hojas"][0]["formato_porcentaje"]["celdas"] == 6
    assert datos[0]["2024"] == 0.304  # el valor guardado no cambia


def test_porcentaje_literal_entre_comillas_no_se_avisa():
    filas = [["Rubro", "2024", "2025"], ["Norte", 30.4, 31], ["Sur", 69.6, 69]]
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("p.xlsx", _xlsx(filas, '0.0" %"'))])
    assert not [a for a in _avisos_hojas(cobertura) if "porcentaje" in a]
    from services.excel_cobertura import es_formato_porcentaje

    assert es_formato_porcentaje("0.0%") and es_formato_porcentaje("#,##0.00%")
    assert not es_formato_porcentaje('0.0" %"') and not es_formato_porcentaje("0.0\\%") and not es_formato_porcentaje("General")


def test_porcentaje_en_muchas_hojas_no_se_repite_hoja_por_hoja():
    filas = [["Rubro", "2024", "2025"], ["Norte", 0.304, 0.31], ["Sur", 0.696, 0.69]]
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("p.xlsx", _xlsx(filas, "0.0%", hojas=6))])
    avisos = [a for a in _avisos_hojas(cobertura) if "porcentaje" in a]
    assert len(avisos) == 4  # las tres primeras hojas y una línea con las demás
    assert avisos[-1].startswith("Y 3 hojas más") and "«Hoja4»" in avisos[-1]
    assert all(h.get("formato_porcentaje") for h in cobertura["hojas"])


# L17 · «.csv» que es una página web -------------------------------------------------------------


@pytest.mark.parametrize("cabeza", [b"<!DOCTYPE html>\n<html lang='es'>", b"\xef\xbb\xbf  <html>\n<head>",
                                    b"<?xml version='1.0'?>\n<html>"])
def test_pagina_web_guardada_como_csv_se_rechaza_con_motivo(cabeza):
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match="es una página web"):
        carga_archivos.tipo_por_contenido("datos.csv", cabeza + b"<body><p>Tu sesi\xc3\xb3n expir\xc3\xb3</p></body></html>")


def test_csv_que_empieza_con_menor_que_sigue_siendo_csv():
    assert carga_archivos.tipo_por_contenido("d.csv", b"<5;3;4\n6;7;8\n") == "csv"
    assert carga_archivos.tipo_por_contenido("d.csv", b"Provincia;2024\nSalta;<5\n") == "csv"


def test_carga_mixta_declara_la_pagina_web_y_lee_el_resto():
    html = b"<!DOCTYPE html><html><body>Ingresar</body></html>"
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura(
        [("bueno.csv", b"Mes;A;B\nene-25;1;2\nfeb-25;3;4\nmar-25;5;6\n"), ("malo.csv", html)])
    malo = next(x for x in lectura if x["archivo"] == "malo.csv")
    assert malo["estado"] == "fallido" and "página web (HTML)" in malo["motivo"] and len(datos) == 3


# L03 · Open XML estricto ------------------------------------------------------------------------


def _estricto(contenido):
    """Un .xlsx de openpyxl con los espacios de nombres del Open XML estricto."""
    pares = [(t, e) for e, t in carga_archivos._ESTRICTO_A_TRANSICIONAL]
    pares.sort(key=lambda p: -len(p[0]))  # los más largos primero: el de relaciones es prefijo de otros
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(contenido)) as z, zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            d = z.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                for t, e in pares:
                    d = d.replace(t, e)
            zo.writestr(item, d)
    return salida.getvalue()


def test_libro_estricto_se_lee_convirtiendolo():
    original = _xlsx(_cuadro_anual())
    estricto = _estricto(original)
    assert b"purl.oclc.org/ooxml" in zipfile.ZipFile(io.BytesIO(estricto)).read("xl/workbook.xml")
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([("e.xlsx", estricto)])
    assert len(datos) == 4 and datos[0]["2017"] == 1000
    assert any("Open XML estricto" in a for a in lectura[0]["avisos"])
    convertido = carga_archivos.libro_transicional(estricto)
    assert b"purl.oclc.org" not in zipfile.ZipFile(io.BytesIO(convertido)).read("xl/workbook.xml")


def test_libro_comun_y_otros_contenidos_no_se_tocan():
    original = _xlsx(_cuadro_anual())
    assert carga_archivos.libro_transicional(original) is original
    texto = b"a;b\n1;2\n"
    assert carga_archivos.libro_transicional(texto) is texto
    roto = b"PK\x03\x04 no es un zip"
    assert carga_archivos.libro_transicional(roto) is roto


# L06 · .xls binario con cierres -----------------------------------------------------------------


def _rec(rid, datos):
    return struct.pack("<HH", rid, len(datos)) + datos


def _u8(s):
    return struct.pack("<BB", len(s), 1) + s.encode("utf-16-le")


def _u16(s):
    return struct.pack("<HB", len(s), 1) + s.encode("utf-16-le")


def _xls(hoja, filas):
    """Libro BIFF8 mínimo de una hoja dentro de un contenedor OLE2 (como lo escribe Excel 97-2003)."""
    fuentes = [_rec(0x0031, struct.pack("<HHHHHBBBB", 200, 0, 0x7FFF, 400, 0, 0, 0, 0, 0) + _u8("Arial")) for _ in range(5)]
    xfs = [_rec(0x00E0, struct.pack("<HHHBBBBIIH", 0, 0, 0xFFF5, 0x20, 0, 0, 0, 0, 0, 0x20C0)) for _ in range(15)]
    xfs.append(_rec(0x00E0, struct.pack("<HHHBBBBIIH", 0, 0, 0x0001, 0x20, 0, 0, 0, 0, 0, 0x20C0)))

    def glob(pos):
        return b"".join([_rec(0x0809, struct.pack("<HHHHII", 0x0600, 0x0005, 0x0DBB, 0x07CC, 0, 6)), _rec(0x0042, struct.pack("<H", 0x04B0)),
                         _rec(0x003D, struct.pack("<HHHHHHHHH", 0, 0, 0x3000, 0x2000, 0x0038, 0, 0, 1, 0x0258))] + fuentes + xfs
                        + [_rec(0x0293, struct.pack("<HBB", 0x8000, 0, 0xFF)), _rec(0x0085, struct.pack("<IBB", pos, 0, 0) + _u8(hoja)),
                           _rec(0x000A, b"")])
    g = glob(len(glob(0)))
    celdas = []
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            if isinstance(v, str):
                celdas.append(_rec(0x0204, struct.pack("<HHH", i, j, 15) + _u16(v)))
            elif v is not None:
                celdas.append(_rec(0x0203, struct.pack("<HHHd", i, j, 15, float(v))))
    datos = g + b"".join([_rec(0x0809, struct.pack("<HHHHII", 0x0600, 0x0010, 0x0DBB, 0x07CC, 0, 6)),
                          _rec(0x0200, struct.pack("<IIHHH", 0, len(filas), 0, max(len(f) for f in filas), 0))] + celdas
                         + [_rec(0x023E, struct.pack("<HHHHHHHHH", 0x06B6, 0, 0, 0x40, 0, 0, 0, 0, 0)), _rec(0x000A, b"")])
    SEC, FREE, END, FATS, NOS = 512, 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD, 0xFFFFFFFF
    datos = datos + b"\x00" * max(0, 4096 - len(datos))
    n = (len(datos) + SEC - 1) // SEC
    fat = [FATS, END] + [3 + i for i in range(n - 1)] + [END]
    fat += [FREE] * (128 - len(fat))
    hdr = carga_archivos.FIRMA_OLE + b"\x00" * 16 + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6) + b"\x00" * 6
    hdr += struct.pack("<IIIIIIIII", 0, 1, 1, 0, 4096, END, 0, END, 0) + struct.pack("<I", 0) + struct.pack("<I", FREE) * 108

    def ent(nom, tipo, hijo, inicio, tam):
        nb = (nom.encode("utf-16-le") + b"\x00\x00") if nom else b""
        e = nb.ljust(64, b"\x00") + struct.pack("<HBB", len(nb), tipo, 1 if nom else 0)
        e += struct.pack("<III", NOS, NOS, hijo) + b"\x00" * 16 + struct.pack("<I", 0) + b"\x00" * 16
        return e + struct.pack("<I", inicio) + struct.pack("<Q", tam)
    dire = ent("Root Entry", 5, 1, END, 0) + ent("Workbook", 2, NOS, 2, len(datos)) + ent("", 0, NOS, 0, 0) + ent("", 0, NOS, 0, 0)
    return hdr + struct.pack("<128I", *fat) + dire + datos.ljust(n * SEC, b"\x00")


def _controles(nombre, contenido):
    from routers.carga_rapida import _aplicar_alcance_series
    from services import excel_cobertura, sanitizacion

    datos, cobertura, _ = carga_archivos.parse_archivos_con_lectura([(nombre, contenido)])
    saneados, reporte, _ = excel_cobertura.preparar(
        datos, cobertura, lambda f: _aplicar_alcance_series(f, "todas"), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, contenido)], saneados, reporte["excel"]), reporte["excel"]


def test_xls_binario_tiene_libro_para_los_controles():
    filas = [["Cuadro 6. Alumnos por región"], ["En alumnos"], [], *[[str(c) if i == 0 else c for i, c in enumerate(f)] for f in _cuadro_anual()],
             ["=No es una fórmula"]]
    contenido = _xls("Cuadro 6", filas)
    [(archivo, libro)] = controles_excel._libros([("v.xls", contenido)])
    ws = libro.valores["Cuadro 6"]
    assert archivo == "v.xls" and ws.cell(5, 2).value == 1000 and isinstance(ws.cell(5, 2).value, int)
    assert ws.cell(4, 2).value == 2017 and ws.cell(9, 1).data_type == "s"
    _, excel = _controles("v.xls", contenido)
    assert excel["estructura_revision"]["cantidad_cierres"] >= 9


# L10 · fecha constante -------------------------------------------------------------------------


def test_fecha_constante_es_metadato_no_eje():
    datos = [{"provincia": p, "2023": 10 + k, "2024": 11 + k, "fecha_actualizacion": "2026-03-16"}
             for k, p in enumerate(["Salta", "Jujuy", "Chaco", "Misiones"])]
    assert cronista.columna_tiempo(datos) is None
    assert cronista.problemas_periodo(datos) == []
    assert columna_etiqueta_fila(datos) == "provincia"
    con_meses = [{"fecha": f"2025-0{m}-01", "valor": m} for m in range(1, 6)]
    assert cronista.columna_tiempo(con_meses) == "fecha"


# Períodos --------------------------------------------------------------------------------------


def test_bimestres():
    p = P.leer("Ene-feb 24")
    assert (p["frecuencia"], p["anio"], p["sub"]) == ("bimestral", 2024, 1) and P.etiqueta(p) == "2024-B1"
    assert P.leer("Nov-dic 25")["sub"] == 6
    assert P.leer("Feb-mar 24") is None  # no empieza en un mes impar
    assert P.leer("Ene-mar 24") is None  # no son meses consecutivos
    eje = P.leer_eje(["Ene-feb 24", "Mar-abr 24", "May-jun 24", "Jul-ago 24"])
    assert eje["frecuencia"] == "bimestral" and [x["ordinal"] for x in eje["periodos"]] == list(range(2024 * 6, 2024 * 6 + 4))
    assert cronista._periodo_columna("Mar-abr 24") == (2024, 3)


def test_prefijo_de_r_solo_si_todos_lo_traen():
    eje = P.leer_eje([f"X{a}" for a in ANIOS])
    assert eje["frecuencia"] == "anual" and eje["periodos"][0]["texto"] == "X2017" and P.etiqueta(eje["periodos"][0]) == "2017"
    assert P.leer_eje(["X2017", "X2018", "2019", "Xenón"]) is None
    datos = [{"Region": r, **{f"X{a}": 100 + 3 * i + k for i, a in enumerate(ANIOS)}} for k, r in enumerate(["AMBA", "NEA", "Cuyo"])]
    series = cronista._series_formato_ancho(datos)
    assert [s["etiqueta"] for s in series] == ["AMBA", "NEA", "Cuyo"] and series[0]["valores"][0][1] == (2017, 1)
    assert cronista._fmt_periodo_label("X2020") == "2020"
    assert _periodos_por_contenido(datos, {f"X{a}" for a in ANIOS}) == "2017 a 2025"


def test_meses_sin_anio_se_declaran_en_el_periodo_observado():
    meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
             "Noviembre", "Diciembre"]
    datos = [{"Región": r, **{m: 100 + i for i, m in enumerate(meses)}} for r in ("AMBA", "Cuyo", "NEA")]
    assert _periodos_por_contenido(datos, set(meses)) == "enero a diciembre (sin año en el rótulo)"


# L14 · arrastre de página ------------------------------------------------------------------------


def _serie(etiqueta, valores, fila, cuadro=1):
    periodos = ANIOS[: len(valores)]
    return {"archivo": "m.csv", "hoja": "m", "cuadro": cuadro, "etiqueta": etiqueta, "nivel": 0, "periodos": periodos,
            "valores": valores, "celdas": [f"{c}{fila}" for c in "BCDEFGHIJ"[: len(valores)]],
            "periodos_validos": periodos, "unidad": None}


def _componentes():
    return [[100 + i for i in range(5)], [200 + 2 * i for i in range(5)], [50 + i for i in range(5)]]


def test_transporte_al_pie_es_total_de_arrastre():
    a, b, c = _componentes()
    cuadro = [_serie("Buenos Aires", a, 2), _serie("Catamarca", b, 3), _serie("Chaco", c, 4),
              _serie("Transporte", [x + y + z for x, y, z in zip(a, b, c)], 5)]
    [(padre, hijos, _)] = list(conciliacion.relaciones_cuadro(cuadro))
    assert padre["etiqueta"] == "Transporte" and [h["etiqueta"] for h in hijos] == ["Buenos Aires", "Catamarca", "Chaco"]
    assert preparar_estructura([dict(s) for s in cuadro])["cantidad_cierres"] == 5


def test_viene_de_es_componente_y_rubro_transporte_no_cambia():
    a, b, c = _componentes()
    viene = [_serie("Viene de la página anterior", a, 2), _serie("Neuquén", b, 3), _serie("Salta", c, 4),
             _serie("Total", [x + y + z for x, y, z in zip(a, b, c)], 5)]
    [(padre, hijos, _)] = list(conciliacion.relaciones_cuadro(viene))
    assert padre["etiqueta"] == "Total" and hijos[0]["etiqueta"] == "Viene de la página anterior"
    rubro = [_serie("Energía", a, 2), _serie("Transporte", b, 3), _serie("Salud", c, 4),
             _serie("Total", [x + y + z for x, y, z in zip(a, b, c)], 5)]
    [(padre, hijos, _)] = list(conciliacion.relaciones_cuadro(rubro))
    assert padre["etiqueta"] == "Total" and "Transporte" in [h["etiqueta"] for h in hijos]
    assert not any(s.get("arrastre") for s in rubro)


def test_cronista_no_toma_sumas_de_pagina_como_series():
    filas = [("Viene de la página anterior", 500), ("Neuquén", 100), ("Salta", 80), ("Transporte", 680)]
    datos = [{"Jurisdicción": n, **{str(a): b + 5 * i for i, a in enumerate(ANIOS[:6])}} for n, b in filas]
    assert [s["etiqueta"] for s in cronista._series_formato_ancho(datos)] == ["Neuquén", "Salta"]


# L07 · «<5» suprimido -----------------------------------------------------------------------------


def test_menor_que_cinco_es_dato_suprimido():
    assert perfilado.numero("<5") is None and perfilado.numero("< 3") is None and perfilado.numero("s.e.") is None
    assert perfilado.es_sin_dato("<5") and perfilado.es_suprimido("s.e.") and not perfilado.es_suprimido("se")
    assert perfilado.numero("5") == 5.0 and not perfilado.es_suprimido("Salta")


def test_csv_con_suprimidos_avisa_aparte_y_no_como_no_interpretables():
    contenido = ("Provincia;2021;2022;2023;2024\nChaco;42;25;23;45\nSalta;<5;<5;7;<5\nFormosa;13;19;17;8\n"
                 "Total;58;47;47;56\n").encode("utf-8")
    datos, info = carga_archivos._leer_csv_con_lectura(contenido, "h.csv")
    assert info["celdas_suprimidas"]["celdas"] == 3 and not info["celdas_no_interpretables"]
    assert any("confidencialidad estadística" in a for a in info["avisos"])
    assert not any("sin dato («-»" in a for a in info["avisos"])


def test_cierre_con_suprimidos_dice_el_motivo_y_sin_cambio_de_tipo():
    salta = ["<5"] * 5
    cuadro = [_serie("Chaco", [42, 25, 23, 45, 36], 2), _serie("Salta", salta, 3),
              _serie("Formosa", [13, 19, 17, 8, 8], 4), _serie("Total", [57, 47, 42, 55, 46], 5)]
    r = preparar_estructura(cuadro)
    assert r["cantidad_cierres"] == 0 and "suprimidos por confidencialidad" in r["motivo_sin_cierres"]
    corrientes = _serie("Corrientes", [8, 6, 9, "<5", 8, 5, 5], 6)
    assert not [h for h in motor_series.analizar(corrientes) if h["evidencia"].get("regla") == "cambio_tipo"]


# L05 · fila «Promedio» ---------------------------------------------------------------------------


def test_fila_promedio_no_es_total_ni_componente():
    assert conciliacion.FILA_NO_SUMABLE.match("promedio por provincia")
    assert conciliacion.FILA_NO_SUMABLE.match("% del total") and conciliacion.FILA_NO_SUMABLE.match("tasa de variacion")
    assert not conciliacion.FILA_NO_SUMABLE.match("limay medio") and not conciliacion.FILA_NO_SUMABLE.match("mediterranea")
    assert not conciliacion.FILA_NO_SUMABLE.match("media tension") and conciliacion.FILA_NO_SUMABLE.match("media por region")
    a, b, c = _componentes()
    cuadro = [_serie("Norte", a, 2), _serie("Sur", b, 3), _serie("Este", c, 4),
              _serie("Total", [x + y + z for x, y, z in zip(a, b, c)], 5),
              _serie("Promedio por provincia", [round((x + y + z) / 3, 1) for x, y, z in zip(a, b, c)], 6)]
    relaciones = list(conciliacion.relaciones_cuadro(cuadro))
    assert [p["etiqueta"] for p, _, _ in relaciones] == ["Total"]
    assert all("Promedio por provincia" not in [h["etiqueta"] for h in hs] for _, hs, _ in relaciones)


# L12 · celda combinada vertical en PDF ---------------------------------------------------------


def _palabra(texto, top):
    return {"text": texto, "top": top, "x0": 10, "x1": 40}


def test_rotulo_de_celda_combinada_se_completa_en_su_grupo():
    # Bordes horizontales que cruzan la primera columna: encabezado, después de cada grupo y al pie.
    edges = [{"orientation": "h", "top": y, "x0": 0, "x1": 500} for y in (90, 160, 210, 260)]
    pagina = SimpleNamespace(edges=edges)
    tops = [100, 120, 140, 170, 190, 220]
    datos = [[_palabra("x", t)] for t in tops]
    grilla = [["", "Jujuy", "1"], ["NOA", "Salta", "2"], ["", "Tucumán", "3"], ["Cuyo", "Mendoza", "4"],
              ["", "San Juan", "5"], ["Total", "", "15"]]
    carga_archivos._completar_combinadas_pdf(pagina, datos, grilla, (5.0, 45.0))
    assert [f[0] for f in grilla] == ["NOA", "NOA", "NOA", "Cuyo", "Cuyo", "Total"]


def test_sin_bordes_la_primera_columna_no_cambia():
    grilla = [["", "Jujuy", "1"], ["NOA", "Salta", "2"], ["", "Tucumán", "3"]]
    carga_archivos._completar_combinadas_pdf(SimpleNamespace(edges=[]), [[_palabra("x", t)] for t in (1, 2, 3)],
                                             grilla, (5.0, 45.0))
    assert [f[0] for f in grilla] == ["", "NOA", ""]
    carga_archivos._completar_combinadas_pdf(object(), [], [], (5.0, 45.0))  # una franja sin bordes legibles no falla


# L13 · capa oculta ---------------------------------------------------------------------------------


def _pdf_cuadro(invisible=False):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica", 10)
    c.drawString(60, 800, "Cuadro 13. Ingresos de turistas por región. 2023-2025")
    filas = [["Región", "2023", "2024", "2025"], ["AMBA", "154.197", "159.884", "155.684"],
             ["Centro", "70.521", "70.151", "71.985"], ["Cuyo", "50.344", "51.954", "53.170"],
             ["Total", "275.062", "281.989", "280.839"]]
    for i, f in enumerate(filas):
        for j, v in enumerate(f):
            c.drawString(60 + 90 * j, 760 - 18 * i, v)
    if invisible:  # capa de texto invisible (como la del OCR de un escaneo): no es una capa opcional
        t = c.beginText(60, 600)
        t.setTextRenderMode(3)
        t.textLine("Texto invisible de OCR")
        c.drawText(t)
    c.save()
    return buf.getvalue()


def _con_capa(contenido, apagada=True, otra_encendida=False, base_apagada=False):
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(contenido)))
    pg = w.pages[0]
    fuente = next(str(k) for k in pg["/Resources"]["/Font"])
    ops = f"q /OC /oc1 BDC BT {fuente} 10 Tf 1 0 0 1 420 742 Tm (999.999) Tj ET BT {fuente} 10 Tf 1 0 0 1 300 500 Tm (BORRADOR NO PUBLICAR) Tj ET EMC Q"
    s = DecodedStreamObject()
    s.set_data(ops.encode("latin-1"))
    cont = pg["/Contents"]
    arr = ArrayObject(list(cont) if isinstance(cont, ArrayObject) else [pg.raw_get("/Contents")])
    arr.append(w._add_object(s))
    pg[NameObject("/Contents")] = arr
    ocg = w._add_object(DictionaryObject({NameObject("/Type"): NameObject("/OCG"), NameObject("/Name"): TextStringObject("Borrador")}))
    props = {NameObject("/oc1"): ocg}
    ocgs = [ocg]
    if otra_encendida:
        otra = w._add_object(DictionaryObject({NameObject("/Type"): NameObject("/OCG"), NameObject("/Name"): TextStringObject("Notas")}))
        props[NameObject("/oc2")] = otra
        ocgs.append(otra)
    pg["/Resources"][NameObject("/Properties")] = DictionaryObject(props)
    w._root_object[NameObject("/OCProperties")] = DictionaryObject({
        NameObject("/OCGs"): ArrayObject(ocgs),
        NameObject("/D"): DictionaryObject({NameObject("/ON"): ArrayObject(ocgs[1:] if apagada else ocgs),
                                            NameObject("/OFF"): ArrayObject([ocg] if apagada else [])})})
    if base_apagada:  # estado base apagado y lista de encendidas: todo lo que no está en /ON queda apagado
        w._root_object["/OCProperties"][NameObject("/D")] = DictionaryObject(
            {NameObject("/BaseState"): NameObject("/OFF"), NameObject("/ON"): ArrayObject(ocgs[1:])})
    salida = io.BytesIO()
    w.write(salida)
    return salida.getvalue()


def _sin_capas(contenido):
    reader = PdfReader(io.BytesIO(contenido))
    native = reader.pages[0].extract_text()
    with pdfplumber.open(io.BytesIO(contenido)) as pdf:
        pagina, texto, aviso = ocr_local._sin_capas_apagadas(reader, 0, pdf.pages[0], native)
        tablas = ocr_local.tablas_pagina_pdf(pagina, texto)
    return native, texto, aviso, tablas


def test_capa_apagada_no_entra_a_la_lectura_y_se_avisa():
    native, texto, aviso, tablas = _sin_capas(_con_capa(_pdf_cuadro()))
    assert "999.999" in native and "BORRADOR" in native
    assert "999.999" not in texto and "BORRADOR" not in texto and "154.197" in texto
    assert "texto oculto (una capa apagada)" in aviso and "BORRADOR NO PUBLICAR" in aviso
    assert not any("999.999" in c for t in tablas for f in t["filas"] for c in f)


def test_capa_apagada_entre_otras_encendidas_solo_se_avisa():
    native, texto, aviso, _ = _sin_capas(_con_capa(_pdf_cuadro(), otra_encendida=True))
    assert texto == native and "no se pudieron separar" in aviso


def test_pdf_sin_capas_o_con_capa_encendida_no_cambia():
    for contenido in (_pdf_cuadro(invisible=True), _con_capa(_pdf_cuadro(), apagada=False)):
        reader = PdfReader(io.BytesIO(contenido))
        native = reader.pages[0].extract_text()
        with pdfplumber.open(io.BytesIO(contenido)) as pdf:
            pagina, texto, aviso = ocr_local._sin_capas_apagadas(reader, 0, pdf.pages[0], native)
            assert pagina is pdf.pages[0] and texto == native and aviso is None
    assert "Texto invisible de OCR" in PdfReader(io.BytesIO(_pdf_cuadro(invisible=True))).pages[0].extract_text()


# PDF en inglés con enteros ----------------------------------------------------------------------


def _filas_en():
    return [["Month", "AMBA", "Centro", "Total"], ["Jan 2025", "24,328", "13,453", "37,781"],
            ["Feb 2025", "24,279", "13,154", "37,433"], ["Mar 2025", "25,001", "13,500", "38,501"]]


def test_pdf_en_ingles_con_coma_de_miles():
    texto = "Table 1. Heat pumps installed by region and month\nIn units\nSource: national registry. Total of the year"
    t = ocr_local._armar_tabla(_filas_en(), "texto_pdf", texto=texto)
    assert t["formato_numerico"] == "en"


def test_texto_en_castellano_o_punto_de_miles_sigue_ambiguo():
    texto_es = "Cuadro 1. Bombas de calor instaladas por región y mes\nEn unidades\nFuente: registro nacional de la provincia"
    assert ocr_local._armar_tabla(_filas_en(), "texto_pdf", texto=texto_es)["formato_numerico"] == "ambiguo"
    texto_en = "Table 1. Heat pumps installed by region and month\nIn units\nSource: national registry. Total of the year"
    mezcla = _filas_en() + [["Apr 2025", "1.234", "13,600", "14,834"]]
    assert ocr_local._armar_tabla(mezcla, "texto_pdf", texto=texto_en)["formato_numerico"] == "ambiguo"
    assert ocr_local._armar_tabla(_filas_en(), "texto_pdf")["formato_numerico"] == "ambiguo"


# Revisión independiente de 0.9.19 -------------------------------------------------------------------


def _cuadro_csv(texto):
    """Filas de un CSV de rótulo + años como series de conciliación (archivo «r.csv»)."""
    lineas = [x.split(";") for x in texto.strip().splitlines()]
    anios = [int(a) for a in lineas[0][1:]]
    out = []
    for k, f in enumerate(lineas[1:], start=2):
        out.append({"archivo": "r.csv", "hoja": "r", "cuadro": 1, "etiqueta": f[0], "nivel": 0, "periodos": anios,
                    "valores": [perfilado.numero(v) if perfilado.numero(v) is not None else v for v in f[1:]],
                    "celdas": [f"{c}{k}" for c in "BCDEFGHIJ"[: len(anios)]], "periodos_validos": anios, "unidad": None})
    return out


PIB = ("Componente;2019;2020;2021;2022;2023\nConsumo;700;690;720;750;770\nInversion;200;180;210;220;230\n"
       "Variacion de existencias;10;-20;15;5;8\nExportaciones;150;140;160;170;180\nImportaciones;-160;-130;-170;-180;-190\n"
       "Total;900;860;935;965;998\n")
TASAS = ("Recurso;2019;2020;2021;2022;2023\nTasa por servicios generales;100;110;120;130;140\nTasa vial;20;22;24;26;28\n"
         "Coparticipacion;500;520;540;560;580\nTotal;620;652;684;716;748\n")
TRANSPORTE_ANCHO = ("Rubro;2019;2020;2021;2022;2023\nSalud;100;110;120;130;140\nEducacion;200;210;220;230;240\n"
                    "Vivienda;50;55;60;65;70\nTransporte;80;85;90;95;300\n")


@pytest.mark.parametrize("texto", [PIB, TASAS])
def test_componentes_que_suman_no_se_excluyen_de_antemano(texto):
    r = preparar_estructura(_cuadro_csv(texto))
    assert r["cantidad_cierres"] == 5 and r["motivo_sin_cierres"] is None


def test_lista_acotada_de_filas_que_no_suman():
    for e in ("Promedio", "Mediana", "Media", "Media por provincia", "% del total", "Porcentaje", "Indice de precios",
              "Tasa de desocupacion", "Tasa de variacion anual"):
        assert conciliacion.FILA_NO_SUMABLE.match(normalizar(e)), e
    for e in ("Tasa vial", "Tasa por servicios generales", "Tasa de seguridad e higiene", "Variación de existencias",
              "Media tensión", "Media y alta tensión", "Tasas", "Indicadores"):
        assert not conciliacion.FILA_NO_SUMABLE.match(normalizar(e)), e


def test_promedio_entre_los_componentes_sale_solo_si_con_el_no_cierra():
    a, b, c = _componentes()
    cuadro = [_serie("Norte", a, 2), _serie("Sur", b, 3), _serie("Este", c, 4),
              _serie("Promedio", [round((x + y + z) / 3, 1) for x, y, z in zip(a, b, c)], 5),
              _serie("Total", [x + y + z for x, y, z in zip(a, b, c)], 6)]
    [(padre, hijos, _)] = list(conciliacion.relaciones_cuadro(cuadro))
    assert padre["etiqueta"] == "Total" and [h["etiqueta"] for h in hijos] == ["Norte", "Sur", "Este"]


def test_transporte_sin_evidencia_de_arrastre_es_un_rubro():
    cuadro = _cuadro_csv(TRANSPORTE_ANCHO)
    assert list(conciliacion.relaciones_cuadro(cuadro)) == [] and not any(s.get("arrastre") for s in cuadro)
    lineas = [x.split(";") for x in TRANSPORTE_ANCHO.strip().splitlines()]
    datos = [dict(zip(lineas[0], [f[0]] + [float(v) for v in f[1:]])) for f in lineas[1:]]
    assert conciliacion.arrastre_en_filas(datos, "Rubro") is None
    assert "Transporte" in [s["etiqueta"] for s in cronista._series_formato_ancho(datos)]
    hallazgos = cronista.analizar(datos)
    assert any("Transporte" in h["descripcion"] and "95" in h["descripcion"] and "300" in h["descripcion"] for h in hallazgos)


def test_transporte_con_viene_de_es_arrastre():
    a, b, c = _componentes()
    cuadro = [_serie("Viene de la página anterior", a, 2), _serie("Salta", b, 3), _serie("Jujuy", c, 4),
              _serie("Transporte", [x + y for x, y in zip(b, c)], 5)]  # no suma, pero la página dice que arrastra
    list(conciliacion.relaciones_cuadro(cuadro))
    assert cuadro[-1].get("arrastre")


def test_transporte_en_formato_largo_no_es_total():
    datos = [{"rubro": r, "anio": str(y), "valor": float(v)} for r in ("Salud", "Educacion", "Transporte")
             for y, v in zip(range(2019, 2024), [10, 11, 12, 13, 14])]
    assert conciliacion.arrastre_en_filas(datos, columna_etiqueta_fila(datos)) is None
    assert conciliacion.arrastre_en_filas(datos, "rubro") is None


def test_idioma_no_decide_con_meses_compartidos_bilingue_o_coma_decimal():
    txt = "Ventas mensuales\nProducto Ene Feb Mar Abr May Jun Jul Ago Sep Oct Nov Dic Total\nA 1,234 2,500 ..."
    assert not ocr_local._texto_en_ingles(txt)
    assert ocr_local._formato_por_idioma([["A", "1,234", "2,500", "3,125"], ["B", "0,875", "1,250", "2,375"]], txt) == "ambiguo"
    bilingue = ("Cuadro 3. Exportaciones / Exports. Fuente / Source: instituto de estadística. Total exports by region and country, "
                "in thousand dollars. Región / Region")
    assert not ocr_local._texto_en_ingles(bilingue)
    assert ocr_local._formato_por_idioma([["Asia", "24,328"], ["Europa", "1,250"]], bilingue) == "ambiguo"
    filas = [["Rubro", "Ene", "Feb", "Mar", "Abr", "May", "Jun"],
             ["Alimentos", "1,052", "1,071", "1,089", "1,104", "1,118", "1,131"],
             ["Vivienda", "0,987", "1,013", "1,032", "1,048", "1,061", "1,077"]]
    texto = ("Indices de precios. Base 2023=1\nRubro Ene Feb Mar Abr May Jun\nAlimentos 1,052 1,071 1,089 1,104 1,118 1,131\n"
             "Vivienda 0,987 1,013 1,032 1,048 1,061 1,077")
    assert ocr_local._armar_tabla(filas, "texto_pdf", "Indices", texto=texto)["formato_numerico"] in ("ambiguo", "es")
    ingles = "Table 1. Heat pumps installed by month\nIn units\nSource: the national registry of the year"
    assert ocr_local._formato_por_idioma([["A", "24,328"], ["B", "0,875"]], ingles) == "ambiguo"  # «0,875»: coma decimal


def test_marcas_de_dato_no_publicado():
    for v in ("<5", "< 3", "<10", "s.e.", "S.E."):
        assert perfilado.es_suprimido(v) and perfilado.numero(v) is None, v
    for v in ("<0,1", "<0.1", "<1"):
        assert perfilado.es_menor_unidad(v) and not perfilado.es_suprimido(v) and perfilado.numero(v) is None, v
    for v in ("<100", "<2020", "s/e", "S/E", "< 5 años"):
        assert not perfilado.es_no_publicado(v), v
    _, informe = perfilado.interpretar_columna(["10", "<0,1", "3,5", "<0,1", "2,2"])
    assert informe["menores"] == [1, 3] and informe["suprimidos"] == []


def test_csv_con_menores_que_la_unidad_avisa_y_cita_la_marca():
    texto = ("Provincia;2019;2020;2021;2022;2023\nA;12,5;13,1;<0,1;14,2;15,0\nB;3,2;<0,1;4,1;4,4;4,9\n"
             "C;7,7;7,9;8,1;8,3;8,6\nTotal;23,4;21,0;12,2;26,9;28,5\n")
    _, info = carga_archivos._leer_csv_con_lectura(texto.encode("utf-8"), "m.csv")
    assert info["celdas_menores_que_unidad"]["celdas"] == 2 and not info.get("celdas_suprimidas")
    assert any("menor que la unidad mínima publicada" in a and "«<0,1»" in a for a in info["avisos"])
    r = preparar_estructura(_cuadro_csv(texto))
    assert r["cantidad_cierres"] == 0 and "«<0,1»" in r["motivo_sin_cierres"] and "«<5»" not in r["motivo_sin_cierres"]


def test_suprimidos_no_tapan_un_total_que_no_cierra():
    texto = ("Provincia;2018;2019;2020;2021;2022;2023\nA;10;12;<5;14;15;16\nB;20;21;22;23;24;25\nC;30;31;32;33;34;35\n"
             "Total;70;74;60;80;83;86\n")
    r = preparar_estructura(_cuadro_csv(texto))
    assert r["cantidad_cierres"] == 0 and r["motivo_sin_cierres"].startswith("los totales no cierran")
    assert "«<5»" in r["motivo_sin_cierres"]


def test_formato_porcentaje_con_valores_en_puntos_y_anios_en_el_encabezado():
    wb = Workbook()
    ws = wb.active
    ws.append(["Provincia", 2019, 2020, 2021, 2022, 2023])
    for n, b in [("A", 30.4), ("B", 12.1), ("C", 8.7)]:
        ws.append([n] + [b + k for k in range(5)])
    for fila in ws.iter_rows(min_row=1, min_col=2):
        for c in fila:
            c.number_format = "0.0%"
    buf = io.BytesIO()
    wb.save(buf)
    _, cobertura = carga_archivos.parse_archivos_con_cobertura([("p.xlsx", buf.getvalue())])
    [aviso] = [a for a in _avisos_hojas(cobertura) if "porcentaje" in a]
    assert aviso.startswith("15 celdas") and "multiplica por 100" in aviso and "guardan la fracción" not in aviso
    assert cobertura["hojas"][0]["formato_porcentaje"]["fraccion"] is False


@pytest.mark.parametrize("nombre,contenido,motivo", [
    ("ssml.xls", b'<?xml version="1.0"?>\n<?mso-application progid="Excel.Sheet"?>\n<Workbook '
                 b'xmlns="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet><Table></Table></Worksheet></Workbook>',
     "Hoja de cálculo XML 2003"),
    ("export.xls", b"<html><body><table><tr><td>Rubro</td><td>2020</td></tr></table></body></html>",
     "tabla HTML guardada con extensión .xls"),
    ("solo.xls", b"<table border=1><tr><td>Rubro</td></tr></table>", "tabla HTML guardada con extensión .xls"),
    ("login.csv", b"<!DOCTYPE html><html><body><form>Ingresar</form></body></html>", "descarga fallida"),
])
def test_html_o_xml_con_extension_de_datos(nombre, contenido, motivo):
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match=motivo):
        carga_archivos.tipo_por_contenido(nombre, contenido)


def test_xls_sin_formulas_no_corre_reglas_de_valor_tipeado():
    from services.controles_excel import transversal
    from services.controles_excel.modelo import Cuadro

    filas = [["Provincia", *ANIOS[:4]], ["Norte", 1234567, 2345678, 3456789, 4567890],
             ["Sur", 1234567, 1111111, 2222222, 3333333], ["Este", 5, 6, 7, 8]]
    [(_, libro)] = controles_excel._libros([("v.xls", _xls("H", filas))])
    assert libro.formulas_desconocidas
    assert transversal.repetidos(Cuadro(formulas_desconocidas=True)) == []


def test_capa_con_estado_base_apagado_solo_se_avisa():
    native, texto, aviso, _ = _sin_capas(_con_capa(_pdf_cuadro(), otra_encendida=True, base_apagada=True))
    assert texto == native and "no se pudieron separar" in aviso
