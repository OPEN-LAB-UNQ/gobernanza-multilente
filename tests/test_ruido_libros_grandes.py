"""Ruido en los libros grandes de un ministerio que entran desde 0.10.56 (749 y 1.675 señales altas siendo correctos).

Casi todo salía de la misma estructura: cada hoja apila un bloque anual, uno trimestral y uno mensual, con una
fila de códigos de serie («45.2_IT_0_T_13», «364.1_TOTALTAL__5») antes de cada bloque. Los códigos se leían como
valores («cambio de tipo»), los trimestres «I.16» y las fechas como texto como años mal escritos («año no entero»),
la marca «n,d,» como un valor extraño y el año sin encuesta, vacío en todas las series, como un faltante de cada una
(«vacío interior»). Libros sintéticos con esa estructura; cada caso tiene su control con un error de verdad.
"""
import datetime as dt
import io

from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion
from services.controles_excel import periodos as P
from services.controles_excel import temporal
from services.excel_estructura import CODIGO_SERIE
from services.perfilado import es_sin_dato

SERIES = ["Producción", "Exportaciones", "Importaciones", "Consumo"]


def _analizar(wb, nombre="libro.xlsx"):
    b = io.BytesIO()
    wb.save(b)
    contenido = b.getvalue()
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, contenido)], d, r["excel"]), r["excel"]


def _reglas(hs, *reglas):
    return sorted((h["evidencia"].get("regla"), h["evidencia"].get("celda")) for h in hs
                  if h["evidencia"].get("regla") in reglas)


def _codigos(n):
    return [f"45.2_IT_0_T_{n}"] + [f"45.2_SERIE{j}__{n + j}" for j in range(len(SERIES))]


def _valor(j, k):
    return round(100 + 7.5 * j + 1.25 * k + (k % 3) * 0.5, 2)


def _apilado(con_codigos=True, cambiar=None):
    """Una hoja como las de un ministerio: anual 2016-2021, trimestral I.16-IV.18 y mensual 2016-01 a 2017-12 (fechas)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "1.1 Serie"
    ws.append(["", "CUADRO 1.1"])
    ws.append(["", "Indicadores, en millones de pesos"])
    ws.append(["Período"] + SERIES)
    bloques = [list(range(2016, 2022)),
               [f"{r}.{a}" for a in (16, 17, 18) for r in ("I", "II", "III", "IV")],
               [dt.datetime(a, m, 1) for a in (2016, 2017) for m in range(1, 13)]]
    k = 0
    for n, bloque in enumerate(bloques):
        if con_codigos:
            ws.append(_codigos(13 + n))
        for periodo in bloque:
            ws.append([periodo] + [_valor(j, k) for j in range(len(SERIES))])
            k += 1
    for celda, valor in (cambiar or {}).items():
        ws[celda] = valor
    return wb


def test_codigo_de_serie_con_guiones_bajos_dobles():
    assert all(CODIGO_SERIE.match(c) for c in ("364.1_TOTALTAL__5", "327.1_INDICE_TIEMPO__13", "44.2_ECTETCU_0_T_35"))
    assert not any(CODIGO_SERIE.match(c) for c in ("1.1 OyD real", "364.1", "364.1_TOTAL", "Total_2"))


def test_lector_de_periodos_trimestre_con_punto_y_fecha_como_texto():
    assert (P.leer("I.16")["frecuencia"], P.leer("IV.25")["anio"], P.leer("IV.25")["sub"]) == ("trimestral", 2025, 4)
    p = P.leer("2020-08-01 00:00:00")
    assert (p["frecuencia"], p["anio"], p["sub"], p["dia"]) == ("mensual", 2020, 8, 1)
    assert P.leer("2020-13-01") is None and P.leer("V.16") is None


def test_bloques_apilados_por_frecuencia_son_un_formato_oficial():
    assert temporal.formato_oficial([2016, 2017, "I.16", "II.16", "2016-01-01 00:00:00", "2016-02-01"])[0] == \
        "bloques_por_frecuencia"
    # Un rótulo que no es período, un bloque de uno solo o una frecuencia que vuelve no lo son.
    assert temporal.formato_oficial([2016, 2017, "2O18", "I.16", "II.16"])[0] is None
    assert temporal.formato_oficial([2016, 2017, 2018, "I.16"])[0] is None
    assert temporal.formato_oficial([2016, 2017, "I.16", "II.16", 2018, 2019])[0] is None


def test_hoja_apilada_con_filas_de_codigos_no_da_ruido():
    hs, excel = _analizar(_apilado())
    assert _reglas(hs, "cambio_tipo", "anio_no_entero", "periodo_fuera_de_secuencia") == []
    hoja = next(h for h in excel["hojas"] if h["hoja"] == "1.1 Serie")
    assert {4, 11, 24} <= set(hoja["filas_codigos_serie"])  # las tres filas de códigos son metadato


def test_hoja_apilada_sin_filas_de_codigos_tampoco():
    hs, _ = _analizar(_apilado(con_codigos=False))
    assert _reglas(hs, "anio_no_entero") == []


def test_un_codigo_suelto_entre_los_valores_sigue_siendo_cambio_de_tipo():
    # Control: un código en una celda de datos, en una fila con números, no es una fila de códigos.
    hs, _ = _analizar(_apilado(cambiar={"C8": "45.2_SERIE1__14"}))
    assert ("cambio_tipo", "C8") in _reglas(hs, "cambio_tipo")


def test_n_d_con_comas_es_la_marca_de_sin_dato_de_la_hoja():
    assert es_sin_dato("n,d,") and es_sin_dato("N,D,") and es_sin_dato("s,d")
    hs, _ = _analizar(_apilado(cambiar={"C7": "n,d,", "D9": "n,d,", "E10": "n,d,"}))
    assert _reglas(hs, "cambio_tipo") == []


def _vertical(vacias, series=SERIES):
    """Años en filas (2008-2015) con una columna por serie; `vacias` = celdas que se dejan vacías."""
    wb = Workbook()
    ws = wb.active
    ws.title = "EAHU"
    ws.append(["Período"] + series)
    for k, anio in enumerate(range(2008, 2016)):
        ws.append([anio] + [_valor(j, k) for j in range(len(series))])
    for celda in vacias:
        ws[celda] = None
    return wb


def test_anio_sin_datos_en_todo_el_cuadro_no_es_un_faltante_de_cada_serie():
    # 2012 (fila 6) sin dato en las cuatro series: la fuente no lo publicó.
    hs, _ = _analizar(_vertical(["B6", "C6", "D6", "E6"]))
    assert _reglas(hs, "vacio_interior") == []


def test_un_faltante_en_una_sola_serie_se_sigue_senalando():
    hs, _ = _analizar(_vertical(["C6"]))
    assert _reglas(hs, "vacio_interior") == [("vacio_interior", "C6")]


def test_con_dos_series_no_alcanza_para_decir_que_el_anio_no_tiene_datos():
    hs, _ = _analizar(_vertical(["B6", "C6"], series=SERIES[:2]))
    assert _reglas(hs, "vacio_interior") == [("vacio_interior", "B6"), ("vacio_interior", "C6")]


def test_fila_de_codigos_con_algunos_incompletos():
    # Algunos códigos vienen sin prefijo o con puntos en un tramo; la fila sigue siendo de códigos si la mayoría son
    # completos. Una fila de rótulos o de tokens sueltos no lo es.
    from services.excel_cobertura import es_fila_codigos_serie

    assert es_fila_codigos_serie(["462.2_INDICE_TIEMPO_0_0_13", "462.2_TASA_ACT_0_0_45", "462.2_TASA_EMP_0_0_46",
                                  "_TASA_DE_DEANO_0_0_42_71", "364.1_DISTRIB._GGAS__30"])
    assert not es_fila_codigos_serie(["_TOTAL_GENERAL_A", "_SUB_TOTAL_B", "462.2_TASA_ACT_0_0_45"])
    assert not es_fila_codigos_serie(["Total", "462.2_TASA_ACT_0_0_45", "462.2_TASA_EMP_0_0_46"])


def test_fila_de_unidades_fisicas_es_metadato():
    wb = Workbook()
    ws = wb.active
    ws.title = "Energía"
    ws.append(["Período", "Demanda", "Potencia", "Temperatura", "Producción"])
    ws.append(["", " GWh", "MW", " °C", "Miles Tns."])
    for k, anio in enumerate(range(2014, 2022)):
        ws.append([anio] + [_valor(j, k) for j in range(4)])
    hs, _ = _analizar(wb, "energia.xlsx")
    assert _reglas(hs, "cambio_tipo") == []


def test_s_d_sin_punto_final_es_sin_dato():
    assert es_sin_dato("s.d") and es_sin_dato("n.d") and not es_sin_dato("s.e")


# Casos de la revisión independiente: el ruido no puede llevarse errores de verdad.
def _filas(filas, titulo="Cuadro"):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    for f in filas:
        ws.append(f)
    return wb


def _v(j, k):
    return round(1000 + 37.5 * j + 11.25 * k + (k % 4) * 3.5, 2)


TRIMESTRES = [f"{r}.{a}" for a in (16, 17, 18) for r in ("I", "II", "III", "IV")]


def test_bloques_apilados_conservan_los_controles_de_cada_bloque():
    # Años y trimestres en la misma columna, sin filas de códigos: B vacía en 2018 y un texto en C en III.17.
    filas = [["Período", "Serie A", "Serie B", "Serie C"]]
    for k, p in enumerate(list(range(2016, 2021)) + TRIMESTRES):
        fila = [p] + [_v(j, k) for j in range(3)]
        if p == 2018:
            fila[2] = None
        if p == "III.17":
            fila[3] = "1.234,5x"
        filas.append(fila)
    hs, _ = _analizar(_filas(filas))
    assert _reglas(hs, "anio_no_entero") == []
    assert ("vacio_interior", "C4") in _reglas(hs, "vacio_interior")
    assert ("cambio_tipo", "D13") in _reglas(hs, "cambio_tipo")


def test_un_periodo_mal_escrito_en_el_borde_de_un_bloque_se_senala():
    # El primer trimestre escrito «2016» en lugar de «I.16»: queda como un año repetido del bloque anual.
    filas = [["Período", "Serie A", "Serie B", "Serie C"]]
    for k, p in enumerate(list(range(2016, 2021)) + [2016] + TRIMESTRES[1:]):
        filas.append([p] + [_v(j, k) for j in range(3)])
    hs, _ = _analizar(_filas(filas))
    assert [c for _, c in _reglas(hs, "anio_duplicado", "periodo_duplicado")] == ["A7"]


def test_faltante_antes_de_que_empiecen_las_otras_series_se_senala():
    # «Serie A» desde 2000 con 2005 vacío; B, C y D empiezan en 2010: 2005 no es un año sin datos del cuadro.
    filas = [["Año", "Serie A", "Serie B", "Serie C", "Serie D"]]
    for k, a in enumerate(range(2000, 2025)):
        filas.append([a] + [None if (j == 0 and a == 2005) or (j > 0 and a < 2010) else _v(j, k) for j in range(4)])
    hs, _ = _analizar(_filas(filas))
    assert ("vacio_interior", "B7") in _reglas(hs, "vacio_interior")


def test_periodo_sin_datos_pide_que_otras_series_lo_abarquen():
    from services.controles_excel.ejes import periodos_con_datos

    si, no = True, False
    # Cuatro series; el período 2 está vacío en todas y las otras lo abarcan: sin datos.
    assert periodos_con_datos([[si, si, no, si]] * 4) == [[si, si, no, si]] * 4
    # La primera serie tiene el período 1 vacío antes de que empiecen las demás: cuenta como con datos.
    marcas = periodos_con_datos([[si, no, si, si], [no, no, si, si], [no, no, si, si], [no, no, si, si]])
    assert marcas[0] == [si, si, si, si]
    # Con dos series no se decide nada.
    assert periodos_con_datos([[si, no, si], [si, no, si]]) == [[si, si, si], [si, si, si]]


def test_los_bloques_conservan_sus_controles_con_la_interpretacion_de_la_revision_rapida():
    # Por la ruta real la carga declara la interpretación (en la rápida, la del plan aritmético): los bloques «1.2»
    # tienen que llevar la del cuadro 1, si no sus series quedan fuera de los controles.
    from services import interpretacion_previa as ip

    filas = [["Período", "Serie A", "Serie B", "Serie C"]]
    for k, p in enumerate(list(range(2016, 2021)) + TRIMESTRES):
        fila = [p] + [_v(j, k) for j in range(3)]
        if p == 2018:
            fila[2] = None
        if p == "III.17":
            fila[3] = "1.234,5x"
        filas.append(fila)
    b = io.BytesIO()
    _filas(filas).save(b)
    contenido = b.getvalue()
    d, c = carga_archivos.parse_archivos_con_cobertura([("libro.xlsx", contenido)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    with ip.contexto(None):
        ip.declarar(ip.revision_rapida(d, r, [("libro.xlsx", contenido)], {"dominio": "general", "ambito": "publico"}, 1))
        hs = controles_excel.analizar([("libro.xlsx", contenido)], d, r["excel"])
    assert ("vacio_interior", "C4") in _reglas(hs, "vacio_interior")
    assert ("cambio_tipo", "D13") in _reglas(hs, "cambio_tipo")
