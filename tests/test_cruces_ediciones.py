"""Cruces entre ediciones: ficha de comparabilidad, emparejamiento y una sola señal por cuadro.

Casos tomados de la evaluación adversarial del 2026-10-03 (cruces E1–E7d), reconstruidos acá.
"""
import io

import pytest
from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion
from services.casos_revision import agrupar, es_nota
from services.controles_excel import ediciones

RUBROS = ["Salud", "Educación", "Promoción social", "Defensa", "Seguridad interior", "Transporte"]
BASE = {"Salud": 4100.0, "Educación": 3650.0, "Promoción social": 5200.0, "Defensa": 1450.0,
        "Seguridad interior": 1730.0, "Transporte": 2280.0}
A23 = list(range(2018, 2024))
A24 = list(range(2018, 2025))
NOTAS = {"cruce_no_comparable", "cobertura_cruce_ediciones"}


def serie(rubro, anios, factor=1.0):
    return {a: round(BASE[rubro] * (1.32 ** i) * factor, 1) for i, a in enumerate(anios)}


def libro(hoja, titulo, unidad, anios, filas):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws["A1"] = titulo
    ws["A2"] = unidad
    ws.cell(4, 1, "Concepto")
    for j, a in enumerate(anios):
        ws.cell(4, 2 + j, a)
    r = 5
    for etiqueta, vals in filas:
        ws.cell(r, 1, etiqueta)
        for j, a in enumerate(anios):
            ws.cell(r, 2 + j, vals.get(a))
        r += 1
    ws.cell(r, 1, "Total")
    for j, a in enumerate(anios):
        ws.cell(r, 2 + j, round(sum(v.get(a) or 0 for _, v in filas), 1))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def filas(anios, factor=1.0, etiquetas=None):
    return [((etiquetas or RUBROS)[k], serie(r, anios, factor)) for k, r in enumerate(RUBROS)]


def analizar(anterior, nueva):
    leidos = [("edicion2023.xlsx", anterior), ("edicion2024.xlsx", nueva)]
    d, c = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar(leidos, d, r["excel"])


def de_ediciones(hs):
    return [h for h in hs if "edicion" in str(h["evidencia"].get("regla")) or h["evidencia"].get("regla") in NOTAS]


def senales(hs):
    return [h for h in de_ediciones(hs) if h["evidencia"]["regla"] not in NOTAS]


def notas(hs, regla):
    return [h for h in hs if h["evidencia"].get("regla") == regla]


# 1. Ficha de comparabilidad ------------------------------------------------------------

INCOMPATIBLES = [
    ("concepto", ("Crédito vigente por finalidad", "En millones de pesos"), ("Gasto devengado por finalidad", "En millones de pesos"), 0.91),
    ("concepto", ("Recursos por rubro", "En millones de pesos"), ("Gastos por rubro", "En millones de pesos"), 1.07),
    ("período o corte", ("Gasto por finalidad - ejecución acumulada al 30 de junio", "En millones de pesos"),
     ("Gasto por finalidad - ejecución anual", "En millones de pesos"), 2.2),
    ("cobertura", ("Gasto por finalidad - Total país", "En millones de pesos"),
     ("Gasto por finalidad - Provincia de Córdoba", "En millones de pesos"), 0.083),
    ("cobertura", ("Gasto por finalidad - Administración Pública Nacional", "En millones de pesos"),
     ("Gasto por finalidad - Administración Central", "En millones de pesos"), 0.71),
    ("moneda", ("Gasto por finalidad", "En millones de pesos corrientes"),
     ("Gasto por finalidad", "En millones de pesos constantes de 2024"), 1.6),
]


@pytest.mark.parametrize("dimension,a,b,factor", INCOMPATIBLES)
def test_ficha_distinta_deja_una_nota_y_ninguna_senal(dimension, a, b, factor):
    hs = analizar(libro("Cuadro 1", *a, A23, filas(A23)), libro("Cuadro 1", *b, A24, filas(A24, factor)))
    assert not senales(hs), [h["descripcion"] for h in senales(hs)]
    nota = notas(hs, "cruce_no_comparable")
    assert len(nota) == 1
    assert dimension in nota[0]["evidencia"]["dimensiones_no_comparables"]
    assert f"No comparable por {dimension}" in nota[0]["descripcion"]
    assert "error de carga" not in nota[0]["descripcion"]
    # La nota no es un caso de control ni contexto temporal.
    r = agrupar(hs)
    assert r["notas_alcance"] >= 1 and all(not es_nota(c["primary_finding"]) for c in r["casos"])


def test_ficha_compatible_conserva_la_diferencia_real():
    nueva = filas(A24)
    nueva[1][1][2020] = round(nueva[1][1][2020] * 1.18, 1)
    hs = analizar(libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A23, filas(A23)),
                  libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A24, nueva))
    s = senales(hs)
    assert [h["evidencia"]["regla"] for h in s] == ["cambio_entre_ediciones"]
    assert s[0]["evidencia"]["indicador"] == "Educación" and s[0]["evidencia"]["periodo"] == 2020


def test_corte_parcial_de_un_anio_excluye_solo_ese_anio():
    a = libro("Cuadro 1", "Gasto por finalidad (2021: acumulado al 30/06, provisorio)", "En millones de pesos", A23, filas(A23))
    nueva = filas(A24)
    nueva[0][1][2019] = round(nueva[0][1][2019] * 1.3, 1)
    hs = analizar(a, libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A24, nueva))
    s = senales(hs)
    assert "Salud" in [h["evidencia"]["indicador"] for h in s] and {h["evidencia"]["periodo"] for h in s} == {2019}
    assert notas(hs, "cobertura_cruce_ediciones")[0]["evidencia"]["anios_excluidos"] == [2021]


def test_cuadro_en_otra_unidad_es_una_senal_de_cuadro_y_no_siete():
    # «miles de millones» frente a «millones»: todo el cuadro es la milésima parte (con redondeo).
    hs = analizar(libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A23, filas(A23)),
                  libro("Cuadro 1", "Gasto por finalidad", "En miles de millones de pesos", A24, filas(A24, 0.001)))
    s = senales(hs)
    assert len(s) == 1
    assert s[0]["evidencia"]["regla"] == "cambio_de_unidad_entre_ediciones"
    assert s[0]["evidencia"]["factor_escala"] == 0.001 and s[0]["evidencia"]["filas_cambiadas"] == 7
    assert len(s[0]["evidencia"]["celdas_participantes"]) > 7  # conserva todas las celdas


def test_revision_global_de_un_cuadro_es_una_sola_senal():
    vieja = [{"archivo": "a.xlsx", "hoja": "C", "cuadro": 1, "etiqueta": r, "fila": i + 2, "periodos": A23,
              "valores": [serie(r, A23)[a] for a in A23], "celdas": [f"{chr(66 + j)}{i + 2}" for j in range(6)]}
             for i, r in enumerate(RUBROS)]
    nueva = [{"archivo": "b.xlsx", "hoja": "C", "cuadro": 1, "etiqueta": r, "fila": i + 2, "periodos": A24,
              "valores": [serie(r, A24)[a] * (1.1 + 0.07 * (j % 3)) for j, a in enumerate(A24)],
              "celdas": [f"{chr(66 + j)}{i + 2}" for j in range(7)]} for i, r in enumerate(RUBROS)]
    s = [h for h in ediciones.analizar(vieja + nueva) if h["evidencia"]["regla"] not in NOTAS]
    assert [h["evidencia"]["regla"] for h in s] == ["cuadro_revisado_entre_ediciones"]


# 2. Cruces omitidos: emparejar por estructura y rótulos, e informar la cobertura -----------

@pytest.mark.parametrize("hoja_nueva,etiquetas", [
    ("Cuadro 1 (2024)", None),
    ("Cuadro 1", ["Servicios de salud", "Educación y cultura", "Promoción y asistencia social", "Defensa",
                  "Seguridad interior", "Transporte"]),
    ("Cuadro 1", ["SALUD", "Educacion", "Promocion  Social", "DEFENSA", "Seguridad Interior", "transporte"]),
])
def test_cruce_emparejado_por_hoja_o_rotulo_equivalente(hoja_nueva, etiquetas):
    nueva = filas(A24, etiquetas=etiquetas)
    nueva[1][1][2020] = round(nueva[1][1][2020] * 1.18, 1)
    hs = analizar(libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A23, filas(A23)),
                  libro(hoja_nueva, "Gasto por finalidad", "En millones de pesos", A24, nueva))
    s = senales(hs)
    assert len(s) == 1 and s[0]["evidencia"]["periodo"] == 2020
    cob = notas(hs, "cobertura_cruce_ediciones")
    assert len(cob) == 1 and cob[0]["evidencia"]["filas_emparejadas"] == 7 and cob[0]["evidencia"]["filas_totales"] == 7
    assert "7 de 7 filas emparejadas" in cob[0]["descripcion"]


def test_cobertura_del_emparejamiento_cuando_no_se_empareja_nada():
    otras = [("Rubro A", serie("Salud", A24, 3.3)), ("Rubro B", serie("Defensa", A24, 2.1))]
    hs = analizar(libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A23, filas(A23)),
                  libro("Cuadro 1", "Gasto por finalidad", "En millones de pesos", A24, otras))
    cob = notas(hs, "cobertura_cruce_ediciones")
    ev = cob[0]["evidencia"]
    assert ev["filas_emparejadas"] < ev["filas_totales"] == 7  # sólo «Total» coincide de rótulo
    assert f"{ev['filas_emparejadas']} de 7 filas emparejadas" in cob[0]["descripcion"]
    assert "Rubro A" in cob[0]["descripcion"]  # las filas sin pareja se nombran
