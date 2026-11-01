"""Falsas alarmas con planillas oficiales reales (0922): seis reglas que se afinan sin apagar errores.

Cada caso reproduce en chico la forma de un cuadro público que alarmaba sin error (deuda pública con
enlaces encadenados, cebada sembrada por departamento, PIB trimestral con «2019 4º trimestre», ventas
de supermercados con la variación porcentual bajo un rótulo repetido, salarios por rama y provincia) y,
al lado, el mismo cuadro con un error real que se sigue detectando.
"""

import io

from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion
from services.controles_excel import periodos, tablas_planas
from services.lentes import ajedrecista, cronista


def _analizar(w, nombre="libro.xlsx"):
    b = io.BytesIO()
    w.save(b)
    b = b.getvalue()
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, b)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, b)], d, r["excel"])


def _de(hs, regla, severidad=None):
    return [h for h in hs if h["evidencia"]["regla"] == regla and (severidad is None or h["severidad"] == severidad)]


def _csv(nombre, contenido):
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([(nombre, contenido)])
    hs, _ = tablas_planas.analizar([(nombre, contenido)], datos, lectura, lambda f: (f, {}), sanitizacion.sanear)
    return [h for h in hs if h["evidencia"].get("naturaleza") == "control"]


# --- A · valor fijo distinto de la fórmula de su línea ---------------------------------------------

def _deuda(enlace):
    """Columna C: o enlaces a otra celda de la misma columna («=+C12» en C10), o una suma de la fila."""
    w = Workbook()
    s = w.active
    s.title = "A.3.3"
    s.append(["Concepto", "Capital", "Saldo", "Tasa"])
    for r in range(2, 22):
        s.cell(r, 1, f"Rubro {r}")
        s.cell(r, 2, 100.0 + 7 * r)
        s.cell(r, 4, 3.0 + (r % 5))
        if enlace:
            # Como en el cuadro de deuda: cada celda copia la de dos filas más abajo; la fila 12 quedó tipeada.
            s.cell(r, 3, f"=+C{r + 2}" if r != 12 else 4.53)
        else:
            s.cell(r, 3, f"=B{r}+D{r}" if r != 12 else round(100.0 + 7 * 12 + 3.0 + 12 % 5 + 25, 2))
    if enlace:
        for r in (22, 23):
            s.cell(r, 3, 1.5 * r)
    return w


def test_enlace_encadenado_en_la_columna_no_es_valor_desactualizado_alto():
    hs = _analizar(_deuda(enlace=True))
    fijos = _de(hs, "valor_fijo_distinto_de_formula")
    assert not [h for h in fijos if h["severidad"] == "alta"]
    assert all(h["evidencia"].get("formula_es_enlace") for h in fijos)


def test_suma_de_la_fila_con_valor_tipeado_sigue_alta():
    hs = _analizar(_deuda(enlace=False))
    assert [h["evidencia"]["celda"] for h in _de(hs, "valor_fijo_distinto_de_formula", "alta")] == ["C12"]


# --- B · período faltante en datos de panel ---------------------------------------------------------

def _cebada(sin_anio_en=None, sin_anio_en_todos=None):
    filas = ["provincia,departamento,campania,superficie_sembrada"]
    deptos = [("Buenos Aires", "Azul"), ("Buenos Aires", "Tandil"), ("Buenos Aires", "Balcarce"),
              ("La Pampa", "Capital"), ("La Pampa", "Toay"), ("Córdoba", "Unión")]
    for k, (p, d) in enumerate(deptos):
        for a in range(2010, 2022):
            if ((p, d) == sin_anio_en and a in (2014, 2015)) or a == sin_anio_en_todos:
                continue
            filas.append(f"{p},{d},{a},{1200 + 150 * k + 13 * (a - 2010) + (a * 7 + k) % 11}")
    return ("\n".join(filas) + "\n").encode()


def test_departamento_sin_un_anio_es_panel_un_solo_aviso_medio():
    hs = _csv("cebada.csv", _cebada(sin_anio_en=("La Pampa", "Toay")))
    faltantes = _de(hs, "periodo_faltante")
    assert len(faltantes) == 1 and faltantes[0]["severidad"] == "media"


def test_anio_ausente_en_todo_el_cuadro_sigue_alto():
    hs = _csv("cebada.csv", _cebada(sin_anio_en_todos=2016))
    altas = _de(hs, "periodo_faltante", "alta")
    assert len(altas) == 1 and "falta en todo el cuadro" in altas[0]["descripcion"]


# --- C · encabezados de período en una sola celda ---------------------------------------------------

def test_rotulos_trimestrales_en_una_celda_son_periodos():
    for rotulo, anio, sub in (("2019 4º trimestre", 2019, 4), ("1er trim. 2020", 2020, 1), ("2019 T4", 2019, 4),
                              ("2024T4*", 2024, 4), ("2023 (1) 4º trimestre", 2023, 4)):
        p = periodos.leer(rotulo)
        assert p and (p["anio"], p["sub"], p["frecuencia"]) == (anio, sub, "trimestral"), rotulo
    assert periodos.leer("Año y trimestre") is None
    assert periodos.leer("2024T4*")["marca"] == "*"


TRIM = ["2019 1º trimestre", "2019 2º trimestre", "2019 3º trimestre", "2019 4º trimestre", "2019 Total",
        "2020 1º trimestre", "2020 2º trimestre"]
RUBROS = [("Consumo privado", 600.0), ("Consumo público", 150.0), ("Formación bruta de capital", 140.0),
          ("Exportaciones", 120.0), ("Importaciones", 130.0), ("Variación de existencias", 8.0)]


def _pib(rubro_con_signo=None):
    w = Workbook()
    s = w.active
    s.title = "cuadro 1"
    s.append(["Oferta y demanda globales a precios constantes"])
    s.append([])
    s.append(["Concepto", *TRIM])
    for k, (n, base) in enumerate(RUBROS):
        vals = [round(base * (1 + 0.01 * j) + (k * 3 + j) % 4, 2) for j in range(len(TRIM))]
        if n == "Variación de existencias":
            vals = [8.0, -5.2, 3.1, -9.4, -0.8, 6.3, -2.2]
        if n == rubro_con_signo:
            vals[2] = -vals[2]
        s.append([n, *vals])
    return w


def test_columnas_trimestrales_no_se_controlan_como_corte():
    hs = _analizar(_pib())
    assert not _de(hs, "signo_opuesto_columna") and not _de(hs, "valor_fuera_de_escala_columna")
    assert not _de(hs, "valor_negativo_en_magnitud_no_negativa")


def test_signo_invertido_en_la_serie_trimestral_sigue_detectado():
    hs = _analizar(_pib(rubro_con_signo="Consumo público"))
    altas = [h for h in hs if h["severidad"] == "alta" and h["evidencia"].get("celda") == "D5"]
    assert altas, sorted({(h["evidencia"]["regla"], h["evidencia"].get("celda")) for h in hs})


# --- D · negativos en una magnitud no negativa ------------------------------------------------------

def _ventas(negativo_en_total=False):
    datos = []
    for k in range(24):
        corriente = 6000.0 + 400 * k
        var = round(-30.0 + 4 * k, 1) if 3 <= k <= 6 else round(5.0 + k % 4, 1)
        datos.append({"Período": f"{2020 + k // 12}-{k % 12 + 1:02d}",
                      "a precios corrientes Total": -corriente if negativo_en_total and k == 7 else corriente,
                      "a precios corrientes Gran Buenos Aires": round(corriente * 0.6, 1),
                      "a precios corrientes Total_2": var})
    return datos


def test_rotulo_repetido_con_variacion_no_es_magnitud_no_negativa_alta():
    hs = [h for h in ajedrecista.imposibilidad(_ventas()) if h["evidencia"]["regla"] == "valor_negativo_en_magnitud_no_negativa"]
    assert hs and all(h["severidad"] == "media" for h in hs)
    assert "repite el de otra columna" in hs[0]["descripcion"]


def test_total_de_ventas_negativo_sigue_alto():
    hs = [h for h in ajedrecista.imposibilidad(_ventas(negativo_en_total=True))
          if h["evidencia"]["regla"] == "valor_negativo_en_magnitud_no_negativa" and h["severidad"] == "alta"]
    assert [h["evidencia"]["columna"] for h in hs] == ["a precios corrientes Total"]


# --- E · serie acumulada que retrocede --------------------------------------------------------------

VARIACION_ACUMULADA = "Variación porcentual acumulada del año respecto a igual período del año anterior"


def _mayoristas(nivel_retrocede=False):
    datos, acum = [], 0.0
    tasas = [5.1, 4.8, 4.2, 4.9, 3.7, 3.9, 4.4, 4.0, 3.8, 3.6, 3.9, 4.1]
    for m in range(1, 13):
        acum += 100.0 + m
        datos.append({"Período": f"2023-{m:02d}-01", VARIACION_ACUMULADA: tasas[m - 1],
                      "Importe acumulado": acum - (250.0 if nivel_retrocede and m == 6 else 0.0)})
    return datos


def test_variacion_acumulada_puede_bajar():
    assert ajedrecista.retrocesos_acumulados(_mayoristas(), [VARIACION_ACUMULADA, "Importe acumulado"]) == []


def test_importe_acumulado_que_baja_sigue_detectado():
    r = ajedrecista.retrocesos_acumulados(_mayoristas(nivel_retrocede=True), [VARIACION_ACUMULADA, "Importe acumulado"])
    assert [(i, col) for i, col, *_ in r] == [(5, "Importe acumulado")]


# --- F · período duplicado en CSV largo con varias dimensiones --------------------------------------

def _salarios(duplicado=False):
    datos = []
    for prov in ("Buenos Aires", "Córdoba", "Santa Fe"):
        for rama in ("Alimentos", "Textil", "Metalmecánica"):
            for m in range(1, 13):
                datos.append({"fecha": f"2022-{m:02d}-01", "provincia": prov, "rama": rama,
                              "salario": 100000 + 1000 * m + len(prov) * 37 + len(rama) * 11})
    if duplicado:
        datos.insert(5, dict(datos[4]))
    return datos


def test_fecha_repetida_en_clave_incompleta_no_es_duplicado():
    assert cronista.problemas_periodo(_salarios()) == []


def test_fecha_repetida_en_la_clave_completa_sigue_duplicada():
    assert [p["tipo"] for p in cronista.problemas_periodo(_salarios(duplicado=True))] == ["periodo_duplicado"]


# --- Ramas de formulas_fijas -------------------------------------------------------------------------------

def _serie_formulas(fijo_en, valor):
    """Fila 2: «Total» = suma de las filas 3 a 6 en cada año; un año queda tipeado."""
    w = Workbook()
    s = w.active
    s.title = "Serie"
    s.append(["Rubro", *range(2015, 2025)])
    s.append(["Total"])
    for k in range(4):
        s.append([f"Rubro {k}", *[1000.0 + 37 * k + 11 * j for j in range(10)]])
    for j in range(10):
        col = chr(ord("B") + j)
        s[f"{col}2"] = valor if col == fijo_en else f"=SUM({col}3:{col}6)"
    return w


def test_anio_viejo_tipeado_antes_del_tramo_calculado_queda_media():
    hs = _analizar(_serie_formulas("B", 4500.0))
    fijos = _de(hs, "valor_fijo_distinto_de_formula")
    assert fijos and all(h["severidad"] == "media" for h in fijos)
    assert not fijos[0]["evidencia"]["formulas_hermanas_a_ambos_lados"]


def test_formula_trasladada_sin_relacion_con_el_valor_queda_media():
    hs = _analizar(_serie_formulas("F", 3.2))
    fijos = _de(hs, "valor_fijo_distinto_de_formula")
    assert fijos and all(h["severidad"] == "media" for h in fijos)
    assert not fijos[0]["evidencia"]["resultado_plausible"]


def test_anio_tipeado_dentro_del_tramo_y_desactualizado_sigue_alto():
    hs = _analizar(_serie_formulas("F", 4300.0))
    assert [h["evidencia"]["celda"] for h in _de(hs, "valor_fijo_distinto_de_formula", "alta")] == ["F2"]


# === Revisión de la fase B (hallazgos 1-11) ================================================================

# --- 1 · columna total sin rótulo: filas de tasas y desvíos grandes ---------------------------------------

def _regiones(error=None):
    w = Workbook()
    s = w.active
    s.title = "regiones"
    s.append(["Indicador", "GBA", "Noreste", "Patagonia", "Conjunto"])
    for k in range(14):
        a, b, c = 5210.0 + 113 * k, 1830.0 + 41 * k, 977.0 + 29 * k
        conjunto = a + b + c
        if error == "chico" and k == 4:
            conjunto += 60
        if error == "grande" and k == 4:
            conjunto *= 1.4
        s.append([f"Población {k}", a, b, c, conjunto])
    # Una tasa del conjunto no es la suma de las tasas regionales.
    s.append(["Tasa de actividad (%)", 47.1, 44.3, 46.0, 46.4])
    s.append(["Índice de salarios", 213.4, 198.2, 230.9, 214.1])
    return w


def test_fila_de_tasa_no_se_toma_como_total_que_no_cierra():
    hs = _analizar(_regiones())
    assert not _de(hs, "columna_total_no_cierra")


def test_conjunto_que_no_cierra_por_poco_sigue_alto_y_el_desvio_grande_queda_medio():
    assert [h["evidencia"]["celda"] for h in _de(_analizar(_regiones("chico")), "columna_total_no_cierra", "alta")] == ["E6"]
    grande = _de(_analizar(_regiones("grande")), "columna_total_no_cierra")
    assert [(h["evidencia"]["celda"], h["severidad"]) for h in grande] == [("E6", "media")]


# --- 2, 6 y 3 · bloques paralelos anidados ----------------------------------------------------------------

COLS_BLOQUES = ["Asalariados", "No asalariados", "Patrones", "Familiares", "Cooperativistas", "Monotributistas"]


def _bloques(error=None, total_omite=False, provincias=None):
    """Regiones con «Varones» y «Mujeres», cada una seguida de sus provincias (también con «Varones» y
    «Mujeres»), y el total del país al pie."""
    w = Workbook()
    s = w.active
    s.title = "Cuadro 3"
    s.append(["Región y sexo", *COLS_BLOQUES])
    provincias = provincias or {"Norte": [("Salta", 3.1), ("Jujuy", 1.7)], "Sur": [("Chubut", 2.3), ("Neuquén", 2.9)]}

    def bloque(nombre, v, m, ajuste=None):
        fila = [round(a + b, 1) for a, b in zip(v, m)]
        if ajuste:
            fila[ajuste[0]] += ajuste[1]
        s.append([nombre, *fila])
        s.append(["Varones", *v])
        s.append(["Mujeres", *m])

    totales_v, totales_m = [0.0] * 6, [0.0] * 6
    for region, provs in provincias.items():
        vs = [[round(100 * f * (1 + 0.3 * j), 1) for j in range(6)] for _, f in provs]
        ms = [[round(90 * f * (1 + 0.2 * j), 1) for j in range(6)] for _, f in provs]
        rv = [round(sum(x[j] for x in vs), 1) for j in range(6)]
        rm = [round(sum(x[j] for x in ms), 1) for j in range(6)]
        bloque(region, rv, rm)
        for k, (nombre, _) in enumerate(provs):
            bloque(nombre, vs[k], ms[k], ajuste=(1, 25.0) if error == "provincia" and nombre == "Chubut" else None)
        if not (total_omite and region == "Sur"):
            totales_v = [round(a + b, 1) for a, b in zip(totales_v, rv)]
            totales_m = [round(a + b, 1) for a, b in zip(totales_m, rm)]
    bloque("Total del país", totales_v, totales_m, ajuste=(2, 40.0) if error == "total" else None)
    return w


def test_region_con_sus_provincias_no_se_cuenta_dos_veces():
    hs = _analizar(_bloques())
    assert not _de(hs, "total_omite_bloques") and not _de(hs, "total_bloques_no_cierra")


def test_total_del_pais_que_no_cierra_con_las_regiones_sigue_alto():
    hs = _analizar(_bloques(error="total"))
    assert [h["evidencia"]["celda"] for h in _de(hs, "total_bloques_no_cierra", "alta")] == ["D20"]


def test_total_que_omite_una_region_con_provincias_se_senala():
    hs = _analizar(_bloques(total_omite=True))
    senal = _de(hs, "total_omite_bloques")
    assert senal and senal[0]["evidencia"]["bloques_omitidos"] == ["Sur"]


def test_provincia_que_no_es_la_suma_de_varones_y_mujeres_se_senala():
    hs = _analizar(_bloques(error="provincia"))
    senal = _de(hs, "total_bloques_no_cierra")
    assert [h["evidencia"]["celda"] for h in senal] == ["C14"]
    assert senal[0]["evidencia"]["bloque"] == "Chubut" and "El bloque «Chubut»" in senal[0]["descripcion"]
    assert senal[0]["severidad"] == "media"


def test_total_aparente_con_mas_de_cuarenta_bloques_se_declara_sin_buscar_tramos():
    from services.controles_excel import totales_cruzados

    llamadas = []
    original = totales_cruzados._cierres_de_tramos
    totales_cruzados._cierres_de_tramos = lambda *a: llamadas.append(a) or original(*a)
    try:
        provs = {f"Región {k}": [(f"P{k}a", 1 + k / 10), (f"P{k}b", 2 + k / 7)] for k in range(42)}
        w = _bloques(provincias=provs)
        s = w.active
        for col in range(2, 8):  # el total ya no es la suma de las regiones, pero sigue siendo varones + mujeres
            s.cell(s.max_row - 2, col).value += 500
            s.cell(s.max_row - 1, col).value += 500
        _, r = _analizar_con_estructura(w)
    finally:
        totales_cruzados._cierres_de_tramos = original
    assert not llamadas
    declarados = r["excel"]["estructura_revision"].get("totales_aparentes_sin_comprobar") or []
    assert declarados and "más de los 40" in declarados[0]["motivo"]


def _analizar_con_estructura(w, nombre="libro.xlsx"):
    b = io.BytesIO()
    w.save(b)
    b = b.getvalue()
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, b)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, b)], d, r["excel"]), r


def test_total_omite_bloques_en_lenguaje_llano_nombra_lo_omitido():
    from services.reportes.lenguaje_llano import resumen_llano

    texto = resumen_llano({"lente": "conciliador", "descripcion": "texto técnico", "severidad": "alta", "evidencia": {
        "regla": "total_omite_bloques", "indicador": "Saldo bruto", "bloques_omitidos": ["Gobierno general", "Banco Central"],
        "total_declarado": 100.0, "total_reconstruido": 150.0}})
    assert "omite «Gobierno general» y «Banco Central»" in texto


# --- 4 · períodos fuera de orden en una serie corta ------------------------------------------------------

def _mensual_corta(orden):
    w = Workbook()
    s = w.active
    s.title = "mensual"
    s.append(["Mes", "Ventas", "Compras", "Stock"])
    for m in orden:
        s.append([f"2023-{m:02d}", 1000.0 + 31 * m, 700.0 + 17 * m, 5000.0 + 3 * m])
    return w


def test_dos_pares_invertidos_en_diez_meses_se_informan():
    hs = _analizar(_mensual_corta([1, 2, 4, 3, 5, 6, 8, 7, 9, 10]))
    assert _de(hs, "periodo_fuera_de_secuencia")


def test_lista_larga_de_fechas_de_emision_no_es_una_secuencia():
    import random

    rnd = random.Random(7)
    w = Workbook()
    s = w.active
    s.title = "titulos"
    s.append(["Título", "Fecha de emisión", "Monto emitido", "Monto en circulación"])
    for k in range(30):
        s.append([f"BONO {k}", f"{rnd.randint(2003, 2024)}-{rnd.randint(1, 12):02d}", 1000.0 + 77 * k, 400.0 + 31 * k])
    hs = _analizar(w)
    assert not [h for h in hs if h["evidencia"]["regla"].startswith("periodo_")]


# --- 5 · fila total sin rótulo con una columna lejos -----------------------------------------------------

def _anual_sobre_meses(desvio=None):
    w = Workbook()
    s = w.active
    s.title = "turismo"
    s.append(["Período", "Aérea", "Fluvial", "Terrestre", "Marítima", "Otras"])
    meses = [[1350.5 + 13 * m + 7 * j for j in range(5)] for m in range(6)]
    anual = [round(sum(x[j] for x in meses), 1) for j in range(5)]
    if desvio:
        anual[2] = round(anual[2] * (1 + desvio), 1)
    s.append(["Año 2022", *anual])
    for m, fila in enumerate(meses):
        s.append([["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio"][m], *fila])
    return w


def test_fila_anual_que_cierra_no_alarma():
    assert not _de(_analizar(_anual_sobre_meses()), "total_sin_rotulo_no_cierra")


def test_columna_que_se_aparta_mucho_del_total_anual_se_dice_como_media():
    senal = _de(_analizar(_anual_sobre_meses(desvio=0.3)), "total_sin_rotulo_no_cierra")
    assert [(h["evidencia"]["celda"], h["severidad"]) for h in senal] == [("D2", "media")]


# --- 7 · una columna «Estado» no es clave del formato largo ----------------------------------------------

def _con_estado(duplicado):
    datos = [{"fecha": f"2022-{m:02d}-01", "estado": "definitivo" if m <= 10 else "provisorio", "valor": 100 + m}
             for m in range(1, 13)]
    if duplicado:
        datos.append({"fecha": "2022-12-01", "estado": "definitivo", "valor": 113})
    return datos


def test_estado_provisorio_no_oculta_un_periodo_repetido():
    assert [p["tipo"] for p in cronista.problemas_periodo(_con_estado(duplicado=True))] == ["periodo_duplicado"]


def test_serie_con_estado_sin_repeticion_no_alarma():
    assert cronista.problemas_periodo(_con_estado(duplicado=False)) == []


# --- 8 · dato trimestral repetido en sus meses -----------------------------------------------------------

def _meses_con_trimestral():
    w = Workbook()
    s = w.active
    s.title = "mensual"
    s.append(["Mes", "PIB", "Empleo", "Inversión"])
    for m in range(1, 13):
        q = (m - 1) // 3
        s.append([f"2023-{m:02d}", 41237.5 + 1313 * q, 12143.7 + 211 * q, 8861.3 + 97 * q])
    return w


def _mensual_con_copia():
    w = Workbook()
    s = w.active
    s.title = "mensual"
    s.append(["Mes", "Ventas", "Compras", "Stock"])
    for m in range(1, 13):
        k = 6 if m == 7 else m  # julio repite junio
        s.append([f"2023-{m:02d}", 41000.0 + 1301 * k, 12100.0 + 211 * k, 8800.0 + 97 * k])
    return w


def test_dato_trimestral_repetido_en_sus_meses_no_es_fila_copiada():
    assert not _de(_analizar(_meses_con_trimestral()), "fila_repetida_periodo_anterior")


def test_un_mes_que_repite_el_anterior_sigue_detectado():
    senal = _de(_analizar(_mensual_con_copia()), "fila_repetida_periodo_anterior")
    assert len(senal) == 1 and senal[0]["evidencia"]["rango"].endswith("B8:D8")


# --- 9 · salto de escala con el cambio de unidad declarado -----------------------------------------------

def _agregados(nota):
    w = Workbook()
    s = w.active
    s.title = "Publicación-Agreg"
    if nota:
        s.append([nota])
        s.append([])
    s.append(["Indicador", *range(1994, 2003)])
    s.append(["A.F.8", 297.0, 313.0, 332.0, 322.0, 353.0, 391.0, 419.0, 501.0, 415.0])
    s.append(["A.F.9", 51.2, 40.3, 44.0, 58.4, 61.1, 66.9, 0.0751, 0.112, 0.219])
    s.append(["A.F.10", 159.0, 79.0, 115.0, 41.0, 111.0, 85.0, 50.0, 128.0, 93.0])
    return w


def test_redenominacion_declarada_en_la_nota_queda_como_contexto():
    senal = _de(_analizar(_agregados("Hasta 1999 en miles de pesos; desde 2000 en millones de pesos")), "salto_escala_final")
    assert senal and all(h["evidencia"].get("naturaleza") == "contexto" for h in senal)


def test_salto_sin_nota_sigue_alto():
    senal = _de(_analizar(_agregados(None)), "salto_escala_final")
    assert senal and senal[0]["severidad"] == "alta"


# --- N4 · saldo final con «Bajas» en positivo y una columna de variación ---------------------------------

def _saldos(bajas_positivas=False, error=False):
    w = Workbook()
    s = w.active
    s.title = "stock"
    s.append(["Concepto", "Bonos", "Letras", "Préstamos", "Avales", "Pagarés", "Otros", "Var. %"])
    ini = [5100.0, 2300.0, 7700.0, 1450.0, 980.0, 3310.0]
    altas = [820.0, 410.0, 990.0, 230.0, 160.0, 515.0]
    bajas = [-610.0, -380.0, -450.0, -120.0, -95.0, -270.0]
    fin = [a + b + c for a, b, c in zip(ini, altas, bajas)]
    if error:
        fin[1] += 70
    s.append(["Saldo inicial", *ini, None])
    s.append(["Altas", *altas, None])
    s.append(["Bajas", *([-x for x in bajas] if bajas_positivas else bajas), None])
    s.append(["Saldo final", *fin, 12.4])
    return w


def test_bajas_publicadas_en_positivo_no_se_toman_como_saldo_que_no_cierra():
    assert not _de(_analizar(_saldos(bajas_positivas=True)), "saldo_final_no_cierra")


def test_saldo_final_que_no_cierra_se_senala_y_la_variacion_no_participa():
    senal = _de(_analizar(_saldos(error=True)), "saldo_final_no_cierra")
    assert [h["evidencia"]["celda"] for h in senal] == ["C5"]
