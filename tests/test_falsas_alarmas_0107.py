"""Falsas alarmas que aparecieron al leer bien encabezados de dos niveles (0.10.6), en planillas oficiales correctas.

Casos del banco: perfil de vencimientos de la deuda pública (ministerio, años arriba y meses abajo, con
«Subtotal 2025» y pagos intermitentes), saldos por trimestre de un instrumento que se cancela y una
columna «Diferencia» de la ejecución presupuestaria de Mendoza. Cada prueba conserva además un error
del mismo tipo que se tiene que seguir señalando.
"""
import io
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura
from services.sanitizacion import sanear

MESES = ["Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre", "Enero", "Febrero", "Marzo",
         "Abril", "Mayo", "Junio"]


def _xlsx(filas, combinar=(), hoja="Cuadro"):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    for i, fila in enumerate(filas, 1):
        for j, v in enumerate(fila, 1):
            if v is not None:
                ws.cell(i, j, v)
    for rango in combinar:
        ws.merge_cells(rango)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _controles(contenido, nombre="cuadro.xlsx"):
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    saneados, reporte, _ = excel_cobertura.preparar(datos, cobertura, lambda f: (f, {}), sanear)
    return controles_excel.analizar([(nombre, contenido)], saneados, reporte["excel"])


def _altas(hs, regla=None):
    return [h for h in hs if h["severidad"] == "alta" and h["evidencia"].get("naturaleza") == "control"
            and (regla is None or h["evidencia"]["regla"] == regla)]


def _perfil(filas_datos, meses=MESES):
    """Perfil de vencimientos: «Instrumento» combinado en dos filas, 2025 sobre julio-diciembre y su
    subtotal, 2026 sobre enero-junio; los meses en la fila de abajo."""
    anios = ["Instrumento", "2025"] + [None] * 6 + ["2026"] + [None] * 5
    fila_meses = [None] + meses[:6] + ["Subtotal 2025"] + meses[6:]
    filas = [anios, fila_meses]
    for rotulo, valores in filas_datos:
        filas.append([rotulo] + valores[:6] + [sum(valores[:6])] + valores[6:])
    return _xlsx(filas, ("A1:A2", "B1:H1", "I1:N1"))


BASE = {
    "Títulos públicos": [2700, 820, 3150, 2130, 1600, 1500, 6670, 2770, 7480, 2100, 1900, 2600],
    "Organismos internacionales": [610, 640, 655, 600, 630, 700, 690, 650, 620, 640, 660, 700],
    "Letras del Tesoro": [1200, 1150, 1300, 1250, 1220, 1280, 1310, 1290, 1270, 1240, 1260, 1300],
    "Bonos provinciales": [400, 410, 395, 420, 405, 415, 430, 425, 410, 420, 435, 440],
}


def test_anio_arriba_y_meses_abajo_no_rompen_la_secuencia_al_cambiar_de_anio():
    hs = _controles(_perfil(list(BASE.items())))
    assert not _altas(hs, "periodo_fuera_de_secuencia")


def test_mes_desordenado_dentro_del_anio_sigue_senalandose():
    meses = MESES[:]
    meses[1], meses[2] = meses[2], meses[1]  # Septiembre antes que Agosto
    hs = _controles(_perfil(list(BASE.items()), meses))
    assert _altas(hs, "periodo_fuera_de_secuencia")


def test_mes_chico_en_un_cronograma_con_ceros_no_es_un_pico():
    # Banca comercial: pagos intermitentes, con meses en cero; septiembre con 3 entre 2.711 y 8.943.
    banca = [972, 2711, 3, 8943, 26819, 0, 0, 2690, 0, 1500, 0, 2200]
    hs = _controles(_perfil(list(BASE.items()) + [("Banca comercial", banca)]))
    assert not [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Banca comercial"]


def test_un_valor_mil_veces_menor_en_una_serie_regular_sigue_siendo_un_pico():
    letras = BASE["Letras del Tesoro"][:]
    letras[4] = 1.22  # 1.220 cargado en otra unidad
    filas = [(k, letras if k == "Letras del Tesoro" else v) for k, v in BASE.items()]
    hs = _controles(_perfil(filas))
    assert [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Letras del Tesoro"]


def _saldos(rubros):
    """Saldos por trimestre: año arriba («2024», «2025») e «I»…«IV» abajo."""
    filas = [["Instrumento", "2024", None, None, None, "2025", None, None, None],
             [None, "I", "II", "III", "IV", "I", "II", "III", "IV"]]
    filas += [[k] + v for k, v in rubros]
    return _xlsx(filas, ("A1:A2", "B1:E1", "F1:I1"))


SALDOS = {
    "Letras": [33700, 31900, 36800, 55300, 64300, 66000, 68100, 70200],
    "Bonos": [120000, 121500, 123100, 124000, 126200, 127900, 129000, 131500],
    "Organismos": [42000, 42600, 43100, 43900, 44400, 45000, 45800, 46100],
    "Adelantos": [3800, 3420, 2990, 2800, 2950, 3010, 3100, 3200],
}


def test_saldo_que_se_cancela_no_es_un_valor_aislado():
    rubros = list(SALDOS.items()) + [("Títulos en canje", [25600, 23700, 24100, 23900, 23500, 625, 0, 0])]
    hs = _controles(_saldos(rubros))
    assert not [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Títulos en canje"]


def test_saldo_cien_veces_menor_entre_vecinos_sigue_senalandose():
    bonos = SALDOS["Bonos"][:]
    bonos[4] = 1262  # 126.200 cargado cien veces menor
    rubros = [(k, bonos if k == "Bonos" else v) for k, v in SALDOS.items()]
    hs = _controles(_saldos(rubros))
    assert [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Bonos"]


def test_columna_diferencia_no_tiene_escala_propia():
    filas = [["Concepto", "Presupuesto", "Ejecutado", "Diferencia"]]
    for k, (p, e) in enumerate([(5000, 4800), (3200, 3150), (7100, 7000), (2500, 2499.98), (9000, 8700),
                                (4100, 3900), (6300, 6150), (1800, 1750)]):
        filas.append([f"Rubro {k}", p, e, round(e - p, 2)])
    hs = _controles(_xlsx(filas))
    assert not [h for h in _altas(hs, "valor_fuera_de_escala_columna") if h["evidencia"].get("columna") == "Diferencia"]


def test_serie_que_empieza_en_cero_sigue_senalando_un_valor_mil_veces_menor():
    # Los ceros del comienzo (el rubro todavía no existía) no hacen intermitente a la serie.
    filas = [["Rubro"] + [str(a) for a in range(2018, 2025)],
             ["Nuevo programa", 0, 0, 5400, 5900, 6.1, 6800, 7200]]
    for k, base in enumerate([1200, 3400, 870, 2500]):
        filas.append([f"Programa {k}"] + [round(base * 1.06 ** t) for t in range(7)])
    hs = _controles(_xlsx(filas))
    assert [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Nuevo programa"]


def test_valor_cien_veces_mayor_junto_a_un_cero_sigue_senalandose():
    bonos = SALDOS["Bonos"][:]
    bonos[2], bonos[3] = 0, 12400000  # un cero y, al lado, 124.000 cargado cien veces mayor
    rubros = [(k, bonos if k == "Bonos" else v) for k, v in SALDOS.items()]
    hs = _controles(_saldos(rubros))
    assert [h for h in _altas(hs) if h["evidencia"].get("indicador") == "Bonos"]


def test_meses_acumulados_no_se_leen_como_meses_sueltos():
    from services.controles_excel import modelo

    class Hoja:
        def __init__(self, filas):
            self.filas = filas

        def cell(self, r, c):
            return type("Celda", (), {"value": self.filas[r - 1][c - 1]})()

    meses = ["Acumulado a julio", "Acumulado a agosto", "Acumulado a septiembre", "Acumulado a octubre"]
    ws = Hoja([["Rubro", "2025", None, None, None], [None] + meses])
    c = {"header": 2, "columnas": [{"idx": j, "numerica": True} for j in range(2, 6)]}
    assert not modelo._dos_niveles_en(c, ws, 1, 2)


def test_la_suite_nunca_usa_la_base_del_entorno():
    # tests/conftest.py fija la base de pruebas aunque el entorno traiga otra (la de producción tras `set -a`).
    raiz = Path(__file__).resolve().parents[1]
    entorno = dict(os.environ, DATABASE_URL="postgresql://nadie:nada@127.0.0.1:1/produccion")
    salida = subprocess.run(
        [sys.executable, "-c", "import tests.conftest; from config import settings; print(settings.database_url)"],
        cwd=raiz, env=entorno, capture_output=True, text=True, timeout=60)
    assert salida.stdout.strip().splitlines()[-1] == "sqlite://"
