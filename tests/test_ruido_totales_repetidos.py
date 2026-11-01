"""0.10.58: ruido de totales y valores repetidos en libros oficiales grandes.

Casos reales que lo motivaron (libros correctos del banco, con 86 «columna total que no cierra» y 50 «valor
idéntico en otra fila» de prioridad alta):
- Filas en cero al pie de una serie (años todavía sin datos) contaban como filas en que un «Total» cierra con
  cualquier relación: «Oferta Global Total» quedaba confirmada como la suma de «Consumo privado» y «Consumo
  público» y se señalaban las veinte filas con datos.
- Un total que en un tramo de la fuente suma otra cosa (2006-2008, siempre 0,4-0,6 % arriba) daba una señal por
  fila; ahora es una señal por tramo, de prioridad alta (un mismo error repetido en bloque deja lo mismo), con la
  diferencia de cada fila y un rango que las cubre a todas.
- Promedios de valores escritos con pocos decimales (tasas anuales = promedio de trimestres con tres decimales;
  73,9666… = 221,9 / 3) se leían con cuatro decimales: ni el redondeo de un total ni la coincidencia de dos
  promedios eran raros. Sólo cuenta para decidir si un total cierra; un promedio de enteros no es resolución gruesa.
- Una serie empalmada (índice viejo de dos decimales por un coeficiente) repite valores de diez cifras: con su
  unidad común, dos iguales no son raros en una columna larga; en una corta, sí.
Libros sintéticos y funciones puras; los escenarios de la revisión independiente del 28/10, como pruebas de que el
error se sigue señalando.
"""
import io

from openpyxl import Workbook
from openpyxl.utils import range_boundaries

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion
from services.controles_excel import totales_cruzados as tc
from services.controles_excel import transversal as tv


def _analizar(wb, nombre="cuadro.xlsx"):
    b = io.BytesIO()
    wb.save(b)
    contenido = b.getvalue()
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    return controles_excel.analizar([(nombre, contenido)], d, r["excel"])


def _regla(hs, regla):
    return [h for h in hs if h["evidencia"].get("regla") == regla]


def _cubre(h, celda):
    """La señal ubica la celda: como la compuerta, se expanden los rangos de sus ubicaciones."""
    col, fila = range_boundaries(celda)[:2]
    for u in h["evidencia"].get("ubicaciones_excel") or []:
        c1, f1, c2, f2 = range_boundaries(str(u).rpartition("!")[2])
        if c1 <= col <= c2 and f1 <= fila <= f2:
            return True
    return h["evidencia"].get("celda") == celda


def _alta_en(hs, celda):
    return any(h["severidad"] == "alta" and _cubre(h, celda) for h in hs)


# --- Funciones puras -------------------------------------------------------------------------------------------

def test_promedio_de_valores_con_pocos_decimales():
    assert tv._promedio_de(221.9 / 3) == (3, 1)
    assert tv._promedio_de(0.363 / 4) == (4, 3)  # 0,09075: cinco decimales escritos, tres de origen
    assert tv._promedio_de(0.0715) is None  # cuatro escritos: se respetan (no se sabe si es un promedio)
    assert tv.decimales_efectivos(0.363 / 4) == 3
    assert tv.decimales_efectivos(0.406 / 3) == 3  # 0,13533…
    assert tv.decimales_efectivos(22027.600000000002) == 1  # ruido de punto flotante
    assert tv.decimales_efectivos(57.577055248729536) == 4  # cifra calculada: sin resolución de origen
    assert tv.decimales_efectivos(0.0836) == 4  # escrito con cuatro: se respeta
    assert tv.decimales_efectivos(1234) == 0
    # Un promedio de enteros no es resolución gruesa (límite conocido: uno de 6 o de 12 enteros encaja antes con k = 3);
    # una cifra grande con decimales propios tampoco. Un entero con ruido de punto flotante sigue siendo entero.
    assert tv._promedio_de(301 / 3) is None and tv.decimales_efectivos(301 / 3) == 4
    assert tv._promedio_de(100.00000000000001) == (1, 0)
    assert tv.decimales_efectivos(1234567.890123) == 4
    # _tolerancia (los demás controles) sigue como antes; _tolerancia_cierre usa la resolución efectiva.
    assert tv._tolerancia([22027.600000000002]) < 1e-3 < tv._tolerancia_cierre([22027.600000000002])


def test_cifras_efectivas_de_un_promedio():
    assert tv._sig(221.9 / 3) >= 7
    assert tv._sig_efectivas(221.9 / 3) == 4
    assert tv._sig_efectivas(1915878.76) == 9
    assert tv._sig_efectivas(57.577055248729536) == tv._sig(57.577055248729536)


def test_fila_en_cero_no_es_comparable():
    assert not tc._comparable(0, [0, 0, None])
    assert not tc._comparable(None, [1, 2])
    assert not tc._comparable(5, [None, 2])
    assert tc._comparable(0, [1, -1])
    assert tc._comparable(3, [0, 0])


def test_tramo_parejo_y_filas_sueltas():
    orden = list(range(10, 40))
    fallas = {r: (5.0, 0.5 + 0.01 * (r % 3)) for r in range(14, 20)}  # seis filas seguidas, parejas
    fallas[30] = (900.0, 9.0)  # una suelta
    tramos, sueltas = tc._tramos_parejos(orden, fallas)
    assert tramos == [list(range(14, 20))]
    assert sueltas == [30]


def test_tramo_que_no_es_parejo_queda_en_filas_sueltas():
    orden = list(range(10, 20))
    # Seguidas pero con diferencias dispares (0,5 % y 5 %), o con signo cambiado, o sólo dos.
    assert tc._tramos_parejos(orden, {11: (1, 0.5), 12: (1, 5.0), 13: (1, 0.6)}) == ([], [11, 12, 13])
    assert tc._tramos_parejos(orden, {11: (1, 0.5), 12: (-1, 0.5), 13: (1, 0.5)}) == ([], [11, 12, 13])
    assert tc._tramos_parejos(orden, {11: (1, 0.5), 12: (1, 0.5)}) == ([], [11, 12])
    # Una fila que cierra en el medio corta el tramo; también una fila del cuadro sin comparación (en cero o sin total).
    # Una fila en blanco no es del cuadro: no corta, y el rango de la señal la cubre.
    assert tc._tramos_parejos(orden, {11: (1, 0.5), 12: (1, 0.5), 14: (1, 0.5), 15: (1, 0.5)}) == ([], [11, 12, 14, 15])
    assert tc._tramos_parejos([3, 4, 50, 51], {3: (1, 0.5), 50: (1, 0.5), 51: (1, 0.5)}) == ([], [3, 50, 51])
    # Un desvío grande (un ×1000 o un ÷10 en bloque) no es un tramo de la fuente: fila por fila.
    assert tc._tramos_parejos(orden, {11: (-1, 98.5), 12: (-1, 98.6), 13: (-1, 98.7)}) == ([], [11, 12, 13])
    assert tc._tramos_parejos(orden, {11: (1, 29.0), 12: (1, 30.0), 13: (1, 29.5)}) == ([], [11, 12, 13])


def test_un_borde_distinto_corta_el_tramo():
    # Sin recortes: un borde que se aparta menos (o más) que el resto no entra en el tramo parejo.
    orden = [20, 21, 22, 23]
    assert tc._tramos_parejos(orden, {20: (1, 0.589), 21: (1, 0.536), 22: (1, 0.095)}) == ([], [20, 21, 22])
    assert tc._tramos_parejos(orden, {20: (1, 0.5), 21: (1, 0.55), 22: (1, 4.0)}) == ([], [20, 21, 22])
    assert tc._tramos_parejos(list(range(10)), {3: (1, 0.001), 4: (1, 40), 5: (1, 45)}) == ([], [3, 4, 5])


def test_unidad_comun_de_una_serie_empalmada():
    coef = 1.1256331778419
    columna = [round(40 + (m * 7919 % 12007) / 100, 2) * coef for m in range(200)]
    assert tv._coincidencias_por_unidad(columna[5], columna) > tv.COINCIDENCIA_ESPERABLE
    continuos = [1800000 + 12345.67 * m + (m % 7) * 1013.29 for m in range(200)]
    assert tv._coincidencias_por_unidad(continuos[5], continuos) == 0
    assert tv._coincidencias_por_unidad(columna[5], columna[:8]) == 0  # muy pocos valores para afirmarlo
    # Columna corta de la misma serie: la unidad es común, pero una coincidencia casual sigue siendo improbable.
    corta = [round(100 + (m * 3571 % 5000) / 100, 2) * coef for m in range(13)]
    assert 0 < tv._coincidencias_por_unidad(corta[5], corta) < tv.COINCIDENCIA_ESPERABLE


# --- Filas en cero ---------------------------------------------------------------------------------------------

def _oferta_demanda(ceros=30, alterar=None):
    """«Oferta Total» = PIB + Importación, a la DERECHA de sus partes (el control de panel la compara con las dos
    columnas que la siguen: los consumos); «Exportaciones Total» = Bienes + Servicios, seguida de sus partes. Al
    pie, `ceros` filas en cero (años todavía sin datos)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "OyD"
    ws["B2"] = "Oferta y demanda globales"
    ws["B3"] = "Millones de pesos"
    ws.append([])
    ws.append([None, "Oferta Global", None, None, "Consumo", None, "Exportaciones", None, None])
    ws.append([None, "PIB", "Importación", "Total", "Privado", "Público", "Total", "Bienes", "Servicios"])
    for k in range(20):
        pib, imp = 1000.0 + 37 * k + (k % 4) * 11, 150.0 + 9 * k + (k % 3) * 7
        priv, pub = round(0.62 * (pib + imp), 2), round(0.18 * (pib + imp), 2)
        bienes, servicios = round(0.13 * pib + k, 2), round(0.04 * pib + 2 * k, 2)
        exportaciones = round(bienes + servicios, 2)
        if alterar == k:
            exportaciones = round(exportaciones * 1.07, 2)
        ws.append([1993 + k, pib, imp, pib + imp, priv, pub, exportaciones, bienes, servicios])
    for _ in range(ceros):
        ws.append([None] + [0] * 8)
    return wb


def test_filas_en_cero_no_confirman_un_total_falso():
    # Con treinta filas en cero «Oferta Total = Consumo privado + Consumo público» cerraba en el 60 % de las filas.
    assert _regla(_analizar(_oferta_demanda()), "columna_total_no_cierra") == []


def test_un_total_verdadero_con_filas_en_cero_sigue_controlado():
    hs = _regla(_analizar(_oferta_demanda(alterar=7)), "columna_total_no_cierra")
    assert [h["evidencia"]["celda"] for h in hs] == ["G14"]


# --- Tramo parejo ----------------------------------------------------------------------------------------------

RUBROS = ["Alimentos", "Limpieza", "Indumentaria", "Electrónicos", "Otros"]


def _ventas(tramo=(), extra=0.005, sueltas=None, bloque=None):
    """Dos paneles con su «Total» (índice y miles de pesos), 48 meses (fila 6 + mes). En `tramo` el total en pesos
    queda `extra` arriba de la suma; `sueltas` = {mes: factor} altera un total suelto; `bloque` = (meses, rubro,
    factor) altera un rubro en esos meses después de sumar (un error repetido en bloque)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Super"
    ws["B2"] = "Encuesta de supermercados"
    ws.append([])
    ws.append(["Período", "Índice"] + [None] * 5 + ["Miles de pesos"] + [None] * 5)
    ws.append([None, "Total"] + RUBROS + ["Total"] + RUBROS)
    for m in range(48):
        partes = [round(1000 * (j + 1) * (1 + 0.01 * m) + 37 * ((m * (j + 3)) % 7)) for j in range(5)]
        total = sum(partes)
        if m in tramo:
            total = round(total * (1 + extra + 0.0004 * (m % 3)))
        if sueltas and m in sueltas:
            total = round(total * sueltas[m])
        indice = [round(100 * p / (1000 * (j + 1)), 1) for j, p in enumerate(partes)]
        if bloque and m in bloque[0]:
            partes[bloque[1]] = round(partes[bloque[1]] * bloque[2], 2)
        ws.append([f"{2006 + m // 12}-{m % 12 + 1:02d}", round(100 * total / 15000, 1)] + indice + [total] + partes)
    return wb


def test_tramo_parejo_es_una_sola_senal_de_prioridad_alta():
    hs = _regla(_analizar(_ventas(tramo=range(5, 17))), "columna_total_no_cierra")
    assert len(hs) == 1
    e = hs[0]["evidencia"]
    assert e["tramo"]["filas"] == 12 and e["ubicaciones_excel"][1].endswith("H11:H22")
    assert hs[0]["severidad"] == "alta"  # un mismo error repetido en bloque deja lo mismo: no se baja
    assert len(e["tramo"]["diferencias_con_signo"]) == 12


def test_error_en_bloque_x1000_se_senala_fila_por_fila():
    # Revisión del 28/10 (S1): «Alimentos» ×1000 en tres meses seguidos daba un tramo parejo (98,5-98,7 %).
    hs = _analizar(_ventas(bloque=(range(20, 23), 0, 1000)))
    for fila in (26, 27, 28):
        assert _alta_en(hs, f"H{fila}"), fila
    assert not any("tramo" in h["evidencia"] for h in _regla(hs, "columna_total_no_cierra"))


def test_error_en_bloque_parejo_y_chico_sigue_siendo_alta():
    # S2: un rubro corrido unos pocos por mil en ocho meses deja un tramo parejo: queda una señal alta que cubre
    # cada fila.
    hs = _analizar(_ventas(bloque=(range(20, 28), 1, 1.003)))
    for fila in range(26, 34):
        assert _alta_en(hs, f"H{fila}"), fila


def test_error_pegado_a_un_tramo_queda_cubierto_y_es_el_ancla():
    # S3: un total alterado un 3 % justo después de un tramo legítimo del 2 %: queda en el tramo, de prioridad alta,
    # y la señal se ancla en esa fila (la que más se aparta).
    hs = _analizar(_ventas(tramo=range(5, 17), extra=0.02, sueltas={17: 1.03}))
    assert _alta_en(hs, "H23")
    assert [h["evidencia"]["celda"] for h in _regla(hs, "columna_total_no_cierra")] == ["H23"]


def test_error_que_completa_un_tramo_queda_cubierto():
    # S3b: sólo dos filas legítimas del 2 % y el error del 3 % completa las tres.
    hs = _analizar(_ventas(tramo=range(15, 17), extra=0.02, sueltas={17: 1.03}))
    assert _alta_en(hs, "H23")


def test_error_al_lado_de_un_tramo_distinto_tiene_su_senal():
    # S4: al lado de un tramo del 10 %, el 3 % no es parejo y queda con su propia señal.
    hs = _analizar(_ventas(tramo=range(5, 17), extra=0.10, sueltas={17: 1.03}))
    assert _alta_en(hs, "H23")
    assert "H23" in [h["evidencia"]["celda"] for h in _regla(hs, "columna_total_no_cierra")]


def test_error_suelto_fuera_del_tramo_sigue_siendo_alta():
    hs = _regla(_analizar(_ventas(tramo=range(5, 17), sueltas={30: 1.08})), "columna_total_no_cierra")
    sueltas = [h for h in hs if "tramo" not in h["evidencia"]]
    assert [h["evidencia"]["celda"] for h in sueltas] == ["H36"]
    assert sueltas[0]["severidad"] == "alta"


def test_error_suelto_dentro_del_tramo_no_se_esconde():
    # El tramo deja de ser parejo (8 % contra 0,5 %): cada fila queda con su señal, también la alterada.
    hs = _regla(_analizar(_ventas(tramo=range(5, 17), sueltas={10: 1.08})), "columna_total_no_cierra")
    assert "H16" in [h["evidencia"]["celda"] for h in hs]
    assert all("tramo" not in h["evidencia"] for h in hs)


# --- Promedios -------------------------------------------------------------------------------------------------

def _tasas(error=None):
    """Tasas = promedio de tres meses publicados con tres decimales; en cada mes el total es la suma de sus partes
    redondeada (puede diferir en 0,001)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tasas"
    ws["B2"] = "Tasa de subocupación"
    ws.append([])
    ws.append(["Año", "Total país", None, None, "Región A", None, None, "Región B", None, None])
    ws.append([None] + ["Total", "Demandantes", "No demandantes"] * 3)
    for k in range(30):
        fila = [1990 + k]
        for g in range(3):
            meses = []
            for q in range(3):
                d = 0.05 + 0.0013 * ((k * 7 + q * 3 + g) % 11) + 0.0001 * q
                n = 0.02 + 0.0011 * ((k * 5 + q + g * 2) % 9) + 0.00004 * k
                meses.append((round(d + n, 3), round(d, 3), round(n, 3)))
            fila += [sum(t[i] for t in meses) / 3 for i in range(3)]
        if error == k:
            fila[1] = fila[2] + fila[3] + 0.011
        ws.append(fila)
    return wb


def test_promedios_de_tasas_redondeadas_no_son_un_total_que_no_cierra():
    assert _regla(_analizar(_tasas()), "columna_total_no_cierra") == []


def test_un_total_de_tasas_que_no_cierra_de_verdad_se_senala():
    hs = _regla(_analizar(_tasas(error=12)), "columna_total_no_cierra")
    assert [h["evidencia"]["celda"] for h in hs] == ["B18"]


def _promedios_de_enteros(columnas, cambiar=None):
    """Promedios de tres enteros (100,333…) por rubro, 12 meses, con la columna de la suma primero. `cambiar` =
    (mes, columna, función) altera una celda después de sumar."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro"
    ws["A1"] = "Promedios trimestrales"
    ws.append([])
    ws.append(["Mes"] + columnas)
    for m in range(12):
        partes = [(300 + 37 * j + 7 * m + (m * (j + 2)) % 5) / 3 for j in range(len(columnas) - 1)]
        fila = [sum(partes)] + partes
        if cambiar and cambiar[0] == m:
            fila[cambiar[1]] = cambiar[2](fila[cambiar[1]])
        ws.append([f"Mes {m + 1}"] + fila)
    return wb


def test_promedios_de_enteros_no_agrandan_la_tolerancia():
    # Revisión del 28/10 (S6): con partes 100,333… un total corrido en una unidad pasaba (tolerancia de ±2).
    hs = _analizar(_promedios_de_enteros(["Total", "A", "B", "C", "D"], cambiar=(4, 0, lambda v: v + 1)))
    assert _alta_en(hs, "B8")


def test_total_sin_rotulo_de_promedios_chicos_se_sigue_reconociendo():
    # S7: «Personas» = Varones + Mujeres con promedios de enteros entre 20 y 40, sin la palabra «Total».
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro"
    ws["A1"] = "Personas por mes"
    ws.append([])
    ws.append(["Mes", "Personas", "Varones", "Mujeres"])
    for m in range(10):
        varones, mujeres = (40 + 5 * m + m % 3) / 3, (55 + 4 * m + (m * 2) % 5) / 3
        if m == 6:
            varones *= 1.3
            ws.append([f"Mes {m + 1}", (40 + 5 * m + m % 3) / 3 + mujeres, varones, mujeres])
        else:
            ws.append([f"Mes {m + 1}", varones + mujeres, varones, mujeres])
    hs = _analizar(wb)
    assert _alta_en(hs, "B10") or _alta_en(hs, "C10")


# --- Valores repetidos -----------------------------------------------------------------------------------------

def _serie(valores, nombre="Alimentos y bebidas"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Serie"
    ws["B2"] = "Estimador mensual, base 2012"
    ws.append([])
    ws.append(["Período", nombre, "Otro rubro"])
    for m, v in enumerate(valores):
        ws.append([f"{1994 + m // 12}-{m % 12 + 1:02d}", v, 1000 + 13.37 * m])
    return wb


def test_dos_promedios_iguales_no_son_una_copia():
    # Promedios trimestrales de meses con un decimal: 67,4333… (= 202,3 / 3) dos veces.
    valores = [(150.0 + 3.1 * m + (m % 5) * 0.7) / 3 for m in range(60)]
    valores[40] = valores[12] = 202.3 / 3
    assert _regla(_analizar(_serie(valores)), "valor_identico_otra_fila") == []


def test_serie_empalmada_repite_valores_de_diez_cifras():
    # Índice viejo de dos decimales por un coeficiente: dos meses con el mismo índice viejo dan el mismo valor de
    # diez cifras.
    coef = 1.1256331778419
    viejo = [round(40 + (m * 7919 % 12007) / 100, 2) for m in range(260)]
    viejo[150], viejo[200] = viejo[20], viejo[33]
    assert _regla(_analizar(_serie([v * coef for v in viejo])), "valor_identico_otra_fila") == []


def test_copia_en_una_serie_empalmada_corta_se_sigue_senalando():
    # Revisión del 28/10 (S5): 13 meses de la misma serie empalmada con una copia; la coincidencia casual es
    # improbable (unas 0,02 esperadas).
    coef = 1.1256331778419
    viejo = [round(100 + (m * 3571 % 5000) / 100, 2) for m in range(13)]
    viejo[12] = viejo[10]
    hs = _regla(_analizar(_serie([v * coef for v in viejo])), "valor_identico_otra_fila")
    assert [h["evidencia"]["celda"] for h in hs] == ["B17"]


def test_copia_en_una_serie_empalmada_de_indice_alto_se_sigue_senalando():
    # Verificación del revisor (A): índice viejo de 100,00 a 400,00; con denominadores de hasta 20.000 la unidad salía
    # tres veces mayor y la copia del 317,37 se callaba.
    coef = 1.1256331778419
    viejo = [round(100 + (m * 7919 % 30001) / 100, 2) for m in range(43)]
    viejo[20] = viejo[40] = 317.37
    hs = _regla(_analizar(_serie([v * coef for v in viejo])), "valor_identico_otra_fila")
    assert [h["evidencia"]["celda"] for h in hs] == ["B45"]


def test_copia_en_un_tramo_empalmado_de_una_columna_mixta_se_sigue_senalando():
    # Verificación del revisor (B): 22 meses empalmados y 178 continuos; las coincidencias esperadas se cuentan
    # entre los 22, no con toda la columna.
    coef = 1.1256331778419
    viejo = [round(60 + (m * 3571 % 10001) / 100, 2) for m in range(22)]  # 231 pares en 10.000 niveles: 0,023
    viejo[21] = viejo[10]
    continuos = [55 + 0.2718281828459 * m + (m % 7) * 0.0314159265358 for m in range(178)]
    hs = _regla(_analizar(_serie([v * coef for v in viejo] + continuos)), "valor_identico_otra_fila")
    assert [h["evidencia"]["celda"] for h in hs] == ["B26"]


def test_un_valor_copiado_de_otra_fila_se_sigue_senalando():
    valores = [1800000 + 12345.67 * m + (m % 7) * 1013.29 for m in range(60)]
    valores[58] = valores[56]
    hs = _regla(_analizar(_serie(valores, "Remuneración")), "valor_identico_otra_fila")
    assert [h["evidencia"]["celda"] for h in hs] == ["B63"]
