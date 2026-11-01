"""Encabezado de dos niveles por combinaciones verticales, total arriba y cruce fila × columna del lector.

Caso que lo motivó (prueba de punta a punta del 22/10/2026): un censo nacional, cuadro 4 de viviendas colectivas
(instituto de estadística). «Código», «Jurisdicción» y «Total» combinados en dos filas; los tipos, debajo de «Tipo de vivienda
colectiva»; el total del país ARRIBA de las jurisdicciones. Con un componente alterado, los tipos quedaban como
«columna_5…», el motor decía «Sin fila de total» y el lector señalaba el total de la fila y el de la columna por
separado, sin nombrar la celda en la que se cruzan.
"""
import io
import re

import pytest
from openpyxl import Workbook

from services import carga_archivos, casos_revision, controles_excel, excel_cobertura
from services import lector_mapa as lm
from services.sanitizacion import sanear

TIPOS = ["Hotel", "Hospital", "Geriátrico", "Prisión"]
JURISDICCIONES = ["Catamarca", "Chaco", "Córdoba", "Jujuy", "Mendoza", "Salta"]
COLS = "DEFG"
NOMBRE = "c2022_vivienda_c4.xlsx"


def _valores(error=None):
    orig = [[(13 * i + 7 * j) % 50 + 10 + 3 * i for j in range(len(TIPOS))] for i in range(len(JURISDICCIONES))]
    vals = [list(v) for v in orig]
    for i, j, cambio in error or []:
        vals[i][j] += cambio
    return orig, vals


def _celdas(error=None, total_arriba=None):
    """Grilla del cuadro (fila 4-5 encabezado, fila 6 total del país, filas 7-12 jurisdicciones)."""
    orig, vals = _valores(error)
    tot = [sum(o[j] for o in orig) for j in range(len(TIPOS))]
    celdas = {"A1": "Censo Nacional de Población, Hogares y Viviendas 2022. Cuadro 4", "A4": "Código",
              "B4": "Jurisdicción", "C4": "Total de viviendas colectivas", "D4": "Tipo de vivienda colectiva",
              "B6": "Total", "C6": total_arriba if total_arriba is not None else sum(tot)}
    celdas.update({f"{c}5": t for c, t in zip(COLS, TIPOS)})
    celdas.update({f"{c}6": (v if total_arriba is None else 7 * v) for c, v in zip(COLS, tot)})
    for i, (nombre, o, v) in enumerate(zip(JURISDICCIONES, orig, vals)):
        f = 7 + i
        celdas.update({f"A{f}": f"{10 + 4 * i}", f"B{f}": nombre, f"C{f}": sum(o)})
        celdas.update({f"{c}{f}": x for c, x in zip(COLS, v)})
    return celdas


def _xlsx(celdas, combinadas=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro 4"
    for ref, v in celdas.items():
        ws[ref] = v
    if combinadas:
        for rango in ("A4:A5", "B4:B5", "C4:C5", "D4:G4"):
            ws.merge_cells(rango)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _abajo(celdas):
    """El mismo cuadro con la fila «Total» debajo de las jurisdicciones (fila 13)."""
    return {(f"{k[0]}13" if re.fullmatch(r"[A-G]6", k) else k): v for k, v in celdas.items()}


def _preparar(contenido, nombre=NOMBRE):
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    saneados, reporte, _ = excel_cobertura.preparar(datos, cobertura, lambda f: (f, {}), sanear)
    return saneados, reporte


def _controles(contenido, nombre=NOMBRE):
    saneados, reporte = _preparar(contenido, nombre)
    return controles_excel.analizar([(nombre, contenido)], saneados, reporte["excel"])


def _reglas(hs, regla):
    return [h for h in hs if h["evidencia"]["regla"] == regla]


# Encabezado -----------------------------------------------------------------------------------------------

def test_combinaciones_verticales_suman_la_fila_de_los_tipos_al_encabezado():
    saneados, _ = _preparar(_xlsx(_celdas()))
    columnas = set().union(*(set(r) for r in saneados))
    assert {f"Tipo de vivienda colectiva {t}" for t in TIPOS} <= columnas
    assert {"Código", "Jurisdicción", "Total de viviendas colectivas"} <= columnas
    assert not any(c.startswith("columna_") for c in columnas)
    # La fila de los tipos es encabezado, no un dato de texto en las columnas de cifras.
    assert not any(r.get("Tipo de vivienda colectiva Hotel") == "Hotel" for r in saneados)


def test_combinacion_vertical_no_suma_una_fila_con_cifras():
    from services.excel_encabezados import combinar

    filas = [["Código", "Jurisdicción", "Total", "Tipo", None], [None, None, None, "Hotel", "Hospital"],
             ["10", "Chaco", 30, 10, 20]]
    # «Código» combinado con la fila de abajo, que tiene cifras: esa fila no es encabezado.
    assert combinar(filas[::2], [(1, 1, 2, 1)], [0]) is None
    niveles, nombres = combinar(filas, [(1, 1, 2, 1), (1, 4, 1, 5)], [0])
    assert niveles == [0, 1] and nombres[3:] == ["Tipo Hotel", "Tipo Hospital"]


def test_combinacion_vertical_no_suma_una_fila_de_unidades():
    from services.excel_encabezados import combinar

    # La fila de unidades la lee la cobertura como metadato; el rótulo de la columna sigue siendo «Total».
    filas = [["Jurisdicción", "Total", "Hotel"], [None, "pesos", "pesos"], ["Chaco", 30, 10]]
    assert combinar(filas, [(1, 1, 2, 1)], [0]) is None


# Motor: total arriba ----------------------------------------------------------------------------------------

@pytest.mark.parametrize("combinadas", [True, False])
def test_total_arriba_ubica_el_componente_en_la_interseccion(combinadas):
    hs = _controles(_xlsx(_celdas(error=[(2, 2, 63)]), combinadas))
    tri = _reglas(hs, "componente_no_cierra_fila_y_columna")
    assert len(tri) == 1 and tri[0]["evidencia"]["celda"] == "F9"
    assert tri[0]["evidencia"]["total_columna"] == "Cuadro 4!F6"
    assert tri[0]["evidencia"]["valor_que_cierra"] == _valores()[0][2][2]
    assert not _reglas(hs, "fila_no_cierra_total")


def test_total_arriba_correcto_no_da_alarmas():
    hs = _controles(_xlsx(_celdas()))
    assert not [h for h in hs if h["severidad"] == "alta" and h["evidencia"].get("naturaleza") == "control"]


def test_total_arriba_que_no_suma_sus_filas_sigue_senalando_la_fila():
    hs = _controles(_xlsx(_celdas(error=[(2, 2, 63)], total_arriba=10 ** 6)))
    assert not _reglas(hs, "componente_no_cierra_fila_y_columna")
    fila = _reglas(hs, "fila_no_cierra_total")
    assert fila and "no alcanza para ubicar una sola celda" in fila[0]["descripcion"]


def test_total_arriba_que_cierra_nombra_el_total_de_la_fila():
    celdas = _celdas()
    celdas["C9"] += 63  # el total de Córdoba, no un componente
    hs = _controles(_xlsx(celdas))
    assert not _reglas(hs, "componente_no_cierra_fila_y_columna")
    fila = [h for h in _reglas(hs, "fila_no_cierra_total") if h["evidencia"]["celda"] == "C9"]
    assert fila and fila[0]["evidencia"]["ubicaciones_excel"][0] == "Cuadro 4!C9"
    assert "el total de esta fila" in fila[0]["descripcion"]


@pytest.mark.parametrize("abajo", [False, True])
@pytest.mark.parametrize("marca", ["s", None])
def test_interseccion_sin_numero_no_rompe_y_no_se_nombra(marca, abajo):
    # Un dato no publicado («s», vacío) explica la diferencia de los dos totales: no es un error de la celda.
    celdas = _celdas()
    celdas["F9"] = marca
    hs = _controles(_xlsx(_abajo(celdas) if abajo else celdas))
    assert not _reglas(hs, "componente_no_cierra_fila_y_columna")
    fila = _reglas(hs, "fila_no_cierra_total")
    assert fila and "no tiene un número publicado" in fila[0]["descripcion"]
    assert fila[0]["severidad"] == "media"


# Lector: cruce de dos cuentas -------------------------------------------------------------------------------

FILA = {"tipo": "columnas", "destino": "C", "expresion": " + ".join(COLS), "filas": [6, 12],
        "fundamento": "el total es la suma de los tipos"}
COLUMNA = {"tipo": "filas", "destino": 6, "expresion": " + ".join(str(f) for f in range(7, 13)),
           "columnas": ["C", "G"], "fundamento": "el total del país es la suma de las jurisdicciones"}


def _grupos(celdas):
    mapas = {lectora: {"relaciones": [FILA, COLUMNA]} for lectora in ("gpt", "gemini")}
    return lm.acordar(mapas, celdas, archivo=NOMBRE, hoja="Cuadro 4")


def _senales(grupos, confirmaciones=None):
    return lm.senales([{"archivo": NOMBRE, "hoja": "Cuadro 4", "grupos": grupos}], confirmaciones or {})


def test_lector_nombra_la_celda_comun_a_las_dos_cuentas():
    grupos = _grupos(_celdas(error=[(2, 2, 63)]))
    assert [g["estado"] for g in grupos] == ["confirmada", "confirmada"]
    ss = _senales(grupos)
    assert {s["evidencia"]["celda"] for s in ss} == {"C9", "F6"}
    for s in ss:
        e = s["evidencia"]
        assert e["celda_probable"] == "F9" and e["valor_que_cierra"] == _valores()[0][2][2]
        assert e["ubicaciones_excel"][0] == "Cuadro 4!F9" and e["grupo_relacion"] == "celda:Cuadro 4!F9"
        assert "Cuadro 4!F9 interviene en esta cuenta" in s["descripcion"]


def test_lector_no_nombra_celda_si_hay_dos_intersecciones_posibles():
    # F9 y G11 alterados por lo mismo: F9+G11 o G9+F11 explican igual las cuatro cuentas.
    grupos = _grupos(_celdas(error=[(2, 2, 5), (4, 3, 5)]))
    assert all("cruce" not in f for g in grupos for f in g["fallas"])
    assert not any("celda_probable" in s["evidencia"] for s in _senales(grupos))


def test_motor_y_lector_sobre_la_misma_celda_son_un_solo_caso():
    celdas = _celdas(error=[(2, 2, 63)])
    hs = _controles(_xlsx(celdas)) + _senales(_grupos(celdas))
    casos = casos_revision.agrupar(hs)["casos"]
    assert len(casos) == 1 and casos[0]["cantidad_senales"] == 3
    assert casos[0]["primary_finding"]["evidencia"]["celda"] == "F9"
    # El mismo cuadro en otro libro es otro caso.
    otro = _senales(_grupos(celdas))
    for s in otro:
        s["evidencia"]["archivo_origen"] = "otro.xlsx"
    assert len(casos_revision.agrupar(hs + otro)["casos"]) == 2


def test_lector_no_cruza_con_una_cuenta_que_la_persona_rechazo():
    grupos = _grupos(_celdas(error=[(2, 2, 63)]))
    columna = next(g for g in grupos if g["tipo"] == "filas")
    ss = _senales(grupos, {columna["id"]: {"confirma": False}})
    assert ss and not any("celda_probable" in s["evidencia"] for s in ss)
