from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook

import models
from scripts.ensayo_estructuras import VARIANTES, evaluar, libro
from services import carga_archivos, controles_excel, excel_cobertura, jobs, sanitizacion
from services.controles_excel import lectura
from services.reportes.markdown import _periodo_observado, render_comprobante_controles_markdown
from tests.flujo_confirmado import cargar_confirmado

pytest_plugins = ["tests.test_ejemplos_zip"]


@pytest.mark.parametrize("semilla", [17, 853])
@pytest.mark.parametrize("variante", VARIANTES)
@pytest.mark.parametrize("alterado", [False, True])
def test_misma_aritmetica_y_control_en_distintas_presentaciones(semilla, variante, alterado):
    data, verdad = libro(semilla, variante, alterado)
    resultado = evaluar(data, verdad)
    assert not resultado["errores"], resultado


def procesar(w):
    b = BytesIO()
    w.save(b)
    leidos = [("prueba.xlsx", b.getvalue())]
    d, c = carga_archivos.parse_archivos_con_cobertura(leidos)
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    ss, _ = controles_excel._cuadros(leidos, d, r["excel"])
    return leidos, d, r, ss


def dos_cuadros():
    w = Workbook()
    s = w.active
    for row, label, years in [(1, "Concepto", [2017, 2018, 2019, 2020]), (7, "Descripción", [2021, 2022])]:
        for j, v in enumerate([label] + years, 1):
            s.cell(row, j, v)
        for i, (label, v) in enumerate([("Salud", 10), ("Educación", 20), ("Total", 30)], 1):
            for j, x in enumerate([label] + [v * (1 + k * 0.1) for k in range(len(years))], 1):
                s.cell(row + i, j, x)
    return w


def test_otra_serie_no_compensa_la_omitida_aunque_los_conteos_coincidan():
    leidos, d, r, ss = procesar(dos_cuadros())
    cov = lectura.resumen(leidos, ss + deepcopy(ss))
    assert len(ss + ss) == 6
    assert cov["series_leidas"] == 3 and cov["series_aparentes"] == 6
    assert cov["estado"] == "parcial" and cov["proporcion"] == 0.5
    assert any("2 períodos" in k for k in cov["archivos"][0]["motivos"])


def test_serie_sin_rotulo_queda_contada_como_no_incorporada():
    w = dos_cuadros()
    w.active["A3"] = None
    leidos, d, r, ss = procesar(w)
    cov = lectura.resumen(leidos, ss)
    assert cov["series_aparentes"] == 6 and cov["series_leidas"] == 2
    assert any("rótulo" in k for k in cov["archivos"][0]["motivos"])


@pytest.mark.parametrize("n", [3, 6])
@pytest.mark.parametrize("blanco", [False, True])
def test_vertical_cobertura_y_periodos_con_fila_vacia(n, blanco):
    w = Workbook()
    s = w.active
    s.append(["Año", "Importe"])
    for j in range(n):
        s.append([2017 + j, 100 + j * 10])
    if blanco:
        s.insert_rows(3)
    leidos, d, r, ss = procesar(w)
    cov = lectura.resumen(leidos, ss)
    assert cov["series_aparentes"] == cov["series_leidas"] == 1 and cov["estado"] == "completa"
    assert _periodo_observado(d) == f"2017 a {2017 + n - 1}"


def test_importes_parecidos_a_anios_no_se_convierten_en_encabezados():
    w = Workbook()
    s = w.active
    s.append(["Concepto", 2017, 2018, 2019])
    s.append(["Salud", 1900, 2000, 2100])
    s.append(["Educación", 3800, 4000, 4200])
    s.append(["Total", 5700, 6000, 6300])
    leidos, d, r, ss = procesar(w)
    cov = lectura.resumen(leidos, ss)
    assert len(d) == 3 and len(ss) == 3 and cov["series_aparentes"] == 3
    assert not controles_excel.analizar(leidos, d, r["excel"])


def test_cobertura_parcial_se_propaga_al_resultado_y_comprobante(flujo):
    client, factory, _ = flujo
    b = BytesIO()
    dos_cuadros().save(b)
    r = cargar_confirmado(client, "/api/analisis-rapido",
        data={"periodo": "2026-09", "dominio": "general", "notificar": "false"},
        files={
            "archivos": ("dos.xlsx", b.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        },
    )
    assert r.status_code == 200, r.text
    assert jobs._process_one()
    with factory() as db:
        i = db.get(models.Informe, r.json()["informe_id"])
        assert i.devolucion.cobertura["estado"] == "COBERTURA_PARCIAL"
        cov = i.devolucion.cobertura["excel"]["lectura"]
        assert cov["series_leidas"] == 3 and cov["series_aparentes"] == 6
        md = render_comprobante_controles_markdown(i)
        assert "LECTURA PARCIAL DEL ARCHIVO (50 %)" in md


def test_comprobante_no_redondea_lectura_parcial_a_cien():
    i = models.Informe(
        id=1,
        area_id=1,
        modo="revision_final",
        periodo="2026-09",
        dominio="general",
        datos=[],
        hallazgos=[],
        saneamiento={},
        devolucion=models.Devolucion(
            cobertura={
                "excel": {
                    "lectura": {
                        "estado": "parcial",
                        "series_leidas": 689,
                        "series_aparentes": 690,
                        "proporcion": 0.999,
                        "archivos": [],
                    }
                }
            }
        ),
    )
    md = render_comprobante_controles_markdown(i)
    assert "99,9 %" in md and "100 %" not in md
