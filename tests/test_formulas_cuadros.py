import io

from openpyxl import Workbook

from services import formula_integridad


def analizar(w, inferir=False):
    b = io.BytesIO()
    w.save(b)
    return formula_integridad.analizar_archivo("prueba.xlsx", b.getvalue(), inferir_patrones=inferir)


def test_encabezados_anuales_y_patrones_distintos_no_se_mezclan():
    w = Workbook()
    s = w.active
    s.title = "Clases"
    for row in [
        ["Clase", 2018, 2019, 2020],
        ["A", 1, 2, "=B2+C2"],
        ["B", 3, 4, "=B3+C3"],
        ["C", 5, 6, "=B4+C4"],
        [],
        ["Otro nivel"],
        ["Clase", 2018, 2019, 2020],
        ["A", 10, 11, 12],
        ["B", 20, 21, "=B9-C9"],
        ["C", 30, 31, "=B10-C10"],
        ["D", 40, 41, "=B11-C11"],
    ]:
        s.append(row)
    assert analizar(w) == []


def test_hardcode_real_dentro_del_mismo_cuadro_sigue_detectado():
    w = Workbook()
    s = w.active
    s.append(["Concepto", "Base", "Resultado"])
    for r in range(2, 8):
        s.append([f"Caso {r}", r, 999 if r == 4 else f"=B{r}*2"])
    hs = analizar(w)
    assert any(
        h["evidencia"]["celda"] == "C4" and h["evidencia"]["regla"] == "valor_fijo_en_serie_de_formulas" for h in hs
    )


def test_formula_corrida_en_segundo_cuadro_se_detecta():
    w = Workbook()
    s = w.active
    s.append(["Concepto", "Base", "Resultado"])
    for r in range(2, 7):
        s.append([f"Caso {r}", r, f"=B{r}*2"])
    s.append(["Concepto", "Base", "Resultado"])
    for r in range(8, 14):
        s.append([f"Caso {r}", r, f"=B{r - 1 if r == 11 else r}*3"])
    hs = analizar(w)
    assert any(h["evidencia"]["celda"] == "C11" and h["evidencia"]["regla"] == "formula_fuera_de_patron" for h in hs)


def test_total_con_tercer_componente_no_se_reduce_a_un_par():
    w = Workbook()
    s = w.active
    s.append(["Clase", 2016, 2017, 2018, 2019, 2020, 2021, 2022])
    s.append(["A", 100, 110, 120, 130, 140, 150, 160])
    s.append(["B", 40, 41, 42, 43, 44, 45, 46])
    s.append(["C", 58, 0, 0, 0, 0, 0, 0])
    s.append(["Total", 198, 151, 162, 173, 184, 195, 206])
    assert analizar(w, True) == []


def test_inferencia_de_par_con_dos_componentes_sigue_detectando():
    w = Workbook()
    s = w.active
    s.append(["Concepto", 2016, 2017, 2018, 2019, 2020, 2021, 2022])
    s.append(["A", 100, 110, 120, 130, 140, 150, 160])
    s.append(["B", 40, 41, 42, 43, 44, 45, 46])
    s.append(["Total", 999, 151, 162, 173, 184, 195, 206])
    hs = analizar(w, True)
    assert any(h["evidencia"]["regla"] == "formula_implicita_rota" for h in hs)
