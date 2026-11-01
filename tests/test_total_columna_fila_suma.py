"""0.10.56: un total de columna que difiere más del 15 % se señala si la fila es claramente una suma.

Caso real que lo motivó: un cuadro mensual de recaudación con «TOTAL» anual en el que un dígito corrido en
diciembre (3.310,8 → 331,08) dejaba el total un 16 % arriba de la suma; el control descartaba las
diferencias mayores al 15 % aunque en la misma fila cerraran las otras doce columnas. Sólo entra un total MAYOR
que la suma de partes presentes y no negativas: promedios, saldos, máximos y personas únicas declaran menos, y una
parte reservada («s») o vacía no se toma como error. Libros sintéticos; los últimos casos, de la revisión
independiente de 0.10.56.
"""
import io

from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura, sanitizacion

COLUMNAS = ["INGRESOS BRUTOS", "INMOBILIARIO", "AUTOMOTORES", "SELLOS", "OTROS"]
MESES = range(1, 13)


def _valor(j, mes):
    # Diciembre triplica (aguinaldo, vencimientos anuales): un dígito corrido ahí mueve el total más del 15 %.
    return round((1000 + 100 * j + 7.5 * mes) * (3 if mes == 12 else 1), 2)


def _libro(columnas=COLUMNAS, cambiar=None, total=None):
    """Recaudación mensual con «TOTAL» anual; `cambiar` = (columna, mes, factor) altera un mes después de sumar;
    `total` = {columna: función de las partes} reemplaza la suma de esa columna."""
    wb = Workbook()
    ws = wb.active
    ws.title = "2025"
    ws["B3"] = "RECURSOS TRIBUTARIOS"
    ws["B4"] = "2025 EN MILLONES DE PESOS"
    ws.cell(5, 2, "AÑO")
    ws.cell(5, 3, "MES")
    for j, nombre in enumerate(columnas):
        ws.cell(5, 4 + j, nombre)
    for i, mes in enumerate(MESES):
        if i == 0:
            ws.cell(6, 2, 2025)
        ws.cell(6 + i, 3, mes)
        for j in range(len(columnas)):
            ws.cell(6 + i, 4 + j, _valor(j, mes))
    ws.cell(18, 2, "TOTAL")
    for j, nombre in enumerate(columnas):
        partes = [_valor(j, mes) for mes in MESES]
        ws.cell(18, 4 + j, round((total or {}).get(nombre, sum)(partes), 2))
    if cambiar:
        nombre, mes, factor = cambiar
        celda = ws.cell(5 + mes, 4 + columnas.index(nombre))
        celda.value = round(celda.value * factor, 4)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _totales(contenido, nombre="recaudacion.xlsx"):
    d, c = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    d, r, _ = excel_cobertura.preparar(d, c, lambda x: (x, {}), sanitizacion.sanear)
    hs = controles_excel.analizar([(nombre, contenido)], d, r["excel"])
    return [h for h in hs if h["evidencia"].get("regla") == "total_columna_no_cierra"]


def test_el_cuadro_original_no_da_senales():
    assert _totales(_libro()) == []


def test_digito_corrido_en_un_mes_se_senala_aunque_pase_el_15():
    (h,) = _totales(_libro(cambiar=("SELLOS", 12, 0.1)))
    e = h["evidencia"]
    assert (e["hoja"], e["celda"]) == ("2025", "G18")
    assert 15 < e["diferencia_porcentual"] < 50 and e["fila_suma_por_otras_columnas"] is True
    assert "la fila es una suma: cierran 4 columnas" in h["descripcion"]


def test_dentro_del_15_sigue_como_antes():
    (h,) = _totales(_libro(cambiar=("SELLOS", 3, 0.1)))
    assert h["evidencia"]["celda"] == "G18" and h["evidencia"]["fila_suma_por_otras_columnas"] is False


def test_con_pocas_columnas_que_cierran_no_alcanza():
    # Dos columnas que cierran no prueban que la fila sea una suma: la diferencia grande no se señala.
    assert _totales(_libro(columnas=["INGRESOS BRUTOS", "INMOBILIARIO", "SELLOS"], cambiar=("SELLOS", 12, 0.1))) == []


def test_mas_de_la_mitad_de_diferencia_no_se_toma_como_suma():
    # «Saldo» al cierre del año: el total es el último mes, no la suma; difiere en más de la mitad.
    assert _totales(_libro(columnas=COLUMNAS + ["SALDO"], total={"SALDO": lambda p: p[-1]})) == []


def test_total_dentro_del_rango_de_sus_partes_no_es_una_suma():
    # Cuadro por región: cuatro columnas cierran y «Referencia» declara el máximo (300 de 100, 120 y 300), un 42 %
    # abajo de la suma. No puede ser una suma de partes no negativas: no se señala.
    wb = Workbook()
    ws = wb.active
    ws.title = "Regiones"
    ws.append(["Región", "Ingresos", "Gastos", "Personal", "Bienes", "Referencia"])
    filas = [["Norte", 410.5, 300.25, 120.5, 80.75, 100], ["Centro", 520.25, 410.5, 150.25, 90.5, 120],
             ["Sur", 610.75, 505.25, 170.75, 95.25, 300]]
    for f in filas:
        ws.append(f)
    ws.append(["Total"] + [round(sum(f[k] for f in filas), 2) for k in range(1, 5)] + [300])
    b = io.BytesIO()
    wb.save(b)
    assert _totales(b.getvalue(), "regiones.xlsx") == []


def _cuadro(filas, nombre="cuadro.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro"
    for f in filas:
        ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return [h["evidencia"]["celda"] for h in _totales(b.getvalue(), nombre)]


PROGRAMAS = [["Programa", "Prestaciones", "Monto", "Expedientes", "Beneficiarios"],
             ["Programa A", 1500, 120000.5, 300, 1200], ["Programa B", 1100, 90000.25, 210, 900],
             ["Programa C", 800, 64000.75, 160, 700], ["Programa D", 600, 48000.5, 120, 500]]


def test_personas_unicas_no_son_una_suma():
    # Quien cobra en dos programas se cuenta una vez: 2.400 contra 3.300 de suma (27 % menos).
    assert _cuadro(PROGRAMAS + [["Total", 4000, 322002, 790, 2400]]) == []


def test_partes_reservadas_o_vacias_no_se_toman_como_error():
    # El total publicado incluye las dos provincias reservadas: las partes visibles suman un 20 % menos.
    filas = [["Provincia", "Empresas", "Puestos", "Remuneraciones", "Exportaciones"],
             ["Buenos Aires", 1200, 35000, 98000.5, 5200.5], ["Córdoba", 800, 21000, 61000.25, 3100.25],
             ["Santa Fe", 700, 19000, 55000.75, 2900.75], ["Mendoza", 400, 9000, 26000.5, "s"],
             ["Tucumán", 300, 7000, 19000.25, "s"]]
    total = [["Total", 3400, 91000, 259002.25, 14001.5]]
    assert _cuadro(filas + total) == []
    assert _cuadro([f[:4] + [None if f[4] == "s" else f[4]] for f in filas] + total) == []


def test_columnas_de_ceros_no_prueban_que_la_fila_sea_una_suma():
    # Sólo «Prestaciones» cierra con valores; «Bajas» y «Suspensiones» son ceros. «Altas» tiene un mes ÷ 10.
    filas = [["Programa", "Prestaciones", "Bajas", "Suspensiones", "Altas"],
             ["Programa A", 1500, 0, 0, 40], ["Programa B", 1100, 0, 0, 300],
             ["Programa C", 800, 0, 0, 200], ["Programa D", 600, 0, 0, 100]]
    assert _cuadro(filas + [["Total", 4000, 0, 0, 1000]]) == []


def test_la_mitad_cuenta_tambien_las_columnas_que_difieren_mas_del_50():
    # Cierran 3 de 7 comparadas: las localidades, municipios y organismos únicos difieren más del 50 %.
    filas = [["Programa", "Prestaciones", "Monto", "Expedientes", "Localidades", "Municipios", "Organismos", "Altas"],
             ["Programa A", 1500, 120000.5, 300, 40, 30, 12, 40], ["Programa B", 1100, 90000.25, 210, 38, 28, 11, 300],
             ["Programa C", 800, 64000.75, 160, 35, 25, 10, 200], ["Programa D", 600, 48000.5, 120, 30, 20, 9, 100]]
    assert _cuadro(filas + [["Total", 4000, 322002, 790, 50, 35, 14, 1000]]) == []


def test_con_tres_columnas_que_cierran_el_mismo_cuadro_si_se_senala():
    # Control de los dos anteriores: con «Prestaciones», «Monto» y «Expedientes» cerrando y sin las columnas de conteos
    # únicos, el «Altas» con un programa ÷ 10 (36 % de diferencia, total mayor que la suma) se señala.
    filas = [["Programa", "Prestaciones", "Monto", "Expedientes", "Altas"],
             ["Programa A", 1500, 120000.5, 300, 40], ["Programa B", 1100, 90000.25, 210, 300],
             ["Programa C", 800, 64000.75, 160, 200], ["Programa D", 600, 48000.5, 120, 100]]
    assert _cuadro(filas + [["Total", 4000, 322002, 790, 1000]]) == ["E6"]
