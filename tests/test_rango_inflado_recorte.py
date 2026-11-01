"""0.10.56: un libro que sólo se rechazaría por el rango usado se recorta al bloque con datos de cada hoja.

Caso real que lo motivó: un libro oficial de 64 hojas y 346.731 celdas con datos abarcaba 25,8 millones de
celdas por una hoja con formato hasta la fila 1.048.446 y una celda suelta en la columna XFD de otra, y se
rechazaba entero. El recorte se declara en la lectura; primero se quita sólo el formato y las celdas sueltas con
datos quedan fuera únicamente si con eso no alcanza. Nada de lo que la lectura completa vuelve a crear (un
comentario, un vínculo) puede quedar fuera del bloque. Los libros de estas pruebas son sintéticos; los casos de
comentarios, vínculos, celdas sin referencia y topes salen de la revisión independiente de 0.10.56.
"""
import io
import json
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

import models
from services import carga_archivos, jobs
from tests.test_ejemplos_zip import flujo  # noqa: F401 — fixture compartida
from tests.test_revision_rapida import XLSX, sin_ia  # noqa: F401

CUADRO = [["Concepto", "2023", "2024"], ["Exportaciones", 100, 120], ["Importaciones", 80, 90], ["Total", 180, 210]]


def _libro(formato_hasta=None, datos_lejos=(), notas_hasta=None, titulo=None, comentario=None, vinculo=None):
    """Hoja «Cuadro» con CUADRO; `datos_lejos` = [(celda, valor)]; `notas_hasta`: otra hoja con formato hasta esa celda;
    `titulo`: un título en A1 combinado hasta esa columna, con el cuadro debajo; `comentario` y `vinculo`: celdas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro"
    if titulo:
        ws["A1"] = "Cuadro 1. Comercio exterior"
        ws.merge_cells(f"A1:{titulo}1")
    for f in CUADRO:
        ws.append(f)
    if formato_hasta:
        ws[formato_hasta].font = Font(bold=True)  # celda vacía con formato: amplía el rango usado
    for celda, valor in datos_lejos:
        ws[celda] = valor
    if comentario:
        ws[comentario].comment = Comment("revisar", "autor")
    if vinculo:
        ws[vinculo].hyperlink = "https://example.org/fuente"
    if notas_hasta:
        notas = wb.create_sheet("Notas")
        notas["A1"] = "Fuente: elaboración propia"
        notas[notas_hasta].font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cambiar_hoja(contenido, cambio, hoja="xl/worksheets/sheet1.xml"):
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(contenido)) as z, zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            datos = z.read(item)
            if item.filename == hoja:
                nuevo = cambio(datos)
                assert nuevo != datos
                datos = nuevo
            zo.writestr(item, datos)
    return salida.getvalue()


def _abrir(contenido, hoja="Cuadro"):
    return load_workbook(io.BytesIO(contenido))[hoja]


def test_un_libro_que_entra_no_se_toca():
    libro = _libro(formato_hasta="Z40")
    nuevo, avisos = carga_archivos.recortar_rango_inflado("chico.xlsx", libro)
    assert nuevo is libro and avisos == []


def test_formato_hasta_el_final_de_la_hoja_se_recorta_y_se_avisa():
    libro = _libro(formato_hasta="C1048000")  # 3 columnas × 1.048.000 filas: más que el tope de rango
    with pytest.raises(carga_archivos.ArchivoNoSoportado):
        carga_archivos.validar_tamano_libro("largo.xlsx", libro, "xlsx")
    nuevo, (aviso,) = carga_archivos.recortar_rango_inflado("largo.xlsx", libro)
    carga_archivos.validar_tamano_libro("largo.xlsx", nuevo, "xlsx")
    assert aviso == ("«largo.xlsx», hoja «Cuadro»: el rango usado llegaba hasta la fila 1.048.000 y la columna C; se "
                     "leyó el bloque con datos, hasta la celda C4. Fuera del bloque había celdas sin datos (sólo formato).")
    ws = _abrir(nuevo)
    assert (ws.max_row, ws.max_column) == (4, 3)
    assert [[c.value for c in fila] for fila in ws.iter_rows()] == CUADRO


def test_celda_suelta_queda_fuera_si_sin_eso_no_entra_y_se_nombra():
    # 200 filas × XFD superan los 3 millones de celdas sólo por un dato suelto: hay que dejarlo fuera.
    libro = _libro(datos_lejos=[("XFD200", 1581511)])
    nuevo, avisos = carga_archivos.recortar_rango_inflado("pesca.xlsx", libro)
    assert avisos == [
        "«pesca.xlsx», hoja «Cuadro»: el rango usado llegaba hasta la fila 200 y la columna XFD; se leyó el bloque con "
        "datos, hasta la celda C4. Quedó fuera de la lectura 1 celda con datos separada del cuadro por columnas o filas "
        "vacías: XFD200."]
    carga_archivos.validar_tamano_libro("pesca.xlsx", nuevo, "xlsx")
    ws = _abrir(nuevo)
    assert ws.max_column == 3 and ws["C4"].value == 210


def test_si_alcanza_con_quitar_el_formato_la_celda_suelta_se_conserva():
    # El exceso viene de «Notas»: recortar su formato alcanza y el dato lejano de «Cuadro» se lee.
    libro = _libro(datos_lejos=[("XFD40", 1581511)], notas_hasta="C1048000")
    nuevo, avisos = carga_archivos.recortar_rango_inflado("pesca.xlsx", libro)
    assert [a.split("»", 2)[1] for a in avisos] == [", hoja «Notas"]
    assert _abrir(nuevo)["XFD40"].value == 1581511


def test_las_celdas_sueltas_salen_solo_de_las_hojas_que_hace_falta():
    # «A» sólo entra sin su dato en XFD200; «B» tiene 10 datos lejos (A5000:A5009) pero recortando «A» el libro
    # ya entra: los de «B» se leen.
    wb = Workbook()
    a = wb.active
    a.title = "A"
    for f in CUADRO:
        a.append(f)
    a["XFD200"] = 99
    b = wb.create_sheet("B")
    for r in range(1, 501):
        b.append([r * 1.5 + c for c in range(26)])
    for r in range(5000, 5010):
        b.cell(r, 1, r)
    buf = io.BytesIO()
    wb.save(buf)
    nuevo, avisos = carga_archivos.recortar_rango_inflado("dos.xlsx", buf.getvalue())
    assert [x.split("»", 2)[1] for x in avisos] == [", hoja «A"]
    assert _abrir(nuevo, "B")["A5009"].value == 5009


def test_muchos_datos_lejos_del_cuadro_son_otro_cuadro_y_se_conservan():
    lejos = [(f"XFD{r}", r) for r in range(1, 31)]  # 30 datos en la misma columna: no es un resto suelto
    nuevo, avisos = carga_archivos.recortar_rango_inflado("dos.xlsx", _libro(datos_lejos=lejos, notas_hasta="C1048000"))
    ws = _abrir(nuevo)
    assert ws.max_column == 16384 and ws["XFD30"].value == 30
    assert [a.split("»", 2)[1] for a in avisos] == [", hoja «Notas"]


def test_combinacion_hasta_xfd_se_recorta_al_ancho_del_cuadro():
    libro = _libro(titulo="XFD", formato_hasta="A12", notas_hasta="C1048000")
    nuevo, _ = carga_archivos.recortar_rango_inflado("titulo.xlsx", libro)
    ws = _abrir(nuevo)
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:C1"]
    assert (ws.max_row, ws.max_column) == (5, 3) and ws["A1"].value == "Cuadro 1. Comercio exterior" and ws["C5"].value == 210


def test_un_comentario_lejos_del_cuadro_impide_el_recorte():
    # La lectura completa vuelve a crear la celda del comentario: recortar el resto no achica el rango real.
    libro = _libro(formato_hasta="C1048000", comentario="XFD5000")
    nuevo, avisos = carga_archivos.recortar_rango_inflado("nota.xlsx", libro)
    assert nuevo is libro and avisos == []
    with pytest.raises(carga_archivos.ArchivoNoSoportado):
        carga_archivos.validar_tamano_libro("nota.xlsx", nuevo, "xlsx")


def test_un_comentario_dentro_del_cuadro_se_conserva():
    nuevo, (aviso,) = carga_archivos.recortar_rango_inflado("nota.xlsx", _libro(formato_hasta="C1048000", comentario="B2"))
    ws = _abrir(nuevo)
    assert ws["B2"].comment.text == "revisar" and (ws.max_row, ws.max_column) == (4, 3)


def test_vinculos_que_salen_del_bloque_se_quitan_en_sus_dos_formas():
    # Un vínculo de rango que empieza en el cuadro y otro con etiqueta de cierre en XFD9000: la lectura crearía sus celdas.
    def vinculos(xml):
        xml = re.sub(rb'(<hyperlink\b[^>]*?\bref=")A2(")', rb"\1A2:A30000\2", xml)
        return xml.replace(b"</hyperlinks>", b'<hyperlink ref="XFD9000" display="lejos"></hyperlink></hyperlinks>')

    libro = _cambiar_hoja(_libro(formato_hasta="C1048000", vinculo="A2"), vinculos)
    nuevo, _ = carga_archivos.recortar_rango_inflado("vinculos.xlsx", libro)
    with zipfile.ZipFile(io.BytesIO(nuevo)) as z:
        assert b"<hyperlink " not in z.read("xl/worksheets/sheet1.xml")
    ws = _abrir(nuevo)
    assert (ws.max_row, ws.max_column) == (4, 3)


def test_celdas_sin_referencia_no_se_recortan():
    # Celdas «<c>» sin «r» (válido; la lectura las ubica una tras otra): no se puede recortar por posición.
    wb = Workbook()
    ws = wb.active
    ws.append(["Provincia", "2023", "2024"])
    for i in range(1, 31):
        ws.append([f"P{i}", 100 + i, 200 + i])
    b = io.BytesIO()
    wb.save(b)

    def sin_r(xml):
        xml = re.sub(rb'(<c) r="[A-Z]+[0-9]+"', rb"\1", xml)
        return re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:Z200000"', xml)

    libro = _cambiar_hoja(b.getvalue(), sin_r)
    nuevo, avisos = carga_archivos.recortar_rango_inflado("sin_r.xlsx", libro)
    assert nuevo is libro and avisos == []


def test_un_libro_con_demasiadas_celdas_escritas_no_se_recorre(monkeypatch):
    # El recorte recorre cada «<c>»: por encima del tope, el libro se rechaza como antes, sin recortarlo.
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_RECORTE", 10)
    libro = _libro(formato_hasta="C1048000")
    assert carga_archivos.recortar_rango_inflado("largo.xlsx", libro) == (libro, [])


def test_un_libro_que_se_rechazaria_por_sus_datos_no_se_recorta(monkeypatch):
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_CON_DATOS", 5)
    libro = _libro(formato_hasta="C1048000")
    assert carga_archivos.recortar_rango_inflado("largo.xlsx", libro) == (libro, [])


def test_la_lectura_declara_el_recorte_y_lee_el_cuadro():
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([("largo.xlsx", _libro(formato_hasta="C1048000"))])
    assert len(datos) == 3 and lectura[0]["estado"] == "leido"
    assert any("se leyó el bloque con datos, hasta la celda C4" in a for a in lectura[0]["avisos"])


def test_un_libro_que_sigue_siendo_grande_despues_del_recorte_se_rechaza(monkeypatch):
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_RANGO", 10)
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match="abarca"):
        carga_archivos.parse_archivos_con_cobertura([("largo.xlsx", _libro(formato_hasta="C1048000"))])


def test_la_carga_recorta_una_vez_a_la_entrada_y_el_informe_lo_dice(sin_ia):  # noqa: F811
    # Por la ruta real: la lectura, los controles y el informe trabajan con el libro recortado.
    client, factory, _ = sin_ia
    r = client.post("/api/analisis-rapido", data={"periodo": "2026-09", "notificar": "false", "modo_revision": "rapida"},
                    files={"archivos": ("largo.xlsx", _libro(formato_hasta="C1048000"), XLSX)})
    assert r.status_code == 200, r.text
    while jobs._process_one():
        pass
    with factory() as db:
        assert [j.estado for j in db.query(models.AnalysisJob).all()] == ["finalizado"]
        saneamiento = json.dumps(db.get(models.Informe, r.json()["informe_id"]).saneamiento, ensure_ascii=False)
    assert saneamiento.count("se leyó el bloque con datos, hasta la celda C4") >= 1
