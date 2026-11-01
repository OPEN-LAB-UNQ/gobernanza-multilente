"""0.9.21: un libro cuya lectura agotaría la memoria del servicio se rechaza al cargarlo, con el motivo.

Los libros de un ministerio con hojas formateadas hasta la columna XFD (16,6 millones de celdas de rango) o con
un millón de celdas con datos superaban 3 GB al leerse; el servicio tiene 2 GB y se reiniciaba sin
responder. El tamaño se mide en el XML de cada hoja, sin abrir el libro.

0.10.56: un libro que sólo se rechazaría por celdas sin datos o por la dimensión declarada se recorta antes al
bloque con datos (tests/test_rango_inflado_recorte.py); los datos de verdad lejos del cuadro siguen contando.
"""
import importlib.util
import io
import time
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from services import carga_archivos

_spec = importlib.util.spec_from_file_location("decimocuarta", Path(__file__).with_name("test_decimocuarta_medicion.py"))
_decimocuarta = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_decimocuarta)


def _libro(filas, hojas=1, formato_hasta=None, combinar=None):
    wb = Workbook()
    for n in range(hojas):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = f"Cuadro {n + 1}"
        for f in filas:
            ws.append(f)
        if formato_hasta:
            ws[formato_hasta].font = Font(bold=True)  # celda vacía con formato: amplía el rango usado
        if combinar:
            ws.merge_cells(combinar)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CUADRO = [["Concepto", "2023", "2024"], ["Exportaciones", 100, 120], ["Importaciones", 80, 90], ["Total", 180, 210]]


def test_libro_normal_se_acepta_con_los_topes_reales():
    datos = carga_archivos.parse_archivos([("cuadro.xlsx", _libro(CUADRO))])
    assert len(datos) == 3


def test_medicion_liviana_rango_y_celdas_con_datos():
    hojas = carga_archivos.medir_libro(_libro(CUADRO, hojas=2, formato_hasta="Z40"), "xlsx")
    assert [h["hoja"] for h in hojas] == ["Cuadro 1", "Cuadro 2"]
    assert all(h["filas"] == 40 and h["columnas"] == 26 and h["con_datos"] == 12 for h in hojas)


def test_celdas_combinadas_amplian_el_rango_como_en_la_lectura():
    (hoja,) = carga_archivos.medir_libro(_libro(CUADRO, combinar="A10:AD12"), "xlsx")
    assert (hoja["filas"], hoja["columnas"]) == (12, 30)


def test_rango_grande_por_formato_se_rechaza_con_motivo(monkeypatch):
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_RANGO", 1000)
    libro = _libro(CUADRO, formato_hasta="XFD3")
    with pytest.raises(carga_archivos.ArchivoNoSoportado) as exc:
        carga_archivos.parse_archivos_con_cobertura([("infoeco.xlsx", libro)])
    motivo = str(exc.value)
    assert motivo.startswith("infoeco.xlsx: el libro abarca 65.536 celdas")
    assert "el máximo es 1.000" in motivo and "«Cuadro 1» llega hasta la fila 3" not in motivo
    assert "«Cuadro 1» llega hasta la fila 4 y la columna XFD" in motivo and "borrá las filas y columnas sobrantes" in motivo


def test_muchas_celdas_con_datos_se_rechaza_con_motivo(monkeypatch):
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_CON_DATOS", 20)
    libro = _libro(CUADRO, hojas=2)
    with pytest.raises(carga_archivos.ArchivoNoSoportado,
                       match=r"el libro tiene 24 celdas con datos en 2 hojas; el máximo es 20\. Dividilo"):
        carga_archivos.parse_archivos_con_cobertura([("grande.xlsx", libro)])


def test_con_varios_archivos_el_grande_se_declara_y_el_resto_se_analiza(monkeypatch):
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_RANGO", 1000)
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura(
        [("grande.xlsx", _libro(CUADRO, formato_hasta="XFD3")), ("chico.xlsx", _libro(CUADRO))])
    assert len(datos) == 3
    grande = next(x for x in lectura if x["archivo"] == "grande.xlsx")
    assert grande["estado"] == "fallido" and "abarca" in grande["motivo"]


def test_xls_binario_tambien_se_mide(monkeypatch):
    libro = _decimocuarta._xls("Cuadro", CUADRO)
    (hoja,) = carga_archivos.medir_libro(libro, "xls")
    assert (hoja["filas"], hoja["columnas"], hoja["con_datos"]) == (4, 3, 12)
    monkeypatch.setattr(carga_archivos, "MAX_CELDAS_CON_DATOS", 10)
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match="12 celdas con datos en 1 hoja;"):
        carga_archivos.parse_archivos_con_cobertura([("viejo.xls", libro)])


def test_el_rechazo_es_rapido_aunque_el_rango_sea_enorme(monkeypatch):
    # Un dato en la columna XFD de cada fila: no son celdas sueltas (son más que MAX_CELDAS_SUELTAS), así que no
    # se recortan; el rango supera el tope real aunque el libro pese poco.
    filas = carga_archivos.MAX_CELDAS_RANGO // 16384 + 10
    wb = Workbook()
    ws = wb.active
    for f in CUADRO:
        ws.append(f)
    for r in range(5, filas + 1):
        ws.cell(row=r, column=16384, value=r)
    buf = io.BytesIO()
    wb.save(buf)
    t0 = time.time()
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match=f"abarca {filas * 16384:,} celdas".replace(",", ".")):
        carga_archivos.parse_archivos_con_cobertura([("ancho.xlsx", buf.getvalue())])
    assert time.time() - t0 < 5


def test_libro_que_no_se_puede_medir_sigue_a_la_lectura_normal():
    with pytest.raises(carga_archivos.ArchivoNoSoportado, match="no es un Excel"):
        carga_archivos.parse_archivos_con_cobertura([("roto.xlsx", b"PK\x03\x04 esto no es un zip")])


def test_dimension_declarada_cuenta_y_se_recorta_al_bloque_con_datos():
    import re
    import zipfile

    origen = zipfile.ZipFile(io.BytesIO(_libro(CUADRO)))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for item in origen.infolist():
            datos = origen.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                datos = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:XFD5000"', datos)
            z.writestr(item, datos)
    (hoja,) = carga_archivos.medir_libro(salida.getvalue(), "xlsx")
    assert (hoja["filas"], hoja["columnas"]) == (5000, 16384)
    # La lectura liviana rellenaría cada fila hasta ese ancho: antes se rechazaba; ahora se corrige la dimensión.
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([("dim.xlsx", salida.getvalue())])
    assert len(datos) == 3
    (aviso,) = [a for a in lectura[0]["avisos"] if "rango usado" in a]
    assert aviso == ("«dim.xlsx», hoja «Cuadro 1»: el rango usado llegaba hasta la fila 5.000 y la columna XFD; "
                     "se leyó el bloque con datos, hasta la celda C4.")
