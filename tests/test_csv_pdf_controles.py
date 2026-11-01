"""Controles por celda sobre tablas de CSV y de PDF con texto, y lectura de encabezados dudosos.

Archivos construidos en el test, con temas variados; las reglas se apoyan sólo en la estructura.
"""

import io
import random

import pytest
from openpyxl import Workbook

from services import carga_archivos, excel_cobertura
from services.controles_excel import tablas_planas
from services.excel_estructura import anios_plausibles, es_encabezado_anual
from services.sanitizacion import sanear
from tests.flujo_confirmado import cargar_confirmado

pytest_plugins = ["tests.test_ejemplos_zip"]

ANIOS = list(range(2016, 2025))


def _es(n, dec=1):
    return f"{n:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _filas(bases, semilla=5, g=0.03):
    rnd = random.Random(semilla)
    return {k: [round(b * (1 + g) ** i * (1 + rnd.uniform(-0.01, 0.01)), 1) for i in range(len(ANIOS))] for k, b in bases.items()}


def _csv_ancho(filas, sep=";", es=True, rotulo="Rubro", total=None):
    fmt = _es if es else (lambda v: f"{v:.1f}")
    out = [sep.join([rotulo] + [str(a) for a in ANIOS])]
    out += [sep.join([k] + [fmt(v) for v in vs]) for k, vs in filas.items()]
    if total is not None:
        out.append(sep.join(["Total"] + [fmt(v) for v in total]))
    return ("\n".join(out) + "\n").encode("utf-8")


def _controles(nombre, contenido):
    datos, _, lectura = carga_archivos.parse_archivos_con_lectura([(nombre, contenido)])
    hs, resumen = tablas_planas.analizar([(nombre, contenido)], datos, lectura, lambda f: (f, {}), sanear)
    return [h for h in hs if h["evidencia"].get("naturaleza") == "control"], resumen


def _altas(hs):
    return [h for h in hs if h["severidad"] == "alta"]


MATRICULA = {"Inicial": 1450.0, "Primario": 4620.0, "Secundario": 3880.0, "Superior": 910.0, "Adultos": 520.0}
ENERGIA = {"Hidráulica": 3120.5, "Térmica": 8820.2, "Nuclear": 1010.7, "Eólica": 640.3, "Solar": 210.9, "Biomasa": 150.4}


def test_csv_signo_invertido_se_ubica_en_fila_y_columna():
    filas = _filas(MATRICULA)
    total = [round(sum(v[i] for v in filas.values()), 1) for i in range(len(ANIOS))]
    filas["Secundario"][5] = -filas["Secundario"][5]
    hs, resumen = _controles("matricula.csv", _csv_ancho(filas, total=total))
    assert resumen and resumen[0]["tipo"] == "csv"
    signo = [h for h in hs if h["evidencia"]["regla"] in {"signo_aislado", "signo_opuesto_columna"}]
    assert signo, [h["evidencia"]["regla"] for h in hs]
    ev = signo[0]["evidencia"]
    # Fila 4 del archivo (encabezado = fila 1), columna G = sexto año (2021).
    assert ev["celda"] == "G4" and ev["hoja"] == "matricula.csv"
    assert ev["ubicacion_legible"] == "matricula.csv · fila 4 («Secundario»), columna G («2021»)"
    assert "fila 4" in signo[0]["descripcion"] and "matricula.csv!G4" not in signo[0]["descripcion"]


def test_csv_valor_por_mil_y_total_que_no_cierra():
    filas = _filas(ENERGIA, semilla=11, g=0.04)
    filas["Nuclear"][4] = round(filas["Nuclear"][4] * 1000, 1)
    hs, _ = _controles("energia.csv", _csv_ancho(filas, sep=",", es=False, rotulo="Fuente"))
    assert any(h["evidencia"]["celda"] == "F4" for h in _altas(hs))

    filas = _filas(MATRICULA, semilla=2)
    total = [round(sum(v[i] for v in filas.values()), 1) for i in range(len(ANIOS))]
    filas["Primario"][3] += 700
    hs, _ = _controles("matricula.csv", _csv_ancho(filas, total=total))
    conc = [h for h in hs if h["evidencia"]["regla"] == "conciliacion_aritmetica"]
    assert conc and conc[0]["evidencia"]["celda"] == "E7"


def _largo(periodos, sep=","):
    base = {"Norte": 820, "Centro": 2400, "Sur": 610, "Cuyo": 700}
    out = [sep.join(["Trimestre", "Region", "Consultas"])]
    for k, t in enumerate(periodos):
        for r, b in base.items():
            out.append(sep.join([t, r, str(int(b * (1 + 0.01 * k) + (k * 7) % 13))]))
    return ("\n".join(out) + "\n").encode()


TRIMESTRES = [f"{y}-Q{q}" for y in (2022, 2023, 2024) for q in (1, 2, 3, 4)]


@pytest.mark.parametrize("periodos,regla", [
    (TRIMESTRES[:6] + [TRIMESTRES[5]] + TRIMESTRES[6:], "periodo_duplicado"),
    (TRIMESTRES[:7] + TRIMESTRES[8:], "periodo_faltante"),
])
def test_csv_formato_largo_periodo_duplicado_o_faltante(periodos, regla):
    hs, _ = _controles("consultas.csv", _largo(periodos))
    assert any(h["evidencia"]["regla"] == regla and h["evidencia"]["celda"].startswith("A") for h in hs)


def _mensual(signo_en=None):
    out = ["Mes;Llegadas;Pernoctes"]
    for k, m in enumerate(f"{y}-{mm:02d}" for y in (2023, 2024) for mm in range(1, 13)):
        ll = 12000 + 300 * k + (k % 3) * 150
        out.append(f"{m};{_es(-ll if k == signo_en else ll, 0)};{_es(ll * 2.7 + (k % 4) * 90, 0)}")
    return ("\n".join(out) + "\n").encode()


def test_csv_mensual_vertical_signo():
    hs, _ = _controles("turismo.csv", _mensual(signo_en=9))
    assert any(h["evidencia"]["celda"] == "B11" for h in _altas(hs))


PROGRAMAS = [("Conducción superior", 1600), ("Salud pública", 21000), ("Educación básica", 26000), ("Seguridad", 23500),
             ("Obras viales", 22000), ("Desarrollo social", 19800), ("Ciencia y técnica", 4100)]


def _presupuesto(factor_en=None, factor=1):
    out = ["Programa;Crédito vigente;Devengado;Pagado"]
    for k, (n, c) in enumerate(PROGRAMAS):
        cv = c * (1 + 0.01 * k)
        dv = cv * (0.86 + 0.01 * (k % 3))
        pg = dv * (0.9 - 0.01 * (k % 2))
        out.append(f"{n};{_es(cv)};{_es(dv * (factor if n == factor_en else 1))};{_es(pg)}")
    return ("\n".join(out) + "\n").encode()


def test_rubro_chico_en_toda_su_fila_no_es_valor_fuera_de_escala():
    hs, _ = _controles("ejecucion.csv", _presupuesto())
    assert not [h for h in hs if h["evidencia"]["regla"] == "valor_fuera_de_escala_columna"]


@pytest.mark.parametrize("fila,factor,celda", [("Salud pública", 10, "C3"), ("Obras viales", 1000, "C6")])
def test_celda_fuera_de_escala_de_su_fila_sigue_senalada(fila, factor, celda):
    hs, _ = _controles("ejecucion.csv", _presupuesto(fila, factor))
    assert any(h["evidencia"]["regla"] == "valor_fuera_de_escala_columna" and h["evidencia"]["celda"] == celda for h in hs)


CORRECTOS = [
    ("matricula_es.csv", lambda: _csv_ancho(_filas(MATRICULA), total=[round(sum(v[i] for v in _filas(MATRICULA).values()), 1) for i in range(len(ANIOS))])),
    ("energia_en.csv", lambda: _csv_ancho(_filas(ENERGIA, 11, 0.04), sep=",", es=False)),
    ("energia_tab.csv", lambda: _csv_ancho(_filas(ENERGIA, 13, 0.02), sep="\t", es=False)),
    ("consultas_largo.csv", lambda: _largo(TRIMESTRES)),
    ("turismo_mensual.csv", lambda: _mensual()),
    ("ejecucion.csv", lambda: _presupuesto()),
]


@pytest.mark.parametrize("nombre,construir", CORRECTOS)
def test_csv_correctos_sin_alarmas_altas(nombre, construir):
    hs, _ = _controles(nombre, construir())
    assert not _altas(hs), [h["descripcion"][:160] for h in _altas(hs)]


# --------------------------------------------------------------------------- PDF

def _pdf(encabezado, filas, total=True):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4))
    data = [encabezado] + [[k] + [_es(v) for v in vs] for k, vs in filas]
    if total is not None:
        data.append(["Total"] + [_es(v) for v in total])
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    st = getSampleStyleSheet()
    doc.build([Paragraph("Cuadro 1. Pasajeros por modo (miles)", st["Heading2"]), Spacer(1, 12), t])
    return buf.getvalue()


MODOS = {"Urbano": 5200, "Interurbano": 2100, "Ferrocarril": 870, "Aéreo": 1500, "Fluvial": 320}


def _modos(alterar=None):
    filas = _filas(MODOS, semilla=3)
    total = [round(sum(v[i] for v in filas.values()), 1) for i in range(len(ANIOS))]
    if alterar:
        alterar(filas)
    return _pdf(["Modo"] + [str(a) for a in ANIOS], list(filas.items()), total)


def test_pdf_correcto_sin_alarmas_y_total_que_no_cierra_ubicado_por_pagina():
    hs, resumen = _controles("transporte.pdf", _modos())
    assert resumen and resumen[0]["tipo"] == "pdf" and not _altas(hs)

    hs, _ = _controles("transporte.pdf", _modos(lambda f: f["Ferrocarril"].__setitem__(4, f["Ferrocarril"][4] + 400)))
    conc = [h for h in hs if h["evidencia"]["regla"] == "conciliacion_aritmetica"]
    assert conc
    ev = conc[0]["evidencia"]
    assert ev["pagina_pdf"] == 1 and ev["tabla_pdf"] == 1
    assert ev["ubicacion_legible"].startswith("transporte.pdf · página 1, cuadro 1, fila ")
    assert "columna «2020»" in ev["ubicacion_legible"]


def test_pdf_signo_invertido():
    hs, _ = _controles("transporte.pdf", _modos(lambda f: f["Aéreo"].__setitem__(6, -f["Aéreo"][6])))
    signo = [h for h in _altas(hs) if "signo" in h["evidencia"]["regla"]]
    assert signo and "«Aéreo»" in signo[0]["evidencia"]["ubicacion_legible"]


def test_pdf_trimestres_2025_t1_se_reconocen():
    trims = [f"{y}-T{q}" for y in (2024, 2025) for q in (1, 2, 3, 4)]
    filas = [("Ocupados", [12000 * (1 + 0.004 * k) for k in range(8)]), ("Desocupados", [900 * (1 - 0.01 * k) for k in range(8)]),
             ("Inactivos", [9000 * (1 + 0.002 * k) for k in range(8)])]
    pdf = _pdf(["Condición"] + trims, filas, total=None)
    hs, resumen = _controles("empleo.pdf", pdf)
    assert resumen and not _altas(hs)
    from services.reportes.markdown import _periodo_observado

    assert _periodo_observado(carga_archivos.parse_archivo("empleo.pdf", pdf)) == "2024-T1 a 2025-T4"


def test_periodo_observado_csv_trimestres_y_meses():
    from services.reportes.markdown import _periodo_observado

    assert _periodo_observado(carga_archivos.parse_archivo("c.csv", _largo(TRIMESTRES))) == "2022-T1 a 2024-T4"
    assert _periodo_observado(carga_archivos.parse_archivo("t.csv", _mensual())) == "2023-01 a 2024-12"


# --------------------------------------------------------------------------- recorrido completo

def test_recorrido_csv_suma_casos_de_control(flujo):
    import models
    from services import jobs

    client, factory, _ = flujo
    filas = _filas(MATRICULA)
    filas["Adultos"][6] = round(filas["Adultos"][6] * 1000, 1)
    r = cargar_confirmado(client, "/api/analisis-rapido", data={"periodo": "", "dominio": "general", "notificar": "false"},
                    files={"archivos": ("matricula.csv", _csv_ancho(filas), "text/csv")})
    assert r.status_code == 200, r.text
    while jobs._process_one():
        pass
    with factory() as db:
        informe = db.get(models.Informe, r.json()["informe_id"])
        controles = [h for h in informe.hallazgos if (h.evidencia or {}).get("naturaleza") == "control"]
        assert any(h.evidencia.get("celda") == "H6" and h.evidencia.get("ubicacion_legible") for h in controles)


# --------------------------------------------------------------------------- encabezados dudosos

def test_importes_cercanos_a_2000_no_son_encabezado_de_anios():
    assert not anios_plausibles([1920.5, 1934.2, 1950.0, 1961.3, 1980.4, 1996.7])
    assert not anios_plausibles([1920, 1934, 1961, 1996, 2047])
    assert anios_plausibles([2001, 2002, 2003, 2005, 2006])
    assert anios_plausibles([1990, 2000, 2010, 2020])
    assert anios_plausibles([2010, 2011, 2012.5, 2013, 2014])  # un año corrupto sigue siendo encabezado
    assert not es_encabezado_anual(["Dermatología", 1920.5, 1934.2, 1950.0, 1961.3, 1980.4, 1996.7])


def _turnos():
    wb = Workbook()
    wb.active.title = "Portada"
    ws = wb.create_sheet("Turnos")
    ws["A1"] = "Turnos otorgados por especialidad"
    ws.append([])
    ws.append([])
    ws.append(["Especialidad", "Ene", "Feb", "Mar", "Abr", "May", "Jun"])
    for n, vs in (("Dermatología", [1920.5, 1934.2, 1950.0, 1961.3, 1980.4, 1996.7]),
                  ("Cardiología", [2310.0, 2290.5, 2350.25, 2402.0, 2380.75, 2420.0]),
                  ("Pediatría", [3100.0, 3155.5, 3201.0, 3188.0, 3240.5, 3302.0]),
                  ("Traumatología", [1520.0, 1544.0, 1561.0, 1580.5, 1602.0, 1611.0]),
                  ("Oftalmología", [980.5, 1002.0, 1011.0, 1025.5, 1040.0, 1052.0])):
        ws.append([n] + vs)
    wb.create_sheet("Notas")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_fila_de_datos_con_importes_cercanos_a_2000_no_se_toma_como_encabezado():
    datos, meta = excel_cobertura.leer_libro(_turnos(), "turnos.xlsx")
    hoja = next(h for h in meta["hojas"] if h["hoja"] == "Turnos")
    assert hoja["filas_datos"] == 5
    assert {d.get("Especialidad") for d in datos} >= {"Dermatología", "Oftalmología"}


def test_encabezado_con_importes_se_declara_en_avisos():
    import pandas as pd

    detalle, avisos = {}, []
    crudo = pd.DataFrame([["Dermatología", 1920.5, 1934.2, 1950.0], ["Cardiología", 2310.0, 2290.5, 2350.25]])
    excel_cobertura._avisar_encabezado_con_datos([[0]], crudo, detalle, avisos)
    assert detalle["encabezado_dudoso"] == [1] and "parecen datos" in avisos[0]


# --------------------------------------------------------------------------- registro de lectura

def test_hoja_oculta_leida_no_figura_sin_leer_y_archivo_fallido_sin_controles():
    from types import SimpleNamespace

    from services.reportes import registro

    saneamiento = {
        "archivos_origen": ["libro.xlsx", "roto.xlsx"],
        "lectura_archivos": [
            {"archivo": "libro.xlsx", "estado": "leido", "tipo_detectado": "xlsx", "avisos": []},
            {"archivo": "roto.xlsx", "estado": "fallido", "motivo": "no es un Excel válido", "avisos": ["«roto.xlsx» no se procesó"]},
        ],
        "excel": {"hojas": [
            {"archivo": "libro.xlsx", "hoja": "Resumen", "estado": "incluida", "filas_preparadas": 3},
            {"archivo": "libro.xlsx", "hoja": "Detalle", "estado": "incluida", "filas_preparadas": 3,
             "visibilidad": {"estado_hoja": "hidden"}},
        ], "bloques": []},
    }
    informe = SimpleNamespace(saneamiento=saneamiento, datos=[{"a": 1}], hallazgos=[], periodo=None, devolucion=None,
                              _condiciones_lentes={"ajedrecista": {"estado": "evaluado"}})
    reg = registro.construir(informe, hallazgos=[], cobertura={})
    assert "hojas ocultas" not in " ".join(p["donde"] for p in reg["pendientes"])
    assert any("Detalle" in a and "se leyeron" in a for a in reg["avisos"])
    roto = next(a for a in reg["archivos"] if a["archivo"] == "roto.xlsx")
    assert roto["controles"] == []
    assert "Controles ejecutados" not in "\n".join(
        linea for linea in registro.lineas_markdown(reg) if "roto" in linea)


def _registros(semilla=4):
    import datetime

    rnd = random.Random(semilla)
    out = ["cuit;apellido;localidad;monto;fecha"]
    for i in range(300):
        d = datetime.date(2024, 1, 1) + datetime.timedelta(days=rnd.randint(0, 360))
        out.append(f"20{rnd.randint(10000000, 45000000)}{rnd.randint(0, 9)};Apellido{i};Loc{rnd.randint(1, 6)};"
                   f"{_es(rnd.uniform(5000, 90000), 2)};{d:%d/%m/%Y}")
    return ("\n".join(out) + "\n").encode()


def test_lista_de_registros_no_se_controla_como_serie():
    hs, resumen = _controles("beneficios.csv", _registros())
    assert not hs and not resumen
    filas = [[t, r, "10"] for t in TRIMESTRES for r in ("Norte", "Sur", "Este")]
    assert not tablas_planas.es_lista_de_registros(filas)
    filas_dup = filas[:9] + filas[6:9] + filas[9:]
    assert not tablas_planas.es_lista_de_registros(filas_dup)


def _serie_id(etiqueta, fila, valores, cuadro=1):
    return {"etiqueta": etiqueta, "archivo": "a.xlsx", "hoja": "H", "cuadro": cuadro, "fila": fila, "nivel": 0,
            "valores": valores, "periodos": list(range(2015, 2015 + len(valores))),
            "celdas": [f"{chr(66 + i)}{fila}" for i in range(len(valores))]}


def test_identidades_con_filtro_de_candidatos_igual_que_comparar_todos_los_pares():
    from services.controles_excel import identidades

    rnd = random.Random(8)
    base = [round(rnd.uniform(100, 900), 1) for _ in range(8)]
    otra = list(base)
    otra[5] = round(otra[5] * 1.2, 1)
    series = [_serie_id("Obras sociales", 4, base, cuadro=1), _serie_id("OO.SS.", 20, otra, cuadro=2)]
    series += [_serie_id(f"Rubro {i}", 30 + i, [round(rnd.uniform(100, 900), 1) for _ in range(8)], cuadro=2) for i in range(40)]
    hs = identidades.analizar(series)
    assert [h["evidencia"]["regla"] for h in hs] == ["identidad_entre_cuadros"]
    assert hs[0]["evidencia"]["celda"] == "G4"
    # Las parejas candidatas incluyen a la única identidad real y descartan el resto.
    valores = [identidades._valores(s) for s in series]
    assert identidades._candidatos(series, valores) == {(0, 1)}


def test_valor_igual_por_azar_en_columna_larga_no_se_senala():
    rnd = random.Random(1)
    poblacion = [round(rnd.uniform(100000, 200000), 1) for _ in range(2500)]
    assert len(set(poblacion)) < len(poblacion)  # hay coincidencias por azar
    out = ["Localidad;Población;Hogares"]
    for i, v in enumerate(poblacion):
        out.append(f"L{i};{_es(v, 1)};{_es(v / 3.1 + i % 7, 1)}")
    hs, _ = _controles("localidades.csv", ("\n".join(out) + "\n").encode())
    assert not [h for h in hs if h["evidencia"]["regla"] == "valor_identico_otra_fila"]
