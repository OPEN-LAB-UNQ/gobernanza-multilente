"""Conciliación por series en cuadros que antes no la recibían bien.

Paneles repetidos sin sangría («Total», «Varones», «Mujeres» con los mismos componentes debajo, instituto de estadística
OI_POBL_03): cada encabezado es el total de su panel. Años arriba de una fila que sólo dice la unidad
(«En millones de pesos corrientes», dirección de estadística Mendoza): los años son el eje y el cuadro llega a la
conciliación como series.
"""
import io

from openpyxl import Workbook

from services import carga_archivos, controles_excel, excel_cobertura
from services.controles_excel.conciliacion import analizar, relaciones_cuadro
from services.sanitizacion import sanear
from tests.test_formatos_adicionales import serie

COMPONENTES = ["Trabajan exclusivamente", "Estudian exclusivamente", "Trabajan y estudian", "No trabajan ni estudian"]


def _paneles(total_primero=(100,) * 6):
    filas, fila = [], 2
    for encabezado, partes in (("Total", (40, 20, 25, 15)), ("Varones", (45, 21, 24, 10)), ("Mujeres", (36, 19, 26, 19))):
        valores = list(total_primero) if encabezado == "Total" else [100] * 6
        filas.append(serie(encabezado, fila, valores))
        for rotulo, v in zip(COMPONENTES, partes):
            fila += 1
            filas.append(serie(rotulo, fila, [v] * 6))
        fila += 1
    return filas


def test_paneles_repetidos_arman_un_total_por_panel():
    relaciones = list(relaciones_cuadro(_paneles()))
    assert [(p["etiqueta"], [x["etiqueta"] for x in c], o) for p, c, o in relaciones] == [
        ("Total", COMPONENTES, "total_precede_componentes"),
        ("Varones", COMPONENTES, "panel_repetido"),
        ("Mujeres", COMPONENTES, "panel_repetido"),
    ]
    assert analizar(_paneles()) == []


def test_total_de_panel_alterado_se_concilia():
    hs = analizar(_paneles(total_primero=(100, 100, 100, 103, 100, 100)))
    assert [h["evidencia"]["ubicaciones_excel"][0] for h in hs] == ["Cuadro!E2"]


def test_sin_paneles_iguales_el_total_toma_lo_que_sigue():
    filas = [serie("Total", 2, [90] * 6), serie("Salud", 3, [10] * 6), serie("Educación", 4, [20] * 6),
             serie("Varones", 5, [30] * 6), serie("Vivienda", 6, [15] * 6), serie("Obras", 7, [15] * 6)]
    relaciones = list(relaciones_cuadro(filas))
    assert [(p["etiqueta"], len(c)) for p, c, _ in relaciones] == [("Total", 5)]


def _mendoza(alterar=False):
    anios = list(range(2007, 2015))
    partes = {"Relacionados con la actividad": [900 + 70 * k for k in range(8)],
              "Patrimoniales": [230 + 20 * k for k in range(8)],
              "Varios": [120 + 10 * k for k in range(8)]}
    if alterar:
        partes["Relacionados con la actividad"][5] /= 10  # 1.250 publicado como 125
    total = [sum(v[k] for v in partes.values()) for k in range(8)]
    if alterar:
        total[5] = 1250 + 330 + 170
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos"
    ws.cell(1, 1, "Ingresos tributarios por concepto. Mendoza")
    ws.cell(3, 1, "Ingresos Tributarios")
    for j, a in enumerate(anios, 2):
        ws.cell(3, j, a)
    ws.cell(4, 2, "En millones de pesos corrientes")
    ws.merge_cells("A3:A4")
    ws.merge_cells("B4:I4")
    for i, (rotulo, valores) in enumerate([("Total", total), *partes.items()], 5):
        ws.cell(i, 1, rotulo)
        for j, v in enumerate(valores, 2):
            ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _altas(contenido, nombre="mendoza.xlsx"):
    datos, cobertura = carga_archivos.parse_archivos_con_cobertura([(nombre, contenido)])
    saneados, reporte, _ = excel_cobertura.preparar(datos, cobertura, lambda f: (f, {}), sanear)
    hs = controles_excel.analizar([(nombre, contenido)], saneados, reporte["excel"])
    return {u for h in hs if h["severidad"] == "alta" for u in h["evidencia"].get("ubicaciones_excel") or []}


def test_anios_sobre_la_fila_de_unidad_llegan_a_la_conciliacion():
    assert _altas(_mendoza()) == set()
    assert _altas(_mendoza(alterar=True)) & {"Ingresos!G5", "Ingresos!G6"}
