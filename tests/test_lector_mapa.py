"""Lector del cuadro: la IA lee a ciegas qué relaciones debería cumplir el cuadro y las reglas las verifican."""
import io
from pathlib import Path

from services import lector_mapa as lm

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre",
         "Noviembre", "Diciembre"]


def _cuadro(error_en=None):
    """Altas, bajas, neto (= altas - bajas) y acumulado del neto, filas 3 a 14; contacto@ejemplo.gob en la nota."""
    celdas = {"A1": "Cuadro 1. Altas y bajas mensuales 2019", "A2": "Mes", "B2": "Altas", "C2": "Bajas",
              "D2": "Neto", "E2": "Acumulado neto", "A16": "Consultas: contacto@ejemplo.gob"}
    acumulado = 0.0
    for i, mes in enumerate(MESES):
        f = i + 3
        altas, bajas = 1000.0 + 37 * i, 400.0 + 11 * i
        neto = altas - bajas
        acumulado += neto
        celdas.update({f"A{f}": mes, f"B{f}": altas, f"C{f}": bajas, f"D{f}": neto, f"E{f}": acumulado})
    if error_en:
        celdas[error_en] = celdas[error_en] * 10
    return celdas


NETO = {"tipo": "columnas", "destino": "D", "expresion": "B - C", "filas": [3, 14], "fundamento": "neto = altas - bajas"}
ACUM = {"tipo": "acumulado", "destino": "E", "componente": "D", "eje": "filas", "filas": [3, 14],
        "fundamento": "acumulado del neto"}


def _lectoras(por_lectora, registro=None):
    def fn(lectora, prompt):
        if registro is not None:
            registro.append((lectora, prompt))
        r = por_lectora.get(lectora)
        if isinstance(r, Exception):
            raise r
        return {"que_es": "altas y bajas", "relaciones": r or []}, f"modelo-{lectora}"
    return fn


def _saneamiento(celdas):
    return {"grillas_lectura": [{"archivo": "a.xlsx", "hoja": "H", "celdas": celdas}]}


def test_la_lectura_es_a_ciegas_y_sin_datos_personales():
    registro = []
    lm.analizar(_saneamiento(_cuadro()), _lectoras({}, registro), ["gpt"], "activo")
    prompt = registro[0][1].split("GRILLA:")[1]
    assert "B3:E3=#" in prompt and "B14:E14=#" in prompt and "2019" in prompt
    assert "1000" not in prompt and "600" not in prompt
    assert "contacto@ejemplo.gob" not in prompt
    assert lm.ciego("12,5%") == "#" and lm.ciego(2019.0) == "2019" and lm.ciego(1234.5) == "#"


def test_el_interprete_es_cerrado():
    vals = {"B": 10.0, "C": 4.0, "31": 50.0, "32": 25.0}
    assert lm.calcular("B - C", vals.get, "columnas") == (6.0, False)
    assert lm.calcular("(B + C) * 2", vals.get, "columnas") == (28.0, True)
    assert lm.calcular("32 / 31 * 100", vals.get, "filas") == (50.0, True)
    assert lm.calcular("32 * 100 / 31", vals.get, "filas") == (50.0, True)
    for malo in ("__import__('os').system('x')", "B +", "B C", "1 + 2", "open(B)"):
        assert lm.calcular(malo, vals.get, "columnas") == (None, False), malo
    assert lm.calcular("B + 31", vals.get, "filas") == (None, False)  # letras en una relación por filas
    assert lm.calcular("B / (C - C)", vals.get, "columnas") == (None, False)


def test_dos_lectoras_confirman_y_la_celda_que_no_cierra_es_senal():
    r = lm.analizar(_saneamiento(_cuadro(error_en="D9")), _lectoras({"gpt": [NETO], "gemini": [NETO], "deepseek": []}),
                    ["gpt", "gemini", "deepseek"], "activo")
    celdas = {s["evidencia"]["celda"] for s in r["senales"] if s["modo"] == "relacion_leida_no_cierra"}
    assert celdas == {"D9"}
    s = next(x for x in r["senales"] if x["evidencia"]["celda"] == "D9")
    assert s["severidad"] == "alta" and s["evidencia"]["lectura_confirmada_por"] == "numeros"
    assert s["evidencia"]["calculado"] == 756.0 and s["evidencia"]["declarado"] == 7560.0
    assert sorted(s["evidencia"]["lectoras"]) == ["gemini", "gpt"]
    assert "la cuenta no cierra: con las demás celdas da 756 y la celda dice 7.560" in s["descripcion"]
    assert "Alguna de las celdas que intervienen no está bien." in s["descripcion"] and "debería" not in s["descripcion"]


def test_el_acumulado_tambien_se_verifica():
    r = lm.analizar(_saneamiento(_cuadro(error_en="E9")), _lectoras({"gpt": [ACUM], "gemini": [ACUM]}),
                    ["gpt", "gemini"], "activo")
    # E9 no es E8 + D9, y E10 no es E9 + D10: las dos celdas que tocan el dato.
    assert {s["evidencia"]["celda"] for s in r["senales"]} == {"E9", "E10"}


def test_una_sola_lectora_no_alcanza():
    r = lm.analizar(_saneamiento(_cuadro(error_en="D9")), _lectoras({"gpt": [NETO], "gemini": []}),
                    ["gpt", "gemini"], "activo")
    assert r["senales"] == [] and r["estado"]["hojas"][0]["grupos"] == []


def test_si_falla_en_muchas_es_la_lectura_la_que_esta_mal():
    mal_leido = dict(NETO, expresion="B + C")
    r = lm.analizar(_saneamiento(_cuadro()), _lectoras({"gpt": [mal_leido], "gemini": [mal_leido]}),
                    ["gpt", "gemini"], "activo")
    assert r["senales"] == []
    assert r["estado"]["hojas"][0]["grupos"][0]["estado"] == "no_confirmada"


def test_con_pocas_aplicaciones_se_pregunta_y_solo_el_si_da_senal():
    total = {"tipo": "filas", "destino": 15, "expresion": " + ".join(str(f) for f in range(3, 15)), "columnas": ["B"],
             "fundamento": "la fila 15 es el total del año"}
    celdas = _cuadro()
    celdas.update({"A15": "Total", "B15": sum(celdas[f"B{f}"] for f in range(3, 15)) + 1000})
    fn = _lectoras({"gpt": [total], "gemini": [total]})
    r = lm.analizar(_saneamiento(celdas), fn, ["gpt", "gemini"], "activo")
    assert r["senales"] == [] and len(r["estado"]["preguntas"]) == 1
    pregunta = r["estado"]["preguntas"][0]
    assert pregunta["respuesta"] is None and "total del año" in pregunta["fundamento"]
    san = dict(_saneamiento(celdas), mapa_lectura=dict(r["estado"], confirmaciones={pregunta["id"]: {"confirma": True}}))
    r2 = lm.analizar(san, fn, ["gpt", "gemini"], "activo")
    assert [s["evidencia"]["celda"] for s in r2["senales"]] == ["B15"]
    assert r2["senales"][0]["evidencia"]["lectura_confirmada_por"] == "persona"
    san["mapa_lectura"]["confirmaciones"] = {pregunta["id"]: {"confirma": False}}
    assert lm.analizar(san, fn, ["gpt", "gemini"], "activo")["senales"] == []


def test_en_observacion_guarda_lo_que_encontraria_sin_senales():
    r = lm.analizar(_saneamiento(_cuadro(error_en="D9")), _lectoras({"gpt": [NETO], "gemini": [NETO]}),
                    ["gpt", "gemini"], "observar")
    assert r["senales"] == [] and r["estado"]["senales"] == 1 and r["estado"]["modo"] == "observar"


def test_apagado_o_sin_grilla_no_hace_nada():
    fn = _lectoras({"gpt": [NETO]})
    assert lm.analizar(_saneamiento(_cuadro()), fn, ["gpt"], "apagado") is None
    assert lm.analizar({}, fn, ["gpt"], "activo") is None


def test_si_la_grilla_no_cambio_no_se_vuelve_a_consultar():
    registro = []
    fn = _lectoras({"gpt": [NETO], "gemini": [NETO]}, registro)
    san = _saneamiento(_cuadro(error_en="D9"))
    r = lm.analizar(san, fn, ["gpt", "gemini"], "activo")
    assert len(registro) == 2
    r2 = lm.analizar(dict(san, mapa_lectura=r["estado"]), fn, ["gpt", "gemini"], "activo")
    assert len(registro) == 2 and len(r2["senales"]) == len(r["senales"])


def test_una_lectora_que_falla_no_frena_a_las_demas():
    r = lm.analizar(_saneamiento(_cuadro(error_en="D9")),
                    _lectoras({"gpt": [NETO], "gemini": [NETO], "deepseek": RuntimeError("sin servicio")}),
                    ["gpt", "gemini", "deepseek"], "activo")
    assert [s["evidencia"]["celda"] for s in r["senales"]] == ["D9"]
    assert "RuntimeError" in r["estado"]["hojas"][0]["mapas"]["deepseek"]["error"]


def test_tipos_fuera_del_catalogo_se_descartan():
    raro = {"tipo": "codigo", "expresion": "import os"}
    r = lm.analizar(_saneamiento(_cuadro()), _lectoras({"gpt": [raro, NETO]}), ["gpt"], "activo")
    assert r["estado"]["hojas"][0]["mapas"]["gpt"]["relaciones"] == [NETO]


def test_grilla_desde_un_libro_y_un_csv():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadro"
    for ref, v in _cuadro().items():
        ws[ref] = v
    oculta = wb.create_sheet("Oculta")
    oculta.sheet_state = "hidden"
    for f in range(1, 10):
        oculta[f"A{f}"] = f
    buf = io.BytesIO()
    wb.save(buf)
    texto = "Mes;Altas;Bajas\n" + "\n".join(f"{m};1.0{i:02d};4{i:02d}" for i, m in enumerate(MESES))
    g = lm.grillas([("libro.xlsx", buf.getvalue()), ("serie.csv", texto.encode("utf-8"))])
    hojas = {(x["archivo"], x["hoja"]) for x in g if x["celdas"]}
    assert hojas == {("libro.xlsx", "Cuadro"), ("serie.csv", "serie.csv")}
    oculta = next(x for x in g if x["hoja"] == "Oculta")
    assert oculta["celdas"] == {} and "oculta" in oculta["recorte"]
    cuadro = next(x for x in g if x["hoja"] == "Cuadro")["celdas"]
    assert cuadro["D3"] == 600.0 and cuadro["A3"] == "Enero"
    csv = next(x for x in g if x["hoja"] == "serie.csv")["celdas"]
    assert csv["B2"] == 1000.0 and csv["C3"] == 401.0


def test_el_motor_solo_verifica_lecturas_anteriores_a_la_confirmacion():
    raiz = Path(__file__).resolve().parent.parent
    fuente = (raiz / "services" / "analisis.py").read_text(encoding="utf-8")
    assert "lector_mapa.analizar(" not in fuente
    assert fuente.index("ip.exigir(informe)") < fuente.index("flujo = ejecutar_lentes_stream")
    assert fuente.index("ip.verificar_relaciones(") < fuente.index('origen_excel = (informe.saneamiento or {}).get("origen_excel")')
    from config import Settings

    assert Settings.model_fields["lector_ia"].default == "apagado"


def test_lee_todas_las_hojas_en_orden_y_declara_lo_que_no_lee():
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for i in range(9):
        ws = wb.create_sheet(f"Año {2016 + i}")
        for ref, v in _cuadro().items():
            ws[ref] = v
    formato = wb["Año 2016"]
    formato["AMJ900"].number_format = "0.00"  # vacía con formato: no es contenido sin leer
    ws = wb.create_sheet("Larga")
    ws["A1"], ws["B1"] = "Mes", "Valor"
    for f in range(2, lm.MAX_FILAS + 21):
        ws[f"A{f}"], ws[f"B{f}"] = f"p{f}", float(f)
    buf = io.BytesIO()
    wb.save(buf)
    g = lm.grillas([("libro.xlsx", buf.getvalue())])
    assert [x["hoja"] for x in g] == [f"Año {2016 + i}" for i in range(9)] + ["Larga"]
    larga = g[-1]
    assert f"filas {lm.MAX_FILAS + 1} a {lm.MAX_FILAS + 20}" in larga["recorte"]
    r = lm.analizar({"grillas_lectura": g}, _lectoras({"gpt": [NETO], "gemini": [NETO]}), ["gpt", "gemini"], "activo")
    leidas = {h["hoja"] for h in r["estado"]["hojas"]}
    assert leidas == {x["hoja"] for x in g}
    assert r["estado"]["sin_leer"] == [{"archivo": "libro.xlsx", "hoja": "Larga", "motivo": larga["recorte"]}]


def test_una_hoja_larga_se_lee_por_tramos_con_el_encabezado():
    celdas = {"A1": "Serie mensual", "A2": "Mes", "B2": "Altas"}
    for f in range(3, lm.VENTANA * 2 + 10):
        celdas[f"A{f}"], celdas[f"B{f}"] = f"m{f}", float(f)
    assert lm.tramos(celdas) == [(1, lm.VENTANA), (lm.VENTANA + 1, 2 * lm.VENTANA), (2 * lm.VENTANA + 1, 2 * lm.VENTANA + 9)]
    segundo = lm.texto_tramo(celdas, lm.VENTANA + 1, 2 * lm.VENTANA)
    assert segundo.startswith("fila 1 (encabezado): A1=Serie mensual") and "fila 2 (encabezado): A2=Mes | B2=Altas" in segundo
    assert f"fila {lm.VENTANA + 1}: " in segundo and f"fila {lm.VENTANA}: " not in segundo
    registro = []
    lm.analizar({"grillas_lectura": [{"archivo": "a.csv", "hoja": "a.csv", "celdas": celdas}]},
                _lectoras({"gpt": []}, registro), ["gpt"], "activo")
    assert len(registro) == 3


def test_la_misma_relacion_con_otro_rango_o_en_otro_orden_es_acuerdo():
    corrida = dict(NETO, expresion="B-C", filas=[5, 14])
    assert lm.clave(NETO) == lm.clave(corrida) and lm.clave(dict(NETO, expresion="-C+B")) != lm.clave(NETO)
    suma = {"tipo": "columnas", "destino": "B", "expresion": "C + D + E", "filas": [3, 9]}
    assert lm.clave(suma) == lm.clave(dict(suma, expresion="E+C+D", filas=[4, 30]))
    # Rangos 3-14 y 11-14: pocas celdas en común (Jaccard < 0,5), pero es la misma relación.
    r = lm.analizar(_saneamiento(_cuadro(error_en="D12")),
                    _lectoras({"gpt": [NETO], "gemini": [], "deepseek": [dict(NETO, expresion="B - C", filas=[11, 14])]}),
                    ["gpt", "gemini", "deepseek"], "activo")
    g = r["estado"]["hojas"][0]["grupos"]
    assert len(g) == 1 and sorted(g[0]["lectoras"]) == ["deepseek", "gpt"]
    assert [s["evidencia"]["celda"] for s in r["senales"]] == ["D12"]


def test_gemini_por_vertex_usa_el_archivo_de_credenciales(monkeypatch):
    import os
    import sys
    import types

    from config import settings

    visto = {}

    class Respuesta:
        candidates = [types.SimpleNamespace(content=types.SimpleNamespace(parts=[
            types.SimpleNamespace(text="pensando...", thought=True),
            types.SimpleNamespace(text='{"relaciones": []}', thought=False)]))]
        text = "no se usa"

    class Cliente:
        def __init__(self, **kw):
            visto.update(kw, credenciales=os.environ.get("GOOGLE_APPLICATION_CREDENTIALS"))
            self.models = types.SimpleNamespace(generate_content=lambda **k: Respuesta())

    falso = types.ModuleType("google.genai")
    falso.Client = Cliente
    monkeypatch.setitem(sys.modules, "google.genai", falso)
    import google

    monkeypatch.setattr(google, "genai", falso, raising=False)
    monkeypatch.setattr(settings, "google_genai_use_enterprise", True)
    monkeypatch.setattr(settings, "google_cloud_project", "proyecto")
    monkeypatch.setattr(settings, "google_application_credentials", "/ruta/credenciales.json")
    monkeypatch.delenv("GOOGLE_APPLICATION_CREDENTIALS", raising=False)
    mapa, _ = lm.llamar("gemini", "pedido")
    assert mapa == {"relaciones": []}
    assert visto["vertexai"] is True and visto["credenciales"] == "/ruta/credenciales.json"


def test_relaciones_distintas_no_son_acuerdo_aunque_compartan_celdas():
    celdas = {}
    for f in range(1, 11):
        celdas.update({f"A{f}": 30.0, f"B{f}": 10.0, f"C{f}": 20.0, f"D{f}": 20.0})
    celdas["A10"] = 99.0
    suma = dict(tipo="columnas", destino="A", expresion="B+C", filas=[1, 10])
    distinta = dict(suma, expresion="B+D")
    mapas = {"gpt": {"relaciones": [suma]}, "gemini": {"relaciones": [distinta]}}
    assert lm.acordar(mapas, celdas) == []
    # Un tercer modelo que coincide realmente debe seguir detectando el error.
    mapas["deepseek"] = {"relaciones": [dict(suma, expresion="C + B")]}
    grupos = lm.acordar(mapas, celdas)
    assert len(grupos) == 1
    assert set(grupos[0]["lectoras"]) == {"gpt", "deepseek"}
    assert grupos[0]["estado"] == "confirmada"
    assert grupos[0]["fallas"] == [{"celda": "A10", "declarado": 99.0, "calculado": 30.0}]
    celdas["A10"] = 30.0
    assert lm.acordar(mapas, celdas)[0]["fallas"] == []


def test_historial_sobrevive_a_reanalisis_y_cambios_de_respuesta():
    celdas = {"B3": 10.0, "B4": 20.0, "B15": 99.0}
    total = dict(tipo="filas", destino=15, expresion="3+4", columnas=["B"])
    fn = _lectoras({"gpt": [total], "gemini": [total]})
    san = _saneamiento(celdas)
    san["mapa_lectura"] = lm.analizar(san, fn, ["gpt", "gemini"], "activo")["estado"]
    pregunta_id = san["mapa_lectura"]["preguntas"][0]["id"]
    historial = []
    for respuesta in (True, False, True):
        san, _ = lm.confirmar(san, pregunta_id, respuesta, 7)
        historial.append(san["mapa_lectura"]["historial_confirmaciones"][-1])
        # JSON round-trip: igual que guardar y recuperar el informe.
        import json
        san = json.loads(json.dumps(san))
        for _ in range(2):
            r = lm.analizar(san, fn, ["gpt", "gemini"], "activo")
            assert r["estado"]["historial_confirmaciones"] == historial
            assert r["estado"]["confirmaciones"][pregunta_id]["confirma"] is respuesta
            assert [s["evidencia"]["celda"] for s in r["senales"]] == (["B15"] if respuesta else [])
            san["mapa_lectura"] = r["estado"]
    assert [h["confirma"] for h in historial] == [True, False, True]
    assert all(h["usuario_id"] == 7 and h["fecha"] for h in historial)
