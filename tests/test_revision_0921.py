"""0.9.21: correcciones de la revisión independiente de los arreglos de falsas alarmas.

Cada exención nueva tenía que perdonar sólo lo que el cuadro declara, no tapar errores verdaderos:
un ×10 en una serie de signo mixto, un negativo en una tasa o en un impuesto, un «-» que falta en una
sola serie, semestres «I» y «II» leídos como trimestres.
"""
import importlib.util
from pathlib import Path

from openpyxl import Workbook

from services.controles_excel import modelo
from services.controles_excel.series import admite_signo
from services.lentes import ajedrecista

_spec = importlib.util.spec_from_file_location("reales_0921", Path(__file__).with_name("test_planillas_reales_0921.py"))
reales = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(reales)


def _cuadro(filas, titulo="Cuadro 3"):
    w = Workbook()
    s = w.active
    s.title = titulo
    s.append(["Concepto", *range(2010, 2020)])
    for f in filas:
        s.append(f)
    return w


def test_pico_por_diez_en_serie_de_signo_mixto_entre_vecinos_normales_se_detecta():
    w = _cuadro([
        ["Cuenta corriente", 1200, 1350, -300, 1100, 12500, 1150, -200, 1300, 1250, 1400],
        ["Exportaciones", 8100, 8350, 8420, 8610, 8790, 8900, 9050, 9210, 9400, 9550],
        ["Importaciones", 6900, 7000, 8720, 7510, 7690, 7750, 9250, 7910, 8150, 8150],
    ])
    hs = reales.analizar_libro(w)
    assert any("Cuadro 3!F2" in (h["evidencia"].get("ubicaciones_excel") or []) or h["evidencia"].get("celda") == "F2"
               for h in reales._altas(hs)), [(h["evidencia"]["regla"], h["evidencia"].get("celda")) for h in hs]


def test_filas_que_admiten_negativos_son_solo_resultados_saldos_y_financiamiento():
    admite = ajedrecista.FILA_ADMITE_NEGATIVOS
    for rotulo in ("resultado financiero", "fuentes financieras", "aplicaciones financieras", "financiamiento neto",
                   "saldo", "deficit primario", "variacion de existencias"):
        assert admite.search(rotulo), rotulo
    for rotulo in ("tasa de desocupacion (%)", "hogares bajo la linea de pobreza (%)", "impuesto a las ganancias",
                   "fuentes de agua potable", "margen bruto", "varios"):
        assert not admite.search(rotulo), rotulo


def test_impuesto_a_las_ganancias_negativo_en_columna_total_sigue_siendo_imposible():
    w = reales._esquema()
    w.active.append(["Impuesto a las Ganancias", -400.0, -100.0, -500.0])
    hs = ajedrecista.imposibilidad(reales._datos(w))
    assert any(h["evidencia"]["regla"] == "valor_negativo_en_magnitud_no_negativa" for h in hs)


def test_rotulos_de_fila_que_admiten_signo():
    for rotulo in ("Transacciones financieras", "Préstamo neto", "Saldo de la cuenta corriente", "Resultado primario",
                   "Revalorizaciones durante el período", "Errores y omisiones netos", "Déficit fiscal"):
        assert admite_signo(rotulo), rotulo
    for rotulo in ("Impuesto a las Ganancias", "Margen bruto", "Ajuste por inflación", "Saldo de deuda bruta",
                   "Otros servicios", "Servicios financieros", "Internet"):
        assert not admite_signo(rotulo), rotulo


def test_marca_repetida_en_una_sola_serie_no_es_convencion():
    hs = reales.analizar_libro(reales._onp([["s/i", "s/i", 0.0571], [1.31, 1.38]]))
    assert any(h["evidencia"]["regla"] == "cambio_tipo" for h in reales._altas(hs)), [h["evidencia"]["regla"] for h in hs]


def test_semestres_i_ii_no_se_leen_como_trimestres():
    w = Workbook()
    ws = w.active
    ws.title = "encuesta de hogares"
    ws.append(["Tasas por semestre"])
    ws.append([])
    ws.append([])
    ws.append([None, 2021, None, 2022, None, 2023, None])
    ws.append([None, "I", "II", "I", "II", "I", "II"])
    for rotulo, base in (("Actividad", 46.0), ("Empleo", 42.0), ("Desocupación", 8.0)):
        ws.append([rotulo, *[round(base + k * 0.3, 1) for k in range(6)]])
    c = modelo.armar("libro.xlsx", ws.title, 1, ws, reales._LibroSimple(ws), 4, list(range(6, 9)),
                     [(j, ws.cell(4, j).value or f"columna_{j}") for j in range(1, ws.max_column + 1)])
    assert not c.get("subperiodos")


def _variacion(marcas):
    w = Workbook()
    s = w.active
    s.title = "Cuadro 2.2"
    s.append(["Concepto", *range(2008, 2024)])
    s.append(["Producción", *[1000 + 37 * k for k in range(16)]])
    s.append(["Variación porcentual", *marcas, *[3.7, 3.4, 3.3, 3.2, 3.1, 3.0, 2.9, 2.8, 2.7, 2.7, 2.6, 2.5, 2.5][: 16 - len(marcas)]])
    s.append(["Consumo", *[800 + 21 * k for k in range(16)]])
    return w


def test_marca_tres_veces_en_una_serie_es_convencion():
    # «///» en los primeros años de una columna de variación (instituto de estadística energía, Cuadro 2.2): convención.
    hs = reales.analizar_libro(_variacion(["///", "///", "///"]))
    assert "cambio_tipo" not in {h["evidencia"]["regla"] for h in reales._altas(hs)}


def test_marca_dos_veces_en_una_serie_sigue_siendo_senal():
    hs = reales.analizar_libro(_variacion(["///", 3.9, "///"]))
    assert "cambio_tipo" in {h["evidencia"]["regla"] for h in reales._altas(hs)}, [h["evidencia"]["regla"] for h in hs]
