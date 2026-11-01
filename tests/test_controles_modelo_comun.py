"""Modelo común de cuadro: ejes por contenido y controles sobre cada forma de cuadro."""

import pytest
from openpyxl import Workbook

from services import ambito
from services.controles_excel import periodos
from services.lentes import ajedrecista, conciliador, cronista
from tests.test_controles_tabulares import analizar_libro

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]


def reglas(hs, naturaleza="control"):
    return {(h["evidencia"]["regla"], h["evidencia"]["celda"]) for h in hs if h["evidencia"].get("naturaleza") == naturaleza}


def _mensual(ajuste=None, nota=None):
    w = Workbook()
    s = w.active
    s.title = "Mensual"
    s.append(["Gasto mensual por programa"])
    s.append(["En pesos"])
    s.append([])
    s.append(["Prog.", "Programa", *MESES])
    base = {"Salud": 100.0, "Educación": 230.0, "Obras": 57.0, "Vivienda": 410.0, "Cultura": 33.0}
    for k, (nombre, nivel) in enumerate(base.items()):
        fila = [10 + k, nombre]
        for m in range(12):
            v = round(nivel * (1 + 0.04 * m) * (1.8 if m == 11 else 1.0) * (1 + 0.01 * ((k * 7 + m * 3) % 5)), 2)
            fila.append(v)
        s.append(fila)
    if ajuste:
        ajuste(s)
    if nota:
        s.append([])
        s.append([nota])
    return w


# 1 y 5. Meses en el encabezado: el cuadro se lee aunque no tenga años.
def test_meses_en_encabezado_valor_x10_en_su_celda():
    hs = analizar_libro(_mensual(lambda s: s.__setitem__("H6", s["H6"].value * 10)))
    alertas = [h for h in hs if h["evidencia"].get("celda") == "H6" and h["evidencia"]["naturaleza"] == "control"]
    assert alertas and alertas[0]["severidad"] == "alta"


def test_meses_en_encabezado_sin_error_no_alerta():
    # El salto de diciembre es común a todos los programas (aguinaldo): no es un error de celda.
    hs = analizar_libro(_mensual())
    assert not [h for h in hs if h["evidencia"]["naturaleza"] == "control"]


def test_signo_documentado_en_nota_es_contexto():
    nota = "Nota: en marzo, el programa 12 registra una desafectación por regularización de devengados."
    hs = analizar_libro(_mensual(lambda s: s.__setitem__("E7", -40.0), nota))
    assert ("signo_aislado", "E7") in reglas(hs, "contexto")
    assert not reglas(hs)
    sin_nota = analizar_libro(_mensual(lambda s: s.__setitem__("E7", -40.0)))
    assert ("signo_aislado", "E7") in reglas(sin_nota)


# 2. Formato largo: secuencia controlada dentro de cada combinación de la clave.
def _largo(ajuste=None):
    w = Workbook()
    s = w.active
    s.title = "Datos"
    s.append(["Provincia", "Año", "Indicador", "Valor"])
    for p, nivel in (("Norte", 10000.0), ("Sur", 25000.0), ("Este", 7000.0), ("Oeste", 18000.0)):
        for a in range(2017, 2026):
            s.append([p, a, "Matrícula", round(nivel * (1 + 0.02 * (a - 2017)), 1)])
            s.append([p, a, "Tasa (%)", 40 + (a - 2017) * 0.3])
    if ajuste:
        ajuste(s)
    return w


def test_formato_largo_correcto_no_repite_periodos_por_fila():
    hs = analizar_libro(_largo())
    assert not [h for h in hs if h["severidad"] == "alta"]


def test_formato_largo_x10_dentro_de_la_clave():
    # Fila 2 + 2 * (año - 2017) + 0 en la primera provincia: 2021 está en la fila 10.
    hs = analizar_libro(_largo(lambda s: s.__setitem__("D10", s["D10"].value * 10)))
    assert any(h["evidencia"]["celda"] == "D10" and h["severidad"] == "alta" for h in hs)


def test_formato_largo_periodo_duplicado_en_una_clave():
    hs = analizar_libro(_largo(lambda s: s.__setitem__("B12", 2021)))
    assert any(h["evidencia"]["regla"] == "periodo_duplicado" and h["evidencia"]["celda"] == "B12" for h in hs)
    assert len([h for h in hs if h["evidencia"]["regla"].startswith("periodo_")]) <= 2


# 3. Marcas y encabezados de dos niveles.
def test_anios_con_marca_son_metadato():
    w = Workbook()
    s = w.active
    s.append(["Concepto", "2019", "2020", "2021", "2022", "2023 (p)", "2024*"])
    for k in range(5):
        s.append([f"Rubro {k}", *[100 * (k + 1) * 1.1 ** i for i in range(6)]])
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] in ("anio_no_entero", "anio_duplicado")]
    assert periodos.leer("2023 (p)")["marca"] == "p" and periodos.leer("2024*")["anio"] == 2024
    assert cronista._periodo_columna("2023 (p)") == (2023, 1)


def test_encabezado_de_dos_niveles_no_es_anio_duplicado():
    w = Workbook()
    s = w.active
    s.append(["Concepto", "Crédito vigente", None, None, "Devengado", None, None])
    s.merge_cells("B1:D1")
    s.merge_cells("E1:G1")
    s.append([None, 2022, 2023, 2024, 2022, 2023, 2024])
    for k in range(5):
        s.append([f"Rubro {k}", *[100 * (k + 1) * 1.2 ** i for i in range(3)], *[90 * (k + 1) * 1.2 ** i for i in range(3)]])
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] == "anio_duplicado"]


def test_dos_niveles_declara_cronista_no_evaluado():
    from services.lentes import _motivo_periodos_repetidos
    motivo = _motivo_periodos_repetidos([{"Concepto": "x", "Crédito vigente 2022": 1, "Devengado 2022": 1}])
    assert "dos niveles" in motivo


# 4. Período duplicado en la columna de tiempo: alerta propia y análisis temporal no evaluable.
def _vertical_mensual(periodos_txt):
    w = Workbook()
    s = w.active
    s.append(["Período", "IVA", "Ganancias", "Total"])
    for i, p in enumerate(periodos_txt):
        a, b = 100 + 3 * i, 50 + 2 * i
        s.append([p, a, b, a + b])
    return w


MESES_ISO = [f"2024-{m:02d}" for m in range(1, 13)]


def test_periodo_duplicado_en_columna_de_tiempo():
    hs = analizar_libro(_vertical_mensual(MESES_ISO[:7] + ["2024-07"] + MESES_ISO[8:]))
    assert ("periodo_duplicado", "A9") in reglas(hs)
    datos = [{"Período": p, "IVA": 1.0 + i} for i, p in enumerate(MESES_ISO[:7] + ["2024-07"] + MESES_ISO[8:])]
    problemas = cronista.problemas_periodo(datos)
    assert problemas and problemas[0]["fila"] == 7
    alerta = cronista.alertas_periodo(datos)[0]
    assert alerta["severidad"] == "alta" and alerta["evidencia"]["filas"] == [7]


def test_periodo_mal_rotulado_fuera_de_secuencia():
    hs = analizar_libro(_vertical_mensual(MESES_ISO[:9] + ["2023-10"] + MESES_ISO[10:]))
    assert any(h["evidencia"]["celda"] == "A11" and h["evidencia"]["regla"].startswith("periodo_") for h in hs)


def test_validez_marca_cronista_no_evaluable_con_motivo():
    from types import SimpleNamespace

    from services.validez import aplicabilidad_lentes
    datos = [{"Período": p, "IVA": 10.0 + i, "Ganancias": 5.0 + i} for i, p in enumerate(MESES_ISO[:7] + ["2024-07"] + MESES_ISO[8:])]
    estados = aplicabilidad_lentes(SimpleNamespace(datos=datos, dominio="general"), [])
    assert estados["cronista"]["estado"] == "no_evaluable"
    assert "2024-07" in estados["cronista"]["motivo"]


# 5. Cambio de escala simultáneo y estacionalidad.
def _gasto_anual(ajuste):
    w = Workbook()
    s = w.active
    s.append(["Concepto", *range(2016, 2025)])
    for k in range(6):
        s.append([f"Rubro {k}", *[round(1000 * (k + 1) * 1.3 ** i * (1 + 0.02 * ((k + i) % 3)), 1) for i in range(9)]])
    ajuste(s)
    return w


def test_cambio_de_escala_sostenido_es_un_caso():
    def dividir(s):
        for fila in s.iter_rows(min_row=2, min_col=7, max_col=10):
            for c in fila:
                c.value = round(c.value / 1000, 3)
    hs = analizar_libro(_gasto_anual(dividir))
    escala = [h for h in hs if h["evidencia"]["regla"] == "cambio_escala_sostenido"]
    assert len(escala) == 1 and escala[0]["severidad"] == "alta"
    assert "Hoja!G2" in escala[0]["evidencia"]["celdas_participantes"] or any(
        u.endswith("!G2") for u in escala[0]["evidencia"]["celdas_participantes"])


def test_periodo_en_otra_escala_que_revierte():
    def dividir(s):
        for fila in s.iter_rows(min_row=2, min_col=6, max_col=6):
            for c in fila:
                c.value = round(c.value / 1000, 3)
    hs = analizar_libro(_gasto_anual(dividir))
    escala = [h for h in hs if h["evidencia"]["regla"] == "periodo_en_otra_escala"]
    assert len(escala) == 1
    assert not [h for h in hs if h["evidencia"]["regla"] == "pico_aislado" and h["evidencia"]["naturaleza"] == "control"]


def test_estacionalidad_del_mismo_mes_no_es_alerta():
    datos = []
    for a in (2024, 2025):
        for m in range(1, 13):
            datos.append({"periodo": f"{a}-{m:02d}", "Ganancias": 100.0 * (1.05 ** m) * (2.8 if m == 5 else 1.0) * (1.3 if a == 2025 else 1)})
    hs = cronista.analizar(datos)
    assert all(h["evidencia"].get("naturaleza") == "contexto" for h in hs)


# 6. Cuadros transversales.
def _provincias(ajuste=None):
    w = Workbook()
    s = w.active
    s.title = "Indicadores"
    s.append(["Código", "Jurisdicción", "Población", "Desocupación (%)", "Pobreza (% personas)", "Indigencia (% personas)",
              "0 a 14 (%)", "15 a 64 (%)", "65 y más (%)"])
    pob = [3100000, 17500000, 430000, 3800000, 1200000, 1100000, 600000, 1400000, 610000, 810000]
    for k, p in enumerate(pob):
        jov = 20 + k % 5
        s.append([2 + 4 * k, f"Provincia {k}", p, 6.0 + (k % 4) * 0.8, 28.0 + k, 4.0 + (k % 3), jov, 100 - jov - 12, 12])
    s.append([None, "Total país", sum(pob), 7.0, 30.0, 5.0, 22, 66, 12])
    if ajuste:
        ajuste(s)
    return w


def test_transversal_correcto_sin_alertas():
    assert not [h for h in analizar_libro(_provincias()) if h["evidencia"]["naturaleza"] == "control"]


@pytest.mark.parametrize("celda,valor,regla", [
    ("F5", 40.0, "indigencia_supera_pobreza"),
    ("D6", 60.0, "valor_fuera_de_escala_columna"),
    ("D7", -7.6, "signo_opuesto_columna"),
    ("C5", 3700000, "total_columna_no_cierra"),
])
def test_transversal_detecta_en_la_celda(celda, valor, regla):
    hs = analizar_libro(_provincias(lambda s: s.__setitem__(celda, valor)))
    encontrados = [h for h in hs if h["evidencia"]["regla"] == regla]
    assert encontrados and encontrados[0]["severidad"] == "alta"
    lugares = {u for h in encontrados for u in h["evidencia"]["celdas_participantes"]}
    assert f"Indicadores!{celda}" in lugares or "Indicadores!C13" in lugares


def test_porcentajes_que_no_suman_100():
    hs = analizar_libro(_provincias(lambda s: s.__setitem__("H4", s["H4"].value + 3)))
    h = [h for h in hs if h["evidencia"]["regla"] == "porcentajes_no_suman_100"]
    assert h and "Indicadores!H4" in h[0]["evidencia"]["celdas_participantes"]


def _balance(activo_extra=0.0, bruto=None):
    w = Workbook()
    er = w.active
    er.title = "Resultados"
    er.append(["Concepto", "Ejercicio 2025", "Ejercicio 2024"])
    er.append(["Ventas de servicios", 4800.0, 4500.0])
    er.append(["Costo de los servicios prestados", -3400.0, -3380.0])
    er.append(["Resultado bruto", 1400.0 if bruto is None else bruto, 1120.0])
    er.append(["Gastos de administración", -510.0, -490.0])
    er.append(["Gastos de comercialización", -180.0, -170.0])
    er.append(["Excedente del ejercicio", 710.0, 460.0])
    bal = w.create_sheet("Balance")
    bal.append(["Rubro", "30/06/2025", "30/06/2024"])
    filas = [("Caja", 180.0, 140.0), ("Créditos", 610.0, 580.0), ("Bienes de uso", 2800.0 + activo_extra, 2700.0),
             ("Total del activo", 3590.0 + activo_extra, 3420.0), ("Deudas comerciales", 500.0, 510.0),
             ("Deudas fiscales", 70.0, 60.0), ("Total del pasivo", 570.0, 570.0), ("Capital", 1250.0, 1250.0),
             ("Resultados acumulados", 1770.0, 1600.0), ("Total patrimonio neto", 3020.0, 2850.0)]
    for f in filas:
        bal.append(list(f))
    return w


def test_identidades_contables_por_rotulo():
    assert not [h for h in analizar_libro(_balance()) if h["evidencia"]["naturaleza"] == "control"]
    hs = analizar_libro(_balance(activo_extra=140.0))
    assert ("activo_distinto_pasivo_mas_pn", "B5") in reglas(hs)
    hs = analizar_libro(_balance(bruto=1500.0))
    assert ("resultado_bruto_no_cierra", "B4") in reglas(hs)


# 7. Ubicación de partes_no_suman_total en coordenadas del cuadro.
def test_partes_no_suman_total_informa_filas_y_columna():
    datos = [{"Período": f"2024-{m:02d}", "A": 10.0 + m, "B": 5.0 + m, "Total": 15.0 + 2 * m} for m in range(1, 9)]
    datos[5]["Total"] += 3
    h = next(h for h in ajedrecista.imposibilidad(datos) if h["evidencia"]["regla"] == "partes_no_suman_total")
    assert h["evidencia"]["filas"] == [5] and h["evidencia"]["columnas"] == ["Total"]
    assert "registro" not in h["descripcion"]
    from services.excel_cobertura import ubicar
    bloque = {"archivo": "a.xlsx", "hoja": "H", "bloque": 1, "inicio": 0, "filas": list(range(8, 16)),
              "columnas": {"Período": "A", "A": "B", "B": "C", "Total": "D"}}
    assert ubicar(h, bloque)["evidencia"]["ubicaciones_excel"] == ["H!D13"]


# 8. Acumulados sobre el eje temporal real.
def test_acumulado_del_anio_se_reinicia_y_retroceso_se_ubica():
    datos, acum = [], 0.0
    for a in (2024, 2025):
        acum = 0.0
        for m in range(1, 13):
            acum += 10.0 + m
            datos.append({"Período": f"{a}-{m:02d}", "Total": 10.0 + m, "Acumulado del año": acum})
    assert not ajedrecista.retrocesos_acumulados(datos, ["Total", "Acumulado del año"])
    datos[17]["Acumulado del año"] = datos[16]["Acumulado del año"] - 50
    r = ajedrecista.retrocesos_acumulados(datos, ["Total", "Acumulado del año"])
    assert [x[0] for x in r] == [17]
    h = [h for h in ajedrecista.imposibilidad(datos) if h["evidencia"]["regla"] == "serie_acumulada_retrocede"]
    assert h and h[0]["evidencia"]["filas"] == [17]


def test_metas_distintas_no_son_periodos_del_acumulado():
    datos = [{"Meta": f"Meta {k}", "Ejec. acum. T1": 100.0 * (5 - k), "Ejec. acum. T2": 200.0 * (5 - k),
              "Ejec. acum. T3": 300.0 * (5 - k), "Ejec. acum. T4": 400.0 * (5 - k)} for k in range(5)]
    assert conciliador._acumulados_que_retroceden(datos, list(datos[0])[1:]) == []
    datos[2]["Ejec. acum. T3"] = 150.0
    r = ajedrecista.retrocesos_acumulados(datos, list(datos[0])[1:])
    assert [(x[0], x[1]) for x in r] == [(2, "Ejec. acum. T3")]
    datos[2]["Ejec. acum. T3"] = 900.0
    datos[3]["Ejec. acum. T2"] = datos[3]["Ejec. acum. T1"]
    e = ajedrecista.estancamientos_acumulados(datos, list(datos[0])[1:])
    assert [x[0] for x in e] == [3]


# 9. Valores fijos donde el cuadro calcula y severidad de las relaciones presupuestarias.
def test_total_general_fijo_en_la_esquina_con_diferencia_chica():
    w = Workbook()
    s = w.active
    s.append(["Programa", "Ene", "Feb", "Mar", "Total anual"])
    for r in range(2, 7):
        s.append([f"P{r}", 1e9 * r, 1.1e9 * r, 1.2e9 * r, f"=SUM(B{r}:D{r})"])
    s.append(["Total general", "=SUM(B2:B6)", "=SUM(C2:C6)", "=SUM(D2:D6)", 6.6e10 + 2.5e7])
    hs = analizar_libro(w)
    h = [h for h in hs if h["evidencia"]["celda"] == "E7" and h["severidad"] == "alta"]
    assert h, [x["evidencia"]["regla"] for x in hs]


def test_devengado_mayor_vigente_aislado_es_alta():
    ambito.declarar("publico")
    datos = [{"Programa": f"P{k}", "Crédito vigente": 1000.0 + k, "Devengado": 900.0 + k} for k in range(8)]
    datos[3]["Devengado"] = 1060.0
    h = next(h for h in ajedrecista.imposibilidad(datos) if h["evidencia"]["regla"] == "devengado_supera_vigente")
    assert h["severidad"] == "alta" and h["evidencia"]["filas"] == [3]


# 10. Quiebre documentado en una nota: contexto declarado.
def test_quiebre_documentado_en_nota_es_contexto():
    datos = [{"periodo": f"{a}-{m:02d}", "Impuesto PAIS": (500.0 + 10 * m if a == 2024 else 0.0)}
             for a in (2024, 2025) for m in range(1, 13)]
    notas = [{"texto": "Notas: el Impuesto PAIS rigió hasta el 23/12/2024.", "celda": "A30", "hoja": "R"}]
    hs = cronista.analizar(datos, notas)
    assert hs and all(h["evidencia"].get("naturaleza") == "contexto" for h in hs)
    assert hs[0]["evidencia"].get("contexto_declarado", {}).get("celda") == "A30" or hs[0]["evidencia"].get("motivo_contexto")


def test_anio_que_clasifica_filas_no_es_anio_duplicado():
    # Formato largo: el año se repite en cada fila (caso real: IDB Data Wall 2026, F8.5 y F1.8).
    w = Workbook()
    s = w.active
    s.append(["Year", "Country", "Spread", "Yield"])
    for anio in (2022, 2023, 2024):
        for k, pais in enumerate(["AR", "BR", "CL", "MX"]):
            s.append([anio, pais, 300 + 10 * k + anio % 10, 150 + 5 * k])
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] == "anio_duplicado"]


def test_anio_repetido_una_vez_sigue_marcado():
    w = Workbook()
    s = w.active
    s.append(["Año", "Recaudación", "Gasto"])
    for anio in (2018, 2019, 2020, 2020, 2022, 2023):
        s.append([anio, 100 + anio % 100, 90 + anio % 100])
    hs = analizar_libro(w)
    assert [h for h in hs if h["evidencia"]["regla"] == "anio_duplicado"]


def test_serie_diaria_no_repite_periodo_por_mes():
    # Caso real: IDB Data Wall 2026, F8.5 (spread diario, varias fechas por mes).
    import datetime

    w = Workbook()
    s = w.active
    s.append(["Year", "Date", "Spread", "Yield"])
    dia = datetime.datetime(2019, 12, 2)
    for k in range(40):
        s.append([dia.year, dia, 300 + (k * 7) % 11, 180 - (k * 3) % 5])
        dia += datetime.timedelta(days=1)
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] in ("periodo_duplicado", "anio_duplicado")]


def test_fecha_diaria_repetida_sigue_marcada():
    from services.lentes.cronista import problemas_periodo

    datos = [{"Fecha": f"2020-01-{d:02d}", "Valor": 10 + d} for d in (1, 2, 3, 3, 4, 5)]
    assert [p["tipo"] for p in problemas_periodo(datos)] == ["periodo_duplicado"]


def test_anio_combinado_sobre_dos_columnas_no_es_anio_no_entero():
    # Caso real: IDB Data Wall 2026, F4.2 (cada año abarca «Men» y «Women»).
    w = Workbook()
    s = w.active
    s.append(["País", 1990, None, 2000, None, 2010, None, 2020, None])
    s.merge_cells("B1:C1")
    s.merge_cells("D1:E1")
    s.merge_cells("F1:G1")
    s.merge_cells("H1:I1")
    s.append([None, "Men", "Women", "Men", "Women", "Men", "Women", "Men", "Women"])
    for k, pais in enumerate(["AR", "BR", "CL", "MX", "PE"]):
        s.append([pais, *[60 + k + i for i in range(8)]])
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] == "anio_no_entero"]


def test_anio_como_dato_de_cada_pais_no_es_serie():
    # Caso real: IDB Data Wall 2026, F7.7 (año en que cada país llega al 10 % de población mayor).
    w = Workbook()
    s = w.active
    s.append(["Region", "ISO", "Año 10 %", "Año 20 %", "Años entre hitos"])
    filas = [("Europe", "EMU", 1966, 2022), ("Asia", "EAS", 2023, 2054), ("USA", "USA", 1966, 2028),
             ("LAC", "LCN", 2023, 2052), ("Haiti", "HTI", 2049, 2088), ("Guatemala", "GTM", 2046, 2071),
             ("Honduras", "HND", 2041, 2065), ("Bolivia", "BOL", 2037, 2068), ("Peru", "PER", 2028, 2054)]
    for region, iso, a, b in filas:
        s.append([region, iso, a, b, b - a])
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"].startswith("periodo_")]


def test_columnas_auxiliares_sin_encabezado_no_son_valores_identicos():
    # Caso real: IDB Data Wall 2026, F2.11 (puntos de una línea de referencia del gráfico).
    w = Workbook()
    s = w.active
    s.append(["Country", "Math", "Expenditure", None, None])
    for fila in [("Serbia", 439.87666, 45963.4307, 1000, 102611.9185), ("Netherlands", 492.67563, 119584.16, 0, 102611.9185),
                 ("Moldova", 414.19806, 35686.3408, 300000, 472.3576499), ("Palestine", 365.74667, 30891.3623, 0, 472.3576499),
                 ("Germany", 474.82645, 121061.54, None, None), ("Finland", 484.13922, 126777.01, None, None),
                 ("Brunei", 442.0931, 207324.225, None, None)]:
        s.append(list(fila))
    hs = analizar_libro(w)
    assert not [h for h in hs if h["evidencia"]["regla"] == "valor_identico_otra_fila"]
